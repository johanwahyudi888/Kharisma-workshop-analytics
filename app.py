import io,re,base64,os,json,hashlib,hmac,platform,uuid,datetime,shutil,calendar
import html as _html
import pandas as pd
import numpy as np
import streamlit as st
import plotly.express as px
import plotly.graph_objects as go
import ref_ui as RK   # komponen kartu & perhitungan halaman V17

st.set_page_config(page_title='Bengkel Analytics • Kharisma', page_icon=':material/build:', layout='wide', initial_sidebar_state='expanded')

# ---- Tema grafik global (hanya tampilan) ----
import plotly.io as pio
_kh_axis=dict(gridcolor='#eef1f6',linecolor='#e7eaf3',zerolinecolor='#e7eaf3',tickfont=dict(size=11,color='#6b7280'),title=dict(font=dict(size=11,color='#6b7280')))
pio.templates['kharisma']=go.layout.Template(layout=dict(
    font=dict(family="Inter, 'Segoe UI', Arial, sans-serif",size=12,color='#374151'),
    colorway=['#2f6fed','#16a34a','#8b5cf6','#f59e0b','#ec4899','#0d9488','#64748b'],
    paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',
    xaxis=_kh_axis,yaxis=_kh_axis,
    hoverlabel=dict(bgcolor='#111827',font=dict(color='#ffffff',size=12),bordercolor='#111827'),
    legend=dict(font=dict(size=11,color='#4b5563'),bgcolor='rgba(0,0,0,0)'),
    margin=dict(l=8,r=8,t=10,b=8)))
pio.templates.default='kharisma'
DEFAULT_BRANCHES=['Kharisma Makassar','Kharisma Kendari','Kharisma Gowa']
# Branch classification for the three branches that still share one source system.
# Service Advisor is the authoritative branch key when the SA is known.
SA_BRANCH_MAP={
    'JOTHA024':'Kharisma Kendari',
    'REMI002':'Kharisma Gowa',
    'FATHA001':'Kharisma Gowa',
    'JOHAN123':'Kharisma Makassar',
    'RYANF017':'Kharisma Makassar',
    'WIRA006':'Kharisma Makassar',
    'GUGUN013':'Kharisma Makassar',
    'MARIF015':'Kharisma Makassar',
}
EXTERNAL_BRANCHES=set()  # dulu menyaring Ambon sehingga cabang tidak bisa ditambah; kini semua cabang boleh ditambahkan
# Cabang yang masih satu sistem sumber (laporan gabungan): cabang per baris ditentukan dari SA / PKB / kode customer.
# Cabang di luar daftar ini (mis. Ambon) memakai sistem terpisah: seluruh isi laporan = cabang itu.
SHARED_SYSTEM_BRANCHES={'Kharisma Makassar','Kharisma Kendari','Kharisma Gowa'}
SA_BRANCH_HISTORY_PATH=os.path.join(os.path.expanduser('~'),'.kharisma_workshop_sa_branch_history.json')

_CUST_K=re.compile(r'[-\u2010\u2011\u2012\u2013\u2014\u2015\u2212]\s*K\s*$',re.IGNORECASE)
def branch_from_customer(name):
    """Penjualan langsung (tanpa PKB): kode unik di AKHIRAN nama customer.
    'CUSTOMER - K' / 'CUSTOMER-K' / 'CUSTOMER – K' -> Kendari. Selain itu (termasuk tanpa kode) -> Makassar."""
    t=re.sub(r'[\s\.]+$','',str(name or '')).strip()
    return 'Kharisma Kendari' if _CUST_K.search(t) else 'Kharisma Makassar'

DEDUP_STATS={}
def _load_sa_branch_history():
    try:
        with open(SA_BRANCH_HISTORY_PATH,'r',encoding='utf-8') as fh:
            data=json.load(fh)
        if not isinstance(data,list): return []
        out=[]
        for r in data:
            if not isinstance(r,dict): continue
            sa=norm_sa(r.get('sa','')) if 'norm_sa' in globals() else re.sub(r'\s+','',str(r.get('sa','')).strip().upper())
            branch=clean(r.get('branch',''))
            effective=clean(r.get('effective_from',''))
            if sa and branch and effective:
                dt=pd.to_datetime(effective,errors='coerce')
                if pd.notna(dt): out.append({'sa':sa,'branch':branch,'effective_from':dt.strftime('%Y-%m-%d')})
        return out
    except Exception:
        return []

def _save_sa_branch_history(rows):
    tmp=SA_BRANCH_HISTORY_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(rows,fh,ensure_ascii=False,indent=2)
    os.replace(tmp,SA_BRANCH_HISTORY_PATH)

def branch_from_sa(sa, fallback='Kharisma Makassar', when=None):
    key=norm_sa(sa) if 'norm_sa' in globals() else re.sub(r'\s+','',str(sa or '').strip().upper())
    if not key: return fallback
    rules=[r for r in _load_sa_branch_history() if r.get('sa')==key]
    if rules:
        dt=pd.to_datetime(when,errors='coerce') if when is not None else pd.NaT
        if pd.notna(dt):
            eligible=[r for r in rules if pd.to_datetime(r['effective_from'],errors='coerce')<=dt]
            if eligible:
                eligible.sort(key=lambda r:r['effective_from'])
                return eligible[-1]['branch']
            # Before the first mutation date, keep the original/static assignment.
            return SA_BRANCH_MAP.get(key, fallback)
        # No usable row date: use the latest configured assignment as current.
        rules.sort(key=lambda r:r['effective_from'])
        return rules[-1]['branch']
    return SA_BRANCH_MAP.get(key, fallback)

def _sa_branch_rows():
    rows=[]
    history=_load_sa_branch_history()
    for sa,branch in SA_BRANCH_MAP.items():
        rules=[r for r in history if r.get('sa')==sa]
        if rules:
            rules.sort(key=lambda r:r['effective_from'],reverse=True)
            rows.append({'Cabang':rules[0]['branch'],'Service Advisor':sa.lower(),'Berlaku Mulai':rules[0]['effective_from'],'Status':'Aktif'})
        else:
            rows.append({'Cabang':branch,'Service Advisor':sa.lower(),'Berlaku Mulai':'Default','Status':'Aktif'})
    for r in history:
        if r.get('sa') not in SA_BRANCH_MAP:
            rows.append({'Cabang':r.get('branch',''),'Service Advisor':str(r.get('sa','')).lower(),'Berlaku Mulai':r.get('effective_from',''),'Status':'Aktif'})
    return rows

def branch_mapping_table():
    return pd.DataFrame(_sa_branch_rows())
BRANCH_STATE_PATH=os.path.join(os.path.expanduser('~'),'.kharisma_workshop_branches.json')
def load_branches():
    try:
        with open(BRANCH_STATE_PATH,'r',encoding='utf-8') as fh:
            data=json.load(fh)
        vals=[str(x).strip() for x in data if str(x).strip() and str(x).strip().upper() not in EXTERNAL_BRANCHES]
        return vals or DEFAULT_BRANCHES.copy()
    except Exception:
        return DEFAULT_BRANCHES.copy()
def save_branches(branches):
    vals=[]
    for x in branches:
        x=str(x).strip()
        if x and x not in vals: vals.append(x)
    tmp=BRANCH_STATE_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(vals,fh,ensure_ascii=False,indent=2)
    os.replace(tmp,BRANCH_STATE_PATH)
BRANCHES=load_branches()
LOGO_PATH=os.path.join(os.path.dirname(__file__),'kharisma_logo.png')
LICENSE_STATE_PATH=os.path.join(os.path.expanduser('~'),'.kharisma_workshop_license.json')
# Persistent local data storage: survives browser refresh and application restart.
DATA_DIR=os.path.join(os.path.expanduser('~'),'.kharisma_workshop_data')
os.makedirs(DATA_DIR, exist_ok=True)

# Compact offline license format: KHS-XXXX-XXXX-XXXX (60 bits).
# The same secret is used by the separate KHARISMA License Manager.
# Note: this is an offline/local licensing mechanism; a determined reverse-engineer
# can extract the secret from the executable/source. For server-grade licensing,
# use a remote verification service.
LICENSE_SECRET=b'KHARISMA-OFFLINE-LICENSE-V12-2026'
_LICENSE_ALPHABET='ABCDEFGHIJKLMNOPQRSTUVWXYZ234567'
_LICENSE_EPOCH=datetime.date(2020,1,1)

def _machine_fingerprint():
    parts=[platform.system(),platform.release(),platform.machine(),platform.node(),hex(uuid.getnode())]
    return hashlib.sha256('|'.join(parts).encode('utf-8')).hexdigest()[:24].upper()

def _base32_encode_int(n, length=12):
    out=[]
    for _ in range(length):
        out.append(_LICENSE_ALPHABET[n & 31]); n >>= 5
    return ''.join(reversed(out))

def _base32_decode_int(text):
    n=0
    for ch in text:
        n=(n<<5)|_LICENSE_ALPHABET.index(ch)
    return n

def _license_tag(data):
    raw=hmac.new(LICENSE_SECRET,data.encode('ascii'),hashlib.sha256).digest()
    return int.from_bytes(raw[:3],'big') >> 4  # 20-bit tag

def _normalize_license_key(key):
    return re.sub(r'[^A-Z2-7]','',str(key or '').upper())

def _license_payload_from_key(key):
    clean=_normalize_license_key(key)
    if len(clean)!=15 or not clean.startswith('KHS'):
        raise ValueError('Format license tidak valid')
    body=clean[3:]
    n=_base32_decode_int(body)
    # 60-bit layout: start day 14 | expiry day 14 | max devices 4 | nonce 8 | tag 20
    tag=n & ((1<<20)-1); n >>= 20
    nonce=n & 0xff; n >>= 8
    max_devices=n & 0x0f; n >>= 4
    expiry_day=n & ((1<<14)-1); n >>= 14
    start_day=n & ((1<<14)-1)
    unsigned=_base32_encode_int((start_day<<(14+4+8)) | (expiry_day<<(4+8)) | (max_devices<<8) | nonce,12)
    expected=_license_tag('KHS'+unsigned)
    if tag!=expected: raise ValueError('Signature license tidak cocok')
    starts=_LICENSE_EPOCH+datetime.timedelta(days=start_day)
    expires=_LICENSE_EPOCH+datetime.timedelta(days=expiry_day)
    return {'starts':starts.isoformat(),'expires':expires.isoformat(),'max_devices':max_devices or 1,'version':'V12'}

def verify_license_key(key):
    try:
        normalized='KHS-'+'-'.join([_normalize_license_key(key)[3:][i:i+4] for i in range(0,12,4)])
        payload=_license_payload_from_key(normalized)
        starts=datetime.date.fromisoformat(payload['starts']); exp=datetime.date.fromisoformat(payload['expires']); today=datetime.date.today()
        if exp<starts: return False,'Periode license tidak valid',payload
        if today<starts: return False,'Lisensi belum mulai berlaku ('+starts.strftime('%d/%m/%Y')+')',payload
        if today>exp: return False,'Lisensi expired pada '+exp.strftime('%d/%m/%Y'),payload
        return True,'Lisensi valid',payload
    except Exception:
        return False,'License key tidak valid atau rusak',None

def load_license_state():
    try:
        with open(LICENSE_STATE_PATH,'r',encoding='utf-8') as f: return json.load(f)
    except Exception: return {}

def save_license_state(state):
    tmp=LICENSE_STATE_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as f: json.dump(state,f,indent=2)
    os.replace(tmp,LICENSE_STATE_PATH)

def activate_license(key,customer_name=''):
    ok,msg,payload=verify_license_key(key)
    if not ok: return False,msg
    fp=_machine_fingerprint(); state=load_license_state()
    if state.get('license_key') and state.get('license_key')!=_normalize_license_key(key):
        return False,'Komputer ini sudah terikat ke license lain. Gunakan RESET BINDING terlebih dahulu.'
    customer_name=(customer_name or state.get('customer_name') or '').strip()
    save_license_state({'license_key':_normalize_license_key(key),'fingerprint':fp,'customer_name':customer_name,'activated_at':datetime.datetime.now().isoformat(timespec='seconds')})
    return True,'Aktivasi berhasil'

def current_license():
    state=load_license_state(); key=state.get('license_key','')
    if not key: return False,'Belum diaktivasi',None
    if state.get('fingerprint')!=_machine_fingerprint(): return False,'Lisensi terikat ke komputer lain',None
    ok,msg,payload=verify_license_key(key)
    if ok and payload is not None: payload['customer']=state.get('customer_name') or ''
    return ok,msg,payload

def logo_data_uri():
    try:
        with open(LOGO_PATH,'rb') as fh:
            return 'data:image/png;base64,'+base64.b64encode(fh.read()).decode('ascii')
    except Exception:
        return ''

def read_excel_auto(data, **kwargs):
    """Robust reader for .xlsx/.xls uploaded bytes.
    Detects legacy OLE .xls by file signature. Uses xlrd first; when xlrd is
    unavailable, tries a local LibreOffice conversion as a fallback.
    """
    raw = data.getvalue() if hasattr(data, 'getvalue') else data
    if isinstance(raw, memoryview):
        raw = raw.tobytes()
    raw = bytes(raw)
    is_legacy_xls = raw[:8].startswith(b'\xd0\xcf\x11\xe0')
    if not is_legacy_xls:
        return pd.read_excel(io.BytesIO(raw), engine='openpyxl', **kwargs)
    if is_legacy_xls:
        # Legacy .xls: use xlrd when available. If xlrd is missing OR cannot
        # parse a particular workbook, try LibreOffice conversion when present.
        xls_error = None
        try:
            import xlrd  # noqa: F401
            return pd.read_excel(io.BytesIO(raw), engine='xlrd', **kwargs)
        except Exception as e:
            xls_error = e
        import tempfile, subprocess, shutil
        soffice = shutil.which('libreoffice') or shutil.which('soffice')
        if soffice:
            with tempfile.TemporaryDirectory() as td:
                src=os.path.join(td,'input.xls'); open(src,'wb').write(raw)
                out=os.path.join(td,'out'); os.makedirs(out,exist_ok=True)
                cp=subprocess.run([soffice,'--headless','--convert-to','xlsx','--outdir',out,src],capture_output=True,text=True,timeout=90)
                converted=os.path.join(out,'input.xlsx')
                if cp.returncode==0 and os.path.isfile(converted):
                    return pd.read_excel(converted, engine='openpyxl', **kwargs)
        detail = str(xls_error) if xls_error else 'xlrd tidak tersedia'
        raise RuntimeError(f'File .xls terdeteksi tetapi belum berhasil dibaca. Pastikan xlrd 2.0.1 terpasang. Detail: {detail}') from xls_error
    return pd.read_excel(io.BytesIO(raw), engine='openpyxl', **kwargs)

def norm_sa(x):
    # Satukan Service Advisor yang hanya berbeda kapitalisasi/spasi.
    s=clean(x).upper()
    return re.sub(r'\s+','',s) if s and s not in {'NAN','NONE','-'} else ''

# =========================
# UI THEME
# =========================
st.markdown('''
<style>
/* =====================================================================
   KHARISMA WORKSHOP ANALYTICS — LIGHT MODERN THEME (V13)
   Hanya lapisan tampilan. Tidak ada logika, data, atau fitur yang diubah.
   Kompatibel Streamlit lama: setiap aturan penting punya versi non-:has()
   sebagai cadangan, dan teks menu sidebar TIDAK PERNAH disembunyikan
   lewat seleksi struktural (mis. :first-child) — hanya lewat data-testid.
   ===================================================================== */
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');
@import url('https://fonts.googleapis.com/css2?family=Material+Symbols+Rounded:opsz,wght,FILL,GRAD@20..48,300..500,0..1,-25..0');

:root{
  --bg:#eef1f8; --surface:#ffffff; --border:#e7eaf3; --border-strong:#d7dce8;
  --text:#111827; --muted:#6b7280; --faint:#9aa3b2;
  --accent:#2f6fed; --accent-hover:#1e54c7; --accent-soft:#e9f0fe;
  --brand:#e31b23; --ok:#16a34a; --ok-soft:#e8f9ef;
  --blue:#2f6fed; --blue-soft:#e9f0fe;
  --green:#16a34a; --green-soft:#e6f8ec;
  --purple:#8b5cf6; --purple-soft:#f2ecfe;
  --orange:#f59e0b; --orange-soft:#fff2e0;
  --pink:#ec4899; --pink-soft:#fdecf3;
  --teal:#0d9488; --teal-soft:#e1f7f4;
  --radius:16px; --radius-sm:10px;
  --shadow:0 1px 2px rgba(17,24,39,.04),0 6px 16px rgba(17,24,39,.05);
  --font:'Inter','Segoe UI',system-ui,-apple-system,Roboto,Arial,sans-serif;
}

/* ---------- Dasar ---------- */
.stApp{ background:var(--bg); color:var(--text); }
.stApp, .stApp p, .stApp label, .stApp li, .stApp input, .stApp textarea, .stApp button,
.stApp h1, .stApp h2, .stApp h3, .stApp h4, .stApp h5, .stApp h6,
.stApp [data-testid="stMetric"], .stApp [data-baseweb="tab"], .stApp [data-baseweb="select"]{
  font-family:var(--font);
}
.stApp [data-testid="stIconMaterial"], .stApp .material-symbols-rounded, .stApp [class*="material-symbols"]{
  font-family:'Material Symbols Rounded' !important;
}
/* Ikon kustom di dalam HTML kartu (kartu KPI, judul section, insight, dsb).
   Ikon bawaan Streamlit (:material/xxx: di luar HTML mentah) sudah otomatis
   konsisten lewat aturan di atas; ini hanya untuk span yang kita tulis sendiri. */
.msr{
  font-family:'Material Symbols Rounded' !important; font-weight:normal; font-style:normal;
  display:inline-block; line-height:1; vertical-align:middle; white-space:nowrap;
  letter-spacing:normal; text-transform:none; word-wrap:normal; -webkit-font-smoothing:antialiased;
  font-variation-settings:'FILL' 0,'wght' 400,'GRAD' 0,'opsz' 24;
}
.summary-icon .msr{ font-size:21px; }
.chart-card h3 .msr, .dashboard-card h3 .msr, .insight-card h3 .msr{ font-size:17px; color:var(--accent); }
.section-title .msr{ font-size:18px; color:var(--accent); }
.insight-dot .msr{ font-size:13px; }
.upload-note .msr, .upload-progress .msr{ font-size:14px; vertical-align:-2px; }
.topbar .msr{ font-size:13px; vertical-align:-2px; }
.stApp h1,.stApp h2,.stApp h3,.stApp h4{ color:var(--text); letter-spacing:-.01em; }
.stApp h3{ font-size:1.05rem; font-weight:650; margin:1.1rem 0 .4rem; padding:0; }
.stApp h4{ font-size:.95rem; font-weight:650; }
[data-testid="stHeaderActionElements"]{ display:none; }
.stApp [data-testid="stCaptionContainer"], .stApp .stCaption{ color:var(--muted); font-size:.78rem; }
hr{ border-color:var(--border) !important; }

/* Header Streamlit: sembunyikan menu/footer bawaan, TAPI biarkan toolbar
   (tombol buka/tutup sidebar) tetap berfungsi di semua versi Streamlit. */
[data-testid="stHeader"]{ background:transparent !important; height:2.6rem; }
[data-testid="stDecoration"], #MainMenu, footer{ display:none !important; visibility:hidden; }
.block-container{ max-width:1420px !important; padding:1.2rem 2rem 3rem !important; }

/* ---------- Sidebar (terang, ala referensi) ---------- */
[data-testid="stSidebar"]{ background:linear-gradient(180deg,#fbfdff 0%,#f1f6ff 42%,#e8f0fd 100%) !important; border-right:1px solid #d9e5f8 !important; min-width:262px; box-shadow:2px 0 14px rgba(47,111,237,.06); }
[data-testid="stSidebar"][aria-expanded="false"]{ min-width:0 !important; }
[data-testid="stSidebar"] > div:first-child{ background:transparent !important; }
[data-testid="stSidebar"] [data-testid="stSidebarContent"], [data-testid="stSidebar"] .block-container{ padding:1rem .85rem 1.4rem !important; }
[data-testid="stSidebar"] hr{ border-color:#d9e5f8 !important; margin:12px 0; }
[data-testid="stSidebar"] [data-testid="stCaptionContainer"]{ color:var(--faint) !important; font-size:.72rem; }

.k-logo{ display:flex; align-items:center; gap:11px; padding:4px 6px 16px; margin-bottom:4px; border-bottom:1px solid #d9e5f8; }
.k-logo img{ width:40px; height:40px; object-fit:contain; border-radius:10px; flex:0 0 auto; }
.k-logo .brand{ line-height:1.2; }
.k-logo .brand br{ display:none; }
.k-logo .brand b{ display:block; font-size:15px; font-weight:800; letter-spacing:.2px; color:var(--text) !important; }
.k-logo .brand span{ font-size:9.5px; letter-spacing:1.3px; color:var(--faint) !important; font-weight:600; text-transform:uppercase; }
.side-section{ font-size:10px; font-weight:700; letter-spacing:1.1px; text-transform:uppercase; color:var(--faint) !important; margin:16px 6px 6px; }

/* Menu navigasi: TIDAK pernah menyembunyikan teks lewat :first-child.
   Radio circle native disembunyikan hanya lewat data-testid stabil. */
[data-testid="stSidebar"] [role="radiogroup"]{ gap:1px !important; display:flex !important; flex-direction:column !important; }
[data-testid="stSidebar"] [role="radiogroup"] label{
  display:flex !important; align-items:center !important; position:relative;
  min-height:40px; padding:8px 12px !important; margin:0 !important;
  border-radius:10px; background:transparent; cursor:pointer; transition:background .12s ease;
}
[data-testid="stSidebar"] [role="radiogroup"] label:hover{ background:rgba(255,255,255,.75); }
[data-testid="stSidebar"] [role="radiogroup"] label p{ font-size:13.5px !important; font-weight:500 !important; margin:0 !important; }
/* Keterlihatan teks dijamin secara eksplisit (bukan hanya warisan) */
[data-testid="stSidebar"] [role="radiogroup"] label [data-testid="stMarkdownContainer"],
[data-testid="stSidebar"] [role="radiogroup"] label [data-testid="stMarkdownContainer"] p{
  display:block !important; visibility:visible !important; opacity:1 !important; width:auto !important; height:auto !important;
}
[data-testid="stSidebar"] [role="radiogroup"] label [data-testid="stMarkdownContainer"] *{
  color:#374151 !important; -webkit-text-fill-color:#374151 !important;
}
/* Status terpilih: progressive enhancement lewat :has(); jika browser lama
   tidak mendukung :has(), teks tetap terlihat (warna default di atas). */
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked){ background:#ffffff; box-shadow:0 2px 10px rgba(47,111,237,.14), 0 0 0 1px #d3e2fb; }
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) [data-testid="stMarkdownContainer"] *{
  color:var(--accent) !important; -webkit-text-fill-color:var(--accent) !important; font-weight:700 !important;
}
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked):before{
  content:''; position:absolute; left:-9px; top:8px; bottom:8px; width:3px; border-radius:3px; background:var(--accent);
}
/* Sembunyikan hanya elemen input radio itu sendiri (bukan pembungkusnya) */
[data-testid="stSidebar"] [data-testid="stRadioOption"] input,
[data-testid="stSidebar"] label[data-baseweb="radio"] input{ position:absolute; opacity:0; pointer-events:none; width:0; height:0; }
[data-testid="stSidebar"] [data-testid="stRadioOption"] svg,
[data-testid="stSidebar"] label[data-baseweb="radio"] svg{ display:none !important; }
/* Bulatan radio Streamlit terbaru: sebuah <div> dekoratif terpisah dari teks
   (bukan svg/input). Disembunyikan berdasarkan ISI (tidak berisi teks menu),
   bukan berdasarkan posisi :first-child — supaya aman lintas versi.
   Progressive enhancement: browser tanpa dukungan :has() akan tetap
   menampilkan bulatan kecil, tapi TEKS MENU tidak pernah ikut hilang. */
[data-testid="stSidebar"] [role="radiogroup"] label div:has(+ [data-testid="stMarkdownContainer"]){
  display:none !important;
}

/* Kontrol di sidebar (selectbox cabang, dsb) */
[data-testid="stSidebar"] [data-testid="stSelectbox"] > div > div,
[data-testid="stSidebar"] [data-baseweb="select"] > div{
  background:var(--surface) !important; border:1px solid var(--border-strong) !important; border-radius:10px !important; min-height:38px;
}
[data-testid="stSidebar"] [data-baseweb="select"] *, [data-testid="stSidebar"] [data-testid="stSelectbox"] input{
  color:var(--text) !important; -webkit-text-fill-color:var(--text) !important;
}
[data-testid="stSidebar"] [data-testid="stWidgetLabel"] p{ color:var(--muted) !important; font-weight:600 !important; }
[data-testid="stSidebar"] [data-testid="stAlert"]{ background:var(--orange-soft); border:1px solid #fbd9a5; }
[data-testid="stSidebar"] [data-testid="stAlert"] *{ color:#92620a !important; }

/* ---------- Ikon menu berwarna ----------
   Warna ditetapkan per item (anak ke-n dari radiogroup, urutan = page_labels) lewat variabel CSS yang diwariskan ke ikon; ikon Streamlit berupa <span role="img" translate="no">.
   Hanya tampilan; logika tidak berubah. */
[data-testid="stSidebar"] [role="radiogroup"] label [role="img"][translate="no"],
[data-testid="stSidebar"] [role="radiogroup"] label [data-testid="stIconMaterial"]{
  color:var(--ic,#2f6fed) !important; -webkit-text-fill-color:var(--ic,#2f6fed) !important; background:var(--ic-bg,rgba(47,111,237,.14));
  border-radius:9px; padding:5px; margin-right:10px; font-size:19px !important; line-height:1 !important; vertical-align:middle !important;
  font-variation-settings:'FILL' 1,'wght' 500,'GRAD' 0,'opsz' 24; transition:transform .12s ease, background .12s ease;
}
[data-testid="stSidebar"] [role="radiogroup"] label:hover [role="img"][translate="no"],
[data-testid="stSidebar"] [role="radiogroup"] label:hover [data-testid="stIconMaterial"]{ transform:scale(1.1); }
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) [role="img"][translate="no"],
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) [data-testid="stIconMaterial"]{
  color:#ffffff !important; -webkit-text-fill-color:#ffffff !important; background:var(--ic,#2f6fed); box-shadow:0 3px 8px var(--ic-bg,rgba(47,111,237,.2));
}

[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(1){ --ic:#2563eb; --ic-bg:rgba(37,99,235,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(2){ --ic:#ef4444; --ic-bg:rgba(239,68,68,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(3){ --ic:#7c3aed; --ic-bg:rgba(124,58,237,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(4){ --ic:#db2777; --ic-bg:rgba(219,39,119,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(5){ --ic:#0891b2; --ic-bg:rgba(8,145,178,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(6){ --ic:#16a34a; --ic-bg:rgba(22,163,74,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(7){ --ic:#ea580c; --ic-bg:rgba(234,88,12,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(8){ --ic:#0d9488; --ic-bg:rgba(13,148,136,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(9){ --ic:#ca8a04; --ic-bg:rgba(202,138,4,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(10){ --ic:#e11d48; --ic-bg:rgba(225,29,72,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(11){ --ic:#4f46e5; --ic-bg:rgba(79,70,229,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(12){ --ic:#0284c7; --ic-bg:rgba(2,132,199,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(13){ --ic:#9333ea; --ic-bg:rgba(147,51,234,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(14){ --ic:#d946ef; --ic-bg:rgba(217,70,239,.15); }
[data-testid="stSidebar"] [role="radiogroup"] > :nth-child(15){ --ic:#65a30d; --ic-bg:rgba(101,163,13,.15); }

/* ---------- User & Access Management ---------- */
.um-head{ display:flex; gap:12px; align-items:center; padding:6px 2px; }
.um-head > .msr{ font-size:26px; color:#047857; background:#d1fae5; border-radius:12px; padding:9px; }
.um-head b{ display:block; font-size:17px; color:#0f172a; } .um-head span:not(.msr){ display:block; color:#64748b; font-size:12.5px; }
.um-th{ font-size:10.5px; font-weight:800; letter-spacing:.7px; color:#64748b; padding:6px 2px; border-bottom:1px solid #e2e8f0; }
.um-name{ font-weight:700; font-size:13.5px; color:#0f172a; } .um-mail{ font-family:ui-monospace,monospace; font-size:11px; color:#64748b; }
.um-branch{ font-size:12.5px; color:#334155; } .um-branch .msr{ font-size:15px; vertical-align:-3px; color:#64748b; }
.um-act{ font-size:12px; color:#64748b; }
.um-cards{ display:grid; grid-template-columns:repeat(auto-fit,minmax(230px,1fr)); gap:12px; margin:18px 0 6px; }
.um-card{ background:#fff; border:1px solid #e2e8f0; border-radius:12px; padding:14px 16px; box-shadow:0 1px 4px rgba(15,23,42,.04); }
[class*="st-key-um_stat_on_"] button{ background:#dcfce7 !important; color:#047857 !important; border:1px solid #86efac !important; font-weight:700 !important; min-height:30px !important; padding:2px 10px !important; }
[class*="st-key-um_stat_off_"] button{ background:#f1f5f9 !important; color:#64748b !important; border:1px solid #cbd5e1 !important; font-weight:700 !important; min-height:30px !important; padding:2px 10px !important; }
[class*="st-key-um_perm_"] button{ background:#eff6ff !important; color:#1d4ed8 !important; border:1px solid #bfdbfe !important; font-weight:600 !important; font-size:12px !important; min-height:30px !important; padding:2px 8px !important; }
[class*="st-key-um_edit_"] button, [class*="st-key-um_del_"] button{ min-height:30px !important; padding:2px 6px !important; background:transparent !important; border:none !important; box-shadow:none !important; }
[data-testid="stSidebar"] .user-chip{ display:flex; gap:9px; align-items:center; background:#ffffff; border:1px solid #d9e5f8; border-radius:12px; padding:8px 10px; margin:4px 0 6px; box-shadow:0 1px 6px rgba(47,111,237,.08); }
[data-testid="stSidebar"] .user-chip .msr{ font-size:26px; color:#2f6fed; }
[data-testid="stSidebar"] .user-chip b{ display:block; font-size:12.5px; color:#0f172a; margin-bottom:3px; }
[data-testid="stDialog"] button[kind="primary"], [data-testid="stDialog"] [data-testid="stBaseButton-primary"]{ background:#059669 !important; border-color:#059669 !important; color:#fff !important; }
[data-testid="stDialog"] button[kind="primary"]:hover, [data-testid="stDialog"] [data-testid="stBaseButton-primary"]:hover{ background:#047857 !important; border-color:#047857 !important; }
[data-testid="stDialog"] [data-testid="stCheckbox"] label[data-selected="true"] > div:first-of-type{ background:#059669 !important; border-color:#059669 !important; }
[data-testid="stDialog"] [data-testid="stRadio"] label[data-selected="true"] > div:first-of-type > div:first-of-type{ background:#059669 !important; border-color:#059669 !important; }
[data-testid="stDialog"] [data-testid="stCaptionContainer"]{ margin-top:-10px; margin-bottom:6px; }
.side-footer{ margin:18px 2px 0; padding:12px 13px; border:1px solid #c9ecd6; border-radius:12px; background:var(--ok-soft); }
.side-footer b{ display:flex; align-items:center; gap:6px; font-size:10.5px; letter-spacing:.8px; color:#15803d !important; }
.side-footer b:before{ content:''; width:7px; height:7px; border-radius:50%; background:#22c55e; display:inline-block; }
.side-footer span{ display:block; margin-top:5px; font-size:11.5px; line-height:1.45; color:#166534 !important; font-weight:500; }


/* ---------- Sidebar V16: filter atas, menu, kartu user di bawah ---------- */
.sb-sec{ display:flex; justify-content:space-between; align-items:center; margin:6px 4px 7px; font-size:10.5px; font-weight:700; letter-spacing:1px; text-transform:uppercase; color:var(--faint) !important; }
.sb-sec > span{ display:flex; align-items:center; gap:6px; }
.sb-sec .msr{ font-size:15px; color:#94a3b8; }
.sb-sec em{ font-style:normal; font-size:10px; letter-spacing:.6px; color:#2f6fed; font-weight:700; max-width:55%; overflow:hidden; text-overflow:ellipsis; white-space:nowrap; }
[data-testid="stSidebar"] [data-testid="stSidebarHeader"]{ display:none !important; }
.st-key-flt_mode{ margin-bottom:4px; }
.st-key-flt_mode [data-testid="stButtonGroup"]{ width:100%; display:flex; background:#eef2f8; border:1px solid #dbe4f3; border-radius:10px; padding:3px; gap:2px; }
.st-key-flt_mode [data-testid="stButtonGroup"] button{ flex:1 1 0; border:none !important; background:transparent !important; box-shadow:none !important; border-radius:8px !important; min-height:32px; font-size:12.5px; font-weight:600; color:#475569 !important; }
.st-key-flt_mode [data-testid="stButtonGroup"] button[kind="segmented_controlActive"],
.st-key-flt_mode [data-testid="stButtonGroup"] [data-testid="stBaseButton-segmented_controlActive"]{ background:#ffffff !important; color:#2563eb !important; box-shadow:0 1px 4px rgba(47,111,237,.18) !important; }
[data-testid="stSidebar"] [data-testid="stDateInput"] > div > div, [data-testid="stSidebar"] [data-testid="stDateInput"] [data-baseweb="input"]{ background:var(--surface) !important; border:1px solid var(--border-strong) !important; border-radius:10px !important; }
[data-testid="stSidebar"] [data-testid="stDateInput"] input{ color:var(--text) !important; -webkit-text-fill-color:var(--text) !important; }
/* menu: ikon berwarna tanpa kotak latar */
[data-testid="stSidebar"] [role="radiogroup"] label [role="img"][translate="no"],
[data-testid="stSidebar"] [role="radiogroup"] label [data-testid="stIconMaterial"]{ background:transparent !important; padding:0 !important; box-shadow:none !important; font-size:21px !important; }
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) [role="img"][translate="no"],
[data-testid="stSidebar"] [role="radiogroup"] label:has(input:checked) [data-testid="stIconMaterial"]{ color:var(--ic,#2f6fed) !important; -webkit-text-fill-color:var(--ic,#2f6fed) !important; background:transparent !important; box-shadow:none !important; }
/* kartu user menempel di dasar sidebar */
[data-testid="stSidebarUserContent"]{ min-height:calc(100vh - 2rem); }
[data-testid="stSidebar"] [data-testid="stVerticalBlock"]:has(> [class*="st-key-sb_bottom"]),
[data-testid="stSidebar"] [data-testid="stVerticalBlock"]:has(> * > [class*="st-key-sb_bottom"]){ min-height:calc(100vh - 3rem); }
[class*="st-key-sb_bottom"]{ margin-top:auto; display:flex; flex-direction:column; gap:.6rem; }
[class*="st-key-sb_userbox"]{ margin-top:0; background:#ffffff; border:1px solid #d9e5f8; border-radius:14px; padding:10px; box-shadow:0 1px 6px rgba(47,111,237,.08); gap:.5rem; display:flex; flex-direction:column; }
[data-testid="stSidebar"] :has(> [class*="st-key-sb_bottom"]){ margin-top:auto; }
.sb-user{ display:flex; align-items:center; gap:10px; margin:0 0 10px; }
.sb-avatar{ width:38px; height:38px; border-radius:10px; background:#2563eb; color:#fff; font-weight:800; font-size:17px; display:flex; align-items:center; justify-content:center; flex:0 0 auto; }
.sb-uinfo b{ display:block; font-size:14px; color:#0f172a; line-height:1.2; }
.sb-uinfo span{ display:block; font-size:12px; color:#2563eb; font-weight:500; }
.lic-card{ margin:0 !important; }
[class*="st-key-sb_bottom"] .side-footer{ margin:0 0 2px !important; padding:8px 10px !important; display:block; position:static !important; clear:both; }
[class*="st-key-sb_userbox"] [data-testid="stMarkdownContainer"]{ overflow:visible; }
[class*="st-key-sb_logout"] button{ background:#fee2e2 !important; border:1px solid #fecaca !important; color:#dc2626 !important; font-weight:600 !important; border-radius:10px !important; }
[class*="st-key-sb_logout"] button:hover{ background:#fecaca !important; }
[class*="st-key-sb_logout"] button *{ color:#dc2626 !important; -webkit-text-fill-color:#dc2626 !important; }

/* tab filter waktu: paksa horizontal (menimpa aturan radiogroup menu) */
.st-key-flt_mode, .st-key-flt_mode > div{ width:100% !important; }
.st-key-flt_mode [role="radiogroup"], .st-key-flt_mode [data-testid="stButtonGroup"], .st-key-flt_mode [data-baseweb="button-group"]{ display:flex !important; flex-direction:row !important; flex-wrap:nowrap !important; width:100% !important; }
.st-key-flt_mode button{ flex:1 1 0 !important; width:auto !important; min-width:0 !important; }
/* tab filter waktu: tombol berwarna, yang aktif jelas terlihat */
[data-testid="stSidebar"] [data-testid="stHorizontalBlock"]{ gap:.35rem !important; }
[class*="st-key-flt_btn_"] button{ border-radius:10px !important; min-height:36px; font-size:12.5px !important; font-weight:600 !important; background:#ffffff !important; color:#64748b !important; border:1px solid #cfdbee !important; box-shadow:none !important; padding:0 4px !important; transition:all .12s ease; }
[class*="st-key-flt_btn_"] button *{ color:inherit !important; -webkit-text-fill-color:currentColor !important; }
[class*="st-key-flt_btn_"] button:hover{ border-color:#93b4f5 !important; color:#1d4ed8 !important; }
.st-key-flt_btn_h button[kind="primary"], .st-key-flt_btn_h [data-testid="stBaseButton-primary"]{ background:linear-gradient(135deg,#2563eb,#3b82f6) !important; color:#ffffff !important; border-color:#1d4ed8 !important; box-shadow:0 3px 10px rgba(37,99,235,.38) !important; }
.st-key-flt_btn_b button[kind="primary"], .st-key-flt_btn_b [data-testid="stBaseButton-primary"]{ background:linear-gradient(135deg,#059669,#10b981) !important; color:#ffffff !important; border-color:#047857 !important; box-shadow:0 3px 10px rgba(5,150,105,.38) !important; }
.st-key-flt_btn_t button[kind="primary"], .st-key-flt_btn_t [data-testid="stBaseButton-primary"]{ background:linear-gradient(135deg,#ea580c,#f59e0b) !important; color:#ffffff !important; border-color:#c2410c !important; box-shadow:0 3px 10px rgba(234,88,12,.38) !important; }
/* perpindahan halaman terasa lebih ringan: jangan redupkan konten lama saat rerun */
[data-stale="true"]{ opacity:1 !important; transition:none !important; }
/* ---------- V17: komponen halaman referensi (rk-*) ---------- */
.rk-head{display:flex;align-items:center;gap:12px;background:#fff;border:1px solid #e3eaf5;border-radius:14px;padding:12px 16px;margin:2px 0 12px;box-shadow:0 1px 4px rgba(15,23,42,.04);}
.rk-head > div > b{display:block;font-size:16px;color:#0f172a;line-height:1.25;}
.rk-head small{display:block;font-size:11.5px;color:#64748b;margin-top:2px;}
.rk-hic{width:38px;height:38px;border-radius:11px;display:flex;align-items:center;justify-content:center;flex:0 0 auto;}
.rk-hic .msr{font-size:21px;}
.rk-dark{display:flex;align-items:center;gap:10px;background:#0b1220;border-radius:12px;padding:11px 16px;margin:0 0 12px;color:#e2e8f0;font-size:12.5px;line-height:1.5;}
.rk-dark b{color:#fff;} .rk-dark .c1{color:#93c5fd;font-weight:700;} .rk-dark .c2{color:#fbbf24;font-weight:700;} .rk-dark .c3{color:#34d399;font-weight:700;}
.rk-dark-t{flex:1 1 auto;} .rk-dot{width:8px;height:8px;border-radius:50%;background:#34d399;flex:0 0 auto;}
.rk-tag{font-family:ui-monospace,monospace;font-size:10.5px;color:#94a3b8;background:#1e293b;border-radius:6px;padding:3px 8px;white-space:nowrap;}
.rk-grid.sm .rk-val{font-size:17px;} .rk-grid.sm .rk-card{padding:12px 13px;}
.rk-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(var(--rkw,190px),1fr));gap:12px;margin:0 0 14px;}
.rk-card{background:#fff;border:1px solid #e3eaf5;border-radius:14px;padding:13px 15px;box-shadow:0 1px 4px rgba(15,23,42,.04);display:flex;flex-direction:column;gap:7px;min-width:0;}
.rk-top{display:flex;justify-content:space-between;align-items:center;gap:8px;}
.rk-lab{font-size:10.5px;font-weight:800;letter-spacing:.5px;text-transform:uppercase;color:#475569;}
.rk-val{font-size:22px;font-weight:800;color:#0f172a;line-height:1.15;word-break:break-word;}
.rk-val small{font-size:12px;font-weight:600;color:#64748b;}
.rk-row{display:flex;justify-content:space-between;align-items:center;font-size:12px;font-weight:600;color:#334155;}
.rk-foot{font-size:11.5px;color:#475569;} .rk-foot2{display:flex;justify-content:space-between;font-size:11px;color:#475569;gap:6px;}
.rk-ic{width:30px;height:30px;border-radius:9px;display:flex;align-items:center;justify-content:center;flex:0 0 auto;}
.rk-ic .msr{font-size:18px;}
.rk-bar{height:8px;border-radius:99px;background:#e2e8f0;overflow:hidden;} .rk-bar i{display:block;height:100%;border-radius:99px;}
.rk-badge{display:inline-block;font-size:11px;font-weight:800;border-radius:7px;padding:2px 9px;white-space:nowrap;}
.rk-ic.green,.rk-hic.green,.rk-badge.green{background:#dcfce7;color:#047857;} .rk-bar i.green{background:#10b981;}
.rk-ic.blue,.rk-hic.blue,.rk-badge.blue{background:#dbeafe;color:#1d4ed8;} .rk-bar i.blue{background:#2563eb;}
.rk-ic.orange,.rk-hic.orange,.rk-badge.orange{background:#ffedd5;color:#c2410c;} .rk-bar i.orange{background:#f59e0b;}
.rk-ic.red,.rk-hic.red,.rk-badge.red{background:#fee2e2;color:#b91c1c;} .rk-bar i.red{background:#ef4444;}
.rk-ic.purple,.rk-hic.purple,.rk-badge.purple{background:#ede9fe;color:#6d28d9;}
.rk-ic.teal,.rk-hic.teal,.rk-badge.teal{background:#ccfbf1;color:#0f766e;}
.rk-ic.gray,.rk-badge.gray{background:#f1f5f9;color:#475569;} .rk-bar i.gray{background:#cbd5e1;}
.rk-mini{border:1px solid #e3eaf5;border-radius:12px;padding:10px 14px;background:#f8fafc;}
.rk-mini.blue{background:#eff6ff;border-color:#bfdbfe;} .rk-mini.green{background:#ecfdf5;border-color:#a7f3d0;} .rk-mini.purple{background:#faf5ff;border-color:#e9d5ff;} .rk-mini.yellow{background:#fffbeb;border-color:#fde68a;}
.rk-dc .rk-val{font-size:20px;} .rk-drs{display:flex;flex-direction:column;gap:5px;margin-top:4px;border-top:1px dashed #e2e8f0;padding-top:8px;} .rk-dr{display:flex;justify-content:space-between;gap:8px;font-size:11.5px;color:#475569;} .rk-dr b{color:#0f172a;font-family:ui-monospace,monospace;font-size:11.5px;text-align:right;}
.rk-mval{font-family:ui-monospace,monospace;font-size:18px;font-weight:800;color:#0f172a;margin-top:3px;} .rk-mini small{font-size:11px;color:#475569;}
.rk-tw{overflow-x:auto;border:1px solid #e3eaf5;border-radius:12px;background:#fff;margin:0 0 12px;}
.rk-t{border-collapse:collapse;width:100%;font-size:12.5px;} .rk-t th{background:#f8fafc;color:#475569;font-size:10.5px;letter-spacing:.4px;text-transform:uppercase;text-align:left;padding:9px 10px;border-bottom:1px solid #e3eaf5;white-space:nowrap;}
.rk-t td{padding:9px 10px;border-bottom:1px solid #eef2f8;color:#0f172a;} .rk-t .n{text-align:right;font-variant-numeric:tabular-nums;} .rk-t tr.tot td{font-weight:800;background:#f1f5f9;}
.rk-t .hl1{background:#eff6ff;} .rk-t .hl2{background:#ecfdf5;font-weight:700;}
.rk-band{background:linear-gradient(135deg,#4c1d95,#1e1b4b);border-radius:16px;padding:16px 18px;margin:12px 0 14px;color:#fff;}
.rk-band-t{display:flex;align-items:center;gap:8px;font-weight:800;font-size:14px;margin-bottom:10px;text-transform:uppercase;letter-spacing:.3px;} .rk-band-t .msr{color:#fbbf24;}
.rk-band-g{display:grid;grid-template-columns:repeat(auto-fit,minmax(230px,1fr));gap:10px;}
.rk-ins{background:rgba(255,255,255,.08);border:1px solid rgba(255,255,255,.14);border-radius:12px;padding:10px 13px;} .rk-ins b{display:block;font-size:10.5px;text-transform:uppercase;letter-spacing:.5px;margin-bottom:4px;}
.rk-ins.g b{color:#6ee7b7;} .rk-ins.o b{color:#fcd34d;} .rk-ins.c b{color:#67e8f9;} .rk-ins p{margin:0;font-size:12.5px;line-height:1.5;color:#f1f5f9;}
.rk-wl{display:flex;align-items:center;gap:12px;background:#fff;border:1px solid #eef2f8;border-radius:10px;padding:8px 12px;margin-bottom:7px;}
.rk-wl .d{width:70px;font-weight:700;font-size:13px;color:#0f172a;} .rk-wl .b{flex:1;height:9px;border-radius:99px;background:#eef2f8;overflow:hidden;} .rk-wl .b i{display:block;height:100%;border-radius:99px;} .rk-wl .v{width:78px;text-align:right;font-size:12px;color:#334155;}
.rk-seg{background:#fff;border:1px solid #eef2f8;border-radius:10px;padding:9px 12px;margin-bottom:8px;} .rk-seg .t{display:flex;justify-content:space-between;font-weight:700;font-size:13px;color:#0f172a;margin-bottom:6px;}
[class*="st-key-rk_tabs"] [role="radiogroup"],[class*="st-key-rk_tabs"] [data-testid="stButtonGroup"]{gap:4px;}
/* ---------- V17b: header halaman, legenda donat, kartu modul, tab input ---------- */
.rk-ph{display:flex;justify-content:space-between;align-items:flex-start;gap:16px;flex-wrap:wrap;background:#fff;border:1px solid #e3eaf5;border-radius:16px;padding:18px 22px;margin:0 0 14px;box-shadow:0 1px 4px rgba(15,23,42,.04);}
.rk-ph-l{flex:1 1 360px;min-width:260px;}
.rk-ph-eb{display:block;font-size:10.5px;font-weight:800;letter-spacing:1.2px;color:#64748b;text-transform:uppercase;}
.rk-ph-t{font-size:26px;line-height:1.15;margin:4px 0 6px;color:#0b1220;font-weight:900;letter-spacing:-.3px;}
.rk-ph p{margin:0;color:#475569;font-size:12.5px;max-width:540px;line-height:1.45;}
.rk-ph-r{display:flex;gap:8px;flex-wrap:wrap;justify-content:flex-end;align-items:center;}
.rk-chip{display:inline-flex;align-items:center;gap:6px;border-radius:9px;padding:6px 11px;font-size:11.5px;font-weight:700;white-space:nowrap;}
.rk-chip small{font-size:10px;opacity:.75;font-weight:700;letter-spacing:.4px;}
.rk-chip.dark{background:#0b1220;color:#fff;font-family:ui-monospace,monospace;}
.rk-chip.soft{background:#eef2f9;color:#334155;border:1px solid #dbe4f3;}
.rk-chip.live{background:#ecfdf5;color:#047857;border:1px solid #a7f3d0;} .rk-chip.live i{width:7px;height:7px;border-radius:50%;background:#10b981;display:inline-block;}
.rk-chip.user{background:#fff;border:1px solid #dbe4f3;color:#0f172a;padding:5px 11px 5px 6px;}
.rk-chip.user .av{width:26px;height:26px;border-radius:50%;background:#2563eb;color:#fff;display:flex;align-items:center;justify-content:center;font-size:12px;}
.rk-chip.user .nm{display:flex;flex-direction:column;line-height:1.2;font-size:12px;} .rk-chip.user .nm small{font-size:9.5px;color:#64748b;font-weight:600;letter-spacing:0;opacity:1;}
.rk-ct{display:flex;justify-content:space-between;align-items:center;gap:8px;margin:4px 2px 8px;color:#0f172a;}
.rk-ct b{font-weight:800;font-size:12.5px;letter-spacing:.3px;} .rk-ct b .a{color:#2563eb;} .rk-ct > span{font-size:11px;color:#64748b;font-family:ui-monospace,monospace;}
.rk-lgw{display:flex;flex-direction:column;gap:8px;margin-top:6px;}
.st-key-dn_sa, .st-key-dn_ly{background:#fff !important;border:1px solid #e3eaf5 !important;border-radius:16px !important;box-shadow:0 1px 6px rgba(15,23,42,.05) !important;padding:16px 20px 12px !important;}
.rk-ct{padding-bottom:10px;border-bottom:1px solid #eef2f8;margin-bottom:12px;} .rk-ct b{display:flex;align-items:center;gap:8px;} .rk-ct b .msr{font-size:18px;color:#2563eb;}
.st-key-dn_sa [data-testid="stPlotlyChart"], .st-key-dn_ly [data-testid="stPlotlyChart"]{background:transparent !important;}
/* ===== Data Input (referensi) ===== */
.st-key-di_head{background:#fff;border:1px solid #e3eaf5;border-radius:16px;padding:16px 20px;box-shadow:0 1px 6px rgba(15,23,42,.05);margin-bottom:12px;}
.di-h,.up-h{display:flex;align-items:flex-start;gap:10px;} .di-h .msr,.up-h .msr{font-size:22px;color:#0e7490;margin-top:1px;} .up-h .msr{color:#2563eb;font-size:20px;}
.di-h b,.up-h b{display:block;font-size:15.5px;color:#0f172a;font-weight:800;} .up-h b{font-size:13.5px;} .di-h small,.up-h small{display:block;font-size:11.5px;color:#64748b;margin-top:2px;line-height:1.45;}
[class*="st-key-upload_mode"] [role="radiogroup"]{background:#eef3fb !important;border-radius:14px !important;} 
[class*="st-key-upload_mode"] label:has(input:checked), [class*="st-key-upload_mode"] label:nth-child(3):has(input:checked){background:#1f5fff !important;box-shadow:0 2px 8px rgba(31,95,255,.35) !important;}
.st-key-up_est{background:#fff;border:1px solid #e3eaf5;border-radius:16px;padding:18px 20px 20px;box-shadow:0 1px 6px rgba(15,23,42,.05);margin-bottom:12px;}
.rk-div{height:1px;background:#eef2f8;margin:2px 0 14px;}
.st-key-up_est .stButton > button[kind="primary"], .st-key-up_est [data-testid="stBaseButton-primary"]{background:#1f5fff !important;border-color:#1f5fff !important;color:#fff !important;min-height:36px;font-size:12.5px;}
.st-key-up_est .stDownloadButton > button{background:#eef2f9 !important;border:1px solid #dbe4f3 !important;color:#334155 !important;min-height:36px !important;font-size:12.5px !important;width:100% !important;}
.st-key-up_est [data-testid="stFileUploaderDropzone"]{display:flex;flex-direction:column;align-items:center;justify-content:center;gap:8px;min-height:210px;border:2px dashed #93c5fd;background:#f8fbff;border-radius:16px;padding:26px 20px;text-align:center;}
.st-key-up_est [data-testid="stFileUploaderDropzone"]:hover{background:#eff6ff;border-color:#3b82f6;}
.st-key-up_est [data-testid="stFileUploaderDropzoneInstructions"]{display:flex;flex-direction:column;align-items:center;gap:10px;}
.st-key-up_est [data-testid="stFileUploaderDropzoneInstructions"] > *{display:none !important;}
.st-key-up_est [data-testid="stFileUploaderDropzoneInstructions"]::before{content:'upload';font-family:'Material Symbols Rounded';font-size:24px;color:#2563eb;background:#dbeafe;width:50px;height:50px;border-radius:50%;display:flex;align-items:center;justify-content:center;}
.st-key-up_est [data-testid="stFileUploaderDropzoneInstructions"]::after{content:'Pilih atau Seret File Laporan Estimasi SA ke Sini';font-weight:800;font-size:14px;color:#0f172a;}
.st-key-up_est [data-testid="stFileUploaderDropzone"]::after{content:'Aplikasi otomatis membaca data transaksi, mengonversi nilai rupiah, dan mencegah duplikasi data PKB. Acuan revenue otomatis mengikuti tanggal pada header halaman (bukan tanggal PKB).';order:5;max-width:540px;font-size:12px;color:#64748b;line-height:1.55;}
.st-key-up_est [data-testid="stFileUploaderDropzone"] button{order:4;font-size:12px;}

.rk-lg{display:flex;align-items:center;gap:10px;border:1px solid #e8eef6;border-radius:10px;padding:8px 12px;background:#fff;}
.rk-lg i{width:10px;height:10px;border-radius:50%;flex:0 0 auto;}
.rk-lg .t{flex:1;min-width:0;} .rk-lg .t b{display:block;font-size:12.5px;color:#0f172a;} .rk-lg small{display:block;font-size:10.5px;color:#64748b;}
.rk-lg .v{text-align:right;} .rk-lg .v b{display:block;font-size:12.5px;color:#0f172a;font-family:ui-monospace,monospace;}
.rk-mod{background:#fff;border:1px solid #e3eaf5;border-radius:12px;padding:11px 14px;margin:0 0 6px;}
.rk-mod-t{display:flex;align-items:center;gap:8px;} .rk-mod-t b{font-size:13.5px;color:#0f172a;}
.rk-mod-f{font-size:11px;color:#475569;margin-top:4px;word-break:break-all;}
.rk-mod-f code{font-size:10.5px;background:#f1f5f9;border-radius:4px;padding:1px 5px;color:#334155;}
.rk-t tr.dim td{color:#94a3b8 !important;}
.rk-dot2{display:inline-block;width:7px;height:7px;border-radius:50%;background:#2563eb;margin-right:7px;}
[class*="st-key-upload_mode"] [role="radiogroup"]{display:flex !important;flex-direction:row !important;flex-wrap:wrap;gap:4px !important;background:#eef2f9;border:1px solid #dbe4f3;border-radius:12px;padding:4px;justify-content:flex-end;}
[class*="st-key-upload_mode"] label{padding:7px 14px !important;border-radius:9px;margin:0 !important;cursor:pointer;}
[class*="st-key-upload_mode"] label div:has(+ [data-testid="stMarkdownContainer"]){display:none !important;}
[class*="st-key-upload_mode"] label p{font-size:12.5px !important;font-weight:600 !important;color:#475569 !important;margin:0 !important;}
[class*="st-key-upload_mode"] label:has(input:checked){background:#2563eb;box-shadow:0 2px 8px rgba(37,99,235,.35);}
[class*="st-key-upload_mode"] label:has(input:checked) *{color:#ffffff !important;-webkit-text-fill-color:#ffffff !important;}
[class*="st-key-upload_mode"] label:nth-child(3):has(input:checked){background:#e11d48;box-shadow:0 2px 8px rgba(225,29,72,.35);}
/* ---------- Header halaman ---------- */
.hero, .input-hero{
  position:relative; background:var(--surface); color:var(--text); border:1px solid var(--border);
  border-radius:var(--radius); padding:18px 22px; margin:0 0 10px; min-height:0; box-shadow:var(--shadow); overflow:hidden;
}
.hero:before, .hero:after{ display:none !important; }
.hero .redline{ display:none; }
.hero .status-pill{ display:none; }
.hero .eyebrow, .input-hero .eyebrow{ font-size:10.5px; font-weight:700; letter-spacing:1.3px; text-transform:uppercase; color:var(--accent); opacity:1; }
.hero h1, .input-hero h1{ margin:4px 0 4px !important; padding:0 !important; font-size:21px !important; line-height:1.25; font-weight:800; color:var(--text) !important; letter-spacing:-.01em; }
.hero p, .input-hero p{ margin:0; max-width:900px; font-size:13px; line-height:1.5; color:var(--muted); }

.topbar{
  display:flex; flex-wrap:wrap; align-items:center; gap:8px; padding:8px; margin:0 0 14px;
  background:transparent; border:0; box-shadow:none; color:var(--muted); font-size:12px;
}
.topbar div{ display:flex; align-items:center; gap:7px; padding:7px 14px; background:var(--surface); border:1px solid var(--border); border-radius:999px; box-shadow:var(--shadow); }
.topbar b{ font-size:10px; font-weight:700; letter-spacing:.6px; text-transform:uppercase; color:var(--faint); }
.topbar span{ color:var(--text); font-weight:700; font-size:12px; }
.topbar-live{ margin-left:auto; color:#15803d !important; font-weight:700 !important; background:var(--ok-soft) !important; border-color:#c9ecd6 !important; }
.topbar-user{ color:var(--accent) !important; font-weight:700 !important; background:var(--blue-soft) !important; border-color:#cfe0fd !important; }

.section-title{ display:flex; align-items:center; gap:8px; margin:22px 0 6px; font-size:17px; font-weight:750; color:var(--text); letter-spacing:-.01em; }
.section-sub{ margin:0 0 14px; font-size:12.5px; color:var(--muted); }
.small-note{ color:var(--muted); font-size:12px; }

/* ---------- Kartu ringkasan & KPI berwarna ---------- */
.summary-card{
  position:relative; background:var(--surface); border:1px solid var(--border); border-radius:var(--radius);
  padding:16px 18px; min-height:118px; margin:0 0 12px; box-shadow:var(--shadow); overflow:hidden; transition:box-shadow .15s ease,transform .15s ease;
}
.summary-card:hover{ box-shadow:0 4px 10px rgba(17,24,39,.06),0 10px 24px rgba(17,24,39,.07); transform:translateY(-1px); }
.summary-card:before{ display:none; }
.summary-icon{ float:right; width:42px; height:42px; border-radius:13px; background:var(--blue-soft); color:var(--blue); display:flex; align-items:center; justify-content:center; font-size:19px; }
.summary-card.c-blue .summary-icon{ background:var(--blue-soft); color:var(--blue); }
.summary-card.c-green .summary-icon{ background:var(--green-soft); color:var(--green); }
.summary-card.c-purple .summary-icon{ background:var(--purple-soft); color:var(--purple); }
.summary-card.c-orange .summary-icon{ background:var(--orange-soft); color:var(--orange); }
.summary-card.c-pink .summary-icon{ background:var(--pink-soft); color:var(--pink); }
.summary-card.c-teal .summary-icon{ background:var(--teal-soft); color:var(--teal); }
.summary-label{ font-size:11px; font-weight:700; letter-spacing:.6px; text-transform:uppercase; color:var(--muted); }
.summary-value{ margin-top:9px; font-size:24px; font-weight:800; line-height:1.2; color:var(--text); overflow-wrap:anywhere; letter-spacing:-.01em; }
.summary-sub{ margin-top:6px; font-size:11.5px; color:var(--faint); }

.kpi-wrap{ background:var(--surface); border:1px solid var(--border); border-radius:var(--radius); padding:14px 16px; box-shadow:var(--shadow); height:100%; }
.kpi-label{ font-size:11px; font-weight:700; letter-spacing:.6px; text-transform:uppercase; color:var(--muted); }
.kpi-value{ margin-top:4px; font-size:20px; font-weight:800; color:var(--text); overflow-wrap:anywhere; }
.mini-kpi{ display:flex; justify-content:space-between; align-items:center; margin-top:8px; padding:10px 14px; background:var(--surface); border:1px solid var(--border); border-radius:12px; }
.mini-kpi span{ font-size:10px; font-weight:700; letter-spacing:.8px; color:var(--muted); }
.mini-kpi b{ font-size:13px; color:var(--text); }

/* st.metric sebagai kartu, dengan aksen warna bergantian per posisi kolom */
[data-testid="stMetric"]{ background:var(--surface); border:1px solid var(--border); border-radius:var(--radius); padding:14px 16px; box-shadow:var(--shadow); min-height:92px; border-left:3px solid var(--blue); }
[data-testid="stColumn"]:nth-of-type(4n+2) [data-testid="stMetric"]{ border-left-color:var(--green); }
[data-testid="stColumn"]:nth-of-type(4n+3) [data-testid="stMetric"]{ border-left-color:var(--purple); }
[data-testid="stColumn"]:nth-of-type(4n+4) [data-testid="stMetric"]{ border-left-color:var(--orange); }
[data-testid="stColumn"]:nth-of-type(3n+2) [data-testid="stMetric"]{ border-left-color:var(--green); }
[data-testid="stColumn"]:nth-of-type(3n+3) [data-testid="stMetric"]{ border-left-color:var(--purple); }
[data-testid="stMetricLabel"], [data-testid="stMetricLabel"] *{ color:var(--muted) !important; font-size:11.5px !important; font-weight:700 !important; letter-spacing:.3px; }
[data-testid="stMetricValue"], [data-testid="stMetricValue"] *{ color:var(--text) !important; font-weight:800 !important; }
[data-testid="stMetricValue"]{ font-size:clamp(18px,1.7vw,25px) !important; line-height:1.2 !important; white-space:normal !important; overflow:visible !important; text-overflow:clip !important; overflow-wrap:anywhere; }
[data-testid="stMetricValue"] > div{ white-space:normal !important; overflow:visible !important; text-overflow:clip !important; }

/* ---------- Kartu grafik & tabel ---------- */
.chart-card, .dashboard-card, .insight-card{ background:var(--surface); border:1px solid var(--border); border-radius:var(--radius); box-shadow:var(--shadow); }
.chart-card{ padding:16px 18px 18px; border-bottom:0; border-radius:var(--radius) var(--radius) 0 0; margin:0; }
.chart-card h3{ margin:0 0 2px !important; font-size:14.5px !important; font-weight:700 !important; color:var(--text) !important; display:flex; align-items:center; gap:8px; }
.chart-card p{ margin:0 0 4px; font-size:12px; color:var(--muted); }
[data-testid="stElementContainer"]:has(.chart-card), .element-container:has(.chart-card){ margin-bottom:-1rem; }
[data-testid="stPlotlyChart"]{ background:var(--surface); border:1px solid var(--border); border-radius:var(--radius); padding:8px 10px; box-shadow:var(--shadow); }
[data-testid="stElementContainer"]:has(.chart-card) + [data-testid="stElementContainer"] [data-testid="stPlotlyChart"],
.element-container:has(.chart-card) + .element-container [data-testid="stPlotlyChart"]{ border-top:0; border-radius:0 0 var(--radius) var(--radius); }
.dashboard-card{ padding:16px 18px 12px; }
.dashboard-card h3{ margin:0 0 3px; font-size:15px; color:var(--text); }
.dashboard-card p{ margin:0 0 8px; font-size:12px; color:var(--muted); }
.insight-card{ padding:16px 18px; }
.insight-row{ display:flex; gap:10px; padding:10px 0; border-bottom:1px solid #f0f2f7; font-size:12.5px; color:#374151; }
.insight-row:last-child{ border-bottom:0; }
.insight-dot{ width:22px; height:22px; flex:0 0 22px; border-radius:50%; background:var(--green-soft); color:var(--green); display:flex; align-items:center; justify-content:center; font-weight:800; font-size:11px; }
.rank-table{ width:100%; border-collapse:collapse; font-size:12.5px; }
.rank-table th{ text-align:left; padding:10px 8px; color:var(--muted); font-size:10.5px; font-weight:700; text-transform:uppercase; letter-spacing:.5px; background:#f7f9fd; border-bottom:1px solid var(--border); }
.rank-table th:first-child{ border-top-left-radius:10px; } .rank-table th:last-child{ border-top-right-radius:10px; }
.rank-table td{ padding:10px 8px; border-bottom:1px solid #f0f2f7; color:#374151; }
.rank-table tr:hover td{ background:#f9fafc; }
.rank-table td.num{ text-align:right; font-weight:700; color:var(--text); }
.progress-bg{ height:7px; background:#e8ecf3; border-radius:99px; overflow:hidden; }
.progress-fill{ height:100%; background:var(--accent); border-radius:99px; }

[data-testid="stDataFrame"]{ border:1px solid var(--border); border-radius:var(--radius); overflow:hidden; box-shadow:var(--shadow); background:var(--surface); }

/* ---------- Tab ---------- */
.stTabs{ margin-top:6px; }
.stTabs [data-baseweb="tab-list"]{ gap:4px; padding:4px; background:#f2f4f9; border-radius:12px; border-bottom:0; }
.stTabs [data-baseweb="tab"]{ height:40px; padding:0 16px; border-radius:9px; background:transparent; color:var(--muted); font-size:13px; font-weight:600; }
.stTabs [data-baseweb="tab"]:hover{ color:var(--text); }
.stTabs [aria-selected="true"]{ color:var(--accent) !important; font-weight:700; background:var(--surface) !important; box-shadow:0 1px 3px rgba(17,24,39,.08); }
.stTabs [data-baseweb="tab-highlight"]{ background:transparent !important; height:0; }
.stTabs [data-baseweb="tab-border"]{ background:transparent !important; }

/* ---------- Tombol & input ---------- */
.stApp .stButton > button, .stApp .stDownloadButton > button, .stApp [data-testid^="stBaseButton-"]{
  border-radius:var(--radius-sm); font-weight:700; font-size:13px; min-height:38px; transition:background .12s ease,border-color .12s ease,box-shadow .12s ease;
}
.stApp [data-testid="stBaseButton-secondary"], .stApp [data-testid="stBaseButton-secondaryFormSubmit"]{
  background:var(--surface); color:var(--text); border:1px solid var(--border-strong);
}
.stApp [data-testid="stBaseButton-secondary"]:hover, .stApp [data-testid="stBaseButton-secondaryFormSubmit"]:hover{
  border-color:var(--accent); color:var(--accent); background:var(--accent-soft);
}
.stApp [data-testid="stBaseButton-primary"], .stApp [data-testid="stBaseButton-primaryFormSubmit"], .stApp button[kind="primary"], .stApp button[kind="primaryFormSubmit"]{
  background:var(--accent); color:#fff; border:1px solid var(--accent); box-shadow:0 1px 2px rgba(47,111,237,.25);
}
.stApp [data-testid="stBaseButton-primary"]:hover, .stApp [data-testid="stBaseButton-primaryFormSubmit"]:hover, .stApp button[kind="primary"]:hover, .stApp button[kind="primaryFormSubmit"]:hover{
  background:var(--accent-hover); border-color:var(--accent-hover); color:#fff;
}
/* Tombol download = minimalis (ghost, kecil) */
.stApp .stDownloadButton > button, .stApp .stDownloadButton > button[kind]{ background:transparent !important; color:#475569 !important; border:1px solid #e2e8f0 !important; box-shadow:none !important; min-height:30px !important; padding:2px 12px !important; font-size:12px !important; font-weight:600 !important; width:auto !important; }
.stApp .stDownloadButton > button:hover{ background:#f1f5f9 !important; border-color:#cbd5e1 !important; color:#0f172a !important; }
.stApp .stDownloadButton > button p, .stApp .stDownloadButton > button span{ color:inherit !important; font-size:12px !important; }
.stApp .stDownloadButton{ display:flex; justify-content:flex-end; }
.stApp [data-baseweb="input"] > div, .stApp [data-baseweb="select"] > div, .stApp [data-baseweb="textarea"]{
  background:var(--surface); border-radius:var(--radius-sm) !important; border-color:var(--border-strong) !important;
}
.stApp [data-baseweb="input"]:focus-within > div, .stApp [data-baseweb="select"]:focus-within > div{ border-color:var(--accent) !important; box-shadow:0 0 0 3px rgba(47,111,237,.14); }
.stApp [data-testid="stTextInputRootElement"], .stApp [data-testid="stNumberInputContainer"], .stApp [data-testid="stTextAreaRootElement"], .stApp [data-testid="stDateInputField"]{
  background:var(--surface) !important; border:1px solid var(--border-strong) !important; border-radius:var(--radius-sm) !important; box-shadow:none;
}
.stApp [data-testid="stTextInputRootElement"]:focus-within, .stApp [data-testid="stNumberInputContainer"]:focus-within, .stApp [data-testid="stTextAreaRootElement"]:focus-within{
  border-color:var(--accent) !important; box-shadow:0 0 0 3px rgba(47,111,237,.14) !important;
}
.stApp [data-testid="stMain"] [data-testid="stSelectbox"] > div > div, .stApp [data-testid="stMain"] [data-testid="stMultiSelect"] > div > div{
  background:var(--surface); border:1px solid var(--border-strong); border-radius:var(--radius-sm);
}
.stApp [data-testid="stWidgetLabel"] p, .stApp label p{ font-size:12.5px; font-weight:700; color:#374151; }

/* Form, expander, uploader, alert */
[data-testid="stForm"]{ background:var(--surface); border:1px solid var(--border); border-radius:var(--radius); padding:18px 20px; box-shadow:var(--shadow); }
[data-testid="stExpander"]{ background:var(--surface); border:1px solid var(--border) !important; border-radius:var(--radius) !important; box-shadow:var(--shadow); }
[data-testid="stExpander"] summary{ font-weight:700; }
[data-testid="stFileUploaderDropzone"]{ background:#f8fafc; border:1.5px dashed var(--border-strong); border-radius:12px; }
[data-testid="stFileUploaderDropzone"]:hover{ border-color:var(--accent); background:var(--accent-soft); }
[data-testid="stAlert"]{ border-radius:12px; border:1px solid var(--border); }
.upload-note{ margin:0 0 14px; padding:11px 14px; background:var(--blue-soft); border:1px solid #cfe0fd; border-left:3px solid var(--accent); border-radius:12px; color:#374151; font-size:12.5px; line-height:1.5; }
.uploaded-ok{ margin-top:6px; padding:6px 10px; border-radius:8px; background:var(--ok-soft); color:#15803d; font-size:11.5px; font-weight:700; }
.upload-progress{ display:flex; justify-content:space-between; align-items:center; margin-top:8px; padding:12px 16px; background:var(--surface); color:var(--text); border:1px solid var(--border); border-left:3px solid var(--accent); border-radius:12px; }
.upload-progress span{ color:var(--muted); font-size:12px; }
.report-download{ margin:4px 0 12px; }
.footer-band{ margin-top:26px; padding:12px 0 4px; background:transparent; color:var(--faint); text-align:center; font-size:11px; letter-spacing:.3px; border-top:1px solid var(--border); border-radius:0; }

/* ---------- Halaman login / lisensi ---------- */
.license-shell{ max-width:520px; margin:7vh auto 14px; }
.license-card{ background:var(--surface); border:1px solid var(--border); border-radius:18px; padding:26px 28px; box-shadow:0 10px 30px rgba(17,24,39,.08); }
.license-brand{ display:flex; align-items:center; gap:14px; margin-bottom:14px; }
.license-brand img{ width:52px; height:52px; object-fit:contain; border-radius:13px; }
.license-brand b{ font-size:18px; font-weight:800; color:var(--text); letter-spacing:.2px; }
.license-brand span{ display:block; margin-top:3px; font-size:10.5px; letter-spacing:1.1px; color:var(--muted); }
.license-card p{ margin:0; }
[data-testid="stElementContainer"]:has(.license-shell), .element-container:has(.license-shell){ margin-bottom:-1rem; }
.license-shell{ margin-bottom:0; }
.license-card{ border-bottom:0; border-radius:18px 18px 0 0; box-shadow:none; padding-bottom:8px; }
.stApp:has(.license-shell) [data-testid="stForm"]{ max-width:520px; margin:0 auto; border-top:0; border-radius:0 0 18px 18px; padding:6px 28px 24px; box-shadow:0 10px 30px rgba(17,24,39,.08); }
.stApp:has(.license-shell) [data-testid="stAlert"]{ max-width:520px; margin:0 auto; }

/* ---------- V21b: donat Dashboard (latar putih, sama tinggi, legenda scroll) ---------- */
.st-key-dn_sa, .st-key-dn_ly{ min-height:650px; background:#ffffff !important; }
.st-key-dn_sa [data-testid="stPlotlyChart"], .st-key-dn_ly [data-testid="stPlotlyChart"]{ background:#ffffff !important; border:0 !important; box-shadow:none !important; padding:0 !important; }
.st-key-dn_sa .rk-lgw, .st-key-dn_ly .rk-lgw{ max-height:300px; overflow-y:auto; padding-right:6px; margin-top:0; }
.st-key-dn_sa .rk-lgw::-webkit-scrollbar, .st-key-dn_ly .rk-lgw::-webkit-scrollbar{ width:6px; }

/* ---------- V21b: jarak antar blok (global) & kartu minimalis ---------- */
/* V23: jarak antar blok dilonggarkan di SEMUA halaman. Blok tingkat atas 2rem, blok di dalam container 1.75rem;
   isi kartu (container ber-border / ber-key) tetap rapat 1rem. */
[data-testid="stMain"] [data-testid="stVerticalBlock"]{ --vgap:2rem; gap:var(--vgap); }
[data-testid="stMain"] [data-testid="stVerticalBlock"] [data-testid="stVerticalBlock"]{ --vgap:1.75rem; gap:var(--vgap); }
[data-testid="stMain"] [data-testid="stVerticalBlockBorderWrapper"] [data-testid="stVerticalBlock"],
[data-testid="stMain"] [data-testid="stVerticalBlock"][class*="st-key-"]{ --vgap:1rem; gap:var(--vgap); }
[data-testid="stMain"] [data-testid="stHorizontalBlock"]{ gap:1.25rem; }
[data-testid="stElementContainer"]:has(.chart-card), .element-container:has(.chart-card){ margin-bottom:calc(-1 * var(--vgap,2rem)) !important; }
[data-testid="stElementContainer"]:has(.pg-wrap){ margin-bottom:calc(.55rem - var(--vgap,2rem)) !important; }
.rk-head, .rk-grid, .rk-dark, .rk-tw, .rk-band, .rk-mg{ margin-top:0 !important; margin-bottom:0 !important; }
.rk-grid{ align-items:stretch; }
.rk-card, .rk-mini{ height:100%; box-sizing:border-box; }
.rk-mg{ display:grid; grid-template-columns:repeat(auto-fit,minmax(190px,1fr)); gap:14px; }
.rk-mn{ background:transparent; border-left:3px solid #cbd5e1; padding:2px 0 2px 12px; min-width:0; }
.rk-mn .l{ display:block; font-size:10.5px; font-weight:700; letter-spacing:.4px; text-transform:uppercase; color:#64748b; }
.rk-mn .v{ display:block; font-size:19px; font-weight:800; color:#0f172a; line-height:1.25; margin-top:2px; }
.rk-mn .s{ display:block; font-size:11px; color:#94a3b8; margin-top:1px; }
.rk-mn.blue{ border-color:#2563eb; } .rk-mn.amber{ border-color:#f59e0b; } .rk-mn.green{ border-color:#10b981; } .rk-mn.purple{ border-color:#8b5cf6; }
.rk-tw{ max-height:620px; overflow:auto; box-shadow:0 1px 4px rgba(15,23,42,.04); }
.rk-t th{ position:sticky; top:0; z-index:1; }
.rk-t tbody tr:nth-child(even) td{ background:#fbfdff; }
.rk-t tbody tr:hover td{ background:#f1f6ff; }
.rk-t tr.tot td{ background:#eaf0fa !important; }
[data-testid="stMain"] [data-testid="stCaptionContainer"]{ line-height:1.5; }
[data-testid="stMain"] h3{ font-size:16px !important; font-weight:800 !important; margin:0 !important; padding:0 !important; color:#0f172a; }
[data-testid="stMain"] hr{ margin:4px 0 !important; border-color:#e8eef6 !important; }
[data-testid="stPlotlyChart"]{ padding:10px 12px; }
[data-testid="stMain"] [data-testid="stButtonGroup"]{ background:#eef2f9; border:1px solid #dbe4f3; border-radius:12px; padding:3px; gap:3px; width:fit-content; max-width:100%; flex-wrap:wrap; }
[data-testid="stMain"] [data-testid="stButtonGroup"] button{ border-radius:9px !important; border:0 !important; min-height:32px !important; font-size:12.5px !important; font-weight:600 !important; background:transparent !important; box-shadow:none !important; }
[data-testid="stMain"] [data-testid="stButtonGroup"] button[aria-checked="true"], [data-testid="stMain"] [data-testid="stButtonGroup"] button[aria-pressed="true"], [data-testid="stMain"] [data-testid="stButtonGroup"] button[kind$="Active"]{ background:#ffffff !important; color:#1d4ed8 !important; box-shadow:0 1px 4px rgba(15,23,42,.12) !important; }
@media (max-width:900px){
  .block-container{ padding:1rem .9rem 2.5rem !important; }
  .rk-ph-t{ font-size:21px; } .rk-val{ font-size:19px; }
  .st-key-dn_sa, .st-key-dn_ly{ min-height:0; }
}


/* ---------- V23: blok header + kontrol (satu kartu), tabel berhalaman, kartu donat UE, strategi ---------- */
[data-testid="stMain"] [class*="st-key-blk_"]{ background:#fff; border:1px solid #e3eaf5; border-radius:14px; padding:12px 16px; box-shadow:0 1px 4px rgba(15,23,42,.04); }
[class*="st-key-blk_"] .rk-head{ border:0 !important; box-shadow:none !important; padding:0 !important; background:transparent !important; margin:0 !important; }
[class*="st-key-blk_"] [data-testid="stSelectbox"] > div > div{ min-height:40px; }
[class*="st-key-blk_"] .stButton > button{ min-height:40px; white-space:nowrap; }
[class*="st-key-pgn_"] [data-testid="stHorizontalBlock"]{ align-items:center !important; flex-wrap:nowrap !important; gap:.3rem !important; justify-content:flex-end; }
[class*="st-key-pgn_"] [data-testid="stColumn"]{ flex:0 0 auto !important; width:auto !important; min-width:0 !important; }
[class*="st-key-pgn_"] [data-testid="stColumn"]:first-child{ flex:1 1 auto !important; }
[class*="st-key-pgn_"] .stButton > button{ min-height:32px !important; min-width:34px; padding:0 10px !important; border-radius:8px !important; font-size:12.5px !important; font-weight:700 !important; }
[class*="st-key-pgn_"] .stButton > button:disabled{ opacity:.45; }
.rk-pginfo{ font-size:12px; color:#64748b; padding:2px 2px; }
.rk-pginfo b{ color:#0f172a; }
.rk-chips{ display:flex; flex-wrap:wrap; gap:6px; }
.rk-chp{ font-family:ui-monospace,monospace; font-size:11px; background:#f1f5f9; border:1px solid #e2e8f0; border-radius:6px; padding:2px 8px; color:#334155; }
.rk-mapcol{ border:1px solid #e3eaf5; border-radius:12px; padding:12px 14px; background:#fafcff; }
.rk-mapcol .t{ display:flex; justify-content:space-between; align-items:center; margin-bottom:9px; font-weight:800; font-size:13px; color:#0f172a; }
.rk-mapcol small{ display:block; margin-top:8px; font-size:11px; color:#64748b; line-height:1.45; }
.st-key-dn_ue_sa, .st-key-dn_ue_bk, .st-key-ue_c1, .st-key-ue_c2{ background:#ffffff !important; border:1px solid #e3eaf5 !important; border-radius:16px !important; box-shadow:0 1px 6px rgba(15,23,42,.05) !important; padding:16px 20px 14px !important; }
.st-key-dn_ue_sa [data-testid="stPlotlyChart"], .st-key-dn_ue_bk [data-testid="stPlotlyChart"], .st-key-ue_c1 [data-testid="stPlotlyChart"], .st-key-ue_c2 [data-testid="stPlotlyChart"]{ background:#ffffff !important; border:0 !important; box-shadow:none !important; padding:0 !important; }
.st-key-dn_ue_sa .rk-lgw, .st-key-dn_ue_bk .rk-lgw{ max-height:300px; overflow-y:auto; padding-right:6px; margin-top:0; }
.rk-ins p b{ display:inline; font-size:inherit; text-transform:none; letter-spacing:0; margin:0; color:#fff; }
/* Management Insight V23 */
.mi-sec{ display:flex; align-items:center; gap:10px; margin:0; }
.mi-sec .n{ width:28px; height:28px; border-radius:9px; background:#2563eb; color:#fff; display:flex; align-items:center; justify-content:center; font-weight:800; font-size:13px; flex:0 0 auto; }
.mi-sec b{ font-size:16px; color:#0f172a; font-weight:800; } .mi-sec small{ display:block; color:#64748b; font-size:12px; font-weight:500; margin-top:1px; }
.mi-g{ display:grid; grid-template-columns:repeat(auto-fit,minmax(var(--mw,290px),1fr)); gap:14px; }
.mi-card{ background:#fff; border:1px solid #e3eaf5; border-radius:14px; padding:15px 17px; box-shadow:0 1px 4px rgba(15,23,42,.04); display:flex; flex-direction:column; gap:9px; min-width:0; border-top:3px solid #cbd5e1; }
.mi-card.blue{ border-top-color:#2563eb; } .mi-card.green{ border-top-color:#10b981; } .mi-card.orange{ border-top-color:#f59e0b; } .mi-card.red{ border-top-color:#ef4444; } .mi-card.purple{ border-top-color:#8b5cf6; } .mi-card.teal{ border-top-color:#0d9488; }
.mi-card h4{ margin:0; font-size:14px; font-weight:800; color:#0f172a; display:flex; align-items:center; gap:8px; }
.mi-card h4 .msr{ font-size:18px; }
.mi-card .ev{ font-size:12px; color:#475569; line-height:1.5; background:#f8fafc; border-radius:9px; padding:8px 10px; }
.mi-card ul{ margin:0; padding-left:18px; font-size:12.5px; color:#1e293b; line-height:1.55; } .mi-card li{ margin-bottom:3px; }
.mi-card .kp{ font-size:11.5px; color:#475569; border-top:1px dashed #e2e8f0; padding-top:8px; } .mi-card .kp b{ color:#0f172a; }
.mi-big{ font-size:26px; font-weight:900; color:#0f172a; letter-spacing:-.4px; line-height:1.15; } .mi-big small{ font-size:13px; color:#64748b; font-weight:600; }
.mi-ach{ display:grid; grid-template-columns:repeat(auto-fit,minmax(300px,1fr)); gap:14px; }
.mi-kv{ display:flex; justify-content:space-between; gap:8px; font-size:12px; color:#475569; padding:4px 0; border-bottom:1px solid #f1f5f9; } .mi-kv b{ color:#0f172a; text-align:right; }
.mi-kv:last-child{ border-bottom:0; }
.mi-sc{ background:#f8fafc; border:1px solid #e3eaf5; border-radius:12px; padding:12px 14px; } .mi-sc .h{ font-size:10.5px; font-weight:800; letter-spacing:.5px; text-transform:uppercase; color:#475569; } .mi-sc .v{ font-size:19px; font-weight:800; color:#0f172a; margin:4px 0 2px; } .mi-sc small{ color:#64748b; font-size:11.5px; line-height:1.45; display:block; }

/* Scrollbar halus */
*::-webkit-scrollbar{ width:10px; height:10px; }
*::-webkit-scrollbar-thumb{ background:#ccd2de; border-radius:10px; border:2px solid transparent; background-clip:content-box; }
*::-webkit-scrollbar-thumb:hover{ background:#aab2c2; background-clip:content-box; border:2px solid transparent; }

</style>
''', unsafe_allow_html=True)

def section(title, subtitle=''):
    """Judul seksi memakai kartu header yang sama dengan Dashboard."""
    _m=re.match(r'\s*<span class="msr">([^<]+)</span>\s*(.*)',str(title),re.S)
    _ic,_t=(_m.group(1),_m.group(2)) if _m else ('insights',str(title))
    st.markdown(RK.head_card(re.sub(r'<[^>]+>','',_t),subtitle,_ic,'blue'),unsafe_allow_html=True)

# ---- Kartu & tabel seragam (desain Dashboard): st.metric -> kartu KPI, st.dataframe kecil -> tabel RK ----
if '_rk_orig_df' not in st.__dict__: st.__dict__['_rk_orig_df']=st.dataframe
_ORIG_DATAFRAME=st.__dict__['_rk_orig_df']
_NUMKEEP=re.compile(r'(tahun|year|^no\b|\bno\b|^id\b|kode|phone|tlp|hp|polisi|pkb|invoice|dokumen|referensi|rank|periode)',re.I)
def _rk_cell(v,col,intcol):
    try:
        if v is None or (not isinstance(v,(str,list,dict,tuple)) and pd.isna(v)): return '-'
    except Exception: pass
    if isinstance(v,(pd.Timestamp,datetime.datetime)): return v.strftime('%d/%m/%Y' if (v.hour==0 and v.minute==0) else '%d/%m/%Y %H:%M')
    if isinstance(v,datetime.date): return v.strftime('%d/%m/%Y')
    if isinstance(v,(bool,np.bool_)): return 'Ya' if v else 'Tidak'
    if isinstance(v,(int,np.integer)): return str(v) if _NUMKEEP.search(str(col)) else f'{int(v):,}'
    if isinstance(v,(float,np.floating)):
        if _NUMKEEP.search(str(col)) and float(v).is_integer(): return str(int(v))
        return f'{v:,.0f}' if intcol else f'{v:,.2f}'
    return _html.escape(str(v))
def _rk_table_from(data):
    """DataFrame / Styler kecil -> HTML tabel desain Dashboard. None bila tidak cocok."""
    sty=None
    if isinstance(data,pd.DataFrame): df=data
    elif hasattr(data,'to_html') and isinstance(getattr(data,'data',None),pd.DataFrame): sty=data; df=data.data
    else: return None
    if isinstance(df.columns,pd.MultiIndex) or len(df)>60 or len(df.columns)>18 or len(df.columns)==0: return None
    if sty is not None:
        try: h=sty.hide(axis='index').to_html(table_attributes='class="rk-t"')
        except Exception: return None
        h=re.sub(r'<td([^>]*?)class="([^"]*)"([^>]*)>\s*([^<]*?)\s*</td>',lambda m:f'<td{m.group(1)}class="{m.group(2)}{" n" if re.fullmatch(r"[-+]?\s*(Rp\s*)?[\d.,]+\s*(%|UE|Jt)?",m.group(4)) else ""}"{m.group(3)}>{m.group(4)}</td>',h)
        return f'<div class="rk-tw">{h}</div>'
    num=[i for i,c in enumerate(df.columns) if pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c]) and not _NUMKEEP.search(str(c))]
    rows=[]
    for _,r in df.iterrows():
        cells=[]
        for i,c in enumerate(df.columns):
            col=df[c]; ic=bool(pd.api.types.is_numeric_dtype(col) and (col.dropna()%1==0).all()) if pd.api.types.is_numeric_dtype(col) else False
            cells.append(_rk_cell(r[c],c,ic))
        rows.append(cells)
    return RK.table_html([str(c) for c in df.columns],rows,num_cols=tuple(num))
def _rk_dataframe(data=None,*a,**k):
    if not (k.get('column_config') or k.get('height') or k.get('on_select') not in (None,'ignore') or k.get('column_order')):
        try:
            h=_rk_table_from(data)
            if h:
                st.markdown(h,unsafe_allow_html=True); return None
        except Exception: pass
    return _ORIG_DATAFRAME(data,*a,**k)
def _rk_metric_html(label,value,delta=None):
    d=(f'<div class="rk-foot">{_html.escape(str(delta))}</div>' if delta not in (None,'') else '')
    return f'<div class="rk-card"><div class="rk-top"><span class="rk-lab">{_html.escape(str(label))}</span></div><div class="rk-val">{_html.escape(str(value))}</div>{d}</div>'
def _rk_metric(label,value,delta=None,*a,**k): st.markdown(_rk_metric_html(label,value,delta),unsafe_allow_html=True)
st.dataframe=_rk_dataframe; st.metric=_rk_metric
try:
    from streamlit.delta_generator import DeltaGenerator as _DG
    if not hasattr(_DG,'_rk_orig_df'): _DG._rk_orig_df=_DG.dataframe
    _DG_df=_DG._rk_orig_df
    def _dg_dataframe(self,data=None,*a,**k):
        if not (k.get('column_config') or k.get('height') or k.get('on_select') not in (None,'ignore') or k.get('column_order')):
            try:
                h=_rk_table_from(data)
                if h: self.markdown(h,unsafe_allow_html=True); return None
            except Exception: pass
        return _DG_df(self,data,*a,**k)
    def _dg_metric(self,label,value,delta=None,*a,**k): self.markdown(_rk_metric_html(label,value,delta),unsafe_allow_html=True)
    if not getattr(_DG,'_rk_patched',False): _DG.dataframe=_dg_dataframe; _DG.metric=_dg_metric; _DG._rk_patched=True
except Exception: pass

# ---- V23: tabel seragam berhalaman (maks 10 baris per halaman + nomor halaman) ----
def _pg_goto(pk,n): st.session_state[pk]=int(n)
def paged_rows(headers,rows,key,per=10,num_cols=(),hl_cols=None,row_cls=None,bold_last=False,unit='baris'):
    """Tabel desain Dashboard (RK) dengan paging. `key` HARUS unik per tabel. rows: list[list[str-html]]."""
    n=len(rows); pages=max(1,-(-n//per))
    pk=f'{key}__page'; sk=f'{key}__sig'
    sig=(n,tuple(headers),str(rows[0]) if n else '',str(rows[-1]) if n else '')
    if st.session_state.get(sk)!=sig: st.session_state[sk]=sig; st.session_state[pk]=1   # data/filter berubah -> kembali ke halaman 1
    try: cur=int(st.session_state.get(pk,1))
    except Exception: cur=1
    cur=min(max(cur,1),pages); st.session_state[pk]=cur
    a=(cur-1)*per; b=min(a+per,n)
    rc=(row_cls[a:b] if row_cls else None)
    _th=RK.table_html(headers,rows[a:b],num_cols=num_cols,hl_cols=hl_cols,row_cls=rc,bold_last=(bold_last and pages==1))
    info=f'<div class="rk-pginfo">Menampilkan <b>{a+1 if n else 0}–{b}</b> dari <b>{n:,}</b> {unit}'+(f' • Halaman <b>{cur}</b> / {pages}' if pages>1 else '')+'</div>'
    if pages==1:
        st.markdown(f'<div>{_th}{info}</div>',unsafe_allow_html=True); return
    st.markdown(f'<div class="pg-wrap">{_th}</div>',unsafe_allow_html=True)
    nums=sorted({1,pages,cur,cur-1,cur+1,cur-2,cur+2} & set(range(1,pages+1)))
    seq=[]
    for i,v in enumerate(nums):
        if i and v-nums[i-1]>1: seq.append(None)
        seq.append(v)
    with st.container(key=f'pgn_{key}'):
        cols=st.columns([1]+[0.001]*(len(seq)+2),gap='small')
        cols[0].markdown(info,unsafe_allow_html=True)
        cols[1].button('‹',key=f'{key}__prev',on_click=_pg_goto,args=(pk,cur-1),disabled=(cur<=1),help='Halaman sebelumnya')
        for ci,v in enumerate(seq):
            if v is None: cols[2+ci].button('…',key=f'{key}__gap{ci}',disabled=True)
            else: cols[2+ci].button(str(v),key=f'{key}__p{v}',on_click=_pg_goto,args=(pk,v),type=('primary' if v==cur else 'secondary'))
        cols[2+len(seq)].button('›',key=f'{key}__next',on_click=_pg_goto,args=(pk,cur+1),disabled=(cur>=pages),help='Halaman berikutnya')
def paged_df(df,key,per=10,money=(),pct=(),code=(),hl_cols=None,unit='baris'):
    """DataFrame -> paged_rows. money: kolom Rupiah; pct: kolom persen (nilai 0-100); code: kolom kode (monospace)."""
    if not isinstance(df,pd.DataFrame): return
    heads=[str(c) for c in df.columns]
    isnum={c:bool(pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c])) for c in df.columns}
    isint={c:bool(isnum[c] and (df[c].dropna()%1==0).all()) for c in df.columns}
    numidx=[i for i,c in enumerate(df.columns) if isnum[c] and (c in money or c in pct or not _NUMKEEP.search(str(c)))]
    rows=[]
    for rec in df.itertuples(index=False,name=None):
        cells=[]
        for c,v in zip(df.columns,rec):
            if v is None or (isnum[c] and pd.isna(v)): cells.append('-')
            elif c in money and isnum[c]: cells.append(fmt(v))
            elif c in pct and isnum[c]: cells.append(f'{float(v):.1f}%')
            elif c in code: cells.append(f'<span style="font-family:ui-monospace,monospace">{_html.escape(str(v))}</span>')
            else: cells.append(_rk_cell(v,c,isint[c]))
        rows.append(cells)
    paged_rows(heads,rows,key,per=per,num_cols=tuple(numidx),hl_cols=hl_cols,unit=unit)
def donut_card_body(labels,values,colors,center_title,center_main,center_sub,legend_items,height=300):
    """Donat + legenda bergaya Dashboard (dipakai UE per SA & UE per Layanan)."""
    fig=go.Figure(go.Pie(labels=labels,values=values,hole=.62,sort=False,textinfo='none',marker=dict(colors=colors,line=dict(color='#ffffff',width=2)),hovertemplate='%{label}<br>%{value:,.0f} (%{percent})<extra></extra>'))
    fig.update_layout(height=height,margin=dict(l=4,r=4,t=4,b=4),hoverlabel=dict(bgcolor='#0f172a',font_color='#fff',font_size=12),paper_bgcolor='#ffffff',plot_bgcolor='#ffffff',showlegend=False,
        annotations=[dict(text=f'<span style="font-size:11px;color:#64748b">{center_title}</span><br><span style="font-size:17px"><b>{center_main}</b></span><br><span style="font-size:11px;color:#2563eb">{center_sub}</span>',x=.5,y=.5,showarrow=False)])
    q1,q2=st.columns([1,1.15])
    with q1: st.plotly_chart(fig,use_container_width=True,config=PCFG)
    with q2: st.markdown(RK.legend_list(legend_items),unsafe_allow_html=True)

def clean(x):
    if pd.isna(x): return ''
    return re.sub(r'\s+',' ',str(x)).strip()
def norm(x):
    s=clean(x).upper()
    return re.sub(r'\s+','',s) if s and s not in {'NAN','NONE','-'} else ''
def num(x):
    if pd.isna(x) or x=='': return 0.0
    if isinstance(x,(int,float,np.number)): return float(x)
    s=re.sub(r'[^0-9.\-]','',str(x).replace(',',''))
    try:return float(s) if s else 0.0
    except:return 0.0
def header_row(raw, need, limit=100):
    for i in range(min(len(raw),limit)):
        vals=[clean(v).upper() for v in raw.iloc[i].tolist()]
        if sum(any(n in v for v in vals) for n in need)>=max(2,len(need)-1): return i
    return None

def fmt(x): return 'Rp '+f'{x:,.0f}'.replace(',','.')
def make_excel_download(master_df, sales_df, parts_df, issue_df, cons_df, opl_df, cust_df):
    # Export ke Excel asli (.xlsx), bukan CSV, sehingga setiap field menjadi kolom terpisah.
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    sheets = {
        'Master': master_df,
        'Sales': sales_df,
        'Parts': parts_df,
        'Issue To PKB': issue_df,
        'Consumable': cons_df,
        'OPL': opl_df,
        'Customer': cust_df,
    }
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
        for name, df in sheets.items():
            out = df.copy() if isinstance(df, pd.DataFrame) else pd.DataFrame()
            # Excel tidak suka nama kolom duplikat; beri nama unik bila ada.
            seen = {}
            cols = []
            for col in out.columns:
                base = str(col)
                n = seen.get(base, 0)
                seen[base] = n + 1
                cols.append(base if n == 0 else f'{base}_{n+1}')
            out.columns = cols
            out.to_excel(writer, sheet_name=name[:31], index=False)

    output.seek(0)
    wb = load_workbook(output)
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 22
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        # Format tanggal, angka, dan lebar kolom agar hasil siap dipakai.
        for col_idx, cell in enumerate(ws[1], start=1):
            header = str(cell.value or '').lower()
            max_len = len(str(cell.value or ''))
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                c = row[0]
                if c.value is not None:
                    max_len = min(max(max_len, len(str(c.value))), 45)
                    if hasattr(c.value, 'year') and hasattr(c.value, 'month'):
                        c.number_format = 'dd/mm/yyyy'
                    if any(k in header for k in ['jasa','sparepart','bahan','revenue','sales','cost','gp','harga','discount','amount','profit']):
                        if isinstance(c.value, (int, float)):
                            c.number_format = '#,##0'
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(max_len + 2, 45))
        # Nomor PKB/polisi/reference diperlakukan sebagai text agar tidak berubah oleh Excel.
        for col_idx, cell in enumerate(ws[1], start=1):
            header = str(cell.value or '').lower()
            if any(k in header for k in ['pkb','polisi','invoice','billing','referensi','dokumen','equipment','no part','opl no']):
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    row[0].number_format = '@'
    final = io.BytesIO()
    wb.save(final)
    final.seek(0)
    return final.getvalue()



def make_template_unit_service(master_df):
    """Export data ke format Template_Unit_Service.xlsx milik user.
    Satu baris = satu UE/PKB, mengikuti 9 kolom template asli.
    """
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    template_path=os.path.join(os.path.dirname(__file__),'Template_Unit_Service.xlsx')
    if not os.path.isfile(template_path):
        raise FileNotFoundError('Template_Unit_Service.xlsx tidak ditemukan di folder aplikasi.')

    wb=load_workbook(template_path)
    ws=wb['Template'] if 'Template' in wb.sheetnames else wb[wb.sheetnames[0]]

    # Bersihkan data lama, tetapi pertahankan header dan format template.
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row-1)

    df=master_df.copy()
    if 'Tanggal' in df.columns:
        df['Tanggal']=pd.to_datetime(df['Tanggal'],errors='coerce')
    # Mapping persis dengan template contoh.
    rows=[]
    for _,r in df.iterrows():
        tanggal=r.get('Tanggal',pd.NaT)
        no_pol=clean(r.get('No Polisi',''))
        pkb=clean(r.get('PKB',''))
        no_wo=no_pol or pkb
        branch=clean(r.get('Cabang','Kharisma Workshop')) or 'Kharisma Workshop'
        tipe=clean(r.get('Type Kendaraan',''))
        entry=clean(r.get('Booking Type',''))
        service=clean(r.get('Jenis Pekerjaan',''))
        rows.append([tanggal if pd.notna(tanggal) else None,no_wo,branch,tipe,entry,service,
                     float(r.get('Jasa',0) or 0),float(r.get('Sparepart',0) or 0),float(r.get('Bahan',0) or 0)])

    for row in rows:
        ws.append(row)

    # Rapikan format mengikuti template dan tetap mempertahankan nama kolom persis.
    for cell in ws[1]:
        cell.font=cell.font.copy(bold=True)
        cell.alignment=cell.alignment.copy(horizontal='center',vertical='center')
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for row in ws.iter_rows(min_row=2,max_row=ws.max_row,min_col=1,max_col=9):
        row[0].number_format='dd.mm.yyyy'
        for c in row[6:9]: c.number_format='#,##0'
        row[1].number_format='@'
    # Lebar kolom mengikuti template, tetapi tidak dibiarkan terlalu sempit.
    widths=[14,22,18,20,14,38,16,16,16]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.row_dimensions[1].height=24

    out=io.BytesIO(); wb.save(out); out.seek(0)
    return out.getvalue()

def make_single_excel_download(df, sheet_name='Data', filename='export.xlsx'):
    """Buat file XLSX untuk tabel tertentu, lengkap dengan filter, freeze header, dan format dasar."""
    if df is None:
        df = pd.DataFrame()
    out_df = df.copy()
    output = io.BytesIO()
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(output, engine='openpyxl', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
        seen = {}
        cols = []
        for col in out_df.columns:
            base = str(col)
            n = seen.get(base, 0)
            seen[base] = n + 1
            cols.append(base if n == 0 else f'{base}_{n+1}')
        out_df.columns = cols
        out_df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
    output.seek(0)
    wb = load_workbook(output)
    ws = wb[wb.sheetnames[0]]
    ws.freeze_panes = 'A2'
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 24
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill(fill_type='solid', fgColor='102A43')
            cell.alignment = Alignment(horizontal='center', vertical='center')
        for col_idx, cell in enumerate(ws[1], start=1):
            header = str(cell.value or '').lower()
            max_len = len(str(cell.value or ''))
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                c = row[0]
                if c.value is not None:
                    max_len = min(max(max_len, len(str(c.value))), 45)
                    if hasattr(c.value, 'year') and hasattr(c.value, 'month'):
                        c.number_format = 'dd/mm/yyyy'
                    if any(k in header for k in ['jasa','sparepart','bahan','revenue','sales','cost','gp','harga','discount','amount','profit','net']):
                        if isinstance(c.value, (int, float)):
                            c.number_format = '#,##0'
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(max_len + 2, 45))
            if any(k in header for k in ['pkb','polisi','invoice','billing','referensi','dokumen','equipment','no part','opl no']):
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    row[0].number_format = '@'
    final = io.BytesIO()
    wb.save(final)
    final.seek(0)
    return final.getvalue()

def bahan_issue_by_branch(master_df, issue_df, branches=None):
    """Rekonsiliasi Bahan Issue To PKB per cabang: nilai ACUAN (master = Estimasi SA / Service Penjualan, sama dengan Dashboard)
    dibandingkan dengan total mentah laporan Issue To PKB. Baris 'Belum terpetakan' = Issue yang PKB-nya belum ditemukan."""
    m=master_df if isinstance(master_df,pd.DataFrame) else pd.DataFrame()
    ref=(m.groupby('Cabang',as_index=False).agg(**{'Jumlah PKB':('PKB','nunique'),'Bahan (acuan Dashboard)':('Bahan','sum')})
         if (not m.empty and 'Cabang' in m.columns and 'Bahan' in m.columns) else pd.DataFrame(columns=['Cabang','Jumlah PKB','Bahan (acuan Dashboard)']))
    raw=pd.Series(dtype=float); unm=0.0
    if isinstance(issue_df,pd.DataFrame) and not issue_df.empty and 'Harga Net' in issue_df.columns and 'Cabang' in issue_df.columns:
        ok=issue_df if 'Sumber Cabang' not in issue_df.columns else issue_df[issue_df['Sumber Cabang']!='Tidak terpetakan']
        raw=ok.groupby('Cabang')['Harga Net'].sum()
        if 'Sumber Cabang' in issue_df.columns: unm=float(issue_df.loc[issue_df['Sumber Cabang']=='Tidak terpetakan','Harga Net'].sum())
    out=ref.copy(); out['Laporan Issue To PKB (mentah)']=out['Cabang'].map(raw).fillna(0.0)
    for cb in (branches or []):
        if cb not in set(out['Cabang']) and cb not in set(raw.index): out=pd.concat([out,pd.DataFrame([{'Cabang':cb,'Jumlah PKB':0,'Bahan (acuan Dashboard)':0.0,'Laporan Issue To PKB (mentah)':0.0}])],ignore_index=True)
    for cb in raw.index:
        if cb not in set(out['Cabang']): out=pd.concat([out,pd.DataFrame([{'Cabang':cb,'Jumlah PKB':0,'Bahan (acuan Dashboard)':0.0,'Laporan Issue To PKB (mentah)':float(raw[cb])}])],ignore_index=True)
    if unm: out=pd.concat([out,pd.DataFrame([{'Cabang':'Belum terpetakan','Jumlah PKB':0,'Bahan (acuan Dashboard)':0.0,'Laporan Issue To PKB (mentah)':unm}])],ignore_index=True)
    out['Selisih (acuan - mentah)']=out['Bahan (acuan Dashboard)']-out['Laporan Issue To PKB (mentah)']
    if not out.empty:
        tot={'Cabang':'TOTAL','Jumlah PKB':int(out['Jumlah PKB'].sum())}
        for c in ['Bahan (acuan Dashboard)','Laporan Issue To PKB (mentah)','Selisih (acuan - mentah)']: tot[c]=float(out[c].sum())
        out=pd.concat([out,pd.DataFrame([tot])],ignore_index=True)
    return out

def make_multi_sheet_excel(sheets):
    """sheets: dict {nama_sheet: DataFrame}. Header berwarna, filter, freeze panes, format angka."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl',datetime_format='dd/mm/yyyy',date_format='dd/mm/yyyy') as w:
        for name,df in sheets.items():
            d=df.copy() if isinstance(df,pd.DataFrame) else pd.DataFrame()
            if d.empty: d=pd.DataFrame({'Keterangan':['Tidak ada data']})
            for c in d.columns:
                if pd.api.types.is_datetime64_any_dtype(d[c]): d[c]=d[c].dt.strftime('%d/%m/%Y').where(d[c].notna(),'')
            d.to_excel(w,sheet_name=str(name)[:31],index=False)
    out.seek(0); wb=load_workbook(out)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill(fill_type='solid',fgColor='102A43'); cell.alignment=Alignment(horizontal='center',vertical='center')
        for ci,cell in enumerate(ws[1],start=1):
            h=str(cell.value or '').lower(); ml=len(str(cell.value or ''))
            for row in ws.iter_rows(min_row=2,min_col=ci,max_col=ci):
                c=row[0]
                if c.value is not None:
                    ml=min(max(ml,len(str(c.value))),45)
                    if isinstance(c.value,(int,float)) and any(k in h for k in ['jasa','sparepart','bahan','revenue','sales','cost','gp','harga','net','selisih']): c.number_format='#,##0'
                    if any(k in h for k in ['pkb','polisi','referensi','dokumen']): c.number_format='@'
            ws.column_dimensions[get_column_letter(ci)].width=max(10,min(ml+2,45))
    fin=io.BytesIO(); wb.save(fin); return fin.getvalue()

def download_table(df, label, filename, key):
    """Tombol download XLSX untuk setiap tabel yang tampil."""
    if isinstance(df, pd.DataFrame):
        dl_button(':material/download: Excel', make_single_excel_download(df, filename.rsplit('.',1)[0]), filename, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', key=key, help=re.sub(r'^:material/\w+:\s*','',label))


@st.cache_data(show_spinner=False)
def parse_pkb(b):
    raw=read_excel_auto(b,header=None); h=None
    for i in range(min(len(raw),30)):
        vals=[clean(x).upper().rstrip('.') for x in raw.iloc[i].tolist()]
        has_pkb=('PKB NO' in vals or 'PKB NO.' in vals)
        has_pol=('POLICE REG NO' in vals or 'POLICE REG NO.' in vals)
        has_type=('TYPE KENDARAAN' in vals or 'TIPE KENDARAAN' in vals)
        if has_pkb and has_pol and has_type: h=i; break
    if h is None: raise ValueError('Header MaintainPKB tidak ditemukan')
    df=raw.iloc[h+1:].copy(); df.columns=[clean(x) or f'COL_{i}' for i,x in enumerate(raw.iloc[h])]
    rename={}
    for c in df.columns:
        u=clean(c).upper(); z=u.rstrip('.')
        if z=='PKB NO':rename[c]='PKB'
        elif z in ('POLICE REG NO','POLICE REG NO.') :rename[c]='No Polisi'
        elif u=='PKB DATE':rename[c]='Tanggal'
        elif u in ('TYPE KENDARAAN','TIPE KENDARAAN'):rename[c]='Type Kendaraan'
        elif u=='TYPE':rename[c]='Booking Type'
        elif u=='JENIS JASA PEKERJAAN':rename[c]='Jenis Jasa Pekerjaan'
        elif u=='JOB TYPE OPERATION':rename[c]='Job Type Operation'
        elif u=='SERVICE CATEGORY':rename[c]='Service Category'
        elif u=='OPERATION DESCRIPTION':rename[c]='Operation Description'
        elif u=='PKB STATUS':rename[c]='PKB Status'
        elif u=='CUSTOMER NAME':rename[c]='Customer'
        elif u=='SERVICE ADVISOR':rename[c]='Service Advisor'
        elif u=='MECHANIC':rename[c]='Mechanic'
        elif u=='EQUIPMENT NO.':rename[c]='Equipment No'
        elif u=='TAHUN PRODUKSI':rename[c]='Tahun Produksi'
    df=df.rename(columns=rename)
    for c in ['PKB','No Polisi','Tanggal','Type Kendaraan','Booking Type','Jenis Jasa Pekerjaan','Job Type Operation','Service Category','Operation Description','PKB Status','Customer','Service Advisor','Mechanic','Equipment No','Tahun Produksi']:
        if c not in df: df[c]=''
    df['PKB']=df['PKB'].map(norm); df['Tanggal']=pd.to_datetime(df['Tanggal'],errors='coerce',dayfirst=True)
    for c in ['No Polisi','Type Kendaraan','Booking Type','Jenis Jasa Pekerjaan','Job Type Operation','Service Category','Operation Description','PKB Status','Customer','Service Advisor','Mechanic','Equipment No'] : df[c]=df[c].map(clean)
    df['Service Advisor']=df['Service Advisor'].map(norm_sa)
    df=df[df.PKB!=''].copy(); df['_fill']=df.notna().sum(axis=1); df=df.sort_values(['PKB','_fill'],ascending=[True,False]).drop_duplicates('PKB').drop(columns='_fill')
    def job(r):
        for c in ['Jenis Jasa Pekerjaan','Job Type Operation','Service Category','Operation Description']:
            if clean(r.get(c,'')): return clean(r[c])
        return 'Tidak diketahui'
    df['Jenis Pekerjaan']=df.apply(job,axis=1); df['UE']=1
    return df.reset_index(drop=True)

@st.cache_data(show_spinner=False)
def parse_service(b):
    raw=read_excel_auto(b,header=None); h=None
    for i in range(min(len(raw),100)):
        v=[clean(x).upper() for x in raw.iloc[i].tolist()]
        if 'NO.' in v and 'TGL.' in v and 'NO. PKB' in v: h=i; break
    if h is None: raise ValueError('Header transaksi service tidak ditemukan')
    rec=[]
    for i in range(h+1,len(raw)):
        r=raw.iloc[i].tolist()
        if len(r)<44 or not re.fullmatch(r'\d+(\.0+)?',clean(r[0])): continue
        pkb=norm(r[3]);
        if not pkb: continue
        rec.append({'Tanggal Sales':pd.to_datetime(r[1],errors='coerce'),'Invoice':clean(r[2]),'PKB':pkb,'Customer Sales':clean(r[4]),'Jenis Pembayaran':clean(r[5]),'Type Kendaraan Sales':clean(r[6]),'Jasa':num(r[17]),'Sparepart':num(r[31]),'Bahan':num(r[43])})
    s=pd.DataFrame(rec)
    if s.empty: raise ValueError('Tidak ada transaksi service')
    a=s.groupby('PKB',as_index=False).agg(**{'Tanggal Sales':('Tanggal Sales','min'),'Invoice':('Invoice',lambda x:', '.join(pd.Series(list(x),dtype='string').dropna().unique().tolist())),'Customer Sales':('Customer Sales','first'),'Jenis Pembayaran':('Jenis Pembayaran','first'),'Type Kendaraan Sales':('Type Kendaraan Sales','first'),'Jasa':('Jasa','sum'),'Sparepart':('Sparepart','sum'),'Bahan':('Bahan','sum'),'Sales Row Count':('PKB','size')})
    a['Revenue']=a[['Jasa','Sparepart','Bahan']].sum(axis=1)
    return a,s

def _looks_like_docno(k):
    """Nomor dokumen/PKB: huruf+angka tanpa spasi/tanda baca (mis. 3203SPK26000123). Tanggal & angka murni diabaikan."""
    return bool(k) and len(k)>=6 and bool(re.fullmatch(r'[A-Z0-9]+',k)) and bool(re.search(r'[A-Z]',k)) and bool(re.search(r'\d',k))

def resolve_issue_branch(ref_keys, customer, pkb_map, fallback):
    """Cabang bahan Issue To PKB: (1) nomor PKB pada baris header dokumen -> cabang PKB; (2) akhiran '-K' pada nama customer -> Kendari;
    (3) tidak terpetakan -> cabang fallback, ditandai agar muncul di Quality."""
    for k in str(ref_keys or '').split('|'):
        k=norm(k)
        if not k: continue
        # Nomor referensi Issue To PKB berawalan SQA (mis. 3203SQA26007752); nomor PKB di MaintainPKB berawalan SPA
        # dengan nomor urut yang sama (3203SPA26007752). Coba nomor apa adanya, lalu versi SPA.
        for cand in (k, re.sub(r'^(\d{4}S)QA',r'\1PA',k)):
            if cand in pkb_map and clean(pkb_map[cand]): return pkb_map[cand],'PKB'
    if branch_from_customer(customer)=='Kharisma Kendari': return 'Kharisma Kendari','Customer (akhiran -K)'
    return fallback,'Tidak terpetakan'

def resolve_issue_sa(ref_keys, ref, sa_map):
    """Kode SA dari nomor PKB pada header dokumen Issue To PKB (nomor SQA -> SPA). Hanya pencocokan nomor penuh, tanpa tebakan."""
    for k in [norm(x) for x in str(ref_keys or '').split('|')]+[norm(ref)]:
        if not k: continue
        for cand in (k, re.sub(r'^(\d{4}S)QA',r'\1PA',k)):
            v=sa_map.get(cand)
            if v: return v
    return ''

def _cust_key(x):
    """Kunci nama customer untuk pencocokan antar laporan (huruf besar, tanpa tanda baca/spasi, tanpa akhiran -K cabang)."""
    t=_CUST_K.sub('',clean(x)).upper()
    return re.sub(r'[^A-Z0-9]','',t)

def fill_issue_sa_by_customer(issue, detail, pkb, max_days=14):
    """Isi 'SA Code' Issue To PKB yang kosong lewat jalur: nama customer (Issue) -> Service Penjualan (customer, PKB, tanggal)
    -> MaintainPKB (PKB -> SA). Bila customer punya beberapa PKB, dipilih PKB dengan tanggal terdekat dari tanggal Issue
    (maks max_days hari); bila tidak ada yang cukup dekat, dipakai hanya jika semua PKB customer itu memiliki SA yang sama."""
    try:
        if not (isinstance(issue,pd.DataFrame) and not issue.empty and 'SA Code' in issue.columns and 'Customer' in issue.columns): return issue
        if not (isinstance(detail,pd.DataFrame) and not detail.empty and {'Customer Sales','PKB'}.issubset(detail.columns)): return issue
        if not (isinstance(pkb,pd.DataFrame) and not pkb.empty and {'PKB','Service Advisor'}.issubset(pkb.columns)): return issue
        sa_of={}
        for _pk,_sa in zip(pkb['PKB'],pkb['Service Advisor']):
            k=norm(_pk)
            if k and clean(_sa) and k not in sa_of: sa_of[k]=norm_sa(_sa)
        dcol='Tanggal Sales' if 'Tanggal Sales' in detail.columns else None
        by_c={}
        for i,(c,pk) in enumerate(zip(detail['Customer Sales'],detail['PKB'])):
            ck=_cust_key(c); k=norm(pk)
            if not ck or not k or k not in sa_of: continue
            dt=pd.to_datetime(detail[dcol].iloc[i],errors='coerce') if dcol else pd.NaT
            by_c.setdefault(ck,[]).append((k,dt))
        out=issue['SA Code'].tolist(); idate=pd.to_datetime(issue['Tanggal'],errors='coerce').tolist() if 'Tanggal' in issue.columns else [pd.NaT]*len(out)
        for i,(cur,cust) in enumerate(zip(out,issue['Customer'])):
            if clean(cur): continue
            cand=by_c.get(_cust_key(cust))
            if not cand: continue
            best=None
            if pd.notna(idate[i]):
                dd=[(abs((dt-idate[i]).days),k) for k,dt in cand if pd.notna(dt)]
                dd=[x for x in dd if x[0]<=max_days]
                if dd: best=sa_of[min(dd)[1]]
            if best is None:
                sas={sa_of[k] for k,_ in cand}
                if len(sas)==1: best=next(iter(sas))
            if best: out[i]=best
        issue=issue.copy(); issue['SA Code']=out
    except Exception:
        pass
    return issue

def _first_text(r,idx,exclude=''):
    """Nama customer pada laporan Crystal Reports bisa bergeser 1 kolom (sel gabungan / No. Referensi kosong).
    Ambil teks pertama yang tidak kosong pada kolom kandidat; '-' dan teks yang sama dengan No. Referensi diabaikan."""
    for j in idx:
        if j<len(r):
            v=clean(r[j])
            if v and v not in {'-','nan','NaN','None'} and v!=exclude: return v
    return ''

@st.cache_data(show_spinner=False)
def parse_parts(b):
    raw=read_excel_auto(b,header=None); rows=[]; bill=''; dt=None; ref=''; cust=''
    # Header/detail layout: metadata rows have a date in col D and billing in col B; detail rows have numeric NO. in col B.
    for i in range(18,len(raw)):
        r=raw.iloc[i].tolist()
        if len(r)<22: continue
        parsed_dt=pd.to_datetime(r[3],errors='coerce',dayfirst=True)
        if pd.notna(parsed_dt) and not re.fullmatch(r'\d+(\.0+)?',clean(r[1])):
            bill=clean(r[1]); dt=parsed_dt; ref=clean(r[7]); cust=_first_text(r,[8,9,10,11],exclude=ref)
        if re.fullmatch(r'\d+(\.0+)?',clean(r[1])) and clean(r[2]):
            rows.append({'No Billing':bill,'Tanggal':dt,'No Referensi':ref,'Customer':cust,'No Part':clean(r[2]),'Deskripsi':clean(r[3]),'Qty':num(r[4]),'Cost':num(r[6]),'Total Cost':num(r[8]),'Harga':num(r[10]),'Discount':num(r[12]),'Total Harga':num(r[14]),'Total Discount':num(r[16]),'Net Sales':num(r[17]),'PPN':num(r[18]),'Sales Amount':num(r[20]),'Gross Profit':num(r[21])})
    return pd.DataFrame(rows)

@st.cache_data(show_spinner=False)
def parse_material(b):
    raw=read_excel_auto(b,header=None); rows=[]; doc=''; ref=''; dt=None; cust=''; refkeys=''
    for i in range(18,len(raw)):
        r=raw.iloc[i].tolist()
        if len(r)<25: continue
        if pd.notna(r[1]) and pd.notna(r[5]) and not re.fullmatch(r'\d+(\.0+)?',clean(r[1])):
            dt=pd.to_datetime(r[1],errors='coerce'); doc=clean(r[5]); ref=clean(r[8]); cust=_first_text(r,[17,18,19,20])
            refkeys='|'.join(dict.fromkeys(k for k in (norm(v) for v in r[2:17]) if _looks_like_docno(k)))
        if re.fullmatch(r'\d+(\.0+)?',clean(r[1])) and clean(r[3]):
            rows.append({'Tanggal':dt,'No Dokumen':doc,'No Referensi':ref,'Ref Keys':refkeys,'Customer':cust,'Kode Bahan':clean(r[3]),'Bahan':clean(r[6]),'Qty':num(r[9]),'Satuan':clean(r[11]),'Harga/Sat':num(r[18]),'Harga Gross':num(r[19]),'Discount':num(r[21]),'Harga Net':num(r[23]),'Cost':num(r[24])})
    return pd.DataFrame(rows)

@st.cache_data(show_spinner=False)
def parse_opl(b):
    raw=read_excel_auto(b,header=None); rows=[]
    for r in raw.iloc[10:].itertuples(index=False,name=None):
        if len(r)>=15 and pd.notna(r[0]) and clean(r[1]):
            rows.append({'OPL Date':pd.to_datetime(r[0],errors='coerce'),'OPL No':clean(r[1]),'Billing No':clean(r[2]),'Bill Date':pd.to_datetime(r[3],errors='coerce'),'OPL Type':clean(r[4]),'Operation Description':clean(r[5]),'Cost OPL':num(r[6]),'Sales OPL':num(r[8]),'Vendor':clean(r[10]),'PKB':norm(r[11]),'No Polisi':clean(r[12]),'Type Kendaraan':clean(r[13]),'Service Advisor':norm_sa(r[14])})
    d=pd.DataFrame(rows)
    if d.empty: return d
    d['Gross Profit OPL']=d['Sales OPL']-d['Cost OPL']; d['Margin OPL %']=np.where(d['Sales OPL']!=0,d['Gross Profit OPL']/d['Sales OPL'],0); return d

@st.cache_data(show_spinner=False)
def parse_customer(b):
    raw=read_excel_auto(b,header=None); rows=[]
    # Actual report header is row 14 (0-based); data starts row 18 in this report, but scan for numeric No. rows.
    for r in raw.iloc[15:].itertuples(index=False,name=None):
        if len(r)>=24 and re.fullmatch(r'\d+(\.0+)?',clean(r[0])):
            rows.append({'No':int(float(clean(r[0]))),'Customer':clean(r[1]),'Alamat':clean(r[2]),'Service Advisor':norm_sa(r[3]),'Equipment No':norm(r[4]),'No Polisi':clean(r[5]),'Type Kendaraan':clean(r[14]),'Tgl Service':pd.to_datetime(r[15],errors='coerce'),'Tgl Selesai':pd.to_datetime(r[16],errors='coerce'),'Tahun Kendaraan':num(r[17]),'KM':num(r[18]),'Tgl Delivery':pd.to_datetime(r[20],errors='coerce'),'Tanggal Kendaraan':clean(r[21]),'Job Type':clean(r[22])})
    return pd.DataFrame(rows)

@st.cache_data(show_spinner=False)
def parse_estimasi(b):
    """Laporan Estimasi Service Advisor (Crystal Reports): estimasi vs actual biaya & waktu per PKB.
    Posisi kolom: biaya est/act=7/8, waktu est/act=11/12, rincian estimasi jasa/part/bahan=16/18/19,
    rincian actual jasa/part/bahan=21/22/24. Blok per Service Advisor ('nama / code')."""
    raw=read_excel_auto(b,header=None)
    if raw.shape[1]<25: raise ValueError('Format Laporan Estimasi Service Advisor tidak dikenali')
    def _dt(x):
        if pd.isna(x) or clean(x)=='': return pd.NaT
        if isinstance(x,(pd.Timestamp,datetime.datetime)): return pd.Timestamp(x)
        s=clean(x); d=pd.to_datetime(s,format='%d.%m.%Y %H:%M',errors='coerce')
        return d if pd.notna(d) else pd.to_datetime(s,errors='coerce',dayfirst=True)
    # Periode laporan dibaca dari HEADER (mis. '01.09.2026 s/d 30.09.2026'), bukan dari tanggal PKB.
    hdr=[pd.to_datetime(d,format='%d.%m.%Y',errors='coerce') for i in range(min(len(raw),20)) for v in raw.iloc[i].tolist() for d in re.findall(r'\d{2}\.\d{2}\.\d{4}',str(v))]
    hdr=[d for d in hdr if pd.notna(d)]
    rec=[]; sa_name=''; sa_code=''; found_header=False
    for i in range(len(raw)):
        r=raw.iloc[i].tolist()
        first=clean(r[1]).lower()
        if first.startswith('service advisor'):
            found_header=True
            val=next((clean(x) for x in r[2:] if clean(x)),'')
            nm,_,cd=val.partition(' / ')
            sa_name=clean(nm); sa_code=norm_sa(cd) or norm_sa(nm)
            continue
        if not re.fullmatch(r'\d+(\.0+)?',clean(r[1])): continue
        pkb=norm(r[2])
        if not pkb or not sa_code: continue
        e_j,e_p,e_b=num(r[16]),num(r[18]),num(r[19]); a_j,a_p,a_b=num(r[21]),num(r[22]),num(r[24])
        est_biaya=num(r[7]) or (e_j+e_p+e_b); act_biaya=num(r[8]) or (a_j+a_p+a_b)
        rec.append({'PKB':pkb,'Tanggal':pd.to_datetime(r[4],errors='coerce'),'Chassis':clean(r[6]),
            'Service Advisor':sa_code,'SA Nama':sa_name,
            'Est Biaya':est_biaya,'Act Biaya':act_biaya,
            'Est Jasa':e_j,'Est Part':e_p,'Est Bahan':e_b,'Act Jasa':a_j,'Act Part':a_p,'Act Bahan':a_b,
            'Est Waktu':_dt(r[11]),'Act Waktu':_dt(r[12])})
    if not found_header: raise ValueError('Blok Service Advisor tidak ditemukan pada laporan estimasi')
    df=pd.DataFrame(rec)
    if df.empty: return df
    p_awal=hdr[0] if len(hdr)>=2 else df['Act Waktu'].min().normalize()
    p_akhir=hdr[1] if len(hdr)>=2 else df['Act Waktu'].max().normalize()
    df['Periode Awal']=p_awal; df['Periode Akhir']=p_akhir
    # Tanggal acuan revenue = tanggal selesai (actual), dijaga tetap di dalam rentang header laporan.
    df['Tanggal Periode']=df['Act Waktu'].dt.normalize().fillna(p_akhir).clip(lower=p_awal,upper=p_akhir)
    df['Selisih Biaya']=df['Act Biaya']-df['Est Biaya']
    df['Selisih Biaya %']=np.where(df['Est Biaya']>0,df['Selisih Biaya']/df['Est Biaya']*100,0.0)
    df['Status Biaya']=np.select([df['Selisih Biaya']>0,df['Selisih Biaya']<0],['Lebih Tinggi','Lebih Rendah'],'Sesuai')
    df['Selisih Jam']=(df['Act Waktu']-df['Est Waktu']).dt.total_seconds()/3600
    df['Status Waktu']=np.select([df['Selisih Jam']>0,df['Selisih Jam']==0,df['Selisih Jam']<0],['Terlambat','Tepat Waktu','Lebih Cepat'],'Tidak diketahui')
    return df.drop_duplicates(['Service Advisor','PKB']).reset_index(drop=True)

def build_master(pkb,sales,opl,customers,branch=None):
    m=pkb.copy()
    if 'Cabang' not in m.columns:
        m['Cabang']=branch or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')
    if 'Cabang' in sales.columns:
        m=m.merge(sales,on=['Cabang','PKB'],how='left')
    else:
        m=m.merge(sales,on='PKB',how='left')
    for c in ['Jasa','Sparepart','Bahan','Revenue']: m[c]=m[c].fillna(0.0)
    m['Sales Row Count']=m['Sales Row Count'].fillna(0).astype(int); m['Match Status']=np.where(m['Sales Row Count']>0,'MATCH','PKB TANPA SALES'); m['Revenue per Unit']=np.where(m.UE>0,m.Revenue/m.UE,0)
    if not customers.empty:
        c=customers.copy()
        # 1) Prefer latest customer visit per Equipment No, then use plate/customer+date as fallbacks.
        c['Customer Type']=c['Type Kendaraan'].map(clean)
        c['No Polisi']=c['No Polisi'].map(clean)
        c['Customer']=c['Customer'].map(clean)
        by_eq=c[c['Equipment No']!=''].sort_values('Tgl Service').drop_duplicates('Equipment No',keep='last')
        enrich=by_eq[['Equipment No','KM','Job Type','Tgl Service','Customer Type','Customer','Service Advisor']].rename(columns={'Job Type':'Customer Job Type','Tgl Service':'Customer Tgl Service','Service Advisor':'Customer SA','Customer':'Customer Report'})
        if 'Cabang' in c.columns:
            by_eq=c[c['Equipment No']!=''].sort_values(['Cabang','Tgl Service']).drop_duplicates(['Cabang','Equipment No'],keep='last')
            enrich=by_eq[['Cabang','Equipment No','KM','Job Type','Tgl Service','Customer Type','Customer','Service Advisor']].rename(columns={'Job Type':'Customer Job Type','Tgl Service':'Customer Tgl Service','Service Advisor':'Customer SA','Customer':'Customer Report'})
            m=m.merge(enrich,on=['Cabang','Equipment No'],how='left')
        else:
            m=m.merge(enrich,on='Equipment No',how='left')
        # Plate fallback where equipment is absent/unmatched.
        by_plate=c[c['No Polisi']!=''].sort_values('Tgl Service').drop_duplicates('No Polisi',keep='last')
        bp=by_plate[['No Polisi','KM','Job Type','Tgl Service','Customer Type','Customer']].rename(columns={'KM':'KM_plate','Job Type':'Customer Job Type_plate','Tgl Service':'Customer Tgl Service_plate','Customer Type':'Customer Type_plate','Customer':'Customer_plate'})
        if 'Cabang' in c.columns:
            by_plate=c[c['No Polisi']!=''].sort_values(['Cabang','Tgl Service']).drop_duplicates(['Cabang','No Polisi'],keep='last')
            bp=by_plate[['Cabang','No Polisi','KM','Job Type','Tgl Service','Customer Type','Customer']].rename(columns={'KM':'KM_plate','Job Type':'Customer Job Type_plate','Tgl Service':'Customer Tgl Service_plate','Customer Type':'Customer Type_plate','Customer':'Customer_plate'})
            m=m.merge(bp,on=['Cabang','No Polisi'],how='left')
        else:
            m=m.merge(bp,on='No Polisi',how='left')
        m['Customer KM']=m['KM'].where(m['KM'].fillna(0)>0,m['KM_plate'])
        m['Customer Job Type']=m['Customer Job Type'].where(m['Customer Job Type'].map(clean)!='',m['Customer Job Type_plate'])
        m['Customer Type Kendaraan']=m['Customer Type'].where(m['Customer Type'].map(clean)!='',m['Customer Type_plate'])
        m['Customer Name Report']=m['Customer'].where(m['Customer'].map(clean)!='',m['Customer_plate'])
        for col in ['KM_plate','Customer Job Type_plate','Customer Tgl Service_plate','Customer Type_plate','Customer_plate','KM']:
            if col in m: m.drop(columns=col,inplace=True)
        m['Type Kendaraan']=m['Type Kendaraan'].where(m['Type Kendaraan'].map(clean)!='',m['Customer Type Kendaraan'])
    if not opl.empty:
        if 'Cabang' in opl.columns:
            oa=opl.groupby(['Cabang','PKB'],as_index=False).agg(**{'OPL Sales':('Sales OPL','sum'),'OPL Cost':('Cost OPL','sum'),'OPL GP':('Gross Profit OPL','sum'),'OPL Count':('OPL No','nunique')})
            m=m.merge(oa,on=['Cabang','PKB'],how='left')
        else:
            oa=opl.groupby('PKB',as_index=False).agg(**{'OPL Sales':('Sales OPL','sum'),'OPL Cost':('Cost OPL','sum'),'OPL GP':('Gross Profit OPL','sum'),'OPL Count':('OPL No','nunique')})
            m=m.merge(oa,on='PKB',how='left')
        for c in ['OPL Sales','OPL Cost','OPL GP','OPL Count']:m[c]=m[c].fillna(0)
    return m

LONG_REPORT_DAYS=27   # laporan Estimasi SA dengan periode >= 27 hari dianggap laporan bulanan
EST_REF_ACTIVE=False
EST_REF_BRANCHES=set()   # cabang yang punya acuan Estimasi SA (diisi saat data dimuat)
def apply_estimasi_reference(m, est, active=True, est_cov=None):
    """Laporan Estimasi Service Advisor sebagai ACUAN revenue bulanan.
    - Bulan = periode HEADER laporan (tanggal selesai), bukan tanggal PKB.
    - Jasa/Sparepart/Bahan PKB = nilai Actual laporan.
    - PKB di laporan tapi belum ada di MaintainPKB ditambahkan.
    - PKB yang jatuh di bulan yang sudah punya laporan tetapi tidak tercantum di laporan tidak dihitung pada bulan itu.
    Bulan tanpa laporan tetap memakai cara lama (tanggal PKB)."""
    m=m.copy(); m['Tanggal PKB']=pd.to_datetime(m['Tanggal'],errors='coerce'); m['Sumber Revenue']='Service Penjualan (tgl PKB)'
    stats={'matched':0,'moved':0,'excluded':0,'added':0}
    excluded=m.iloc[0:0].copy()
    _cov_src=est_cov if (est_cov is not None and not est_cov.empty) else est
    if not active or _cov_src is None or _cov_src.empty: return m,excluded,stats
    if est is None or est.empty: est=_cov_src.iloc[0:0].copy()
    e=est.sort_values('Tanggal Periode',kind='stable').drop_duplicates('PKB',keep='last').set_index('PKB')
    # Bulan yang 'terliput' laporan dihitung PER CABANG. Laporan gabungan Makassar/Kendari/Gowa meliput ketiganya; cabang sistem
    # terpisah (mis. Ambon) hanya terliput bila punya laporan Estimasi SA sendiri. Tanpa ini, PKB cabang yang belum punya
    # laporan Estimasi ikut terbuang pada tampilan Semua Cabang sehingga total != jumlah revenue per cabang.
    # REVENUE & UE MURNI DARI ESTIMASI SA: sejak tanggal laporan Estimasi pertama sebuah cabang, PKB yang tidak tercantum di
    # laporan Estimasi TIDAK dihitung (termasuk pada hari yang laporannya belum diupload; tetap tampil di Quality untuk evaluasi).
    # Riwayat SEBELUM laporan Estimasi pertama tetap memakai Service Penjualan. Laporan panjang (>= LONG_REPORT_DAYS) dihitung
    # mulai awal BULAN-nya (perilaku lama); laporan pendek (harian) mulai tanggalnya.
    first_by_branch={}
    _ecols=['Cabang','Periode Awal','Periode Akhir'] if 'Cabang' in _cov_src.columns else None
    def _add_first(key,a_,b_):
        if pd.isna(a_) or pd.isna(b_): return
        st_=a_.to_period('M').to_timestamp() if (b_-a_).days>=LONG_REPORT_DAYS else a_.normalize()
        first_by_branch[key]=min(first_by_branch.get(key,st_),st_)
    if _ecols:
        for cab,a_,b_ in _cov_src[_ecols].drop_duplicates().itertuples(index=False): _add_first(str(cab),a_,b_)
    else:
        for a_,b_ in _cov_src[['Periode Awal','Periode Akhir']].drop_duplicates().itertuples(index=False): _add_first('*',a_,b_)
    def _first(cab):
        if not _ecols: return first_by_branch.get('*',pd.Timestamp.max)
        if cab in SHARED_SYSTEM_BRANCHES:
            v=[first_by_branch[c] for c in SHARED_SYSTEM_BRANCHES if c in first_by_branch]; return min(v) if v else pd.Timestamp.max
        return first_by_branch.get(str(cab),pd.Timestamp.max)
    hit=m['PKB'].isin(e.index)
    if hit.any():
        for src,dst in [('Act Jasa','Jasa'),('Act Part','Sparepart'),('Act Bahan','Bahan')]:
            m.loc[hit,dst]=m.loc[hit,'PKB'].map(e[src]).astype(float)
        new_date=m.loc[hit,'PKB'].map(e['Tanggal Periode'])
        stats['matched']=int(hit.sum()); stats['moved']=int((new_date.dt.to_period('M')!=m.loc[hit,'Tanggal PKB'].dt.to_period('M')).sum())
        m.loc[hit,'Tanggal']=new_date; m.loc[hit,'Match Status']='MATCH'; m.loc[hit,'Sumber Revenue']='Laporan Estimasi SA (periode laporan)'
    # PKB di laporan estimasi yang belum ada di MaintainPKB: tetap dihitung agar revenue = laporan
    miss=e.loc[~e.index.isin(m['PKB'])].reset_index()
    if not miss.empty:
        add=pd.DataFrame({'PKB':miss['PKB'],'Cabang':miss['Cabang'],'Service Advisor':miss['Service Advisor'],'Tanggal':miss['Tanggal Periode'],'Tanggal PKB':miss['Tanggal'],
            'Jasa':miss['Act Jasa'],'Sparepart':miss['Act Part'],'Bahan':miss['Act Bahan'],'UE':1,'Match Status':'MATCH','Sumber Revenue':'Laporan Estimasi SA (periode laporan)','Sales Row Count':0})
        m=pd.concat([m,add],ignore_index=True); stats['added']=len(add)
        for c in m.columns:
            if m[c].dtype==object: m[c]=m[c].where(m[c].notna(),'')
    m['Revenue']=m[['Jasa','Sparepart','Bahan']].sum(axis=1)
    m['Revenue per Unit']=np.where(m.UE>0,m.Revenue/m.UE,0)
    # PKB di bulan yang sudah punya laporan tapi tidak tercantum -> belum/ tidak menjadi pencapaian bulan tsb
    own=m['Sumber Revenue'].eq('Service Penjualan (tgl PKB)')
    _pdt=m['Tanggal PKB'].dt.normalize()
    _cab=m['Cabang'].astype(str) if 'Cabang' in m.columns else pd.Series('',index=m.index)
    gone=own & pd.Series([(pd.notna(d) and d>=_first(c)) for d,c in zip(_pdt,_cab)],index=m.index)
    excluded=m[gone].copy(); stats['excluded']=int(gone.sum())
    return m[~gone].reset_index(drop=True),excluded,stats

# =========================
# LOGIN + USER MANAGEMENT + LICENSE ACCESS
# =========================
AUTH_STATE_PATH=os.path.join(DATA_DIR,'user_credentials.json')
LOGIN_SESSION_PATH=os.path.join(DATA_DIR,'login_session.json')
LOGIN_SESSION_HOURS=48
def _default_auth():
    return {'username':'admin','password_hash':hashlib.sha256('admin123'.encode()).hexdigest()}
def load_auth():
    try:
        with open(AUTH_STATE_PATH,'r',encoding='utf-8') as fh: return json.load(fh)
    except Exception:
        stt=_default_auth(); save_auth(stt); return stt
def save_auth(stt):
    tmp=AUTH_STATE_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(stt,fh,indent=2)
    os.replace(tmp,AUTH_STATE_PATH)
def check_login(username,password):
    stt=load_auth(); return username.strip()==stt.get('username','admin') and hashlib.sha256(password.encode()).hexdigest()==stt.get('password_hash','')

def save_login_session(username):
    # Persistent local login session. This survives Streamlit reruns/browser refreshes
    # and is automatically invalid after 48 hours. It is also tied to this computer.
    now=datetime.datetime.now().timestamp()
    session={
        'username': username.strip(),
        'created_at': now,
        'expires_at': now + (LOGIN_SESSION_HOURS * 3600),
        'fingerprint': _machine_fingerprint(),
        'token': secrets.token_urlsafe(32) if 'secrets' in globals() else base64.urlsafe_b64encode(os.urandom(24)).decode('ascii')
    }
    tmp=LOGIN_SESSION_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(session,fh,indent=2)
    os.replace(tmp,LOGIN_SESSION_PATH)

def clear_login_session():
    try:
        if os.path.exists(LOGIN_SESSION_PATH): os.remove(LOGIN_SESSION_PATH)
    except Exception:
        pass

def load_login_session():
    try:
        with open(LOGIN_SESSION_PATH,'r',encoding='utf-8') as fh: session=json.load(fh)
        now=datetime.datetime.now().timestamp()
        if session.get('fingerprint') != _machine_fingerprint():
            clear_login_session(); return None
        if not session.get('username') or float(session.get('expires_at',0)) <= now:
            clear_login_session(); return None
        # Pastikan akun sesi masih ada & aktif (admin bawaan atau pengguna di daftar User Management).
        if not _login_exists(session.get('username')):
            clear_login_session(); return None
        return session
    except Exception:
        clear_login_session(); return None

# =========================
# RBAC: PENGGUNA, PERAN & HAK AKSES (Role-Based Access Control + izin khusus per pengguna)
# =========================
USERS_PATH=os.path.join(DATA_DIR,'users.json')
ROLE_ORDER=['super_admin','admin_cabang','sa','foreman','staf']
ROLE_LABELS={'super_admin':'Super Admin','admin_cabang':'Administrator Cabang','sa':'Service Advisor','foreman':'Foreman QC','staf':'Staf (Kasir/Admin/Gudang)'}
ROLE_STYLE={'super_admin':('#6d28d9','#f3e8ff','#e9d5ff'),'admin_cabang':('#1d4ed8','#e0ecff','#bfdbfe'),'sa':('#047857','#dcfce7','#bbf7d0'),
            'foreman':('#b45309','#fef3c7','#fde68a'),'staf':('#475569','#f1f5f9','#e2e8f0')}
# (kunci, nama, kelompok, deskripsi) -- hanya izin yang benar-benar dipakai modul aplikasi ini
PERMISSIONS=[
 ('dashboard_view','Melihat Dashboard Utama','Dashboard & Analitik','Melihat kartu KPI, grafik performa, dan pencapaian per Service Advisor.'),
 ('analytics_view','Melihat Analitik Lanjutan','Dashboard & Analitik','Comparative Analytics, Management Insight, UE, serta Revenue & Profit.'),
 ('all_branches','Akses Semua Cabang','Dashboard & Analitik','Dapat berpindah dan melihat performa seluruh cabang bengkel.'),
 ('ops_view','Melihat Laporan Operasional','Operasional','Sparepart, Bahan & Consumable, OPL, Customer, dan Quality Control.'),
 ('target_view','Melihat Target Perusahaan','Target Perusahaan','Melihat target bulanan & harian serta persentase pencapaian.'),
 ('target_edit','Input & Edit Target Perusahaan','Target Perusahaan','Mengubah nominal target UE, jasa, part, bahan, dan revenue.'),
 ('upload_daily','Upload Laporan Harian (Estimasi SA)','Input Laporan','Mengupload laporan Estimasi Service Advisor harian.'),
 ('upload_monthly','Upload Laporan Bulanan','Input Laporan','MaintainPKB, Service Penjualan, Billing Part/Service, Consumable, Issue To PKB, OPL, Pelanggan.'),
 ('data_view','Melihat Data & Export','Laporan & Ekspor','Membuka halaman Data (filter dan tabel detail).'),
 ('export_excel','Ekspor Laporan Excel','Laporan & Ekspor','Mengunduh tabel dan laporan dalam format Excel/CSV.'),
 ('branch_manage','Mengelola Cabang','Administrasi','Menambah cabang dan memetakan Service Advisor ke cabang.'),
 ('user_manage','Mengelola Pengguna','Administrasi','Menambah/mengubah pengguna dan hak akses (sesuai cakupan cabang).'),
]
PERM_KEYS=[p[0] for p in PERMISSIONS]
ROLE_DEFAULTS={
 'super_admin':list(PERM_KEYS),
 'admin_cabang':['dashboard_view','analytics_view','ops_view','target_view','target_edit','upload_daily','upload_monthly','data_view','export_excel','user_manage'],
 'sa':['dashboard_view','ops_view','target_view','upload_daily'],
 'foreman':['dashboard_view','ops_view'],
 'staf':['ops_view','upload_monthly','export_excel'],
}
ROLE_DESC={
 'super_admin':('Semua Cabang','Akses penuh seluruh cabang, pengaturan master target perusahaan, dan manajemen seluruh pengguna sistem.'),
 'admin_cabang':('Khusus cabangnya','Mengelola operasional cabang, menginput target cabang, upload laporan harian SA & bulanan PKB, ekspor laporan Excel, dan manajemen staf cabang.'),
 'sa':('Khusus cabangnya','Upload laporan Estimasi SA harian, melihat kartu target & realisasi cabang, serta memonitor unit servis.'),
 'foreman':('Khusus cabangnya','Memantau progres unit dan laporan operasional serta Quality Control sebelum unit diserahkan.'),
 'staf':('Khusus cabangnya','Entri administrasi berkas PKB (upload laporan bulanan), memantau laporan sparepart & bahan, serta ekspor laporan ke Excel.'),
}
PAGE_PERM={
 '📊 Dashboard':['dashboard_view'],'🎯 Target Perusahaan':['target_view'],
 '📊 Comparative Analytics':['analytics_view'],'🧠 Management Insight':['analytics_view'],'🚗 UE':['analytics_view'],'💰 Revenue & Profit':['analytics_view'],
 '🔩 Sparepart':['ops_view'],'🧪 Bahan & Consumable':['ops_view'],'🛠️ OPL':['ops_view'],'👥 Customer':['ops_view'],'✅ Quality':['ops_view'],
 '📥 Data Input':['upload_daily','upload_monthly'],'📋 Data':['data_view'],
 '🏢 Branch Management':['branch_manage'],'👤 User Management':['user_manage'],
}

def _hash_pw(pw):
    salt=os.urandom(16); h=hashlib.pbkdf2_hmac('sha256',pw.encode('utf-8'),salt,200000)
    return f'pbkdf2${salt.hex()}${h.hex()}'
def _verify_pw(pw,stored):
    stored=stored or ''
    if stored.startswith('pbkdf2$'):
        try:
            _,s_,h_=stored.split('$')
            return hmac.compare_digest(hashlib.pbkdf2_hmac('sha256',pw.encode('utf-8'),bytes.fromhex(s_),200000).hex(),h_)
        except Exception: return False
    return hmac.compare_digest(hashlib.sha256(pw.encode('utf-8')).hexdigest(),stored)

def _builtin_user(auth):
    return {'id':'builtin-admin','builtin':True,'name':'Administrator Pusat','email':auth.get('username','admin'),'role':'super_admin',
            'branch':'Semua Cabang','active':True,'permissions':list(PERM_KEYS),'last_login':None}
def load_users():
    """Daftar pengguna. Akun admin bawaan (login lama) selalu ada sebagai Super Admin dan tidak bisa dinonaktifkan/dihapus."""
    auth=load_auth()
    try:
        with open(USERS_PATH,'r',encoding='utf-8') as fh: rows=json.load(fh)
        if not isinstance(rows,list): rows=[]
    except Exception: rows=[]
    b=next((r for r in rows if r.get('builtin')),None)
    if b is None: b=_builtin_user(auth); rows.insert(0,b)
    b.update({'email':auth.get('username','admin'),'role':'super_admin','branch':'Semua Cabang','active':True,'permissions':list(PERM_KEYS),'builtin':True})
    return rows
def save_users(rows):
    tmp=USERS_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(rows,fh,ensure_ascii=False,indent=2)
    os.replace(tmp,USERS_PATH)
def find_user(login):
    login=(login or '').strip()
    if not login: return None
    rows=load_users(); auth=load_auth()
    if login==auth.get('username','admin'): return next((r for r in rows if r.get('builtin')),None)
    return next((r for r in rows if not r.get('builtin') and str(r.get('email','')).strip().lower()==login.lower()),None)
def _touch_login(uid):
    rows=load_users()
    for r in rows:
        if r.get('id')==uid: r['last_login']=datetime.datetime.now().isoformat(timespec='seconds')
    try: save_users(rows)
    except Exception: pass
def authenticate_user(login,pw):
    """(user, pesan_error). Admin bawaan memakai kredensial lama; pengguna lain memakai email + password."""
    login=(login or '').strip(); auth=load_auth()
    if login==auth.get('username','admin'):
        if check_login(login,pw):
            u=find_user(login); _touch_login(u['id']); return u,''
        return None,'Username/Email atau password salah.'
    u=find_user(login)
    if not u or not _verify_pw(pw,u.get('pw','')): return None,'Username/Email atau password salah.'
    if not u.get('active',True): return None,'Akun Anda nonaktif. Hubungi administrator.'
    _touch_login(u['id']); return u,''
def resolve_current_user():
    u=find_user(st.session_state.get('logged_in_username',''))
    return u if (u and u.get('active',True)) else None
def _login_exists(login):
    u=find_user(login); return bool(u and u.get('active',True))

def has_perm(key,user=None):
    u=user if user is not None else globals().get('CURRENT_USER')
    if not u: return False
    if u.get('builtin') or u.get('role')=='super_admin': return True
    return key in (u.get('permissions') or [])
def user_branches(user=None):
    """Cabang yang boleh diakses: semua cabang bila punya izin 'Akses Semua Cabang', selain itu cabang penugasannya."""
    u=user if user is not None else globals().get('CURRENT_USER')
    if not u: return []
    if has_perm('all_branches',u): return list(BRANCHES)
    b=u.get('branch')
    return [b] if b in BRANCHES else (list(BRANCHES) if b=='Semua Cabang' else [])
def can_manage_user(editor,target):
    if not editor or not has_perm('user_manage',editor): return False
    if editor.get('role')=='super_admin' or editor.get('builtin'): return True
    return (not target.get('builtin')) and target.get('role') in ('sa','foreman','staf') and target.get('branch')==editor.get('branch')
def assignable_roles(editor): return list(ROLE_ORDER) if (editor.get('role')=='super_admin' or editor.get('builtin')) else ['sa','foreman','staf']
def assignable_branches(editor): return (['Semua Cabang']+list(BRANCHES)) if (editor.get('role')=='super_admin' or editor.get('builtin')) else [editor.get('branch')]
def assignable_perms(editor): return list(PERM_KEYS) if (editor.get('role')=='super_admin' or editor.get('builtin')) else [k for k in PERM_KEYS if has_perm(k,editor)]
def allowed_pages(): return [p for p in page_labels if any(has_perm(k) for k in PAGE_PERM.get(p,[]))]

def dl_button(*a,**k):
    """st.download_button yang mengikuti izin 'Ekspor Laporan Excel'."""
    if has_perm('export_excel'): return st.download_button(*a,**k)
    st.caption(':material/lock: Ekspor tidak diizinkan untuk akun Anda.'); return False

def _fmt_activity(iso):
    if not iso: return 'Belum pernah login'
    try: dt=datetime.datetime.fromisoformat(iso)
    except Exception: return '-'
    d=(datetime.date.today()-dt.date()).days
    return f'Hari ini, {dt:%H:%M}' if d==0 else (f'Kemarin, {dt:%H:%M}' if d==1 else f'{d} hari lalu')

def _rchip(text,role=None,color='#334155',bg='#f1f5f9',border='#e2e8f0'):
    if role: color,bg,border=ROLE_STYLE[role]
    return f'<span style="display:inline-block;padding:3px 10px;border-radius:6px;font-size:11.5px;font-weight:700;color:{color};background:{bg};border:1px solid {border};white-space:nowrap">{_html.escape(str(text))}</span>'

def _clear_dlg():
    for k in [k for k in st.session_state.keys() if str(k).startswith('dlg_')]: st.session_state.pop(k,None)

def _dlg_apply_role(p):
    role=st.session_state.get(p+'role'); editor=CURRENT_USER; allow=set(assignable_perms(editor))
    for k in PERM_KEYS: st.session_state[p+'perm_'+k]=(k in ROLE_DEFAULTS.get(role,[]) and k in allow)
def _dlg_reset_perms(p): _dlg_apply_role(p)

def _user_form(uid):
    """Isi dialog Tambah / Edit Pengguna & Hak Akses (dipanggil dari st.dialog)."""
    editor=CURRENT_USER; users=load_users()
    target=next((u for u in users if u.get('id')==uid),None) if uid else None
    is_new=target is None; builtin=bool(target and target.get('builtin')); is_self=bool(target and editor and target.get('id')==editor.get('id'))
    if target is not None and not can_manage_user(editor,target) and not is_self:
        st.error('Anda tidak berwenang mengubah pengguna ini.'); return
    p=f'dlg_{uid or "new"}_'
    roles=assignable_roles(editor); branches=assignable_branches(editor); aperms=set(assignable_perms(editor))
    if p+'init' not in st.session_state:
        st.session_state[p+'name']=target.get('name','') if target else ''
        st.session_state[p+'email']=target.get('email','') if target else ''
        st.session_state[p+'role']=(target.get('role') if target else 'sa')
        st.session_state[p+'branch']=(target.get('branch') if target else (branches[1] if len(branches)>1 else branches[0]))
        st.session_state[p+'status']='Aktif' if (target.get('active',True) if target else True) else 'Nonaktif'
        base=(target.get('permissions') or []) if target else ROLE_DEFAULTS.get(st.session_state[p+'role'],[])
        for k in PERM_KEYS: st.session_state[p+'perm_'+k]=(k in base)
        st.session_state[p+'init']=True
    sub=('Isi data pengguna dan atur hak aksesnya.' if is_new else f'Mengubah konfigurasi untuk {_html.escape(str(target.get("name","")))}')
    st.markdown(f'<div style="color:#64748b;font-size:12.5px;margin:-8px 0 12px">{sub}</div>',unsafe_allow_html=True)
    locked=builtin or is_self
    c1,c2=st.columns(2)
    name=c1.text_input('Nama Lengkap *',key=p+'name')
    email=c2.text_input('Email Akun *',key=p+'email',disabled=builtin,help='Dipakai untuk login.' if not builtin else 'Akun admin bawaan: ubah username & password di bagian Profil.')
    c3,c4=st.columns(2)
    ropts=list(roles) if not builtin else ['super_admin']
    if st.session_state.get(p+'role') not in ropts: ropts=[st.session_state.get(p+'role')]+ropts
    role=c3.selectbox('Peran / Role *',ropts,key=p+'role',format_func=lambda r:ROLE_LABELS[r],disabled=locked,on_change=_dlg_apply_role,args=(p,))
    bopts=list(branches) if not builtin else ['Semua Cabang']
    if target and target.get('branch') not in bopts: bopts=[target.get('branch')]+bopts
    if st.session_state.get(p+'branch') not in bopts: st.session_state[p+'branch']=bopts[0]
    c4.selectbox('Penugasan Cabang *',bopts,key=p+'branch',disabled=locked or len(bopts)==1)
    st.radio('Status Akun',['Aktif','Nonaktif'],key=p+'status',horizontal=True,disabled=locked)
    pw=st.text_input('Password Baru' if not is_new else 'Password *',type='password',key=p+'pw',disabled=builtin,
                     placeholder='Minimal 6 karakter' if is_new else 'Kosongkan bila tidak diubah',help='Diperlukan agar pengguna dapat login.')
    npick=sum(1 for k in PERM_KEYS if st.session_state.get(p+'perm_'+k))
    h1,h2=st.columns([3,1])
    h1.markdown(f'<div style="font-weight:800;font-size:12px;letter-spacing:.6px;color:#047857;margin-top:6px"><span class="msr" style="font-size:16px;vertical-align:-3px">key</span> CAKUPAN HAK AKSES KHUSUS ({npick} DIPILIH)</div>',unsafe_allow_html=True)
    h2.button(':material/restart_alt: Reset ke Standar Role',key=p+'reset',on_click=_dlg_reset_perms,args=(p,),disabled=locked,use_container_width=True)
    with st.container(height=300,border=True):
        for k,label,grp,desc in PERMISSIONS:
            dis=locked or role=='super_admin' or (k not in aperms)
            if role=='super_admin': st.session_state[p+'perm_'+k]=True
            st.checkbox(f'{label}  ·  {grp}',key=p+'perm_'+k,disabled=dis,help=desc)
            st.caption(desc)
    if role!='super_admin' and not locked and aperms!=set(PERM_KEYS): st.caption('Anda hanya dapat memberikan izin yang Anda miliki sendiri.')
    b1,b2,b3=st.columns([2,1,1.4])
    if b2.button('Batal',key=p+'cancel',use_container_width=True): _clear_dlg(); st.rerun()
    if b3.button(':material/save: Simpan Perubahan',key=p+'save',type='primary',use_container_width=True):
        nm=(name or '').strip(); em=(email or '').strip()
        sel=[k for k in PERM_KEYS if st.session_state.get(p+'perm_'+k)]
        errs=[]
        if not nm: errs.append('Nama lengkap wajib diisi.')
        if not builtin:
            if not em or '@' not in em: errs.append('Email akun wajib diisi dengan format yang benar.')
            low=em.lower(); auth=load_auth()
            if low==str(auth.get('username','admin')).lower() or any(u.get('id')!=(target or {}).get('id') and str(u.get('email','')).strip().lower()==low for u in users): errs.append('Email tersebut sudah dipakai akun lain.')
            if is_new and len(pw or '')<6: errs.append('Password minimal 6 karakter.')
            if pw and len(pw)<6: errs.append('Password minimal 6 karakter.')
        if not locked and role not in roles: errs.append('Peran tidak diizinkan.')
        if errs:
            for e in errs: st.error(e)
        else:
            if role=='super_admin': sel=list(PERM_KEYS)
            sel=[k for k in sel if (k in aperms or builtin or is_self)]
            for u in users:
                if is_new: break
                if u.get('id')==target.get('id'):
                    u['name']=nm
                    if not builtin: u['email']=em
                    if not locked:
                        u['role']=role; u['branch']=st.session_state[p+'branch']; u['active']=(st.session_state[p+'status']=='Aktif'); u['permissions']=sel
                    if pw and not builtin: u['pw']=_hash_pw(pw)
            if is_new:
                users.append({'id':'u-'+uuid.uuid4().hex[:10],'name':nm,'email':em,'role':role,'branch':st.session_state[p+'branch'],'active':st.session_state[p+'status']=='Aktif',
                              'permissions':sel,'pw':_hash_pw(pw),'last_login':None,'created_at':datetime.datetime.now().isoformat(timespec='seconds')})
            save_users(users); _clear_dlg(); st.session_state['um_flash']=('Pengguna berhasil ditambahkan.' if is_new else 'Perubahan pengguna tersimpan.'); st.rerun()

@st.dialog('Tambah Pengguna Baru',width='large')
def _user_dialog_new(): _user_form(None)
@st.dialog('Edit Pengguna & Hak Akses',width='large')
def _user_dialog_edit(uid): _user_form(uid)
def _user_dialog(uid):
    if uid is None: _user_dialog_new()
    else: _user_dialog_edit(uid)

@st.dialog('Hapus Pengguna')
def _delete_dialog(uid):
    users=load_users(); t=next((u for u in users if u.get('id')==uid),None)
    if not t or t.get('builtin') or not can_manage_user(CURRENT_USER,t): st.error('Pengguna ini tidak dapat dihapus.'); return
    st.warning(f'Hapus akun **{t.get("name","")}** ({t.get("email","")})? Tindakan ini tidak dapat dibatalkan.')
    a,b=st.columns(2)
    if a.button('Batal',key='del_cancel',use_container_width=True): st.rerun()
    if b.button(':material/delete: Ya, Hapus',key='del_ok',type='primary',use_container_width=True):
        save_users([u for u in users if u.get('id')!=uid]); st.session_state['um_flash']='Pengguna dihapus.'; st.rerun()

def rbac_admin_area():
    editor=CURRENT_USER; users=load_users()
    visible=[u for u in users if can_manage_user(editor,u) or u.get('id')==editor.get('id')]
    if st.session_state.get('um_flash'): st.success(st.session_state.pop('um_flash'))
    hc1,hc2=st.columns([4,1.3])
    hc1.markdown('<div class="um-head"><span class="msr">groups</span><div><b>User & Access Management</b><span>Pengaturan hak akses staf, Service Advisor, foreman, dan administrator cabang bengkel Kharisma</span></div></div>',unsafe_allow_html=True)
    if hc2.button(':material/person_add: Tambah Pengguna',key='um_add',type='primary',use_container_width=True): _clear_dlg(); _user_dialog(None)
    tab_list,tab_matrix=st.tabs([f':material/groups: Daftar Pengguna ({len(visible)})',':material/key: Matriks Hak Akses Peran'])
    with tab_list:
        f1,f2,f3=st.columns([2.4,1,1])
        q=f1.text_input('Cari',placeholder='Cari berdasarkan nama atau email pengguna...',label_visibility='collapsed',key='um_q')
        rf=f2.selectbox('Peran',['Semua Peran']+ROLE_ORDER,format_func=lambda r:r if r=='Semua Peran' else ROLE_LABELS[r],key='um_role',label_visibility='collapsed')
        bf=f3.selectbox('Cabang',['Semua Cabang']+list(BRANCHES),key='um_branch',label_visibility='collapsed')
        rows=[u for u in visible if (not q or q.lower() in (str(u.get('name',''))+' '+str(u.get('email',''))).lower()) and (rf=='Semua Peran' or u.get('role')==rf) and (bf=='Semua Cabang' or u.get('branch')==bf)]
        W=[2.3,2.1,1.6,2.0,1.0,1.4,0.9]
        hd=st.columns(W)
        for c,t in zip(hd,['PENGGUNA','PERAN / ROLE','PENUGASAN CABANG','CAKUPAN HAK AKSES','STATUS','AKTIVITAS TERAKHIR','AKSI']): c.markdown(f'<div class="um-th">{t}</div>',unsafe_allow_html=True)
        if not rows: st.info('Tidak ada pengguna yang cocok dengan filter.')
        for u in rows:
            uid=u['id']; manage=can_manage_user(editor,u); self_row=(uid==editor.get('id'))
            c=st.columns(W,vertical_alignment='center')
            c[0].markdown(f'<div class="um-name">{_html.escape(str(u.get("name","")))}</div><div class="um-mail">{_html.escape(str(u.get("email","")))}</div>',unsafe_allow_html=True)
            c[1].markdown(_rchip(ROLE_LABELS.get(u.get('role'),u.get('role')),u.get('role')),unsafe_allow_html=True)
            c[2].markdown(f'<span class="um-branch"><span class="msr">apartment</span> {_html.escape(str(u.get("branch","")))}</span>',unsafe_allow_html=True)
            n=len(PERM_KEYS) if (u.get('role')=='super_admin' or u.get('builtin')) else len([k for k in (u.get('permissions') or []) if k in PERM_KEYS])
            if c[3].button(f':material/shield: {n} Hak Akses Aktif',key=f'um_perm_{uid}',disabled=not (manage or self_row),help='Lihat / ubah hak akses'): _clear_dlg(); _user_dialog(uid)
            act=bool(u.get('active',True)); can_toggle=manage and not u.get('builtin') and not self_row
            if c[4].button('Aktif' if act else 'Nonaktif',key=f'um_stat_{"on" if act else "off"}_{uid}',disabled=not can_toggle,help='Klik untuk mengubah status akun'):
                allu=load_users()
                for r in allu:
                    if r.get('id')==uid: r['active']=not act
                save_users(allu); st.rerun()
            c[5].markdown(f'<span class="um-act">{_fmt_activity(u.get("last_login"))}</span>',unsafe_allow_html=True)
            a1,a2=c[6].columns(2)
            if a1.button(':material/edit:',key=f'um_edit_{uid}',disabled=not (manage or self_row),help='Edit'): _clear_dlg(); _user_dialog(uid)
            if a2.button(':material/delete:',key=f'um_del_{uid}',disabled=not (manage and not u.get('builtin') and not self_row),help='Hapus'): _delete_dialog(uid)
        st.caption('Klik tombol Aktif/Nonaktif untuk mengubah status akun; pengguna nonaktif tidak dapat login hingga diaktifkan kembali.')
    with tab_matrix:
        st.markdown('**Izin standar tiap peran** — izin ini diterapkan otomatis saat peran dipilih, dan dapat dikustomisasi per pengguna.')
        th=''.join(f'<th style="padding:8px 6px;text-align:center">{_rchip(ROLE_LABELS[r],r)}</th>' for r in ROLE_ORDER)
        body='';last=None
        for k,label,grp,desc in PERMISSIONS:
            if grp!=last: body+=f'<tr><td colspan="{len(ROLE_ORDER)+1}" style="background:#f1f5f9;font-size:11px;font-weight:800;letter-spacing:.5px;color:#475569;padding:6px 10px">{grp.upper()}</td></tr>'; last=grp
            cells=''.join(f'<td style="text-align:center;padding:8px 6px;color:{"#16a34a" if k in ROLE_DEFAULTS[r] else "#cbd5e1"};font-weight:800">{"✓" if k in ROLE_DEFAULTS[r] else "–"}</td>' for r in ROLE_ORDER)
            body+=f'<tr style="border-bottom:1px solid #e8eef6"><td style="padding:8px 10px"><b style="font-size:13px">{label}</b><div style="color:#64748b;font-size:11.5px">{desc}</div></td>{cells}</tr>'
        st.markdown(f'<div style="overflow-x:auto;background:#fff;border:1px solid #e2e8f0;border-radius:12px"><table style="width:100%;border-collapse:collapse"><thead><tr><th style="text-align:left;padding:8px 10px;color:#475569;font-size:11px;letter-spacing:.6px">HAK AKSES</th>{th}</tr></thead><tbody>{body}</tbody></table></div>',unsafe_allow_html=True)
    cards=''.join(f'<div class="um-card"><div style="margin-bottom:6px">{_rchip(ROLE_LABELS[r],r)} <span style="color:#64748b;font-size:11.5px">· {ROLE_DESC[r][0]}</span></div><div style="font-size:12.5px;color:#334155;line-height:1.5">{ROLE_DESC[r][1]}</div></div>' for r in ROLE_ORDER)
    st.markdown(f'<div class="um-cards">{cards}</div>',unsafe_allow_html=True)

def my_account_box():
    """Pengguna non-bawaan dapat mengganti password sendiri."""
    u=CURRENT_USER
    if not u or u.get('builtin'): return
    with st.expander(':material/key: Ganti password saya'):
        with st.form('my_pw_form'):
            old=st.text_input('Password saat ini',type='password'); n1=st.text_input('Password baru',type='password'); n2=st.text_input('Konfirmasi password baru',type='password')
            if st.form_submit_button('Simpan password',use_container_width=True):
                if not _verify_pw(old,u.get('pw','')): st.error('Password saat ini salah.')
                elif len(n1)<6: st.error('Password baru minimal 6 karakter.')
                elif n1!=n2: st.error('Konfirmasi password tidak sama.')
                else:
                    allu=load_users()
                    for r in allu:
                        if r.get('id')==u['id']: r['pw']=_hash_pw(n1)
                    save_users(allu); st.success('Password diperbarui.')

def user_management_page(license_ok,license_msg,license_payload):
    st.markdown("""<div class="input-hero"><div class="eyebrow">KHARISMA • USER & LICENSE CONTROL</div><h1>User Management</h1><p>Kelola akun lokal, status lisensi, aktivasi license key, dan reset binding komputer.</p></div>""",unsafe_allow_html=True)
    is_super=bool(CURRENT_USER and (CURRENT_USER.get('builtin') or CURRENT_USER.get('role')=='super_admin'))
    c1,c2,c3=st.columns(3)
    if license_ok and license_payload:
        exp=datetime.date.fromisoformat(license_payload['expires']); days=(exp-datetime.date.today()).days
        c1.metric('Status License','AKTIF'); c2.metric('Berlaku Sampai',exp.strftime('%d/%m/%Y')); c3.metric('Sisa Hari',max(0,days))
    else:
        c1.metric('Status License','EXPIRED / BELUM AKTIF'); c2.metric('Keterangan',license_msg); c3.metric('Akses','User Management saja')
    if is_super:
        st.markdown('<div class="section-title"><span class="msr">lock</span> License Activation</div>',unsafe_allow_html=True)
        _lic_state_now=load_license_state()
        with st.form('user_license_activation'):
            cust_name=st.text_input('Nama Customer / Perusahaan',value=_lic_state_now.get('customer_name',''),placeholder='Contoh: PT Kharisma Sentosa')
            key=st.text_input('License Key',placeholder='Paste license key di sini',type='password')
            a,b=st.columns([2,1])
            with a: submitted=st.form_submit_button(':material/lock: AKTIFKAN / PERPANJANG LISENSI',use_container_width=True,type='primary')
            with b: reset=st.form_submit_button(':material/autorenew: RESET BINDING',use_container_width=True)
        if submitted:
            ok,msg=activate_license(key,cust_name)
            (st.success if ok else st.error)(msg)
            if ok: st.rerun()
        if reset:
            save_license_state({}); st.success('Binding license lokal sudah di-reset. Silakan masukkan license key baru.'); st.rerun()
    if has_perm('user_manage'): rbac_admin_area()
    if is_super:
        with st.expander(':material/admin_panel_settings: Akun Administrator & Lisensi (Profil, Admin License)'):
            tabs=st.tabs(['Profil','Admin License'])
            auth=load_auth()
            with tabs[0]:
                st.write(f"**Username aktif:** `{auth.get('username','admin')}`")
                with st.form('change_user'):
                    nu=st.text_input('Ganti Username',value=auth.get('username','admin'))
                    np=st.text_input('Ganti Password',type='password',placeholder='Minimal 6 karakter')
                    cp=st.text_input('Konfirmasi Password',type='password')
                    if st.form_submit_button(':material/save: SIMPAN PERUBAHAN',use_container_width=True):
                        if not nu.strip(): st.error('Username tidak boleh kosong.')
                        elif np and len(np)<6: st.error('Password minimal 6 karakter.')
                        elif np and np!=cp: st.error('Konfirmasi password tidak sama.')
                        else:
                            auth['username']=nu.strip()
                            if np: auth['password_hash']=hashlib.sha256(np.encode()).hexdigest()
                            save_auth(auth)
                            clear_login_session()
                            st.session_state['authenticated']=False
                            st.session_state.pop('logged_in_username',None)
                            st.success('Profil berhasil diperbarui. Silakan login ulang agar sesi 48 jam dibuat ulang.')
                            st.rerun()
                with st.expander(':material/key: Lupa password / reset ke default'):
                    st.warning('Reset lokal akan mengembalikan password menjadi `admin123`.')
                    if st.button('Reset Password ke admin123',key='reset_pw'):
                        auth['password_hash']=hashlib.sha256('admin123'.encode()).hexdigest(); save_auth(auth); st.success('Password sudah di-reset ke admin123.')
            with tabs[1]:
                state=load_license_state()
                st.write(f"**Computer ID:** `{_machine_fingerprint()}`")
                st.write(f"**Binding tersimpan:** `{state.get('fingerprint','Belum ada')}`")
                st.write('Gunakan RESET BINDING di atas untuk menghapus ikatan license lokal sebelum aktivasi ulang.')
    my_account_box()
    st.divider()
    if st.button(':material/logout: LOGOUT',use_container_width=True):
        clear_login_session()
        st.session_state['authenticated']=False; st.session_state.pop('page',None); st.rerun()

# Streamlit Session State is reset when the browser is refreshed. Therefore the
# application uses a small local, machine-bound login session with a 48-hour TTL.
# This is intentionally separate from license state and user credentials.
if 'authenticated' not in st.session_state:
    persisted_login=load_login_session()
    st.session_state['authenticated']=bool(persisted_login)
    if persisted_login:
        st.session_state['logged_in_username']=persisted_login.get('username','')
if not st.session_state['authenticated']:
    logo=logo_data_uri()
    st.markdown(f"""<div class="license-shell"><div class="license-card"><div class="license-brand"><img src="{logo}"><div><b>KHARISMA WORKSHOP ANALYTICS</b><span>LOGIN • USER MANAGEMENT</span></div></div><p style="color:#475569;font-size:13px;line-height:1.6;">Silakan login untuk masuk ke aplikasi. Default awal: <b>admin</b> / <b>admin123</b>.</p></div></div>""",unsafe_allow_html=True)
    with st.form('login_form'):
        u=st.text_input('Username / Email',value='admin')
        pw=st.text_input('Password',type='password')
        login=st.form_submit_button(':material/lock: LOGIN',use_container_width=True,type='primary')
    if login:
        _lu,_lmsg=authenticate_user(u,pw)
        if _lu:
            save_login_session(u)
            st.session_state['authenticated']=True
            st.session_state['logged_in_username']=u.strip()
            st.rerun()
        else: st.error(_lmsg)
    st.stop()

# Pengguna aktif (peran + hak akses). Akun yang dinonaktifkan/dihapus otomatis dikeluarkan.
CURRENT_USER=resolve_current_user()
if CURRENT_USER is None:
    clear_login_session(); st.session_state['authenticated']=False; st.session_state.pop('logged_in_username',None); st.rerun()

license_ok, license_msg, license_payload=current_license()
page_labels=['📊 Dashboard','🎯 Target Perusahaan','📊 Comparative Analytics','🧠 Management Insight','🚗 UE','💰 Revenue & Profit','🔩 Sparepart','🧪 Bahan & Consumable','🛠️ OPL','👥 Customer','✅ Quality','📥 Data Input','🏢 Branch Management','👤 User Management']
# Nilai di atas (termasuk emoji) TIDAK diubah: dipakai di banyak perbandingan
# `if page == '...'` di seluruh file. Yang berubah hanya TAMPILANNYA di menu,
# lewat format_func di st.radio (lihat PAGE_ICON_LABELS di bawah) — nilai
# yang dikembalikan st.radio tetap identik dengan item di page_labels.
PAGE_ICON_LABELS={
    '📊 Dashboard':':material/dashboard: Dashboard',
    '🎯 Target Perusahaan':':material/target: Target & Pencapaian',
    '📊 Comparative Analytics':':material/account_tree: Comparative Analytics',
    '🧠 Management Insight':':material/lightbulb: Management Insight',
    '🚗 UE':':material/directions_car: UE',
    '💰 Revenue & Profit':':material/attach_money: Revenue & Profit',
    '🔩 Sparepart':':material/settings: Sparepart',
    '🧪 Bahan & Consumable':':material/science: Bahan & Consumable',
    '🛠️ OPL':':material/build: OPL',
    '👥 Customer':':material/group: Customer',
    '✅ Quality':':material/verified_user: Quality',
    '📥 Data Input':':material/upload_file: Data Input',
    '📋 Data':':material/table_chart: Data',
    '🏢 Branch Management':':material/apartment: Branch Management',
    '👤 User Management':':material/manage_accounts: User Management',
}
MONTH_NAMES=['Januari','Februari','Maret','April','Mei','Juni','Juli','Agustus','September','Oktober','November','Desember']
TARGET_METRICS=['UE','Jasa','Part','Bahan','Revenue']
TARGET_PATH=os.path.join(DATA_DIR,'targets.json')

BAHAN_MAP_PATH=os.path.join(DATA_DIR,'bahan_mapping.json')
def load_bahan_manual():
    try:
        with open(BAHAN_MAP_PATH,'r',encoding='utf-8') as fh: d=json.load(fh)
        return {RK.norm_kode(k):v for k,v in d.items() if v in RK.BAHAN_KATS} if isinstance(d,dict) else {}
    except Exception: return {}
def save_bahan_manual(d):
    tmp=BAHAN_MAP_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(d,fh,ensure_ascii=False,indent=1)
    os.replace(tmp,BAHAN_MAP_PATH)

def load_targets():
    """Target perusahaan: {'2026-10': {'Kharisma Makassar': {'UE':..,'Jasa':..,'Part':..,'Bahan':..,'Revenue':..}}}"""
    try:
        with open(TARGET_PATH,'r',encoding='utf-8') as fh: d=json.load(fh)
        return d if isinstance(d,dict) else {}
    except Exception: return {}

def save_targets(d):
    tmp=TARGET_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(d,fh,ensure_ascii=False,indent=1)
    os.replace(tmp,TARGET_PATH)

def _available_years():
    ys={datetime.date.today().year}
    try:
        with open(os.path.join(DATA_DIR,'datasets.json'),'r',encoding='utf-8') as fh: man=json.load(fh)
        for r in man:
            for k in ('period_start','period_end'):
                v=pd.to_datetime(r.get(k),errors='coerce')
                if pd.notna(v): ys.add(int(v.year))
    except Exception: pass
    return sorted(ys,reverse=True)

def _reset_time_filter():
    for k in ('flt_year','flt_month','flt_day'): st.session_state[k]='Semua'

if 'page' not in st.session_state: st.session_state['page']='📊 Dashboard'
def _latest_data_date():
    """Tanggal data terbaru (dari manifest) -> default filter harian."""
    best=None
    try:
        with open(os.path.join(DATA_DIR,'datasets.json'),'r',encoding='utf-8') as fh: man=json.load(fh)
        for r in man:
            for k in ('period_end','period_start'):
                v=pd.to_datetime(r.get(k),errors='coerce')
                if pd.notna(v) and (best is None or v>best): best=v
    except Exception: pass
    return best.date() if best is not None else datetime.date.today()

def _sb_head(icon,title,tag=''):
    _t=f'<em>{_html.escape(str(tag))}</em>' if tag else ''
    st.markdown(f'<div class="sb-sec"><span><span class="msr">{icon}</span>{_html.escape(title)}</span>{_t}</div>',unsafe_allow_html=True)

def _logout_cb():
    clear_login_session(); st.session_state.update({'authenticated':False}); st.session_state.pop('page',None)

with st.sidebar:
    _logo=logo_data_uri()
    st.markdown(f"""<div class="k-logo"><img src="{_logo}"><div class="brand"><b>KHARISMA</b><br><span>WORKSHOP ANALYTICS</span></div></div>""",unsafe_allow_html=True)
    if license_ok:
        allowed=allowed_pages()
        if not allowed:
            st.error('Akun Anda belum memiliki hak akses ke menu apa pun. Hubungi administrator.')
    else:
        allowed=['👤 User Management'] if has_perm('user_manage') else []
        st.warning('Lisensi belum aktif / expired. Akses analitik dikunci.')
    if not allowed:
        st.button(':material/logout: Logout',key='sb_logout_empty',on_click=_logout_cb,use_container_width=True)
        st.stop()

    if license_ok:
        # ---------- FILTER WAKTU (tab: Harian / Bulanan / Tahunan) ----------
        _years=_available_years(); _ld=_latest_data_date()
        st.session_state.setdefault('flt_mode','Harian')
        st.session_state.setdefault('flt_date',_ld)
        st.session_state.setdefault('flt_month_sel',MONTH_NAMES[_ld.month-1])
        st.session_state.setdefault('flt_year_m',_ld.year if _ld.year in _years else _years[0])
        st.session_state.setdefault('flt_year_y',_ld.year if _ld.year in _years else _years[0])
        _mode=st.session_state.get('flt_mode') or 'Harian'
        _sb_head('schedule','Filter Waktu',_mode.upper())
        _mc=st.columns(3,gap='small')
        for _col,(_m,_k) in zip(_mc,[('Harian','h'),('Bulanan','b'),('Tahunan','t')]):
            with _col:
                st.button(_m,key=f'flt_btn_{_k}',type='primary' if _mode==_m else 'secondary',use_container_width=True,on_click=lambda m=_m: st.session_state.update({'flt_mode':m}))
        _mode=st.session_state.get('flt_mode') or 'Harian'
        if _mode=='Harian':
            st.date_input('Tanggal',key='flt_date',format='DD/MM/YYYY',label_visibility='collapsed')
            _d=st.session_state.get('flt_date')
            if not isinstance(_d,datetime.date): _d=_ld
            st.session_state['flt_year'],st.session_state['flt_month'],st.session_state['flt_day']=_d.year,MONTH_NAMES[_d.month-1],_d.day
        elif _mode=='Bulanan':
            _c1,_c2=st.columns([3,2])
            with _c1: st.selectbox('Bulan',MONTH_NAMES,key='flt_month_sel',label_visibility='collapsed')
            with _c2: st.selectbox('Tahun',_years,key='flt_year_m',label_visibility='collapsed')
            st.session_state['flt_year'],st.session_state['flt_month'],st.session_state['flt_day']=st.session_state['flt_year_m'],st.session_state['flt_month_sel'],'Semua'
        else:
            st.selectbox('Tahun',_years,key='flt_year_y',label_visibility='collapsed',format_func=lambda v:f'Tahun {v}')
            st.session_state['flt_year'],st.session_state['flt_month'],st.session_state['flt_day']=st.session_state['flt_year_y'],'Semua','Semua'

        # ---------- FILTER CABANG ----------
        _ub=user_branches()
        branch_options=(['Semua Cabang']+list(BRANCHES)) if (has_perm('all_branches') or set(_ub)>=set(BRANCHES)) else (_ub or ['Semua Cabang'])
        if st.session_state.get('selected_branch') not in branch_options:
            st.session_state['selected_branch']=branch_options[0]
        _cur_b=st.session_state.get('global_branch_selector')
        if _cur_b not in branch_options: _cur_b=st.session_state['selected_branch']
        _sb_head('storefront','Filter Cabang',_cur_b)
        selected_branch=st.selectbox('Pilih Cabang',branch_options,index=branch_options.index(st.session_state['selected_branch']),key='global_branch_selector',label_visibility='collapsed',disabled=(len(branch_options)==1))
        st.session_state['selected_branch']=selected_branch

    # ---------- MENU UTAMA ----------
    st.markdown('<div class="side-section">Menu Utama</div>',unsafe_allow_html=True)
    current=st.session_state.get('page','👤 User Management' if not license_ok else '📊 Dashboard')
    if current not in allowed: current=allowed[0]
    page=st.radio('NAVIGATION',allowed,index=allowed.index(current),label_visibility='collapsed',format_func=lambda x: PAGE_ICON_LABELS.get(x,x))
    st.session_state['page']=page

    # ---------- KARTU USER (menempel di bawah) ----------
    _uname=str(CURRENT_USER.get('name','') or CURRENT_USER.get('username','') or '?')
    _role_txt=ROLE_LABELS.get(CURRENT_USER.get('role'),'')
    with st.container(key='sb_bottom'):
        with st.container(key='sb_userbox'):
            st.markdown(f'<div class="sb-user"><div class="sb-avatar">{_html.escape(_uname[:1].upper())}</div><div class="sb-uinfo"><b>{_html.escape(_uname)}</b><span>{_html.escape(_role_txt)}</span></div></div>',unsafe_allow_html=True)
            st.button(':material/logout: Logout',key='sb_logout',on_click=_logout_cb,use_container_width=True)
        if license_payload:
            exp=datetime.date.fromisoformat(license_payload['expires']); days=(exp-datetime.date.today()).days
            _foot_cust=(license_payload.get('customer') or '').strip() or 'Lisensi Aktif'
            st.markdown(f'<div class="side-footer lic-card"><b>LICENSE ACTIVE</b><span>{_html.escape(_foot_cust)} • {max(0,days)} hari tersisa<br>Expired {exp.strftime("%d/%m/%Y")}</span></div>',unsafe_allow_html=True)
        else:
            st.markdown('<div class="side-footer lic-card"><b>LICENSE REQUIRED</b><span>Aktifkan melalui User Management</span></div>',unsafe_allow_html=True)

# =========================
# V11.4 MULTI-FILE HISTORICAL DATABASE
# =========================
FILE_SPECS=[
    ('pkb','1️⃣ MaintainPKB.xls','MaintainPKB'),
    ('sales','2️⃣ RptServicePenjualan.xls','Service Penjualan'),
    ('parts_service','3️⃣ RptPenjualanPart.Billing Service.xls','Penjualan Part • Billing Service'),
    ('parts_spare','4️⃣ RptPenjualanPart.Billing SparePart.xls','Penjualan Part • Billing SparePart'),
    ('issue','5️⃣ RptNotaBahan.Bill Of Material Issue To PKB.xls','Bahan • Issue To PKB'),
    ('cons','6️⃣ RptNotaBahan.Bill Of Sales Order Consumable.xls','Bahan • Consumable'),
    ('opl','7️⃣ Laporan OPL.xls','OPL'),
    ('cust','8️⃣ RptPelanggan.xls','Pelanggan'),
    ('est','9️⃣ RptEstimasiServiceAdvisor.xls','Estimasi Service Advisor')]
REPORT_NAMES={k:n for k,_,n in FILE_SPECS}
HIST_DIR=os.path.join(DATA_DIR,'historical'); MANIFEST_PATH=os.path.join(DATA_DIR,'datasets.json')
os.makedirs(HIST_DIR,exist_ok=True)

def _load_manifest():
    try:
        with open(MANIFEST_PATH,'r',encoding='utf-8') as fh: return json.load(fh)
    except Exception: return []
def _save_manifest(rows):
    tmp=MANIFEST_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(rows,fh,ensure_ascii=False,indent=2)
    os.replace(tmp,MANIFEST_PATH)
def _safe_filename(name):
    name=re.sub(r'[^A-Za-z0-9._-]+','_',str(name)).strip('._') or 'laporan.xlsx'
    return name[:150]
def _safe_branch(branch):
    return _safe_filename(branch)[:80] or 'Cabang'
def _file_hash(data): return hashlib.sha256(data).hexdigest()
def _period_from_df(df,date_cols):
    dates=[]
    for c in date_cols:
        if c in df.columns:
            x=pd.to_datetime(df[c],errors='coerce')
            if x.notna().any(): dates.extend(x.dropna().tolist())
    if not dates:return None,None,'Tidak terdeteksi'
    a,b=pd.Timestamp(min(dates)),pd.Timestamp(max(dates)); label=a.strftime('%b %Y') if a.to_period('M')==b.to_period('M') else f'{a.strftime("%b %Y")} – {b.strftime("%b %Y")}'
    return a.strftime('%Y-%m-%d'),b.strftime('%Y-%m-%d'),label
def _detect_period(rt,data):
    try:
        if rt=='pkb': return _period_from_df(parse_pkb(data),['Tanggal'])
        if rt=='sales': return _period_from_df(parse_service(data)[1],['Tanggal Sales'])
        if rt in ('parts_service','parts_spare'): return _period_from_df(parse_parts(data),['Tanggal'])
        if rt in ('issue','cons'): return _period_from_df(parse_material(data),['Tanggal'])
        if rt=='opl': return _period_from_df(parse_opl(data),['OPL Date','Bill Date'])
        if rt=='cust': return _period_from_df(parse_customer(data),['Tgl Service','Tgl Selesai'])
        if rt=='est':
            e=parse_estimasi(data)
            if e.empty: return None,None,'Tidak terdeteksi'
            a,b=pd.Timestamp(e['Periode Awal'].iloc[0]),pd.Timestamp(e['Periode Akhir'].iloc[0])
            if (b-a).days<LONG_REPORT_DAYS:   # laporan harian / periode pendek: tampilkan tanggalnya
                _lbl=a.strftime('%d %b %Y') if a==b else f'{a.strftime("%d %b")} – {b.strftime("%d %b %Y")}'
                return a.strftime('%Y-%m-%d'),b.strftime('%Y-%m-%d'),_lbl
            return a.strftime('%Y-%m-%d'),b.strftime('%Y-%m-%d'),(a.strftime('%b %Y') if a.to_period('M')==b.to_period('M') else f'{a.strftime("%b %Y")} – {b.strftime("%b %Y")}')
    except Exception: pass
    return None,None,'Tidak terdeteksi'
def _migrate_legacy_files():
    legacy={'pkb':'MaintainPKB.xlsx','sales':'RptServicePenjualan.xlsx','parts_service':'RptPenjualanPart.xlsx','issue':'Issue To PKB.xlsx','cons':'Consumable.xlsx','cust':'RptPelanggan.xlsx'}
    legacy_branch=BRANCHES[0] if BRANCHES else 'Kharisma Makassar'
    manifest=_load_manifest(); known={(r.get('sha256'),r.get('branch') or legacy_branch) for r in manifest}; changed=False
    for r in manifest:
        if not r.get('branch'):
            r['branch']=legacy_branch; changed=True
    for rt,fn in legacy.items():
        path=os.path.join(DATA_DIR,fn)
        if not os.path.isfile(path) or os.path.getsize(path)==0: continue
        try:
            data=open(path,'rb').read(); sha=_file_hash(data)
            if (sha,legacy_branch) in known: continue
            a,b,label=_detect_period(rt,data); folder=os.path.join(HIST_DIR,rt,_safe_branch(legacy_branch)); os.makedirs(folder,exist_ok=True)
            stored=os.path.join(folder,f'{sha[:12]}_{_safe_filename(fn)}'); open(stored,'wb').write(data)
            manifest.append({'id':sha[:16]+'-'+_safe_branch(legacy_branch)[:20],'report_type':rt,'branch':legacy_branch,'original_filename':fn,'stored_path':stored,'sha256':sha,'period_start':a,'period_end':b,'period_label':label,'uploaded_at':datetime.datetime.now().isoformat(timespec='seconds'),'size_bytes':len(data)})
            known.add((sha,legacy_branch)); changed=True
        except Exception: pass
    for r in manifest:
        if r.get('period_label') in (None,'','Tidak terdeteksi') and os.path.isfile(r.get('stored_path','')):
            try:
                blob=open(r['stored_path'],'rb').read(); a,b,label=_detect_period(r.get('report_type',''),blob)
                if label!='Tidak terdeteksi':
                    r['period_start']=a; r['period_end']=b; r['period_label']=label; changed=True
            except Exception: pass
    if changed:_save_manifest(manifest)
def _historical_records(rt=None, branch=None):
    rows=_load_manifest(); out=[]; default_branch=BRANCHES[0] if BRANCHES else 'Kharisma Makassar'
    for r in rows:
        if rt and r.get('report_type')!=rt: continue
        rb=r.get('branch') or default_branch
        if branch and branch!='Semua Cabang' and rb!=branch: continue
        if os.path.isfile(r.get('stored_path','')) and os.path.getsize(r.get('stored_path',''))>0:
            rr=dict(r); rr['branch']=rb; out.append(rr)
    return out
def detect_branch_from_report(data, branches):
    """Baca ~15 baris pertama laporan (header Crystal Reports: nama perusahaan/cabang/alamat) dan cari nama cabang
    sistem terpisah (mis. 'Ambon'). Hanya cabang di luar SHARED_SYSTEM_BRANCHES yang dideteksi, karena laporan gabungan
    Makassar/Gowa/Kendari sudah dipilah per baris. Mengembalikan nama cabang atau None bila tidak ada / ambigu."""
    try:
        raw=read_excel_auto(data,header=None,nrows=15)
    except Exception:
        return None
    txt=' '.join(str(v) for v in raw.fillna('').astype(str).values.ravel() if str(v).strip()).upper()
    hits=[]
    for b in branches:
        if b in SHARED_SYSTEM_BRANCHES: continue
        words=[w for w in re.split(r'\s+',str(b).upper().strip()) if w and w not in {'KHARISMA','CABANG','PT','PT.'}]
        key=' '.join(words)
        if key and re.search(r'(?<![A-Z0-9])'+re.escape(key)+r'(?![A-Z0-9])',txt): hits.append(b)
    return hits[0] if len(hits)==1 else None

def _store_uploaded(rt,uploaded_files,branch):
    branch=clean(branch) or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')
    manifest=_load_manifest(); default_branch=BRANCHES[0] if BRANCHES else 'Kharisma Makassar'
    known={(r.get('report_type'),r.get('sha256')) for r in manifest}; results=[]
    selected_branch=branch
    for up in uploaded_files:
        data=up.getvalue(); sha=_file_hash(data)
        # Cabang otomatis dari isi laporan (mis. header memuat 'Ambon'); bila tidak terdeteksi, pakai cabang terpilih.
        _det=detect_branch_from_report(data,BRANCHES)
        if _det and _det not in user_branches(): _det=None   # pengguna tidak boleh mengarahkan file ke cabang di luar wewenangnya
        branch=_det or selected_branch
        folder=os.path.join(HIST_DIR,rt,_safe_branch(branch)); os.makedirs(folder,exist_ok=True)
        if (rt,sha) in known:
            prev=next((r for r in manifest if r.get('report_type')==rt and r.get('sha256')==sha),{})
            prev_branch=prev.get('branch') or '-'
            if prev and prev_branch!=branch and branch not in SHARED_SYSTEM_BRANCHES:
                # File identik sebelumnya tersimpan di cabang lain (mis. diupload ke Makassar sebelum cabang Ambon dibuat).
                # Cabang sistem terpisah (Ambon) 'mengklaim' file itu: dipindahkan, tidak digandakan.
                try:
                    newpath=os.path.join(folder,os.path.basename(prev.get('stored_path','')) or f'{sha[:12]}_{_safe_filename(up.name)}')
                    if prev.get('stored_path') and os.path.isfile(prev['stored_path']): shutil.move(prev['stored_path'],newpath)
                    else: open(newpath,'wb').write(data)
                    prev['stored_path']=newpath; prev['branch']=branch
                    did=f'{sha[:16]}-{_safe_branch(branch)[:20]}'
                    prev['id']=did if not any(r is not prev and r.get('id')==did for r in manifest) else f'{sha[:16]}-{uuid.uuid4().hex[:8]}'
                    results.append(('ok',up.name,f'{prev.get("period_label","")} • file identik sudah ada di cabang {prev_branch} → DIPINDAHKAN ke {branch} • Cabang: {branch}'))
                except Exception as e: results.append(('err',up.name,f'Gagal memindahkan ke {branch}: {e}'))
                continue
            when=str(prev.get('uploaded_at','') or '').replace('T',' ')
            results.append(('skip',up.name,f'File identik sudah tersimpan di cabang {prev_branch}'+(f' (diupload {when})' if when else '')+(f' sebagai "{prev.get("original_filename")}"' if prev.get('original_filename') and prev.get('original_filename')!=up.name else '')+'. Dilewati agar data tidak dobel; datanya sudah ada — pilih cabang tersebut di sidebar untuk melihatnya. Untuk mengupload ulang, hapus dulu file lama di Kelola histori.')); continue
        a,b,label=_detect_period(rt,data); stored=os.path.join(folder,f'{sha[:12]}_{_safe_filename(up.name)}')
        try:
            open(stored,'wb').write(data)
            did=f'{sha[:16]}-{_safe_branch(branch)[:20]}'
            if any(r.get('id')==did for r in manifest): did=f'{sha[:16]}-{uuid.uuid4().hex[:8]}'
            manifest.append({'id':did,'report_type':rt,'branch':branch,'original_filename':up.name,'stored_path':stored,'sha256':sha,'period_start':a,'period_end':b,'period_label':label,'uploaded_at':datetime.datetime.now().isoformat(timespec='seconds'),'size_bytes':len(data)})
            known.add((rt,sha))
            ov=[r.get('original_filename') for r in manifest[:-1] if r.get('report_type')==rt and a and b and r.get('period_start') and r.get('period_end') and r['period_start']<=b and a<=r['period_end']]
            _dtxt=(f' • cabang terdeteksi otomatis dari isi laporan: {branch}' if _det else '')
            if ov: results.append(('warn',up.name,f'{label}{_dtxt} • Cabang: {branch} • periode TUMPANG TINDIH dengan: {", ".join(ov[:3])}. Baris identik otomatis tidak dihitung dua kali; bila ini export ulang yang isinya berubah, hapus file lama di Kelola histori.'))
            else: results.append(('ok',up.name,label+_dtxt+f' • Cabang: {branch}'))
        except Exception as e: results.append(('err',up.name,str(e)))
    _save_manifest(manifest); return results
def _delete_all_datasets():
    for r in _historical_records():
        try: os.remove(r.get('stored_path',''))
        except Exception: pass
    _save_manifest([])

_dialog=getattr(st,'dialog',None) or getattr(st,'experimental_dialog',None)
def _confirm_delete_body(ids,labels,delete_all):
    st.warning('**Data yang dihapus tidak dapat dikembalikan.**' ,icon=':material/warning:')
    if delete_all: st.write(f'Anda akan menghapus **SEMUA {len(labels)} file histori** (seluruh cabang dan seluruh laporan).')
    else:
        st.write(f'Anda akan menghapus **{len(ids)} file**:')
        for i in ids[:12]: st.markdown(f'- {labels.get(i,i)}')
        if len(ids)>12: st.caption(f'… dan {len(ids)-12} file lainnya')
    ok=True
    if delete_all:
        ok=st.text_input('Ketik **HAPUS** untuk melanjutkan',key='confirm_delete_all_text').strip().upper()=='HAPUS'
    c1,c2=st.columns(2)
    if c1.button('Ya, hapus',type='primary',disabled=not ok,use_container_width=True,key='confirm_delete_yes'):
        if delete_all: _delete_all_datasets(); st.session_state['_del_toast']='Seluruh histori sudah dihapus.'
        else:
            for did in ids: _delete_dataset(did)
            st.session_state['_del_toast']=f'{len(ids)} file histori dihapus.'
        st.rerun()
    if c2.button('Batal',use_container_width=True,key='confirm_delete_no'): st.rerun()
if _dialog:
    @_dialog('Konfirmasi hapus data')
    def confirm_delete_dialog(ids,labels,delete_all=False): _confirm_delete_body(ids,labels,delete_all)
else:
    def confirm_delete_dialog(ids,labels,delete_all=False): _confirm_delete_body(ids,labels,delete_all)

def _delete_dataset(did):
    manifest=_load_manifest(); kept=[]
    for r in manifest:
        if r.get('id')==did:
            try: os.remove(r.get('stored_path',''))
            except Exception: pass
        else: kept.append(r)
    _save_manifest(kept)
def _read_all_bytes(rt, branch=None):
    out=[]
    for r in _historical_records(rt,branch):
        try: out.append(open(r['stored_path'],'rb').read())
        except Exception: pass
    return out
def _read_records_with_bytes(rt, branch=None):
    out=[]
    for r in _historical_records(rt,branch):
        try: out.append((r,open(r['stored_path'],'rb').read()))
        except Exception: pass
    return out
def load_historical_data():
    _migrate_legacy_files()
    km={'pkb':'fpkb_bytes','sales':'fsales_bytes','parts_service':'fparts_service_bytes','parts_spare':'fparts_spare_bytes','issue':'fissue_bytes','cons':'fcons_bytes','opl':'fopl_bytes','cust':'fcust_bytes','est':'fest_bytes'}
    for rt,key in km.items():
        blobs=_read_all_bytes(rt); st.session_state[key]=blobs[-1] if blobs else None
def aggregate_historical(branch=None):
    P=[];SD=[];PAS=[];PAD=[];I=[];C=[];O=[];CU=[]
    parsers={'pkb':(P,parse_pkb),'sales':(SD,lambda b:parse_service(b)[1]),'parts_service':(PAS,parse_parts),'parts_spare':(PAD,parse_parts),'issue':(I,parse_material),'cons':(C,parse_material),'opl':(O,parse_opl),'cust':(CU,parse_customer)}

    # Parse every stored file first.  IMPORTANT: do not filter the manifest by
    # upload branch here. One shared source file can contain Makassar + Gowa +
    # Kendari, so the real branch must be resolved from SA/PKB before filtering.
    parsed=[]
    for rt,(bucket,fn) in parsers.items():
        for rec,b in _read_records_with_bytes(rt,None):
            try:
                df=fn(b)
                if not isinstance(df,pd.DataFrame): continue
                df=df.copy()
                fallback=rec.get('branch') or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')
                df['_Source SHA']=str(rec.get('sha256',''))
                df['_Uploaded Branch']=fallback
                parsed.append((rt,rec,df))
            except Exception:
                pass

    def _date_col_for_report(rt,df):
        candidates={
            'pkb':['Tanggal'], 'sales':['Tanggal Sales'],
            'parts_service':['Tanggal'], 'parts_spare':['Tanggal'],
            'issue':['Tanggal'], 'cons':['Tanggal'],
            'opl':['Bill Date','OPL Date'], 'cust':['Tgl Service','Tgl Selesai']
        }.get(rt,[])
        for c in candidates:
            if c in df.columns: return c
        return None

    def _resolve_sa_branches(rt,df,fallback):
        if 'Service Advisor' not in df.columns:
            return pd.Series([fallback]*len(df),index=df.index,dtype='object')
        dc=_date_col_for_report(rt,df)
        dates=df[dc] if dc else pd.Series([pd.NaT]*len(df),index=df.index)
        return pd.Series([branch_from_sa(sa,fallback,dt) for sa,dt in zip(df['Service Advisor'],dates)],index=df.index,dtype='object')

    # MaintainPKB is the master branch key because it contains Service Advisor + PKB.
    for rt,rec,df in parsed:
        if rt=='pkb':
            fallback=rec.get('branch') or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')
            if 'Service Advisor' in df.columns:
                df['Cabang']=_resolve_sa_branches(rt,df,fallback)
            else:
                df['Cabang']=fallback
            df.drop(columns=['_Uploaded Branch'],inplace=True,errors='ignore')
            P.append(df)

    # Build PKB -> branch map.  This keeps Service Penjualan / OPL / other
    # PKB-linked reports in the same branch even when they do not contain SA.
    pkb_branch_map={}
    for df in P:
        if {'PKB','Cabang'}.issubset(df.columns):
            for _,r in df[['PKB','Cabang']].dropna().iterrows():
                pkb_key=norm(r.get('PKB'))
                if pkb_key: pkb_branch_map[pkb_key]=r.get('Cabang')

    pkb_sa_map={}
    for df in P:
        if {'PKB','Service Advisor'}.issubset(df.columns):
            for _pk,_sa in zip(df['PKB'],df['Service Advisor']):
                _k=norm(_pk)
                if _k and clean(_sa) and _k not in pkb_sa_map: pkb_sa_map[_k]=norm_sa(_sa)
    # Peta PKB->cabang juga diperkaya dari Laporan Estimasi SA (kode SA) agar PKB yang belum ada di MaintainPKB tetap terpetakan.
    try:
        _est_map=aggregate_estimasi(None)
        if isinstance(_est_map,pd.DataFrame) and not _est_map.empty:
            for _pk,_cb in zip(_est_map['PKB'],_est_map['Cabang']):
                _k=norm(_pk)
                if _k and _k not in pkb_branch_map: pkb_branch_map[_k]=_cb
            if 'Service Advisor' in _est_map.columns:
                for _pk,_sa in zip(_est_map['PKB'],_est_map['Service Advisor']):
                    _k=norm(_pk)
                    if _k and clean(_sa) and _k not in pkb_sa_map: pkb_sa_map[_k]=norm_sa(_sa)
    except Exception:
        pass

    for rt,rec,df in parsed:
        if rt=='pkb':
            continue
        fallback=rec.get('branch') or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')
        if rt=='issue':
            _kc=df['Ref Keys'] if 'Ref Keys' in df.columns else df['No Referensi']
            _res=[resolve_issue_branch(k if clean(k) else r_,(c if fallback in SHARED_SYSTEM_BRANCHES else ''),pkb_branch_map,fallback) for k,r_,c in zip(_kc,df['No Referensi'],df['Customer'])]
            df['Cabang']=[x[0] for x in _res]; df['Sumber Cabang']=[x[1] for x in _res]
            df['SA Code']=[resolve_issue_sa(k,r_,pkb_sa_map) for k,r_ in zip(_kc,df['No Referensi'])]
            df.drop(columns=['_Uploaded Branch','Ref Keys'],inplace=True,errors='ignore')
            bucket=dict(parsers)[rt][0]; bucket.append(df)
            continue
        # Prefer PKB branch, then SA branch, then the branch selected during upload.
        if 'PKB' in df.columns:
            branch_key=df['PKB'].map(norm)
        elif 'No Referensi' in df.columns:
            branch_key=df['No Referensi'].map(norm)
        else:
            branch_key=pd.Series(['']*len(df),index=df.index,dtype='object')
        df['Cabang']=branch_key.map(lambda x: pkb_branch_map.get(x, ''))
        if 'Service Advisor' in df.columns:
            sa_branch=_resolve_sa_branches(rt,df,'')
            df['Cabang']=df['Cabang'].where(df['Cabang'].map(clean)!='',sa_branch)
        df['Cabang']=df['Cabang'].where(df['Cabang'].map(clean)!='',fallback)
        if rt in ('parts_spare','cons') and 'Customer' in df.columns and fallback in SHARED_SYSTEM_BRANCHES:
            # Penjualan langsung: laporan masih gabungan dengan Makassar -> cabang dari kode akhiran nama customer.
            # (Cabang sistem terpisah seperti Ambon tidak memakai aturan ini: cabangnya = cabang upload/terdeteksi.)
            df['Cabang']=df['Customer'].map(branch_from_customer)
            df.drop(columns=['Ref Keys'],inplace=True,errors='ignore')
        df.drop(columns=['_Uploaded Branch'],inplace=True,errors='ignore')
        bucket=dict(parsers)[rt][0]
        bucket.append(df)

    # Apply the selected branch only AFTER branch classification. This is what
    # allows one uploaded combined report to be viewed correctly per branch.
    def _filter_branch(df):
        if df is None or df.empty or not branch or branch=='Semua Cabang': return df
        if 'Cabang' not in df.columns: return df
        return df[df['Cabang'].map(clean)==clean(branch)].copy()

    P=[_filter_branch(x) for x in P]
    SD=[_filter_branch(x) for x in SD]
    PAS=[_filter_branch(x) for x in PAS]
    PAD=[_filter_branch(x) for x in PAD]
    I=[_filter_branch(x) for x in I]
    C=[_filter_branch(x) for x in C]
    O=[_filter_branch(x) for x in O]
    CU=[_filter_branch(x) for x in CU]

    def _concat_dedup(frames,name=''):
        """Gabungkan semua file. Baris yang IDENTIK dengan baris dari file lain (periode tumpang tindih / export ulang)
        dibuang sehingga tidak terhitung dua kali. Baris kembar yang memang ada DI DALAM satu file tetap dipertahankan."""
        frames=[x for x in frames if isinstance(x,pd.DataFrame) and not x.empty]
        if not frames: DEDUP_STATS[name]=0; return pd.DataFrame()
        d=pd.concat(frames,ignore_index=True); before=len(d)
        cols=[c for c in d.columns if c!='_Source SHA']
        if cols and len(frames)>1:
            key=pd.util.hash_pandas_object(d[cols].astype(str),index=False).values
            sha=d['_Source SHA'].astype(str).values if '_Source SHA' in d.columns else np.zeros(len(d))
            occ=pd.DataFrame({'s':sha,'k':key}).groupby(['s','k']).cumcount().values
            d=d.assign(_k=key,_o=occ).drop_duplicates(['_k','_o'],keep='first').drop(columns=['_k','_o'])
        DEDUP_STATS[name]=before-len(d)
        return d.reset_index(drop=True)

    pkb=_concat_dedup(P,'pkb')
    detail=_concat_dedup(SD,'detail')
    parts_service=_concat_dedup(PAS,'parts_service')
    parts_direct=_concat_dedup(PAD,'parts_direct')
    issue=_concat_dedup(I,'issue')
    cons=_concat_dedup(C,'cons')
    opl=_concat_dedup(O,'opl')
    cust=_concat_dedup(CU,'cust')
    if not pkb.empty:
        pkb['_fill']=pkb.notna().sum(axis=1); pkb=pkb.sort_values(['Cabang','PKB','_fill'],ascending=[True,True,False]).drop_duplicates(['Cabang','PKB']).drop(columns='_fill').reset_index(drop=True)
    if detail.empty:
        sales=pd.DataFrame(columns=['Cabang','PKB','Tanggal Sales','Invoice','Customer Sales','Jenis Pembayaran','Type Kendaraan Sales','Jasa','Sparepart','Bahan','Sales Row Count','Revenue'])
    else:
        sales=detail.groupby(['Cabang','PKB'],as_index=False).agg(**{'Tanggal Sales':('Tanggal Sales','min'),'Invoice':('Invoice',lambda x:', '.join(pd.Series(list(x),dtype='string').dropna().unique().tolist())),'Customer Sales':('Customer Sales','first'),'Jenis Pembayaran':('Jenis Pembayaran','first'),'Type Kendaraan Sales':('Type Kendaraan Sales','first'),'Jasa':('Jasa','sum'),'Sparepart':('Sparepart','sum'),'Bahan':('Bahan','sum'),'Sales Row Count':('PKB','size')})
        sales['Revenue']=sales[['Jasa','Sparepart','Bahan']].sum(axis=1)
    ps=parts_service.copy()
    pdirect=parts_direct.copy()
    if not ps.empty: ps['Sumber Penjualan']='Billing Service'; ps['Kategori Penjualan']='Sparepart dengan PKB'
    if not pdirect.empty: pdirect['Sumber Penjualan']='Billing Part'; pdirect['Kategori Penjualan']='Sparepart tanpa PKB'
    parts=pd.concat([ps,pdirect],ignore_index=True) if (not ps.empty or not pdirect.empty) else pd.DataFrame()
    issue=fill_issue_sa_by_customer(issue,detail,pkb)   # cadangan: customer Issue -> Service Penjualan -> MaintainPKB
    if not issue.empty: issue['Kategori Penjualan']='Bahan Issue To PKB'
    if not cons.empty: cons['Kategori Penjualan']='Bahan Consumable tanpa PKB'
    for _df in [pkb,sales,parts,issue,cons,opl,cust,detail]:
        if isinstance(_df,pd.DataFrame):
            _df.drop(columns=['_Source SHA'],inplace=True,errors='ignore')
    return pkb,sales,parts,issue,cons,opl,cust,detail

MANUAL_EST_PATH=os.path.join(DATA_DIR,'manual_est.json')
def load_manual_est():
    try:
        with open(MANUAL_EST_PATH,'r',encoding='utf-8') as fh: rows=json.load(fh)
        return rows if isinstance(rows,list) else []
    except Exception: return []
def save_manual_est(rows):
    os.makedirs(DATA_DIR,exist_ok=True); tmp=MANUAL_EST_PATH+'.tmp'
    with open(tmp,'w',encoding='utf-8') as fh: json.dump(rows,fh,ensure_ascii=False,indent=1)
    os.replace(tmp,MANUAL_EST_PATH)
def manual_est_df():
    """Transaksi Estimasi SA yang diinput manual (tanpa file) -> kolom sama dengan parse_estimasi()."""
    rec=[]
    for r in load_manual_est():
        try: d=pd.Timestamp(r['tanggal']).normalize()
        except Exception: continue
        sa=norm_sa(r.get('sa_kode','')); pk=norm(r.get('pkb',''))
        if not sa or not pk: continue
        g=lambda k:float(r.get(k,0) or 0)
        ej,ep,eb,aj,ap,ab=g('est_jasa'),g('est_part'),g('est_bahan'),g('act_jasa'),g('act_part'),g('act_bahan')
        rec.append({'PKB':pk,'Tanggal':d,'Chassis':clean(r.get('no_polisi','')),'Service Advisor':sa,'SA Nama':clean(r.get('sa_nama','')),
            'Est Biaya':ej+ep+eb,'Act Biaya':aj+ap+ab,'Est Jasa':ej,'Est Part':ep,'Est Bahan':eb,'Act Jasa':aj,'Act Part':ap,'Act Bahan':ab,
            'Est Waktu':d+pd.Timedelta(hours=8),'Act Waktu':d+pd.Timedelta(hours=16),'_fb':clean(r.get('cabang',''))})
    if not rec: return pd.DataFrame()
    df=pd.DataFrame(rec)
    df['Periode Awal']=df['Tanggal']; df['Periode Akhir']=df['Tanggal']; df['Tanggal Periode']=df['Tanggal']
    df['Selisih Biaya']=df['Act Biaya']-df['Est Biaya']
    df['Selisih Biaya %']=np.where(df['Est Biaya']>0,df['Selisih Biaya']/df['Est Biaya']*100,0.0)
    df['Status Biaya']=np.select([df['Selisih Biaya']>0,df['Selisih Biaya']<0],['Lebih Tinggi','Lebih Rendah'],'Sesuai')
    df['Selisih Jam']=(df['Act Waktu']-df['Est Waktu']).dt.total_seconds()/3600
    df['Status Waktu']=np.select([df['Selisih Jam']>0,df['Selisih Jam']==0,df['Selisih Jam']<0],['Terlambat','Tepat Waktu','Lebih Cepat'],'Tidak diketahui')
    df['Cabang']=[branch_from_sa(sa,(fb or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')),dt) for sa,fb,dt in zip(df['Service Advisor'],df['_fb'],df['Tanggal'])]
    return df.drop(columns=['_fb']).drop_duplicates(['Service Advisor','PKB'],keep='last').reset_index(drop=True)

@st.dialog('Tambah Transaksi Manual — Estimasi SA',width='large')
def _manual_est_dialog(default_branch):
    st.caption('Isi satu PKB selesai tanpa file. Transaksi ikut dihitung sebagai UE & revenue sama seperti laporan Estimasi SA (nilai Aktual).')
    _c1,_c2=st.columns(2)
    me_tgl=_c1.date_input('Tanggal selesai PKB',value=datetime.date.today(),key='me_tgl')
    _bl=list(BRANCHES) or ['Kharisma Makassar']
    me_cab=_c2.selectbox('Cabang',_bl,index=(_bl.index(default_branch) if default_branch in _bl else 0),key='me_cab')
    _c3,_c4=st.columns(2)
    me_sa=_c3.text_input('Kode / akun SA *',key='me_sa',placeholder='mis. johan123')
    me_nm=_c4.text_input('Nama lengkap SA',key='me_nm',placeholder='mis. Johan Wahyudi')
    _c5,_c6=st.columns(2)
    me_pkb=_c5.text_input('No PKB *',key='me_pkb')
    me_pol=_c6.text_input('No Polisi / Chassis',key='me_pol')
    st.markdown('**Estimasi (Rp)**')
    _e1,_e2,_e3=st.columns(3)
    ej=_e1.number_input('Jasa',min_value=0,step=10000,key='me_ej'); ep=_e2.number_input('Part',min_value=0,step=10000,key='me_ep'); eb=_e3.number_input('Bahan',min_value=0,step=10000,key='me_eb')
    st.markdown('**Aktual (Rp)** — dipakai untuk revenue')
    _same=st.checkbox('Aktual sama dengan Estimasi',value=True,key='me_same')
    _a1,_a2,_a3=st.columns(3)
    aj=_a1.number_input('Jasa ',min_value=0,step=10000,key='me_aj',disabled=_same); ap=_a2.number_input('Part ',min_value=0,step=10000,key='me_ap',disabled=_same); ab=_a3.number_input('Bahan ',min_value=0,step=10000,key='me_ab',disabled=_same)
    if _same: aj,ap,ab=ej,ep,eb
    if st.button(':material/save: Simpan Transaksi',type='primary',key='me_save',use_container_width=True):
        if not norm_sa(me_sa) or not norm(me_pkb): st.error('Kode SA dan No PKB wajib diisi.')
        elif (aj+ap+ab)<=0: st.error('Nilai Aktual (Jasa/Part/Bahan) belum diisi.')
        else:
            _rows=[r for r in load_manual_est() if not (norm(r.get('pkb',''))==norm(me_pkb) and norm_sa(r.get('sa_kode',''))==norm_sa(me_sa))]
            _rows.append({'id':'m-'+uuid.uuid4().hex[:8],'tanggal':str(me_tgl),'cabang':me_cab,'sa_kode':clean(me_sa),'sa_nama':clean(me_nm),'pkb':clean(me_pkb),'no_polisi':clean(me_pol),
                'est_jasa':float(ej),'est_part':float(ep),'est_bahan':float(eb),'act_jasa':float(aj),'act_part':float(ap),'act_bahan':float(ab),
                'created_by':str((CURRENT_USER or {}).get('name','')),'created_at':datetime.datetime.now().isoformat(timespec='seconds')})
            save_manual_est(_rows); st.session_state['_me_saved']=clean(me_pkb)
            for _k in ['me_sa','me_nm','me_pkb','me_pol','me_ej','me_ep','me_eb','me_aj','me_ap','me_ab']: st.session_state.pop(_k,None)
            st.rerun()

def aggregate_estimasi(branch=None):
    """Gabungkan semua laporan Estimasi Service Advisor tersimpan. Cabang ditentukan dari kode SA (sama seperti laporan lain)."""
    frames=[]
    for rec,b in _read_records_with_bytes('est',None):
        try: d=parse_estimasi(b)
        except Exception: continue
        if not isinstance(d,pd.DataFrame) or d.empty: continue
        d=d.copy(); fallback=rec.get('branch') or (BRANCHES[0] if BRANCHES else 'Kharisma Makassar')
        d['Cabang']=[branch_from_sa(sa,fallback,dt) for sa,dt in zip(d['Service Advisor'],d['Tanggal'])]
        frames.append(d)
    _md=manual_est_df()
    if not _md.empty: frames.append(_md)
    if not frames: return pd.DataFrame()
    out=pd.concat(frames,ignore_index=True).sort_values('Tanggal Periode',kind='stable').drop_duplicates('PKB',keep='last')
    if branch and branch!='Semua Cabang': out=out[out['Cabang'].map(clean)==clean(branch)]
    return out.reset_index(drop=True)

load_historical_data()
if page == '👤 User Management':
    user_management_page(license_ok,license_msg,license_payload)
    st.stop()

if page == '🏢 Branch Management':
    # Header Branch Management dibuat aman sebelum dataframe analitik `f` dibentuk.
    # Periode dan UE diambil langsung dari histori cabang aktif agar tidak ada NameError.
    _bm_branch=st.session_state.get('selected_branch','Semua Cabang')
    _bm_pkb,_bm_sales,*_=aggregate_historical(_bm_branch)
    _bm_dates=pd.to_datetime(_bm_pkb['Tanggal'],errors='coerce') if isinstance(_bm_pkb,pd.DataFrame) and 'Tanggal' in _bm_pkb.columns else pd.Series(dtype='datetime64[ns]')
    _bm_dates=_bm_dates.dropna()
    _bm_period=(f"{_bm_dates.min().strftime('%d/%m/%Y')} — {_bm_dates.max().strftime('%d/%m/%Y')}" if not _bm_dates.empty else '-')
    _bm_ue=len(_bm_pkb) if isinstance(_bm_pkb,pd.DataFrame) else 0
    _bm_user=st.session_state.get('logged_in_username') or 'admin'
    st.markdown(RK.page_header(_bm_period,f'{_bm_ue:,} UE',(CURRENT_USER or {}).get('name') or _bm_user,ROLE_LABELS.get((CURRENT_USER or {}).get('role'),'')),unsafe_allow_html=True)
    section('<span class="msr">apartment</span> Branch Management','Kelola daftar cabang yang tersedia untuk upload dan analisa.')
    st.markdown('<div class="upload-note"><span class="msr">push_pin</span> Tambahkan nama cabang di sini. Data lama tetap aman; cabang hanya menjadi metadata pemisah laporan.</div>',unsafe_allow_html=True)
    st.markdown('### Daftar Cabang')
    branch_df=pd.DataFrame({'No':range(1,len(BRANCHES)+1),'Cabang':BRANCHES})
    st.dataframe(branch_df,use_container_width=True,hide_index=True)
    with st.form('add_branch_form'):
        new_branch=st.text_input('Nama Cabang Baru',placeholder='Contoh: Kharisma Palu')
        add=st.form_submit_button(':material/add: Tambah Cabang',use_container_width=True,type='primary')
    if add:
        nb=clean(new_branch)
        if not nb: st.error('Nama cabang tidak boleh kosong.')
        elif nb.lower() in {x.lower() for x in BRANCHES}: st.warning('Nama cabang tersebut sudah ada.')
        else:
            BRANCHES.append(nb); save_branches(BRANCHES); st.session_state['selected_branch']=nb
            st.success(f'Cabang **{nb}** berhasil ditambahkan.'); st.rerun()
    st.markdown('### :material/link: Mapping Service Advisor → Cabang')
    st.dataframe(branch_mapping_table(),use_container_width=True,hide_index=True)
    st.caption('Pencocokan SA tidak membedakan huruf besar/kecil. Mapping dibaca berdasarkan tanggal laporan, sehingga SA yang dimutasi bisa memiliki cabang berbeda mulai tanggal tertentu.')

    st.markdown('### :material/sync: Mutasi Service Advisor')
    st.caption('Tambahkan aturan baru tanpa mengubah histori lama. Contoh: JOTHA024 pindah dari Kendari ke Makassar mulai 2026-10-01. Laporan sebelum tanggal tersebut tetap Kendari; laporan setelahnya menjadi Makassar.')
    with st.form('sa_mutation_form'):
        c1,c2,c3=st.columns([1.2,1.2,1])
        sa_mut=c1.text_input('Service Advisor',placeholder='Contoh: jotha024')
        branch_mut=c2.selectbox('Cabang Baru',BRANCHES,key='mutation_branch')
        effective_mut=c3.date_input('Berlaku Mulai',value=datetime.date.today(),key='mutation_date')
        save_mut=st.form_submit_button(':material/save: Simpan Mutasi SA',use_container_width=True,type='primary')
    if save_mut:
        sa_key=norm_sa(sa_mut)
        if not sa_key:
            st.error('Service Advisor wajib diisi.')
        elif not branch_mut:
            st.error('Cabang baru wajib dipilih.')
        else:
            hist=_load_sa_branch_history()
            hist=[r for r in hist if not (r.get('sa')==sa_key and r.get('effective_from')==effective_mut.strftime('%Y-%m-%d'))]
            hist.append({'sa':sa_key,'branch':branch_mut,'effective_from':effective_mut.strftime('%Y-%m-%d')})
            hist.sort(key=lambda r:(r.get('sa',''),r.get('effective_from','')))
            _save_sa_branch_history(hist)
            st.success(f'Mutasi **{sa_key.lower()} → {branch_mut}** berlaku mulai **{effective_mut.strftime("%d/%m/%Y")}** berhasil disimpan.')
            st.rerun()

    hist=_load_sa_branch_history()
    if hist:
        st.markdown('#### Riwayat Mutasi SA')
        hdf=pd.DataFrame([{'Service Advisor':r['sa'].lower(),'Cabang':r['branch'],'Berlaku Mulai':r['effective_from']} for r in hist])
        st.dataframe(hdf,use_container_width=True,hide_index=True)
    st.info('ℹ️ Cabang dengan sistem terpisah (mis. Ambon) tidak memakai mapping SA di atas. Saat upload, cabangnya dikenali otomatis dari nama cabang pada header laporan; bila tidak ada, dipakai cabang yang dipilih.')
    st.info(':material/lightbulb: Cabang tidak mengubah License Key, login/session, Revenue & Profit, Dashboard, atau rumus analitik yang sudah ada.')
    st.stop()

if page == '📥 Data Input':
    _di_rec=_historical_records()
    _di_a=[pd.to_datetime(r.get('period_start'),errors='coerce') for r in _di_rec]; _di_b=[pd.to_datetime(r.get('period_end'),errors='coerce') for r in _di_rec]
    _di_a=[x for x in _di_a if pd.notna(x)]; _di_b=[x for x in _di_b if pd.notna(x)]
    _di_per=(f"{min(_di_a).strftime('%d/%m/%Y')} - {max(_di_b).strftime('%d/%m/%Y')}" if _di_a and _di_b else '-')
    st.markdown(RK.page_header(_di_per,f'{len(_di_rec)} file',(CURRENT_USER or {}).get('name') or 'admin',ROLE_LABELS.get((CURRENT_USER or {}).get('role'),''),data_label='DATA TERSIMPAN'),unsafe_allow_html=True)
    if st.session_state.get('_del_toast'): st.toast(st.session_state.pop('_del_toast'),icon=':material/check:')
    _modes=[m for m,pk in (('Harian','upload_daily'),('Bulanan','upload_monthly'),('Target','target_edit')) if has_perm(pk)] or ['Bulanan']
    if st.session_state.get('upload_mode') not in _modes: st.session_state.pop('upload_mode',None)
    with st.container(key='di_head'):
        _ti1,_ti2=st.columns([2,3.4])
        with _ti1: st.markdown('<div class="di-h"><span class="msr">upload_file</span><div><b>Data Input — Pilih Jenis Input</b><small>Kelola laporan harian (Estimasi SA), modul bulanan, atau tentukan target perusahaan bulanan per cabang.</small></div></div>',unsafe_allow_html=True)
        with _ti2:
            upload_mode=st.radio('Pilih Jenis Input',_modes,horizontal=True,key='upload_mode',label_visibility='collapsed',format_func=lambda m:{'Harian':':material/schedule: 1. Laporan Harian (Estimasi SA)','Bulanan':':material/inventory_2: 2. Laporan Bulanan (Akhir Bulan)','Target':':material/track_changes: 3. Input Target Perusahaan'}.get(m,m),
                help='Harian: laporan Estimasi Service Advisor per hari. Bulanan: MaintainPKB, Service Penjualan, Billing Part/Service, Consumable, Issue To PKB, OPL, Pelanggan.')
    if upload_mode=='Target':
        _tg_all=load_targets()
        _ref=pd.Timestamp(datetime.date.today())
        try:
            _mf=[pd.to_datetime(r.get('period_end'),errors='coerce') for r in _load_manifest()]; _mf=[x for x in _mf if pd.notna(x)]
            if _mf: _ref=max(_mf)
        except Exception: pass
        _mopts=sorted({f'{(_ref+pd.DateOffset(months=k)).year}-{(_ref+pd.DateOffset(months=k)).month:02d}' for k in range(-11,13)}|set(_tg_all.keys()))
        _fy=st.session_state.get('flt_year'); _fmn=st.session_state.get('flt_month')
        _dk=f'{int(_fy)}-{MONTH_NAMES.index(_fmn)+1:02d}' if isinstance(_fy,int) and _fmn in MONTH_NAMES else f'{_ref.year}-{_ref.month:02d}'
        _tl1,_tl2=st.columns([3,1.4])
        with _tl1: st.markdown(RK.head_card('Input Target Perusahaan Bulanan (Manual per Cabang)','Tentukan target bulanan untuk <b>UE, Jasa, Part, Bahan, dan Revenue</b> masing-masing cabang.','track_changes','red'),unsafe_allow_html=True)
        with _tl2: pk_=st.selectbox('Pilih Bulan Target',_mopts,index=(_mopts.index(_dk) if _dk in _mopts else len(_mopts)//2),format_func=lambda k:f'{MONTH_NAMES[int(k[5:7])-1]} {k[:4]}',key='rk_tgt_month')
        _cur_t=_tg_all.get(pk_,{}) or {}
        _ebr=[b for b in BRANCHES if b in user_branches()]
        _auto=st.checkbox('Revenue otomatis = Jasa + Part + Bahan',value=True,key=f'rk_tgt_auto_{pk_}')
        _vals={}
        for _i in range(0,len(_ebr),2):
            _cc=st.columns(2)
            for _col,b in zip(_cc,_ebr[_i:_i+2]):
                with _col:
                    with st.container(border=True):
                        st.markdown(f'<div class="rk-top"><b style="font-size:13.5px">{_html.escape(b.upper())}</b><span class="rk-badge red">Target Bulan {pk_}</span></div>',unsafe_allow_html=True)
                        _sv=_cur_t.get(b,{}) or {}
                        _k=lambda m,b=b:f'rk_tgt_{pk_}_{b}_{m}'
                        _x1,_x2=st.columns(2)
                        with _x1: _ue=st.number_input('Target UE (unit)',min_value=0,step=1,value=int(float(_sv.get('UE',0) or 0)),key=_k('UE'))
                        with _x2: _js=st.number_input('Target Jasa (Rp)',min_value=0,step=1000000,value=int(float(_sv.get('Jasa',0) or 0)),key=_k('Jasa'))
                        _x3,_x4=st.columns(2)
                        with _x3: _pt=st.number_input('Target Part (Rp)',min_value=0,step=1000000,value=int(float(_sv.get('Part',0) or 0)),key=_k('Part'))
                        with _x4: _bh=st.number_input('Target Bahan (Rp)',min_value=0,step=1000000,value=int(float(_sv.get('Bahan',0) or 0)),key=_k('Bahan'))
                        if _auto: st.number_input('Target Total Revenue (Rp) — otomatis',value=int(_js+_pt+_bh),disabled=True,key=f'{_k("Rev")}_auto_{int(_js+_pt+_bh)}'); _rv=int(_js+_pt+_bh)
                        else: _rv=st.number_input('Target Total Revenue (Rp)',min_value=0,step=1000000,value=int(float(_sv.get('Revenue',0) or 0)),key=_k('Revenue'))
                        _vals[b]={'UE':float(_ue),'Jasa':float(_js),'Part':float(_pt),'Bahan':float(_bh),'Revenue':float(_rv)}
        if _vals:
            _sm={m:sum(v[m] for v in _vals.values()) for m in TARGET_METRICS}
            st.markdown(RK.dark_banner(f'<b>TOTAL TARGET KONSOLIDASI ({pk_})</b> &nbsp;•&nbsp; UE <span class="c1">{_sm["UE"]:,.0f}</span> &nbsp;•&nbsp; Jasa <span class="c3">{RK.fmt_jt(_sm["Jasa"])}</span> &nbsp;•&nbsp; Part <span class="c1">{RK.fmt_jt(_sm["Part"])}</span> &nbsp;•&nbsp; Bahan <span class="c2">{RK.fmt_jt(_sm["Bahan"])}</span> &nbsp;•&nbsp; Revenue <span class="c3">{RK.fmt_jt(_sm["Revenue"])}</span>','Jumlah seluruh cabang'),unsafe_allow_html=True)
            _s1,_s2=st.columns(2)
            def _save_t(pk=pk_,ebr=tuple(_ebr),auto=_auto,cur=_cur_t):
                # baca langsung dari session_state (nilai terbaru yang diketik), bukan dari variabel run sebelumnya
                _t=load_targets(); _out={}
                for b in ebr:
                    _sv=cur.get(b,{}) or {}
                    g=lambda m:float(st.session_state.get(f'rk_tgt_{pk}_{b}_{m}',_sv.get(m,0)) or 0)
                    v={'UE':g('UE'),'Jasa':g('Jasa'),'Part':g('Part'),'Bahan':g('Bahan')}
                    v['Revenue']=(v['Jasa']+v['Part']+v['Bahan']) if auto else g('Revenue')
                    _out[b]=v
                _t.setdefault(pk,{}).update(_out); save_targets(_t); st.session_state['_tgt_saved']=pk
            _s1.button(':material/save: Simpan Target Perusahaan',type='primary',key='rk_tgt_save',on_click=_save_t,use_container_width=True)
            _py,_pm=int(pk_[:4]),int(pk_[5:7]); _prev=f'{_py-1}-12' if _pm==1 else f'{_py}-{_pm-1:02d}'
            def _copy_t():
                _t=load_targets()
                if _prev in _t:
                    _t[pk_]=json.loads(json.dumps(_t[_prev])); save_targets(_t)
                    for _kk in [x for x in st.session_state.keys() if str(x).startswith(f'rk_tgt_{pk_}_')]: st.session_state.pop(_kk,None)
                    st.session_state['_tgt_saved']=pk_
            _s2.button(f':material/content_copy: Salin dari {_prev}',key='rk_tgt_copy',on_click=_copy_t,disabled=_prev not in _tg_all,use_container_width=True)
            if st.session_state.pop('_tgt_saved',None): st.success('Target tersimpan. Hasil pencapaian otomatis diperbarui di halaman Target & Pencapaian.')
        else: st.info('Akun Anda belum memiliki cabang yang boleh diedit.')
        st.stop()
    upload_branch=st.session_state.get('selected_branch','Semua Cabang')
    if upload_branch=='Semua Cabang':
        upload_branch=(st.selectbox(':material/apartment: Cabang untuk Upload',BRANCHES,key='upload_branch_selector') if BRANCHES else 'Kharisma Makassar')
    with st.expander(':material/info: Petunjuk & cabang upload',expanded=False):
        st.markdown(f'<div class="upload-note"><span class="msr">push_pin</span> Cabang upload aktif: <b>{_html.escape(str(upload_branch))}</b> • Format: <b>Excel (.xlsx / .xls)</b>. Pilih banyak file sekaligus untuk jenis laporan yang sama. File identik boleh digunakan di cabang berbeda.</div>',unsafe_allow_html=True)
        st.markdown('<div class="upload-note"><span class="msr">warning</span> Cabang <b>dikenali otomatis dari isi laporan</b> bila header memuat nama cabang sistem terpisah (mis. Ambon). Selain itu file disimpan ke cabang yang dipilih di sidebar / selector. Laporan gabungan Makassar-Gowa-Kendari tetap dipilah otomatis per baris.</div>',unsafe_allow_html=True)
        if upload_mode=='Harian': st.markdown('<div class="upload-note"><span class="msr">today</span> <b>Laporan harian</b> = Estimasi Service Advisor yang berisi <b>transaksi hari itu saja</b>. Tanggal dibaca dari header laporan, boleh upload banyak hari sekaligus. Revenue & UE murni dari laporan ini (tanggal PKB selesai); hari yang belum diupload belum terhitung.</div>',unsafe_allow_html=True)
        else: st.markdown('<div class="upload-note"><span class="msr">calendar_month</span> <b>Laporan bulanan</b> = MaintainPKB, Service Penjualan, Billing Part, Billing Service, Consumable, Issue To PKB, OPL dan Pelanggan. Estimasi Service Advisor diupload lewat mode <b>Harian</b>.</div>',unsafe_allow_html=True)
    _card=None
    if upload_mode=='Harian':
        _specs=[x for x in FILE_SPECS if x[0]=='est']
        _card=st.container(key='up_est')
        with _card:
            _u1,_u2,_u3=st.columns([3.4,1.7,1.6])
            _u1.markdown('<div class="up-h"><span class="msr">upload</span><div><b>Input & Upload Laporan Estimasi SA Harian</b><small>Mendukung upload berkas Excel atau input langsung transaksi baru tanpa file.</small></div></div>',unsafe_allow_html=True)
            with _u2:
                if has_perm('upload_daily') and st.button(':material/add: Tambah Transaksi Manual',type='primary',key='me_open',use_container_width=True): _manual_est_dialog(upload_branch)
            with _u3:
                _tp=os.path.join(os.path.dirname(os.path.abspath(__file__)),'Template_Unit_Service.xlsx')
                if os.path.exists(_tp):
                    with open(_tp,'rb') as _fh: st.download_button(':material/download: Download Template Excel SA',_fh.read(),'Template_Unit_Service.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='dl_tpl_sa',use_container_width=True)
            st.markdown('<div class="rk-div"></div>',unsafe_allow_html=True)
    else:
        _specs=[x for x in FILE_SPECS if x[0]!='est']
    for i in range(0,len(_specs),2):
        cols=[_card] if _card is not None else st.columns(2)
        for col,(rt,label,_) in zip(cols,_specs[i:i+2]):
            with col:
                if rt=='est': label='📅 '+label.split(' ',1)[1]+' (Harian)'
                _mr=sorted(_historical_records(rt),key=lambda r:str(r.get('uploaded_at','')))
                _mn=REPORT_NAMES.get(rt,label)+(' (Harian)' if rt=='est' else '')
                _mf=(f"File: <code>{_html.escape(str(_mr[-1].get('original_filename','')))}</code> • {_html.escape(str(_mr[-1].get('period_label','')))} • {len(_mr)} file tersimpan" if _mr else 'Belum ada file tersimpan untuk jenis laporan ini.')
                if _card is None: st.markdown(RK.module_card_head(_mn,bool(_mr),_mf),unsafe_allow_html=True)
                files=st.file_uploader(label,type=['xlsx','xls'],accept_multiple_files=True,key=f'{rt}_multi_uploader',label_visibility='collapsed')
                if files:
                    # Hasil upload disimpan per sesi supaya pesan 'berhasil' tidak berubah menjadi 'file identik sudah tersimpan'
                    # saat halaman rerun (mis. setelah mengganti widget lain). Cache dipakai hanya bila semua file masih ada di histori.
                    _shas=tuple(_file_hash(f.getvalue()) for f in files)
                    _ck=(rt,_shas,upload_branch); _cache=st.session_state.setdefault('_upload_results',{})
                    _present={(r.get('report_type'),r.get('sha256')) for r in _load_manifest()}
                    if _ck in _cache and all((rt,h) in _present for h in _shas): results=_cache[_ck]
                    else:
                        results=_store_uploaded(rt,files,upload_branch); _cache[_ck]=results
                    for status,name,msg in results:
                        if status=='ok': st.success(f':material/check: {name} • {msg}',icon=':material/save:')
                        elif status=='warn': st.warning(f'{name} • {msg}',icon=':material/warning:')
                        elif status=='skip': st.info(f':material/undo: {name} • {msg}',icon='ℹ')
                        else: st.error(f':material/close: {name} • {msg}')
    if upload_mode=='Harian':
        if st.session_state.pop('_me_saved',None): st.success('Transaksi manual tersimpan dan sudah dihitung di Dashboard.',icon=':material/check:')
        _mrows=load_manual_est()
        with st.expander(f':material/edit_note: Transaksi manual tersimpan ({len(_mrows)})',expanded=False):
            if not _mrows: st.caption('Belum ada transaksi manual. Gunakan tombol Tambah Transaksi Manual di atas.')
            else:
                _mt=[[_html.escape(str(r.get('tanggal',''))),_html.escape(str(r.get('cabang',''))),_html.escape(str(r.get('sa_nama') or r.get('sa_kode',''))),_html.escape(str(r.get('pkb',''))),fmt(float(r.get('act_jasa',0))+float(r.get('act_part',0))+float(r.get('act_bahan',0)))] for r in _mrows]
                st.markdown(RK.table_html(['Tanggal','Cabang','SA','No PKB','Revenue (Aktual)'],_mt,num_cols=(4,)),unsafe_allow_html=True)
                if has_perm('upload_daily'):
                    _mch={r['id']:f"{r.get('tanggal','')} • {r.get('pkb','')} • {r.get('sa_nama') or r.get('sa_kode','')}" for r in _mrows}
                    _msel=st.multiselect('Pilih transaksi manual yang ingin dihapus',list(_mch),format_func=lambda x:_mch[x],key='me_del_sel')
                    if st.button(':material/delete: Hapus terpilih',disabled=not _msel,key='me_del_btn'):
                        save_manual_est([r for r in _mrows if r['id'] not in _msel]); st.session_state.pop('me_del_sel',None); st.rerun()
        # Status kelengkapan laporan harian: tanggal yang sudah terupload dan tanggal kosong di antaranya.
        _days=[]
        for _r in _load_manifest():
            if _r.get('report_type')!='est': continue
            _a=pd.to_datetime(_r.get('period_start'),errors='coerce'); _b=pd.to_datetime(_r.get('period_end'),errors='coerce')
            if pd.isna(_a) or pd.isna(_b) or (_b-_a).days>=LONG_REPORT_DAYS: continue
            _days.extend(pd.date_range(_a.normalize(),_b.normalize(),freq='D'))
        st.markdown(RK.head_card('Status Laporan Harian','Kelengkapan tanggal Estimasi SA yang sudah diupload pada bulan berjalan','event_available','green'),unsafe_allow_html=True)
        if not _days: st.info('Belum ada laporan Estimasi SA harian yang diupload.')
        else:
            _ds=pd.Series(sorted(set(_days))); _last=_ds.max(); _mstart=_last.replace(day=1)
            _inm=_ds[_ds>=_mstart]; _full=pd.date_range(_mstart,_last,freq='D'); _miss=[d for d in _full if d not in set(_inm)]
            c1,c2,c3=st.columns(3); c1.metric('Hari terupload bulan ini',f'{len(_inm)} / {len(_full)}'); c2.metric('Tanggal terakhir',_last.strftime('%d/%m/%Y')); c3.metric('Tanggal kosong',len(_miss))
            if _miss: st.warning('Tanggal belum diupload (sampai tanggal terakhir): '+', '.join(d.strftime('%d/%m') for d in _miss)+'. Revenue & UE hari kosong belum terhitung (acuan murni Laporan Estimasi SA) sampai laporannya diupload.')
            else: st.success('Semua tanggal bulan ini sudah ada, tanpa hari kosong.')
    records=_historical_records()
    total_size=sum(int(r.get('size_bytes',0) or 0) for r in records)
    st.markdown(f'<div class="upload-progress"><b>{len(records)} file histori tersimpan</b><span><span class="msr">save</span> {total_size/1024/1024:.1f} MB • permanen di komputer ini</span></div>',unsafe_allow_html=True)
    st.markdown(RK.head_card('Registry Laporan','File disimpan terpisah sehingga laporan bulan berikutnya tidak menimpa histori sebelumnya.','folder_open','blue'),unsafe_allow_html=True)
    if records:
        reg=pd.DataFrame([{'Cabang':r.get('branch',''),'Jenis':REPORT_NAMES.get(r.get('report_type'),r.get('report_type')),'Periode':r.get('period_label','Tidak terdeteksi'),'File':r.get('original_filename'),'Upload':r.get('uploaded_at'),'Ukuran KB':round(int(r.get('size_bytes',0))/1024,1)} for r in records])
        st.dataframe(reg,use_container_width=True,hide_index=True)
        choices={r['id']:f"{r.get('branch','')} • {REPORT_NAMES.get(r.get('report_type'),r.get('report_type'))} • {r.get('period_label','?')} • {r.get('original_filename')}" for r in records}
        with st.expander(':material/warning: Kelola histori / hapus file tertentu'):
            selected=st.multiselect('Pilih file histori yang ingin dihapus',list(choices),format_func=lambda x:choices[x])
            if st.button(':material/delete: Hapus file histori terpilih',disabled=not selected):
                confirm_delete_dialog(list(selected),choices,False)
            if st.button(':material/delete: Hapus SEMUA histori',type='secondary'):
                confirm_delete_dialog([r['id'] for r in records],choices,True)
    else: st.info('Belum ada laporan tersimpan. Upload MaintainPKB dan RptServicePenjualan untuk membuka analisa utama.')
    st.session_state['est_ref_on']=st.checkbox('Gunakan Laporan Estimasi Service Advisor sebagai acuan revenue bulanan per cabang (bulan mengikuti periode header laporan, bukan tanggal PKB)',value=st.session_state.get('est_ref_on',True))
    st.info(f'Database histori lokal: `{HIST_DIR}`. File Excel asli tetap tersimpan dan dapat ditambah kapan saja.',icon=':material/inventory_2:')


active_branch=st.session_state.get('selected_branch','Semua Cabang')
@st.cache_resource(show_spinner=False)
def _pipe_store(): return {}
_PIPE_CACHE=_pipe_store()   # bertahan antar rerun/pindah halaman (sebelumnya dibuat ulang tiap klik)
def _data_sig():
    sig=[]
    for p in (MANIFEST_PATH,HIST_DIR,MANUAL_EST_PATH):
        try: stt=os.stat(p); sig.append((stt.st_mtime_ns,stt.st_size))
        except Exception: sig.append(None)
    return tuple(sig)+(tuple(map(str,BRANCHES)),)
_PIPE_KEYS=['pkb','sales','parts','issue','cons','opl','cust','sales_detail','est_all','est_cov','master_raw','master_all','est_excluded','est_stats','master','qc_no_sales','EST_REF_ACTIVE','EST_REF_BRANCHES','error','dedup']
def _pipe_compute(active_branch,est_on):
    pkb,sales,parts,issue,cons,opl,cust,sales_detail=aggregate_historical(active_branch)
    est_all=aggregate_estimasi(active_branch)
    est_cov=aggregate_estimasi('Semua Cabang') if active_branch!='Semua Cabang' else est_all   # cakupan laporan Estimasi selalu dari SEMUA cabang
    if pkb.empty and sales.empty and not est_all.empty:
        # Mode harian murni: belum ada laporan bulanan sama sekali -> PKB & penjualan dibentuk dari laporan Estimasi SA (nilai Actual per PKB).
        _e0=est_all.copy()
        pkb=pd.DataFrame({'No Polisi':'','PKB':_e0['PKB'],'Tanggal':pd.to_datetime(_e0['Tanggal'],errors='coerce'),'Delivery Datetime':'','PKB Status':'Closed','Customer':'','No Tlp/HP':'',
            'Operation Description':'','Vendor':'','Booking Type':'','Ch.To':'','Service Advisor':_e0['Service Advisor'],'Mechanic':'','Tahun Produksi':'','Type Kendaraan':'','Age':'','Status':'',
            'Range Jam':'','Periode Jual':'','Jenis Jasa Pekerjaan':'','PIC Foreman':'','Release Date':'','Job Type Operation':'','Service Category':'','Equipment No':'','Decision Maker':'',
            'Decision Maker Phone':'','Contact Person':'','Contact Person Phone':'','Customer Address':'','Email':'','Kilometer':'','Manual PKB No':'','Jenis Pekerjaan':'','UE':1,'Cabang':_e0['Cabang']})
        sales=pd.DataFrame({'Cabang':_e0['Cabang'],'PKB':_e0['PKB'],'Tanggal Sales':pd.to_datetime(_e0['Tanggal Periode']),'Invoice':'','Customer Sales':'','Jenis Pembayaran':'','Type Kendaraan Sales':'',
            'Jasa':_e0['Act Jasa'].astype(float),'Sparepart':_e0['Act Part'].astype(float),'Bahan':_e0['Act Bahan'].astype(float),'Sales Row Count':1,
            'Revenue':(_e0['Act Jasa']+_e0['Act Part']+_e0['Act Bahan']).astype(float)})
    error=None
    if not (pkb.empty or sales.empty):
        try:
            master_raw=build_master(pkb,sales,opl,cust,active_branch)
            master_all,est_excluded,est_stats=apply_estimasi_reference(master_raw,est_all,active=est_on,est_cov=est_cov)
            # UE tanpa penjualan (PKB cancelled / created / released / completed, dst.) TIDAK dihitung sebagai UE maupun revenue.
            master=master_all[master_all['Match Status']=='MATCH'].reset_index(drop=True)
            # PKB tanpa penjualan: SELALU didaftar untuk evaluasi (Quality), termasuk yang jatuh di bulan berlaporan Estimasi SA
            # tetapi tidak tercantum di laporan (yang dikeluarkan dari master_all). Tidak pernah dihitung sebagai UE/revenue.
            _ns=[master_all[master_all['Match Status']!='MATCH']]
            if isinstance(est_excluded,pd.DataFrame) and not est_excluded.empty: _ns.append(est_excluded[est_excluded['Match Status']!='MATCH'])
            qc_no_sales=pd.concat(_ns,ignore_index=True).drop_duplicates(['Cabang','PKB']).reset_index(drop=True)
            EST_REF_ACTIVE=est_on and not est_cov.empty
            EST_REF_BRANCHES=set()
            if EST_REF_ACTIVE and 'Cabang' in est_cov.columns:
                EST_REF_BRANCHES=set(est_cov['Cabang'].astype(str))
                if EST_REF_BRANCHES & SHARED_SYSTEM_BRANCHES: EST_REF_BRANCHES|=SHARED_SYSTEM_BRANCHES
        except Exception as e:
            error=str(e)
    dedup=dict(DEDUP_STATS)
    _l=locals()
    return {k:_l.get(k) for k in _PIPE_KEYS}
def _pipe(active_branch,est_on):
    key=(active_branch,bool(est_on),_data_sig())
    hit=_PIPE_CACHE.get(key)
    if hit is None:
        hit=_pipe_compute(active_branch,bool(est_on))
        while len(_PIPE_CACHE)>=8: _PIPE_CACHE.pop(next(iter(_PIPE_CACHE)))
        _PIPE_CACHE[key]=hit
    DEDUP_STATS.clear(); DEDUP_STATS.update(hit.get('dedup') or {})
    return {k:(v.copy() if isinstance(v,(pd.DataFrame,set,dict,list)) else v) for k,v in hit.items()}
_P=_pipe(active_branch,st.session_state.get('est_ref_on',True))
pkb,sales,parts,issue,cons,opl,cust,sales_detail,est_all,est_cov=[_P[k] for k in ['pkb','sales','parts','issue','cons','opl','cust','sales_detail','est_all','est_cov']]
if pkb.empty or sales.empty:
    if page == '📥 Data Input':
        st.info(f'Belum ada data utama untuk cabang **{active_branch}**. Upload minimal MaintainPKB dan RptServicePenjualan pada cabang tersebut.')
        st.stop()
    if page == '🏢 Branch Management':
        st.stop()
    st.warning(f'Belum ada data utama untuk cabang **{active_branch}**. Buka menu **Data Input** untuk mengupload laporan Excel.')
    st.stop()
if _P.get('error'):
    st.error(f"Gagal membaca data histori: {_P['error']}"); st.stop()
master_raw,master_all,est_excluded,est_stats,master,qc_no_sales,EST_REF_ACTIVE,EST_REF_BRANCHES=[_P[k] for k in ['master_raw','master_all','est_excluded','est_stats','master','qc_no_sales','EST_REF_ACTIVE','EST_REF_BRANCHES']]

# ===== FILTER WAKTU GLOBAL (Tahun / Bulan / Hari dari sidebar) =====
# Salinan tanpa filter dipakai halaman Target Perusahaan (punya pemilih periode sendiri).
UNF={'master':master,'parts':parts,'issue':issue,'cons':cons}
TF_Y=st.session_state.get('flt_year','Semua'); TF_M=st.session_state.get('flt_month','Semua'); TF_D=st.session_state.get('flt_day','Semua')
TF_ACTIVE=any(v!='Semua' for v in (TF_Y,TF_M,TF_D))
def _tf_mask(dates):
    dt=pd.to_datetime(dates,errors='coerce'); ok=pd.Series(True,index=dt.index)
    if TF_Y!='Semua': ok&=(dt.dt.year==int(TF_Y))
    if TF_M!='Semua': ok&=(dt.dt.month==MONTH_NAMES.index(TF_M)+1)
    if TF_D!='Semua': ok&=(dt.dt.day==int(TF_D))
    return ok.fillna(False)
def tf(df,col='Tanggal'):
    if not TF_ACTIVE or not isinstance(df,pd.DataFrame) or df.empty or col not in df.columns: return df
    return df[_tf_mask(df[col]).values].copy()
if TF_ACTIVE:
    master=tf(master); master_all=tf(master_all); qc_no_sales=tf(qc_no_sales); est_excluded=tf(est_excluded)
    parts=tf(parts); issue=tf(issue); cons=tf(cons); sales=tf(sales,'Tanggal Sales')
    if isinstance(opl,pd.DataFrame) and not opl.empty and 'PKB' in opl.columns and isinstance(master_all,pd.DataFrame) and 'PKB' in master_all.columns:
        opl=opl[opl['PKB'].isin(set(master_all['PKB']))].copy()
def add_period_fields(df):
    df=df.copy()
    if 'Tanggal' in df.columns:
        df['Tanggal']=pd.to_datetime(df['Tanggal'],errors='coerce')
        df['Periode Bulan']=df['Tanggal'].dt.to_period('M').astype('string')
        df['Periode Bulan Label']=df['Tanggal'].dt.strftime('%b %Y')
        df['Periode Kuartal']=df['Tanggal'].dt.to_period('Q').astype('string')
        df['Periode Tahun']=df['Tanggal'].dt.strftime('%Y')
    return df

def _growth(curr, prev):
    if prev is None or pd.isna(prev) or float(prev)==0:
        return np.nan if curr in [0, 0.0] else np.inf
    return (float(curr)/float(prev)-1)*100

def _period_key_to_label(key, gran):
    return period_key_label(str(key), gran)

def build_comparative_table(df, field, gran):
    """Comparative KPI per periode tanpa double-count direct sales."""
    if df.empty or field not in df.columns:
        return pd.DataFrame()
    x=df.copy()
    x['Tanggal']=pd.to_datetime(x.get('Tanggal'),errors='coerce')
    x=x[x['Tanggal'].notna()].copy()
    if x.empty: return pd.DataFrame()

    # Master/PKB metrics.
    g=x.groupby(field, dropna=True).agg(
        UE=('UE','sum'), Revenue=('Revenue','sum'), Jasa=('Jasa','sum'),
        Sparepart=('Sparepart','sum'), Bahan=('Bahan','sum')
    ).reset_index()
    g[field]=g[field].astype(str)

    # Direct sales are aggregated independently by transaction period, then added once.
    def period_series(dates):
        dt=pd.to_datetime(dates,errors='coerce')
        if gran=='Bulan': return dt.dt.to_period('M').astype(str)
        if gran=='Kuartal': return dt.dt.to_period('Q').astype(str)
        return dt.dt.year.astype('Int64').astype(str)

    direct_part=pd.Series(dtype=float); direct_bahan=pd.Series(dtype=float)
    if isinstance(parts,pd.DataFrame) and not parts.empty and 'Tanggal' in parts.columns and 'Sales Amount' in parts.columns and 'Sumber Penjualan' in parts.columns:
        z=parts[parts['Sumber Penjualan'].eq('Billing Part')].copy()
        if not z.empty:
            z['_period']=period_series(z['Tanggal'])
            direct_part=z.groupby('_period')['Sales Amount'].sum()
    if isinstance(cons,pd.DataFrame) and not cons.empty and 'Tanggal' in cons.columns and 'Harga Net' in cons.columns:
        z=cons.copy(); z['_period']=period_series(z['Tanggal'])
        direct_bahan=z.groupby('_period')['Harga Net'].sum()

    g['Direct Part']=g[field].map(direct_part).fillna(0.0)
    g['Direct Bahan']=g[field].map(direct_bahan).fillna(0.0)
    g['Sparepart']=g['Sparepart']+g['Direct Part']
    g['Bahan']=g['Bahan']+g['Direct Bahan']
    g['Revenue']=g['Revenue']+g['Direct Part']+g['Direct Bahan']
    g['Revenue/UE']=np.where(g['UE']>0,g['Revenue']/g['UE'],0)

    if gran=='Bulan': g['_sort']=pd.PeriodIndex(g[field],freq='M')
    elif gran=='Kuartal': g['_sort']=pd.PeriodIndex(g[field],freq='Q')
    else: g['_sort']=pd.to_numeric(g[field],errors='coerce')
    g=g.sort_values('_sort').reset_index(drop=True)
    g['Label']=g[field].map(lambda x:_period_key_to_label(x,gran))

    lag_yoy={'Bulan':12,'Kuartal':4,'Tahun':1}[gran]
    for metric in ['UE','Revenue','Revenue/UE','Jasa','Sparepart','Bahan']:
        g[f'{metric} Prev']=g[metric].shift(1)
        g[f'{metric} Prev Growth %']=g.apply(lambda r:_growth(r[metric],r[f'{metric} Prev']),axis=1)
        g[f'{metric} YoY']=g[metric].shift(lag_yoy)
        g[f'{metric} YoY Growth %']=g.apply(lambda r:_growth(r[metric],r[f'{metric} YoY']),axis=1)
    g['Revenue Mix Jasa %']=np.where(g['Revenue']!=0,g['Jasa']/g['Revenue']*100,0)
    g['Revenue Mix Sparepart %']=np.where(g['Revenue']!=0,g['Sparepart']/g['Revenue']*100,0)
    g['Revenue Mix Bahan %']=np.where(g['Revenue']!=0,g['Bahan']/g['Revenue']*100,0)
    return g


def build_management_insights(df):
    """Generate executive insights from the currently selected historical data."""
    out=[]
    if df.empty: return out
    x=df.copy()
    x['Tanggal']=pd.to_datetime(x.get('Tanggal'),errors='coerce')
    x=x[x['Tanggal'].notna()].copy()
    if x.empty: return out
    for c in ['UE','Revenue','Jasa','Sparepart','Bahan','Revenue per Unit']:
        if c not in x.columns: x[c]=0
    total_ue=float(x['UE'].sum()); rev=float(x['Revenue'].sum()); rpu=rev/total_ue if total_ue else 0
    # Current period vs immediately preceding period.
    monthly=x.groupby(x['Tanggal'].dt.to_period('M')).agg(UE=('UE','sum'),Revenue=('Revenue','sum'),Jasa=('Jasa','sum'),Sparepart=('Sparepart','sum'),Bahan=('Bahan','sum')).sort_index()
    if len(monthly)>=2:
        cur=monthly.iloc[-1]; prev=monthly.iloc[-2]
        rg=_growth(cur['Revenue'],prev['Revenue']); ug=_growth(cur['UE'],prev['UE'])
        if not pd.isna(rg):
            direction='naik' if rg>=0 else 'turun'
            out.append(f"Revenue periode terakhir {direction} {abs(rg):.1f}% dibanding bulan sebelumnya ({fmt(cur['Revenue'])} vs {fmt(prev['Revenue'])}).")
        if not pd.isna(ug):
            direction='naik' if ug>=0 else 'turun'
            out.append(f"Volume UE {direction} {abs(ug):.1f}% dibanding bulan sebelumnya ({int(cur['UE']):,} vs {int(prev['UE']):,} UE).")
        cur_rpu=cur['Revenue']/cur['UE'] if cur['UE'] else 0; prev_rpu=prev['Revenue']/prev['UE'] if prev['UE'] else 0
        rgpu=_growth(cur_rpu,prev_rpu)
        if not pd.isna(rgpu):
            if rgpu < -5:
                out.append(f'<span class="msr">warning</span> Revenue/UE turun {abs(rgpu):.1f}%; kenaikan/penurunan revenue perlu dilihat bersama produktivitas per unit.')
            elif rgpu > 5:
                out.append(f'<span class="msr">trending_up</span> Revenue/UE naik {rgpu:.1f}%; kualitas monetisasi per unit membaik.')
    # Revenue mix and concentration.
    mix={c:(float(x[c].sum())/rev*100 if rev else 0) for c in ['Jasa','Sparepart','Bahan']}
    dominant=max(mix,key=mix.get)
    out.append(f"Revenue mix didominasi {dominant} sebesar {mix[dominant]:.1f}% (Jasa {mix['Jasa']:.1f}%, Sparepart {mix['Sparepart']:.1f}%, Bahan {mix['Bahan']:.1f}%).")
    # Matching quality.
    if 'Match Status' in x:
        unmatched=int((x['Match Status']!='MATCH').sum()); rate=(1-unmatched/len(x))*100 if len(x) else 0
        if unmatched:
            out.append(f'<span class="msr">warning</span> {unmatched:,} UE ({100-rate:.1f}%) belum memiliki sales yang matched; ini menjadi prioritas rekonsiliasi data.')
        else:
            out.append('<span class="msr">check_circle</span> Seluruh UE pada data terpilih sudah memiliki transaksi sales yang matched.')
    # Booking mix.
    if 'Booking Type' in x:
        bt=x['Booking Type'].map(clean).str.upper()
        booking=int((bt=='BOOKING').sum()); walk=int((bt=='WALK IN').sum()); denom=booking+walk
        if denom:
            out.append(f"Booking menyumbang {booking/denom*100:.1f}% dari UE dengan entry yang teridentifikasi; Walk In {walk/denom*100:.1f}%.")
    # Top SA with revenue per UE.
    if 'Service Advisor' in x:
        sa=x[x['Service Advisor'].map(clean)!=''].groupby('Service Advisor').agg(UE=('UE','sum'),Revenue=('Revenue','sum')).sort_values('Revenue',ascending=False)
        if not sa.empty:
            top=sa.iloc[0]; topname=sa.index[0]; top_rpu=top['Revenue']/top['UE'] if top['UE'] else 0
            out.append(f"SA dengan revenue terbesar: {topname} ({fmt(top['Revenue'])} dari {int(top['UE']):,} UE; {fmt(top_rpu)}/UE).")
    # Vehicle concentration.
    if 'Type Kendaraan' in x:
        veh=x.groupby('Type Kendaraan').agg(UE=('UE','sum'),Revenue=('Revenue','sum')).sort_values('Revenue',ascending=False)
        veh=veh[veh.index.map(clean)!='']
        if not veh.empty:
            top=veh.iloc[0]; out.append(f"Type kendaraan penyumbang revenue terbesar: {veh.index[0]} ({fmt(top['Revenue'])}, {int(top['UE']):,} UE).")
    # Anomaly flags using monthly history.
    if len(monthly)>=4:
        hist=monthly.iloc[:-1]['Revenue']; latest=float(monthly.iloc[-1]['Revenue']); mean=float(hist.mean()); sd=float(hist.std(ddof=0))
        if sd>0 and abs(latest-mean)>1.5*sd:
            out.append(f'<span class="msr">search</span> Anomali revenue terdeteksi: periode terakhir {fmt(latest)} berbeda signifikan dari pola historis (rata-rata {fmt(mean)}).')
    # Overall executive recommendation.
    if rev>0 and total_ue>0:
        out.append(f"Fokus KPI manajemen: pertahankan UE {int(total_ue):,}, dorong Revenue/UE {fmt(rpu)}, dan selesaikan gap matching sebelum menarik kesimpulan profitabilitas.")
    return out[:12]

def revenue_components(master_df, parts_df, issue_df, cons_df):
    """Satu sumber rumus revenue untuk Dashboard, Comparative, Revenue & Profit, Sparepart dan Bahan.
    Dihitung PER CABANG lalu dijumlahkan, sehingga total 'Semua Cabang' selalu = jumlah revenue masing-masing cabang
    (pilihan sumber Bahan/Part detail vs Service Penjualan ditentukan per cabang, bukan global)."""
    def _cab(df):
        return set(df['Cabang'].fillna('').astype(str)) if isinstance(df,pd.DataFrame) and not df.empty and 'Cabang' in df.columns else set()
    cabs=_cab(master_df)|_cab(parts_df)|_cab(issue_df)|_cab(cons_df)
    def _ea(c): return bool(EST_REF_ACTIVE and (c in EST_REF_BRANCHES or not EST_REF_BRANCHES))
    if len(cabs)<=1: return _revenue_components_single(master_df,parts_df,issue_df,cons_df,_ea(next(iter(cabs),'')))
    def _sub(df,c):
        return df[df['Cabang'].fillna('').astype(str)==c] if isinstance(df,pd.DataFrame) and not df.empty and 'Cabang' in df.columns else df
    tot=None
    for c in sorted(cabs):
        r=_revenue_components_single(_sub(master_df,c),_sub(parts_df,c),_sub(issue_df,c),_sub(cons_df,c),_ea(c))
        tot=r if tot is None else {k:tot[k]+r[k] for k in tot}
    return tot

def revenue_components_by_branch(master_df, parts_df, issue_df, cons_df, branches=None):
    """Tabel komponen revenue per cabang. Semua cabang di `branches` (mis. BRANCHES) SELALU tampil, termasuk yang belum
    punya data (nilai 0) dan cabang baru. Baris TOTAL = jumlah baris cabang = angka Dashboard 'Semua Cabang'."""
    def _cab(df):
        return set(df['Cabang'].fillna('').astype(str)) if isinstance(df,pd.DataFrame) and not df.empty and 'Cabang' in df.columns else set()
    seen=_cab(master_df)|_cab(parts_df)|_cab(issue_df)|_cab(cons_df)
    order=list(branches or [])+[c for c in sorted(seen) if c and c not in set(branches or [])]
    def _sub(df,c):
        return df[df['Cabang'].fillna('').astype(str)==c] if isinstance(df,pd.DataFrame) and not df.empty and 'Cabang' in df.columns else (df.iloc[0:0] if isinstance(df,pd.DataFrame) else df)
    rows=[]
    for c in order:
        m=_sub(master_df,c); r=_revenue_components_single(m,_sub(parts_df,c),_sub(issue_df,c),_sub(cons_df,c),bool(EST_REF_ACTIVE and (c in EST_REF_BRANCHES or not EST_REF_BRANCHES)))
        ue=int(m['UE'].sum()) if isinstance(m,pd.DataFrame) and 'UE' in m.columns else 0
        rows.append({'Cabang':c,'UE':ue,'Jasa':r['jasa'],'Sparepart (dengan PKB)':r['service_sparepart'],'Bahan (dengan PKB)':r['service_bahan'],
                     'Part tanpa PKB':r['direct_part'],'Bahan tanpa PKB':r['direct_bahan'],'Total Revenue':r['total_revenue']})
    out=pd.DataFrame(rows)
    if not out.empty:
        t={'Cabang':'TOTAL SEMUA CABANG'}
        for c in out.columns[1:]: t[c]=out[c].sum()
        out=pd.concat([out,pd.DataFrame([t])],ignore_index=True)
        out.insert(out.columns.get_loc('Total Revenue')+1,'Revenue/UE',np.where(out['UE']>0,out['Total Revenue']/out['UE'].where(out['UE']>0,1),0.0))
    return out

def daily_components(master_df, parts_df, issue_df, cons_df):
    """Realisasi PER TANGGAL: UE, Jasa, Part, Bahan, Revenue. Memakai aturan sumber yang sama dengan revenue_components
    (per cabang: acuan Estimasi aktif / detail Billing Service & Issue To PKB bila ada), sehingga jumlah harian = angka Dashboard."""
    def _cabs(df):
        return set(df['Cabang'].fillna('').astype(str)) if isinstance(df,pd.DataFrame) and not df.empty and 'Cabang' in df.columns else set()
    def _sub(df,c):
        if not isinstance(df,pd.DataFrame) or df.empty: return pd.DataFrame()
        return df[df['Cabang'].fillna('').astype(str)==c] if 'Cabang' in df.columns else df
    def _by(df,col,val):
        if not isinstance(df,pd.DataFrame) or df.empty or col not in df.columns or val not in df.columns: return pd.Series(dtype=float)
        d=pd.to_datetime(df[col],errors='coerce').dt.normalize()
        return df.assign(_d=d).dropna(subset=['_d']).groupby('_d')[val].sum().astype(float)
    cabs=_cabs(master_df)|_cabs(parts_df)|_cabs(issue_df)|_cabs(cons_df) or {''}
    acc=[]
    for c in cabs:
        m=_sub(master_df,c) if len(cabs)>1 or c else master_df; p=_sub(parts_df,c) if len(cabs)>1 or c else parts_df
        i=_sub(issue_df,c) if len(cabs)>1 or c else issue_df; k=_sub(cons_df,c) if len(cabs)>1 or c else cons_df
        est_a=bool(EST_REF_ACTIVE and (c in EST_REF_BRANCHES or not EST_REF_BRANCHES))
        if isinstance(p,pd.DataFrame) and not p.empty and 'Sales Amount' in p.columns:
            if 'Sumber Penjualan' in p.columns:
                p_det=p[p['Sumber Penjualan'].eq('Billing Service')]; p_dir=p[p['Sumber Penjualan'].eq('Billing Part')]
            else: p_det=p.iloc[0:0]; p_dir=p
        else: p_det=p_dir=pd.DataFrame()
        use_pdet=(float(p_det['Sales Amount'].sum()) if len(p_det) else 0.0)>0 and not est_a
        use_iss=(float(i['Harga Net'].sum()) if isinstance(i,pd.DataFrame) and not i.empty and 'Harga Net' in i.columns else 0.0)>0 and not est_a
        ue=_by(m,'Tanggal','UE'); jasa=_by(m,'Tanggal','Jasa')
        part=(_by(p_det,'Tanggal','Sales Amount') if use_pdet else _by(m,'Tanggal','Sparepart')).add(_by(p_dir,'Tanggal','Sales Amount'),fill_value=0)
        bahan=(_by(i,'Tanggal','Harga Net') if use_iss else _by(m,'Tanggal','Bahan')).add(_by(k,'Tanggal','Harga Net'),fill_value=0)
        acc.append(pd.DataFrame({'UE':ue,'Jasa':jasa,'Part':part,'Bahan':bahan}).fillna(0.0))
    if not acc: return pd.DataFrame(columns=['UE','Jasa','Part','Bahan','Revenue'])
    out=pd.concat(acc).groupby(level=0).sum()
    out['Revenue']=out['Jasa']+out['Part']+out['Bahan']
    out.index.name='Tanggal'
    return out.sort_index()

def _revenue_components_single(master_df, parts_df, issue_df, cons_df, est_active=None):
    est_active=EST_REF_ACTIVE if est_active is None else est_active
    x=master_df.copy()
    jasa=float(x['Jasa'].sum()) if 'Jasa' in x else 0.0
    service_spare=float(x['Sparepart'].sum()) if 'Sparepart' in x else 0.0
    service_bahan=float(x['Bahan'].sum()) if 'Bahan' in x else 0.0
    direct_part=0.0
    service_part_detail=0.0
    if isinstance(parts_df,pd.DataFrame) and not parts_df.empty and 'Sales Amount' in parts_df:
        service_part_detail=float(parts_df.loc[parts_df.get('Sumber Penjualan','').eq('Billing Service'),'Sales Amount'].sum()) if 'Sumber Penjualan' in parts_df else 0.0
        direct_part=float(parts_df.loc[parts_df.get('Sumber Penjualan','').eq('Billing Part'),'Sales Amount'].sum()) if 'Sumber Penjualan' in parts_df else float(parts_df['Sales Amount'].sum())
    issue_bahan=float(issue_df['Harga Net'].sum()) if isinstance(issue_df,pd.DataFrame) and not issue_df.empty and 'Harga Net' in issue_df else 0.0
    direct_bahan=float(cons_df['Harga Net'].sum()) if isinstance(cons_df,pd.DataFrame) and not cons_df.empty and 'Harga Net' in cons_df else 0.0
    # Detail reports are the authoritative breakdown for parts/materials when available.
    if service_part_detail>0 and not est_active: service_spare=service_part_detail
    if issue_bahan>0 and not est_active: service_bahan=issue_bahan
    total_spare=service_spare+direct_part
    total_bahan=service_bahan+direct_bahan
    revenue_core=jasa+service_spare+service_bahan
    total_revenue=revenue_core+direct_part+direct_bahan
    return {'jasa':jasa,'service_sparepart':service_spare,'direct_part':direct_part,'total_sparepart':total_spare,
            'service_bahan':service_bahan,'direct_bahan':direct_bahan,'total_bahan':total_bahan,
            'revenue_core':revenue_core,'total_revenue':total_revenue}

def enrich_comparative_revenue(df):
    """Tambahkan revenue total/direct sales per tanggal agar Comparative Analytics tidak hanya membaca PKB."""
    x=df.copy()
    if 'Tanggal' not in x.columns: return x
    x['Tanggal']=pd.to_datetime(x['Tanggal'],errors='coerce')
    x['_date']=x['Tanggal'].dt.normalize()
    for col in ['Direct Part','Direct Bahan']:
        x[col]=0.0
    if isinstance(parts,pd.DataFrame) and not parts.empty and 'Tanggal' in parts.columns and 'Sales Amount' in parts.columns and 'Sumber Penjualan' in parts.columns:
        z=parts[parts['Sumber Penjualan']=='Billing Part'].copy(); z['Tanggal']=pd.to_datetime(z['Tanggal'],errors='coerce'); z=z[z['Tanggal'].notna()]
        if not z.empty:
            g=z.groupby(z['Tanggal'].dt.normalize())['Sales Amount'].sum()
            x['Direct Part']=x['_date'].map(g).fillna(0.0)
    if isinstance(cons,pd.DataFrame) and not cons.empty and 'Tanggal' in cons.columns and 'Harga Net' in cons.columns:
        z=cons.copy(); z['Tanggal']=pd.to_datetime(z['Tanggal'],errors='coerce'); z=z[z['Tanggal'].notna()]
        if not z.empty:
            g=z.groupby(z['Tanggal'].dt.normalize())['Harga Net'].sum()
            x['Direct Bahan']=x['_date'].map(g).fillna(0.0)
    x['Sparepart Total']=x.get('Sparepart',0)+x['Direct Part']
    x['Bahan Total']=x.get('Bahan',0)+x['Direct Bahan']
    x['Revenue Total']=x.get('Revenue',0)+x['Direct Part']+x['Direct Bahan']
    x=x.drop(columns=['_date'],errors='ignore')
    return x

def period_key_label(period_key, granularity):
    try:
        if granularity=='Bulan':
            return pd.Period(period_key,freq='M').strftime('%b %Y')
        if granularity=='Kuartal':
            return str(period_key).replace('Q',' Q')
        return str(period_key)
    except Exception:
        return str(period_key)

if page == '📥 Data Input':
    if st.session_state.get('upload_mode')=='Bulanan':
        _rbm=revenue_components_by_branch(master,parts,issue,cons,BRANCHES)
        if not _rbm.empty:
            _tt=_rbm[_rbm['Cabang']=='TOTAL SEMUA CABANG'].iloc[0]; _body=_rbm[_rbm['Cabang']!='TOTAL SEMUA CABANG']
            _core_t=float(_tt['Total Revenue']-_tt['Part tanpa PKB']-_tt['Bahan tanpa PKB'])
            _pl=(f"{st.session_state.get('flt_month')} {st.session_state.get('flt_year')}" if st.session_state.get('flt_mode')!='Tahunan' else f"Tahun {st.session_state.get('flt_year')}")
            st.markdown(RK.head_card('Hasil Analisis & Integrasi Laporan Bulanan Workshop',f'Konsolidasi revenue Kharisma Workshop Group • periode Filter Waktu: <b>{_html.escape(str(_pl))}</b> (ubah di sidebar)','hub','blue'),unsafe_allow_html=True)
            st.markdown(RK.dark_banner(f'Total Revenue <span class="c3">({fmt(_tt["Total Revenue"])})</span> = <span class="c1">Revenue Core PKB ({fmt(_core_t)})</span> + <span class="c2">Part Tanpa PKB ({fmt(_tt["Part tanpa PKB"])})</span> + <span class="c2">Bahan Tanpa PKB ({fmt(_tt["Bahan tanpa PKB"])})</span>','Dari 8 modul laporan bulanan'),unsafe_allow_html=True)
            st.markdown(RK.grid([RK.kpi_card('1. Revenue Core (ada PKB)',fmt(_core_t),f'Jasa {fmt(_tt["Jasa"])} • Part &amp; Bahan PKB','build','blue'),
                RK.kpi_card('2. Penjualan Part tanpa PKB',fmt(_tt['Part tanpa PKB']),'Counter sales suku cadang retail','settings','orange'),
                RK.kpi_card('3. Penjualan Bahan tanpa PKB',fmt(_tt['Bahan tanpa PKB']),'Chemical & cat non-PKB','science','orange'),
                RK.kpi_card('4. Total Revenue',fmt(_tt['Total Revenue']),f'{int(_tt["UE"]):,} UE • Rev/UE {fmt(_tt["Revenue/UE"])}','payments','green')],220),unsafe_allow_html=True)
            _hr=['Cabang','UE','Jasa','Part (PKB)','Bahan (PKB)','Revenue Core (PKB)','Part tanpa PKB','Bahan tanpa PKB','Total Revenue','Revenue/UE']
            _rr=[]
            for _,r in pd.concat([_body,_rbm[_rbm['Cabang']=='TOTAL SEMUA CABANG']]).iterrows():
                _rr.append([_html.escape(str(r['Cabang'])),f'{int(r["UE"]):,}',fmt(r['Jasa']),fmt(r['Sparepart (dengan PKB)']),fmt(r['Bahan (dengan PKB)']),fmt(r['Total Revenue']-r['Part tanpa PKB']-r['Bahan tanpa PKB']),fmt(r['Part tanpa PKB']),fmt(r['Bahan tanpa PKB']),fmt(r['Total Revenue']),fmt(r['Revenue/UE'])])
            st.markdown(RK.table_html(_hr,_rr,num_cols=tuple(range(1,10)),bold_last=True,hl_cols={5:'hl1',8:'hl2'}),unsafe_allow_html=True)
            st.caption('Nilai part & bahan tanpa PKB berasal dari file Billing Part dan Consumable yang diupload, sehingga angkanya selalu sama dengan Dashboard. Untuk mengoreksinya, upload ulang / hapus file terkait di Registry Laporan.')
    section('<span class="msr">assignment</span> Status Data','Laporan yang berhasil dibaca oleh aplikasi.')
    status_rows=[]
    _dd={'MaintainPKB':'pkb','Service Penjualan':'detail','Penjualan Part':None,'Issue To PKB':'issue','Consumable':'cons','OPL':'opl','Pelanggan':'cust','Estimasi SA':None}
    for label,df in [('MaintainPKB',pkb),('Service Penjualan',sales),('Penjualan Part',parts),('Issue To PKB',issue),('Consumable',cons),('OPL',opl),('Pelanggan',cust),('Estimasi SA',est_all)]:
        k=_dd.get(label); dup=(DEDUP_STATS.get('parts_service',0)+DEDUP_STATS.get('parts_direct',0)) if label=='Penjualan Part' else (DEDUP_STATS.get(k,0) if k else 0)
        status_rows.append((label,len(df),'Siap' if not df.empty else 'Belum tersedia',int(dup)))
    st.dataframe(pd.DataFrame(status_rows,columns=['Laporan','Jumlah Baris','Status','Baris duplikat antar-file dibuang']),use_container_width=True,hide_index=True)
    st.success('Data utama sudah siap. Silakan buka **Dashboard** untuk melihat performa bengkel.')
    st.stop()

f=add_period_fields(master.copy())
if 'Tanggal' in f:
    f['Tanggal']=pd.to_datetime(f['Tanggal'],errors='coerce'); f=f[f['Tanggal'].notna()].copy()
dmin,dmax=(f.Tanggal.min(),f.Tanggal.max()) if not f.empty else (pd.NaT,pd.NaT)
period_text=f"{dmin.strftime('%d/%m/%Y')} — {dmax.strftime('%d/%m/%Y')}" if pd.notna(dmin) and pd.notna(dmax) else '-'
# Header atas halaman (sesuai referensi): judul + periode + jumlah data + status live + user.
def _filter_period_text():
    """Periode yang dipilih di Filter Waktu sidebar, dalam format dd/mm/yyyy."""
    try:
        _mo=st.session_state.get('flt_mode'); _y=int(st.session_state.get('flt_year')); _m=st.session_state.get('flt_month'); _d=st.session_state.get('flt_day')
        if _mo=='Harian': return datetime.date(_y,MONTH_NAMES.index(_m)+1,int(_d)).strftime('%d/%m/%Y')
        if _mo=='Bulanan': _mi=MONTH_NAMES.index(_m)+1; return f'01/{_mi:02d}/{_y} - {calendar.monthrange(_y,_mi)[1]:02d}/{_mi:02d}/{_y}'
        if _mo=='Tahunan': return f'01/01/{_y} - 31/12/{_y}'
    except Exception: pass
    return None
def _filter_period_label():
    try:
        _mo=st.session_state.get('flt_mode'); _y=int(st.session_state.get('flt_year')); _m=st.session_state.get('flt_month'); _d=st.session_state.get('flt_day')
        if _mo=='Harian': return f'{int(_d)} {_m} {_y}'
        if _mo=='Bulanan': return f'{_m} {_y}'
        if _mo=='Tahunan': return f'Tahun {_y}'
    except Exception: pass
    return period_text
PCFG={'displayModeBar':False,'scrollZoom':False,'doubleClick':False,'staticPlot':False}
def fixfig(fig):
    """Grafik tetap (tanpa zoom/geser/scroll), hover tetap aktif."""
    fig.update_layout(dragmode=False)
    try: fig.update_xaxes(fixedrange=True); fig.update_yaxes(fixedrange=True)
    except Exception: pass
    return fig
def _sa_name_map():
    """Kode/akun SA (mis. johan123) -> nama lengkap (mis. Johan Wahyudi). Sumber: akun pengguna, lalu nama SA pada laporan Estimasi."""
    m={}
    try:
        if isinstance(est_cov,pd.DataFrame) and not est_cov.empty and 'SA Nama' in est_cov:
            for c_,n_ in est_cov[['Service Advisor','SA Nama']].drop_duplicates().itertuples(index=False):
                c_=norm_sa(c_); n_=clean(n_)
                if c_ and n_ and norm_sa(n_)!=c_: m[c_]=n_
    except Exception: pass
    try:
        for u_ in load_users():
            nm_=clean(u_.get('name','')); em_=clean(u_.get('email',''))
            if not nm_ or u_.get('builtin'): continue
            for k_ in {norm_sa(em_),norm_sa(em_.split('@')[0]),norm_sa(u_.get('sa_code',''))}:
                if k_: m[k_]=nm_
    except Exception: pass
    return {k:(v.title() if (v.isupper() or v.islower()) else v) for k,v in m.items()}
SA_NAMES=_sa_name_map()
def sa_disp(x):
    """Tampilkan nama lengkap SA; jika belum ada nama, tampilkan kodenya apa adanya."""
    if isinstance(x,pd.Series): return x.map(lambda v:SA_NAMES.get(norm_sa(v),clean(v)))
    return SA_NAMES.get(norm_sa(x),clean(x))
_topbar_user=(CURRENT_USER.get('name') if CURRENT_USER else None) or st.session_state.get('logged_in_username') or 'admin'
_topbar_role=ROLE_LABELS.get((CURRENT_USER or {}).get('role'),'') if CURRENT_USER else ''
if page != '📥 Data Input':
    st.markdown(RK.page_header(_filter_period_text() or period_text,f'{len(f):,} UE',_topbar_user,_topbar_role),unsafe_allow_html=True)
if TF_ACTIVE and f.empty and page not in ('🎯 Target Perusahaan','🧠 Management Insight'):
    st.warning('Tidak ada data pada filter waktu yang dipilih. Ubah filter waktu di sidebar.')
    st.stop()

def render_summary_cards(df_summary, title=':material/bar_chart: Ringkasan Periode', subtitle=None):
    x=df_summary.copy()
    total_ue=int(x['UE'].sum()) if 'UE' in x else len(x)
    rc=revenue_components(x,parts,issue,cons)
    rev=rc['total_revenue']; rpu=rev/total_ue if total_ue else 0
    section(title, subtitle or 'Ringkasan dari seluruh laporan yang tersedia')
    summary_cards=[
        ('TOTAL UE',f'{total_ue:,}','<span class="msr">directions_car</span>','1 PKB = 1 UE','c-blue'),
        ('JASA',fmt(rc['jasa']),'<span class="msr">home_repair_service</span>','Jasa dari PKB / Service Penjualan','c-green'),
        ('SPAREPART',fmt(rc['total_sparepart']),'<span class="msr">settings</span>','Dengan PKB + tanpa PKB','c-purple'),
        ('BAHAN',fmt(rc['total_bahan']),'<span class="msr">science</span>','Issue To PKB + Consumable','c-orange'),
        ('TOTAL REVENUE',fmt(rev),'<span class="msr">payments</span>','Core + penjualan tanpa PKB','c-pink'),
        ('REVENUE / UE',fmt(rpu),'<span class="msr">trending_up</span>','Total Revenue ÷ UE','c-teal'),
    ]
    for row in [summary_cards[:3],summary_cards[3:]]:
        cols=st.columns(3)
        for col,(k,v,ic,sub,cc) in zip(cols,row):
            with col: st.markdown(f'<div class="summary-card {cc}"><div class="summary-icon">{ic}</div><div class="summary-label">{k}</div><div class="summary-value">{v}</div><div class="summary-sub">{sub}</div></div>',unsafe_allow_html=True)

if page == '📊 Dashboard':
    rc=revenue_components(f,parts,issue,cons)
    total_ue=int(f['UE'].sum()) if 'UE' in f.columns else len(f)
    rev=float(rc['total_revenue']); rpu=(rev/total_ue if total_ue else 0.0)
    _d1,_d2=st.columns([5,1.3])
    with _d1:
        st.markdown(RK.head_card('Dashboard Workshop Performance',f'<span class="msr">calendar_month</span> Periode Aktif: <b>{_html.escape(str(_filter_period_label()))}</b> &nbsp;•&nbsp; Cabang: <b>{_html.escape(str(active_branch))}</b>','bar_chart','blue'),unsafe_allow_html=True)
    with _d2:
        if has_perm('export_excel'):
            if st.button(':material/table_view: Siapkan Excel',key='dash_prep_xlsx',use_container_width=True):
                with st.spinner('Menyusun file Excel...'):
                    st.session_state['_dash_xlsx']=make_excel_download(f,sales,parts,issue,cons,opl,cust)
            if st.session_state.get('_dash_xlsx'):
                st.download_button(':material/download: Download Excel',st.session_state['_dash_xlsx'],'dashboard_workshop_performance.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='dash_dl_xlsx',type='primary',use_container_width=True)
    st.markdown(RK.grid([
        RK.kpi_card('Total UE',f'{total_ue:,}',_html.escape(str(active_branch)),'directions_car','blue'),
        RK.kpi_card('Jasa',fmt(rc['jasa']),'Biaya jasa perbaikan','build','green'),
        RK.kpi_card('Sparepart',fmt(rc['total_sparepart']),f'PKB: {fmt(rc["service_sparepart"])} • Non: {fmt(rc["direct_part"])}','settings','purple'),
        RK.kpi_card('Bahan',fmt(rc['total_bahan']),f'PKB: {fmt(rc["service_bahan"])} • Non: {fmt(rc["direct_bahan"])}','science','orange'),
        RK.kpi_card('Total Revenue',fmt(rev),f'Core: {fmt(rc["revenue_core"])} + Non-PKB','payments','red'),
        RK.kpi_card('Revenue / UE',fmt(rpu),'Rata-rata per unit','trending_up','teal')],165,'sm'),unsafe_allow_html=True)
    st.markdown(RK.dark_banner('<b>FORMULA</b>&nbsp; Total Revenue = <span class="c1">Revenue Core (PKB)</span> + <span class="c2">Part Non-PKB</span> + <span class="c2">Bahan Non-PKB</span>',f'Akumulasi: {fmt(rev)} (Core {fmt(rc["revenue_core"])} + Non-PKB {fmt(rc["direct_part"]+rc["direct_bahan"])})'),unsafe_allow_html=True)
    st.caption('Revenue Core = Jasa + sparepart dengan PKB + bahan dengan PKB. Total Revenue = Revenue Core + part tanpa PKB + bahan tanpa PKB.')

    # ===== Revenue per Cabang (rincian Core PKB + Non-PKB) =====
    def _all_branch_frames():
        """Data semua cabang dengan Filter Waktu yang sama (pipeline di-cache) untuk tabel pembanding antar cabang."""
        if active_branch=='Semua Cabang' or not has_perm('all_branches'): return f,parts,issue,cons
        try:
            _Pa=_pipe('Semua Cabang',st.session_state.get('est_ref_on',True))
            if _Pa.get('error') or not isinstance(_Pa.get('master'),pd.DataFrame): return f,parts,issue,cons
            return tf(_Pa['master']),tf(_Pa['parts']),tf(_Pa['issue']),tf(_Pa['cons'])
        except Exception: return f,parts,issue,cons
    _bm_,_bp_,_bi_,_bc_=_all_branch_frames()
    _show_br=list(BRANCHES) if has_perm('all_branches') else [b for b in BRANCHES if b in user_branches()]
    _rb=revenue_components_by_branch(_bm_,_bp_,_bi_,_bc_,_show_br)
    if not _rb.empty:
        _body=_rb[_rb['Cabang']!='TOTAL SEMUA CABANG']; _tot=_rb[_rb['Cabang']=='TOTAL SEMUA CABANG']
        _hd=['Cabang','UE','Jasa','Sparepart (PKB)','Bahan (PKB)','Revenue Core (PKB)','Part tanpa PKB','Bahan tanpa PKB','Total Revenue','Revenue/UE']
        _rws=[]; _cls=[]
        for _,r in pd.concat([_body,_tot]).iterrows():
            is_tot=(r['Cabang']=='TOTAL SEMUA CABANG'); sel=(r['Cabang']==active_branch)
            _nm=_html.escape(str(r['Cabang']))
            _rws.append([('<span class="rk-dot2"></span>' if sel and active_branch!='Semua Cabang' else '')+_nm,f'{int(r["UE"]):,}',f'{r["Jasa"]:,.0f}',f'{r["Sparepart (dengan PKB)"]:,.0f}',f'{r["Bahan (dengan PKB)"]:,.0f}',
                f'{r["Total Revenue"]-r["Part tanpa PKB"]-r["Bahan tanpa PKB"]:,.0f}',f'{r["Part tanpa PKB"]:,.0f}',f'{r["Bahan tanpa PKB"]:,.0f}',f'{r["Total Revenue"]:,.0f}',f'{r["Revenue/UE"]:,.0f}'])
            _cls.append('' if (is_tot or active_branch=='Semua Cabang' or sel) else 'dim')
        st.markdown(f'<div class="rk-ct"><b>REVENUE PER CABANG (RINCIAN CORE PKB + NON-PKB)</b><span>{len(_body)} Cabang Terdata</span></div>',unsafe_allow_html=True)
        st.markdown(RK.table_html(_hd,_rws,num_cols=tuple(range(1,10)),bold_last=True,hl_cols={5:'hl1',8:'hl2'},row_cls=_cls),unsafe_allow_html=True)
        if active_branch=='Semua Cabang':
            _sum_rows=float(_body['Total Revenue'].sum())
            if abs(_sum_rows-float(rc['total_revenue']))<1: st.caption(f'Total semua cabang {fmt(rc["total_revenue"])} = jumlah revenue masing-masing cabang. Cabang baru otomatis tampil di tabel ini; nilai 0 berarti data cabang belum diupload.')
            else: st.warning(f'Selisih {fmt(float(rc["total_revenue"])-_sum_rows)} antara total dan jumlah cabang — periksa data cabang.')
        download_table(_rb,':material/download: Download Excel — Revenue per Cabang','revenue_per_cabang.xlsx','dl_rev_cabang_dash')

    # ===== Revenue per SA & UE berdasarkan Layanan (donat) =====
    _cA,_cB=st.columns(2)
    _pal=['#2563eb','#10b981','#f59e0b','#ef4444','#8b5cf6','#06b6d4','#ec4899','#84cc16','#f97316','#14b8a6','#6366f1','#a855f7']
    with _cA:
        with st.container(border=True,key='dn_sa'):
            _sa=f[f['Service Advisor'].astype(str).str.strip()!=''].copy()
            st.markdown(f'<div class="rk-ct"><b><span class="msr">groups</span><span>REVENUE PER SA — <span class="a">{_html.escape(str(active_branch)).upper()}</span></span></b><span>{_sa["Service Advisor"].nunique() if not _sa.empty else 0} Service Advisor</span></div>',unsafe_allow_html=True)
            if _sa.empty: st.info('Belum ada data Service Advisor pada periode ini.')
            else:
                _sa['Service Advisor']=sa_disp(_sa['Service Advisor'])
                _sat=_sa.groupby('Service Advisor',as_index=False).agg(UE=('UE','sum'),Jasa=('Jasa','sum'),Sparepart=('Sparepart','sum'),Bahan=('Bahan','sum'),Revenue=('Revenue','sum'))
                _sel=st.radio('Tampilkan pencapaian berdasarkan',['Revenue','UE','Jasa','Sparepart','Bahan'],horizontal=True,key='sa_pie_metric')
                _sat=_sat.sort_values(_sel,ascending=False).reset_index(drop=True)
                _sat['% Kontribusi']=(_sat[_sel]/_sat[_sel].sum()*100) if _sat[_sel].sum()>0 else 0.0
                _cs=0.0   # Counter Sales (Non-PKB) tidak ditampilkan
                _top=_sat.head(6); _rest=_sat.iloc[6:]
                _labels=_top['Service Advisor'].tolist()+(['Lainnya'] if len(_rest) else []); _vals=_top[_sel].tolist()+([float(_rest[_sel].sum())] if len(_rest) else [])
                _tot_txt=f"{int(sum(_vals)):,}" if _sel=='UE' else RK.fmt_jt(sum(_vals))
                _fig=go.Figure(go.Pie(labels=_labels,values=_vals,hole=.62,sort=False,textinfo='none',marker=dict(colors=_pal[:len(_labels)],line=dict(color='#ffffff',width=2)),
                    hovertemplate='%{label}<br>%{value:,.0f} (%{percent})<extra></extra>'))
                _fig.update_layout(height=300,margin=dict(l=4,r=4,t=4,b=4),hoverlabel=dict(bgcolor='#0f172a',font_color='#fff',font_size=12),paper_bgcolor='#ffffff',plot_bgcolor='#ffffff',showlegend=False,
                    annotations=[dict(text=f'<span style="font-size:11px;color:#64748b">TOTAL SA</span><br><span style="font-size:17px"><b>{len(_sat)-(1 if _cs>0 else 0)} SA</b></span><br><span style="font-size:11px;color:#2563eb">{_tot_txt}</span>',x=.5,y=.5,showarrow=False)])
                _q1,_q2=st.columns([1,1.15])
                with _q1: st.plotly_chart(_fig,use_container_width=True,config=PCFG)
                with _q2:
                    _tv=float(_sat[_sel].sum())
                    _its=[]
                    for _k,(_,rw) in enumerate(_top.iterrows()):
                        _val=(f'{int(rw[_sel]):,} UE' if _sel=='UE' else RK.fmt_jt(rw[_sel])); _sub=(RK.fmt_jt(rw['Revenue']) if _sel=='UE' else (f'{int(rw["UE"]):,} UE' if int(rw['UE'])>0 else 'Penjualan tanpa PKB'))
                        _its.append((_pal[_k%len(_pal)],str(rw['Service Advisor']),_sub,_val,f'{rw["% Kontribusi"]:.0f}% Share'))
                    if len(_rest): _its.append(('#94a3b8',f'Lainnya ({len(_rest)} SA)','',(f'{int(_rest[_sel].sum()):,} UE' if _sel=='UE' else RK.fmt_jt(_rest[_sel].sum())),f'{_rest["% Kontribusi"].sum():.0f}% Share'))
                    st.markdown(RK.legend_list(_its),unsafe_allow_html=True)
                with st.expander(':material/table_view: Tabel rinci & download'):
                    _show=_sat.copy(); _show['% Kontribusi']=_show['% Kontribusi'].map(lambda v:f'{v:.1f}%')
                    for _c in ['Jasa','Sparepart','Bahan','Revenue']: _show[_c]=_show[_c].map(lambda v:f'{v:,.0f}')
                    _show['UE']=_show['UE'].map(lambda v:f'{int(v):,}')
                    st.dataframe(_show[['Service Advisor','UE','Jasa','Sparepart','Bahan','Revenue','% Kontribusi']],use_container_width=True,hide_index=True)
                    download_table(_sat,':material/download: Download Excel — Pencapaian per SA','pencapaian_per_sa.xlsx','dl_sa_pencapaian')
                st.caption('Revenue per SA = jasa + sparepart + bahan dengan PKB (revenue core). Penjualan part/bahan tanpa PKB tidak terikat ke SA.')
    with _cB:
        with st.container(border=True,key='dn_ly'):
            st.markdown(f'<div class="rk-ct"><b><span class="msr">explore</span><span>UE BERDASARKAN LAYANAN — <span class="a">{_html.escape(str(active_branch)).upper()}</span></span></b><span>Booking vs Walk-In vs Field</span></div>',unsafe_allow_html=True)
            _bt=f.copy(); _bt['_L']=_bt['Booking Type'].map(clean).str.upper().str.replace('-',' ',regex=False).str.strip()
            _lmap={'BOOKING':'Booking','WALK IN':'Walk-In','WALKIN':'Walk-In','FIELD SERVICE':'Field Service','FIELD':'Field Service'}
            _bt['_L']=_bt['_L'].map(lambda x:_lmap.get(x,'Lainnya' if x else 'Tidak diisi'))
            _lz=_bt.groupby('_L',as_index=False).agg(UE=('UE','sum')).sort_values('UE',ascending=False)
            _lt=float(_lz['UE'].sum()); _lc={'Booking':'#2563eb','Walk-In':'#06b6d4','Field Service':'#f59e0b','Lainnya':'#94a3b8','Tidak diisi':'#cbd5e1'}
            if _lz.empty or _lt<=0: st.info('Belum ada data jenis layanan pada periode ini.')
            else:
                _fig2=go.Figure(go.Pie(labels=_lz['_L'].tolist(),values=_lz['UE'].tolist(),hole=.62,sort=False,textinfo='none',marker=dict(colors=[_lc.get(x,'#94a3b8') for x in _lz['_L']],line=dict(color='#ffffff',width=2))))
                _fig2.update_layout(height=300,margin=dict(l=4,r=4,t=4,b=4),hoverlabel=dict(bgcolor='#0f172a',font_color='#fff',font_size=12),paper_bgcolor='#ffffff',plot_bgcolor='#ffffff',showlegend=False,
                    annotations=[dict(text=f'<span style="font-size:11px;color:#64748b">TOTAL UE</span><br><span style="font-size:17px"><b>{int(_lt):,} UE</b></span><br><span style="font-size:11px;color:#64748b">Unit Service</span>',x=.5,y=.5,showarrow=False)])
                _w1,_w2=st.columns([1,1.15])
                with _w1: st.plotly_chart(_fig2,use_container_width=True,config=PCFG)
                with _w2: st.markdown(RK.legend_list([(_lc.get(r['_L'],'#94a3b8'),r['_L'],'Kedatangan Unit',f'{int(r["UE"]):,} UE',RK.share_pct(r['UE'],_lt)) for _,r in _lz.iterrows()]),unsafe_allow_html=True)
                download_table(_lz.rename(columns={'_L':'Layanan'}),':material/download: Download Excel — UE Booking Walk In','ue_booking_walkin_dashboard.xlsx','dl_dash_layanan')

    # Fitur dashboard lama dipertahankan: trend, Top 5, Booking/Walk In, Management Insight.
    booking=int((f['Booking Type'].str.upper()=='BOOKING').sum()); walkin=int((f['Booking Type'].str.upper()=='WALK IN').sum())
    top_vehicle=f.groupby('Type Kendaraan').UE.sum().sort_values(ascending=False).head(5)
    top_job=f.groupby('Jenis Pekerjaan').Revenue.sum().sort_values(ascending=False).head(5)
    sa_rank=f[f['Service Advisor']!=''].assign(**{'Service Advisor':lambda d:sa_disp(d['Service Advisor'])}).groupby('Service Advisor').UE.sum().sort_values(ascending=False).head(10)
    daily=f.groupby('Tanggal',as_index=False).agg(UE=('UE','sum'),Revenue=('Revenue','sum')).sort_values('Tanggal')
    fig=go.Figure(); fig.add_trace(go.Bar(x=daily['Tanggal'],y=daily['UE'],name='UE',marker_color='#2563eb',opacity=.78)); fig.add_trace(go.Scatter(x=daily['Tanggal'],y=daily['Revenue']/1e6,name='Revenue (Juta)',mode='lines+markers',line=dict(color='#e31b23',width=3),marker=dict(size=5),yaxis='y2'))
    fig.update_layout(height=350,margin=dict(l=8,r=12,t=8,b=10),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',hovermode='x unified',legend=dict(orientation='h',y=1.08,x=0),yaxis=dict(title='UE',gridcolor='#e8eef6'),yaxis2=dict(title='Revenue (Juta)',overlaying='y',side='right',showgrid=False))
    st.markdown('<div class="chart-card"><h3><span class="msr">directions_car</span> Trend UE & Revenue</h3><p>Pergerakan harian selama periode laporan.</p>',unsafe_allow_html=True); st.plotly_chart(fixfig(fig),use_container_width=True,config=PCFG)
    a,b=st.columns(2)
    with a:
        st.markdown('<div class="chart-card"><h3><span class="msr">emoji_events</span> Top 5 Type Kendaraan</h3><p>Kontribusi UE terbesar.</p>',unsafe_allow_html=True)
        vdf=top_vehicle.sort_values(); figv=go.Figure(go.Bar(x=vdf.values,y=vdf.index,orientation='h',marker_color='#2563eb',text=[f'{int(x):,}' for x in vdf.values],textposition='outside',cliponaxis=False)); figv.update_layout(height=300,margin=dict(l=5,r=45,t=5,b=5),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',xaxis=dict(title='UE',gridcolor='#edf2f7'),yaxis=dict(title='')); st.plotly_chart(fixfig(figv),use_container_width=True,config=PCFG)
    with b:
        st.markdown('<div class="chart-card"><h3><span class="msr">payments</span> Top 5 Job Type</h3><p>Job type dengan revenue terbesar.</p>',unsafe_allow_html=True)
        jdf=top_job.sort_values(); figj=go.Figure(go.Bar(x=jdf.values/1e6,y=jdf.index,orientation='h',marker_color='#10b981',text=[fmt(x) for x in jdf.values],textposition='outside',cliponaxis=False)); figj.update_layout(height=300,margin=dict(l=5,r=90,t=5,b=5),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',xaxis=dict(title='Revenue (Juta)',gridcolor='#edf2f7'),yaxis=dict(title='')); st.plotly_chart(fixfig(figj),use_container_width=True,config=PCFG)
    d=st.container()
    with d:
        no_sales = int(len(qc_no_sales))
        top_sa = str(sa_rank.index[0]) if len(sa_rank) else '-'
        top_sa_ue = int(sa_rank.iloc[0]) if len(sa_rank) else 0
        if len(top_vehicle):
            vehicle_name = str(top_vehicle.index[0])
            vehicle_ue = int(top_vehicle.iloc[0])
            vehicle_insight = f'{vehicle_name} menjadi kendaraan dengan UE tertinggi ({vehicle_ue:,} UE).'
        else:
            vehicle_insight = 'Belum ada data kendaraan.'
        dashboard_total_ue = int(f['UE'].sum()) if 'UE' in f.columns else len(f)
        dashboard_total_rev = float(revenue_components(f,parts,issue,cons)['total_revenue'])
        safe_rev_ue = (dashboard_total_rev / dashboard_total_ue) if dashboard_total_ue else 0.0
        revenue_insight = f'Revenue / UE periode ini {fmt(safe_rev_ue)}.'
        sales_insight = f'{no_sales:,} PKB tanpa penjualan tidak dihitung sebagai UE (lihat menu Quality).'
        if top_sa != '-':
            sa_insight = f'{top_sa} menjadi SA dengan UE tertinggi ({top_sa_ue:,} UE).'
        else:
            sa_insight = 'Belum ada data Service Advisor.'
        insights = [vehicle_insight, revenue_insight, sales_insight, sa_insight]
        html='<div class="insight-card"><h3 style="margin:0 0 5px;color:#0f172a;"><span class="msr">lightbulb</span> Management Insight</h3><p style="margin:0 0 8px;color:#64748b;font-size:11px;">Insight otomatis dari data.</p>'+''.join(f'<div class="insight-row"><span class="insight-dot"><span class="msr">check</span></span><span>{x}</span></div>' for x in insights)+'</div>'; st.markdown(html,unsafe_allow_html=True)
    st.markdown('<div class="footer-band">Data Lebih Akurat &nbsp; | &nbsp; Analisa Lebih Dalam &nbsp; | &nbsp; Keputusan Lebih Baik</div>',unsafe_allow_html=True)

if page == '🎯 Target Perusahaan':
    st.markdown('<style>[data-testid="stMain"] [data-testid="stMarkdownContainer"] p, [data-testid="stMarkdownContainer"] li, section.main [data-testid="stMarkdownContainer"] p, [data-testid="stMarkdownContainer"] li{color:#0f172a !important;}[data-testid="stMain"] [data-testid="stCaptionContainer"], [data-testid="stCaptionContainer"] *, section.main [data-testid="stCaptionContainer"], [data-testid="stCaptionContainer"] *{color:#1e293b !important;font-size:13.5px !important;font-weight:500 !important;}[data-testid="stMain"] [data-testid="stMetricLabel"], [data-testid="stMetricLabel"] *, section.main [data-testid="stMetricLabel"], [data-testid="stMetricLabel"] *{color:#1e293b !important;font-weight:700 !important;}[data-testid="stMain"] [data-testid="stMetricValue"], [data-testid="stMetricValue"] *, section.main [data-testid="stMetricValue"], [data-testid="stMetricValue"] *{color:#0b1b3a !important;font-weight:800 !important;}[data-testid="stMain"] [data-testid="stMetricDelta"], [data-testid="stMetricDelta"] *, section.main [data-testid="stMetricDelta"], [data-testid="stMetricDelta"] *{color:#1d4ed8 !important;font-weight:700 !important;opacity:1 !important;}[data-testid="stMain"] .stTabs [data-baseweb="tab"] p, section.main .stTabs [data-baseweb="tab"] p{color:#334155 !important;font-weight:700 !important;}[data-testid="stMain"] .stTabs [aria-selected="true"] p, section.main .stTabs [aria-selected="true"] p{color:#1d4ed8 !important;}[data-testid="stMain"] [data-testid="stProgress"] > div > div, section.main [data-testid="stProgress"] > div > div{background:#cbd5e1 !important;}[data-testid="stMain"] h1,h2,h3,h4, section.main h1,h2,h3,h4{color:#0f172a !important;}</style>',unsafe_allow_html=True)
    t_master,t_parts,t_issue,t_cons=UNF['master'],UNF['parts'],UNF['issue'],UNF['cons']   # tanpa filter waktu sidebar
    _today=datetime.date.today()
    _dd=pd.to_datetime(t_master['Tanggal'],errors='coerce').dropna() if isinstance(t_master,pd.DataFrame) and not t_master.empty and 'Tanggal' in t_master.columns else pd.Series(dtype='datetime64[ns]')
    _dy,_dm=((int(_dd.max().year),int(_dd.max().month)) if len(_dd) else (_today.year,_today.month))
    try: ty=int(st.session_state.get('flt_year'))
    except Exception: ty=_dy
    _fm=st.session_state.get('flt_month','Semua'); _tgt_note=''
    if _fm in MONTH_NAMES: tm=MONTH_NAMES.index(_fm)+1
    else:   # mode Tahunan: pakai bulan terakhir yang punya data pada tahun itu
        _yd=_dd[_dd.dt.year==ty]
        tm=int(_yd.max().month) if len(_yd) else (_today.month if ty==_today.year else 12)
        _tgt_note=f' (mode Tahunan menampilkan bulan terakhir yang memiliki data pada {ty})'
    tm_name=MONTH_NAMES[tm-1]
    pkey=f'{ty}-{tm:02d}'; days_in=calendar.monthrange(ty,tm)[1]
    tab_mon=st.container()

    with tab_mon:
        _targets=load_targets(); _cur=_targets.get(pkey,{})
        scope=active_branch
        _bset=list(dict.fromkeys(list(BRANCHES)+[b for b in _cur if b not in BRANCHES]))
        def _tg(m): return float(sum((_cur.get(b,{}).get(m,0) or 0) for b in _bset)) if scope=='Semua Cabang' else float((_cur.get(scope,{}) or {}).get(m,0) or 0)
        T={m:_tg(m) for m in TARGET_METRICS}
        def _inmonth(df,col):
            if not isinstance(df,pd.DataFrame) or df.empty or col not in df.columns: return df
            d=pd.to_datetime(df[col],errors='coerce'); return df[((d.dt.year==ty)&(d.dt.month==tm)).values].copy()
        mm,pp,ii,cc=_inmonth(t_master,'Tanggal'),_inmonth(t_parts,'Tanggal'),_inmonth(t_issue,'Tanggal'),_inmonth(t_cons,'Tanggal')
        _daily=daily_components(mm,pp,ii,cc)
        _idx=pd.date_range(pd.Timestamp(ty,tm,1),periods=days_in,freq='D')
        _daily=_daily.reindex(_idx,fill_value=0.0); _daily.index.name='Tanggal'
        R={m:float(_daily[col].sum()) for m,col in [('UE','UE'),('Jasa','Jasa'),('Part','Part'),('Bahan','Bahan'),('Revenue','Revenue')]}
        if pd.Timestamp(ty,tm,1).date()>_today: elapsed=0
        elif pd.Timestamp(ty,tm,days_in).date()<_today: elapsed=days_in
        else: elapsed=_today.day
        days_left=days_in-elapsed; exp_pct=(elapsed/days_in*100) if days_in else 0.0
        def _fv(m,v): return f'{int(round(v)):,}' if m=='UE' else fmt(v)
        _lab={'UE':'UE','Jasa':'Jasa','Part':'Part (sparepart)','Bahan':'Bahan','Revenue':'Total Revenue'}
        # ================= V17: tampilan ringkas sesuai referensi =================
        _rcm=revenue_components(mm,pp,ii,cc)
        with st.container(key='blk_tgt'):   # judul + tombol input target dalam SATU blok
            _hc1,_hc2=st.columns([5,1.25],vertical_alignment='center')
            with _hc1:
                st.markdown(RK.head_card('Target Perusahaan & Persentase Pencapaian',f'Periode Target: <b>{tm_name} {ty}</b> &nbsp;•&nbsp; Tampilan: <b>{_html.escape(str(scope))}</b>','track_changes','red'),unsafe_allow_html=True)
            with _hc2:
                st.button(':material/edit: Input Target',key='tgt_goto_input',type='primary',use_container_width=True,on_click=lambda:st.session_state.update({'page':'📥 Data Input','upload_mode':'Target'}))
        _kpi_x=pd.DataFrame([{'Metrik':_lab[m],'Target':T[m],'Realisasi':R[m],'% Pencapaian':(R[m]/T[m]*100 if T[m]>0 else 0.0)} for m in TARGET_METRICS])
        st.markdown(RK.dark_banner(f'Realisasi Aktual Revenue dihitung komprehensif: <span class="c1">Revenue Core PKB ({fmt(_rcm["revenue_core"])})</span> + <span class="c2">Part Non-PKB ({fmt(_rcm["direct_part"])})</span> + <span class="c2">Bahan Non-PKB ({fmt(_rcm["direct_bahan"])})</span> = <span class="c3">Total Revenue ({fmt(_rcm["total_revenue"])})</span>.',str(scope)),unsafe_allow_html=True)
        if _tgt_note: st.caption(f'Mode Tahunan menampilkan bulan terakhir yang memiliki data pada {ty}.')
        if sum(T.values())==0: st.info('Target untuk periode & cabang ini belum diinput. Isi lewat menu **Data Input → Input Target Perusahaan**; realisasi tetap ditampilkan.')

        st.markdown(RK.head_card('Metrik per Kartu','Realisasi, sisa target, kebutuhan per hari, dan proyeksi akhir bulan untuk tiap metrik','dashboard','blue'),unsafe_allow_html=True)
        _dcs=[]
        _tone_m={'UE':'blue','Jasa':'green','Part':'purple','Bahan':'orange','Revenue':'red'}; _ico_m={'UE':'directions_car','Jasa':'build','Part':'settings','Bahan':'science','Revenue':'payments'}
        for m in TARGET_METRICS:
            pct=(R[m]/T[m]*100) if T[m]>0 else 0.0; remain=max(T[m]-R[m],0.0)
            need=(remain/days_left) if days_left>0 else remain; proj=(R[m]/elapsed*days_in) if elapsed>0 else 0.0
            if T[m]>0:
                if pct>=100: _st_t,_st_c='Tercapai','green'
                elif pct>=exp_pct: _st_t,_st_c='Di atas pace','green'
                elif pct>=exp_pct*0.9: _st_t,_st_c='Mendekati pace','orange'
                else: _st_t,_st_c='Tertinggal','red'
            else: _st_t,_st_c='Target kosong','gray'
            _dcs.append(RK.detail_card(_lab[m],_ico_m[m],_tone_m[m],_fv(m,R[m]) if m=='UE' else RK.fmt_jt(R[m]),(_fv(m,T[m]) if m=='UE' else RK.fmt_jt(T[m])),pct,_st_t,_st_c,
                [('Sisa target',_fv(m,remain) if T[m]>0 else '-'),(f'Butuh / hari ({days_left} hari)',_fv(m,need) if T[m]>0 else '-'),('Proyeksi akhir bulan',(f'{_fv(m,proj)} ({(proj/T[m]*100):.0f}%)' if T[m]>0 else _fv(m,proj)))],target_set=T[m]>0))
        st.markdown(RK.grid(_dcs,230),unsafe_allow_html=True)

        st.markdown(RK.head_card('Kontrol Harian','Realisasi harian, akumulasi, dan pace target linear per hari kalender','event_note','teal'),unsafe_allow_html=True)
        _dm_sel=st.segmented_control('Metrik grafik',TARGET_METRICS,default='Revenue',format_func=lambda m:_lab[m],key='tgt_chart_metric',label_visibility='collapsed') or 'Revenue'
        _col=_dm_sel; _div=1.0 if _dm_sel=='UE' else 1e6; _unit='UE' if _dm_sel=='UE' else 'Jt'
        _cum=_daily[_col].cumsum(); _tgt_cum=pd.Series([T[_dm_sel]*(i+1)/days_in for i in range(days_in)],index=_idx)
        _fig=go.Figure()
        _fig.add_trace(go.Bar(x=_idx,y=_daily[_col]/_div,name=f'Realisasi harian ({_unit})',marker_color='#93c5fd',yaxis='y2',opacity=.9))
        _shown=_cum.where(_idx<=pd.Timestamp(_today)) if elapsed<days_in else _cum
        _fig.add_trace(go.Scatter(x=_idx,y=_shown/_div,name=f'Kumulatif ({_unit})',mode='lines+markers',line=dict(color='#2563eb',width=3),marker=dict(size=5)))
        if T[_dm_sel]>0: _fig.add_trace(go.Scatter(x=_idx,y=_tgt_cum/_div,name='Target kumulatif',mode='lines',line=dict(color='#ef4444',width=2,dash='dash')))
        _avgT=(T[_dm_sel]/days_in) if days_in else 0.0
        if _avgT>0: _fig.add_trace(go.Scatter(x=_idx,y=[_avgT/_div]*days_in,name='Target harian rata-rata',mode='lines',yaxis='y2',line=dict(color='#f59e0b',width=2.5,dash='dot'),hovertemplate='%{y:,.2f} '+_unit+'<extra>Target harian rata-rata</extra>'))
        _fig.update_layout(font=dict(color='#0f172a',size=12),height=340,margin=dict(l=8,r=8,t=8,b=8),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',hovermode='x unified',legend=dict(orientation='h',y=-.15,x=0),
            yaxis=dict(title=f'Kumulatif ({_unit})',gridcolor='#e8eef6'),yaxis2=dict(title=f'Harian ({_unit})',overlaying='y',side='right',showgrid=False,range=[0,max(float(_daily[_col].max())/_div,_avgT/_div)*1.25 or 1.0]),xaxis=dict(tickformat='%d/%m'))
        st.plotly_chart(fixfig(_fig),use_container_width=True,config=PCFG)
        _act_d=_daily[_daily[_col]>0]; _avgA=float(_act_d[_col].mean()) if len(_act_d) else 0.0
        def _fa(v): return f'{v:,.1f} UE' if _dm_sel=='UE' else RK.fmt_jt(v)
        _pa=(_avgA/_avgT*100) if _avgT>0 else 0.0
        if len(_act_d): _pkd=_act_d[_col].idxmax(); _pk_t=f'Tgl {_pkd.day:02d}/{_pkd.month:02d}'; _pk_s=_fa(float(_act_d.loc[_pkd,_col]))
        else: _pk_t,_pk_s='-',''
        _hit=int((_act_d[_col]>=_avgT).sum()) if _avgT>0 else 0
        st.markdown(RK.min_grid([RK.min_card('Target Harian Rata-rata',_fa(_avgT) if _avgT>0 else '-',f'{_lab[_dm_sel]} ÷ {days_in} hari','amber'),
            RK.min_card('Rata-rata Aktual/Hari',_fa(_avgA),(f'{_pa:.0f}% dari target harian' if _avgT>0 else f'{len(_act_d)} hari aktif'),'blue'),
            RK.min_card('Hari Tertinggi (Peak)',_pk_t,_pk_s,'green'),
            RK.min_card('Hari Target Tercapai',f'{_hit}',f'dari {max(elapsed,len(_act_d))} hari','purple')]),unsafe_allow_html=True)
        _cumR=_daily['Revenue'].cumsum(); _tcR=[T['Revenue']*(i+1)/days_in for i in range(days_in)]
        _trs2=[]
        for i in range(days_in):
            _d=_idx[i]; _rv=float(_daily['Revenue'].iloc[i]); _cv=float(_cumR.iloc[i]); _tv=_tcR[i]; _gap=_cv-_tv
            _pc_=(_cv/T['Revenue']*100 if T['Revenue']>0 else 0.0)
            if i>=elapsed: _sb=RK.status_badge('Belum','gray')
            elif _rv==0: _sb=RK.status_badge('Hari kosong','gray')
            elif T['Revenue']>0: _sb=RK.status_badge('Di atas pace' if _gap>=0 else 'Di bawah pace','green' if _gap>=0 else 'red')
            else: _sb=RK.status_badge('Berjalan','blue')
            _gc=(f'<span style="color:{"#047857" if _gap>=0 else "#b91c1c"};font-weight:700">{_gap:+,.0f}</span>' if T['Revenue']>0 and i<elapsed else '-')
            _trs2.append([f'{_d.day:02d}/{_d.month:02d}/{_d.year}',f'{int(_daily["UE"].iloc[i]):,}',f'{_daily["Jasa"].iloc[i]:,.0f}',f'{_daily["Part"].iloc[i]:,.0f}',f'{_daily["Bahan"].iloc[i]:,.0f}',f'<b>{_rv:,.0f}</b>',f'{_cv:,.0f}',f'{_pc_:.1f}%' if T['Revenue']>0 else '-',_gc,_sb])
        st.markdown(RK.table_html(['Tanggal','UE','Jasa','Part','Bahan','Revenue','Kumulatif','% Target','Selisih vs Pace','Status'],_trs2,num_cols=(1,2,3,4,5,6,7,8)),unsafe_allow_html=True)
        _chk=revenue_components(mm,pp,ii,cc)['total_revenue']
        if abs(_chk-R['Revenue'])<1: st.caption(f'Realisasi {fmt(R["Revenue"])} sama dengan Total Revenue Dashboard untuk periode ini.')
        else: st.warning(f'Selisih {fmt(_chk-R["Revenue"])} antara realisasi harian dan Total Revenue Dashboard — periksa data.')
        st.caption('Realisasi mengikuti acuan Dashboard (Laporan Estimasi SA + penjualan part/bahan tanpa PKB). Target kumulatif dihitung linear per hari kalender. Periode target mengikuti Filter Waktu di sidebar.')

        # ---- rincian per cabang ----
        st.markdown(RK.head_card(f'Rincian Target vs Realisasi per Cabang ({tm_name} {ty})','Evaluasi ketercapaian target bulanan masing-masing cabang','apartment','purple'),unsafe_allow_html=True)
        _rb2=revenue_components_by_branch(mm,pp,ii,cc,BRANCHES)
        _rb2=_rb2[_rb2['Cabang']!='TOTAL SEMUA CABANG'].copy()
        if scope!='Semua Cabang': _rb2=_rb2[_rb2['Cabang']==scope]
        _rb2['Part']=_rb2['Sparepart (dengan PKB)']+_rb2['Part tanpa PKB']; _rb2['Bahan']=_rb2['Bahan (dengan PKB)']+_rb2['Bahan tanpa PKB']; _rb2['Revenue']=_rb2['Total Revenue']
        _hdr=['Cabang','UE (Aktual / Target)','% ACH UE','Jasa (Aktual / Target)','% ACH Jasa','Part (Aktual / Target)','% ACH Part','Bahan (Aktual / Target)','% ACH Bahan','Revenue (Aktual / Target)','% ACH Rev','Status']
        _trs=[]; _sumA={m:0.0 for m in TARGET_METRICS}; _sumT={m:0.0 for m in TARGET_METRICS}
        def _pcell(a,t): 
            p=(a/t*100 if t>0 else 0.0); return f'<span class="rk-badge {RK.pct_tone(p) if t>0 else "gray"}">{p:.1f}%</span>'
        def _acell(m,a,t): return (f'{int(round(a)):,} / {int(round(t)):,}' if m=='UE' else f'{RK.fmt_jt(a)}<br><small style="color:#64748b">{RK.fmt_jt(t)}</small>')
        def _rowcells(name,A,Tt):
            cs=[f'<b>{_html.escape(str(name))}</b>']
            for m in TARGET_METRICS: cs+= [_acell(m,A[m],Tt[m]),_pcell(A[m],Tt[m])]
            _p=(A['Revenue']/Tt['Revenue']*100 if Tt['Revenue']>0 else 0.0); _sl,_st=RK.status_label(_p)
            cs.append(RK.status_badge(_sl,_st) if Tt['Revenue']>0 else RK.status_badge('Target belum diisi','gray'))
            return cs
        for _,r_ in _rb2.iterrows():
            tg_=_cur.get(r_['Cabang'],{}) or {}
            A_={m:float(r_[m]) for m in TARGET_METRICS}; T_={m:float(tg_.get(m,0) or 0) for m in TARGET_METRICS}
            for m in TARGET_METRICS: _sumA[m]+=A_[m]; _sumT[m]+=T_[m]
            _trs.append(_rowcells(r_['Cabang'],A_,T_))
        if _trs:
            _trs.append(_rowcells('TOTAL KONSOLIDASI' if scope=='Semua Cabang' else 'TOTAL',_sumA,_sumT))
            st.markdown(RK.table_html(_hdr,_trs,num_cols=tuple(range(1,11)),bold_last=True),unsafe_allow_html=True)
        st.caption('Target di atas dapat disesuaikan kapan saja melalui menu Data Input → Input Target Perusahaan.')


if page == '📊 Comparative Analytics':
    section('<span class="msr">bar_chart</span> Comparative Analytics','Perbandingan antar bulan, kuartal atau tahun. Dashboard utama sengaja dibuat ringkas; analisa perbandingan berada di halaman ini.')
    # ================= V17: komparasi Periode 1 vs Periode 2 (sesuai referensi) =================
    _cu={k:UNF[k] for k in ('master','parts','issue','cons')}
    _alld=pd.concat([pd.to_datetime(v['Tanggal'],errors='coerce') for v in _cu.values() if isinstance(v,pd.DataFrame) and not v.empty and 'Tanggal' in v.columns]+[pd.Series(dtype='datetime64[ns]')],ignore_index=True)
    _gmap={'Per Bulan':'Bulan','Per Kuartal (3 Bln)':'Kuartal','Per Semester (6 Bln)':'Semester','Per Tahun (YoY)':'Tahun','3 Bulan Terakhir':'LAST3','6 Bulan Terakhir':'LAST6'}
    with st.container(key='blk_cmp'):   # judul + dropdown mode komparasi dalam SATU blok
        _gh1,_gh2=st.columns([3.4,1.4],vertical_alignment='center')
        with _gh1: st.markdown(RK.head_card('Comparative Analytics & Pertumbuhan Workshop',f'Analisis komparasi performa, evaluasi tren periode, dan benchmarking Core vs Non-PKB • Cabang: <b>{_html.escape(str(active_branch))}</b> (mengikuti filter sidebar)','compare_arrows','purple'),unsafe_allow_html=True)
        with _gh2:
            if st.session_state.get('rk_tabs_cmp_gran') not in _gmap: st.session_state['rk_tabs_cmp_gran']='Per Bulan'
            _gsel=st.selectbox('Mode komparasi',list(_gmap.keys()),key='rk_tabs_cmp_gran',label_visibility='collapsed',help='Pilih granularitas komparasi. "3 / 6 Bulan Terakhir" otomatis membandingkan bulan-bulan terakhir yang ada di data.')
    _g=_gmap[_gsel]
    def _cmp_data(key,gran):
        s_,e_=RK.period_bounds(key,gran)
        return [RK.slice_dates(_cu[k],'Tanggal',s_,e_) for k in ('master','parts','issue','cons')]
    def _mx(d,r):
        ue=int(d[0]['UE'].sum()) if isinstance(d[0],pd.DataFrame) and 'UE' in d[0].columns else 0
        tot=float(r['total_revenue']); return dict(ue=ue,core=float(r['revenue_core']),part=float(r['direct_part']),bahan=float(r['direct_bahan']),direct=float(r['direct_part']+r['direct_bahan']),total=tot,rev_ue=(tot/ue if ue else 0.0))
    def _gb(a,b):
        g_=RK.growth_pct(b,a); return RK.fmt_growth(g_),('green' if (g_ or 0)>=0 else 'red')
    if _g in ('LAST3','LAST6'):
        # ================= 3 / 6 bulan terakhir: bandingkan semua bulan secara berurutan =================
        _nm=3 if _g=='LAST3' else 6
        _lastd=_alld.dropna().max() if len(_alld.dropna()) else pd.NaT
        if pd.isna(_lastd):
            st.info('Belum ada data tanggal untuk dibandingkan.')
        else:
            _mk=[pd.Timestamp(_lastd.year,_lastd.month,1)-pd.DateOffset(months=k) for k in range(_nm-1,-1,-1)]
            _mk=[f'{t.year}-{t.month:02d}' for t in _mk]
            _ml=[RK.period_label(k,'Bulan') for k in _mk]; _ms_=[MONTH_NAMES[int(k[5:7])-1][:3]+' '+k[:4] for k in _mk]
            _MD=[_cmp_data(k,'Bulan') for k in _mk]; _MR=[revenue_components(*d) for d in _MD]; _MM=[_mx(d,r) for d,r in zip(_MD,_MR)]
            _mg=[None]+[RK.growth_pct(_MM[i]['total'],_MM[i-1]['total']) for i in range(1,len(_MM))]
            _m1,_m2=_MM[-2],_MM[-1]; _l1,_l2=_ml[-2],_ml[-1]
            st.markdown(RK.dark_banner(f'Membandingkan <b>{_nm} bulan terakhir</b> pada data: <span class="c1">{_html.escape(" • ".join(_ml))}</span>. Kartu di bawah menampilkan <b>{_html.escape(_l2)}</b> dibanding <b>{_html.escape(_l1)}</b>.',_html.escape(str(active_branch))),unsafe_allow_html=True)
            _gu,_tu=_gb(_m1['ue'],_m2['ue']); _gc,_tc=_gb(_m1['core'],_m2['core']); _gt,_tt=_gb(_m1['total'],_m2['total']); _gd,_td=_gb(_m1['direct'],_m2['direct'])
            st.markdown(RK.grid([
                RK.kpi_card('Volume Unit Entry (UE)',f'{_m2["ue"]:,} UE <small><s>({_m1["ue"]:,} UE)</s></small>',f'Perubahan: <b>{_m2["ue"]-_m1["ue"]:+,} Unit</b>','directions_car','blue',_gu),
                RK.kpi_card('Revenue Core (ada PKB)',fmt(_m2['core']),f'Delta: <b>{"+" if _m2["core"]-_m1["core"]>=0 else "-"}{fmt(abs(_m2["core"]-_m1["core"]))}</b>','build','blue',_gc),
                RK.kpi_card('Part & Bahan Non-PKB',fmt(_m2['direct']),f'Baseline: {fmt(_m1["direct"])} ({"+" if _m2["direct"]-_m1["direct"]>=0 else "-"}{fmt(abs(_m2["direct"]-_m1["direct"]))})','storefront','orange',_gd),
                RK.kpi_card('Total Revenue',fmt(_m2['total']),f'Rev/UE: <b>{fmt(_m2["rev_ue"])}</b> ({RK.fmt_growth(RK.growth_pct(_m2["rev_ue"],_m1["rev_ue"]))})','payments','purple',_gt)],230),unsafe_allow_html=True)
            st.markdown(RK.head_card(f'Diagram Tren {_nm} Bulan Terakhir','Batang = nilai per bulan (bulan terakhir berwarna biru), garis oranye = pertumbuhan dibanding bulan sebelumnya (%)','show_chart','blue'),unsafe_allow_html=True)
            _cm=st.segmented_control('Metrik diagram',['Total Revenue','Unit Entry (UE)','Revenue Core','Part & Bahan Non-PKB'],default='Total Revenue',key='rk_cmp_chart_metric',label_visibility='collapsed') or 'Total Revenue'
            _fk={'Total Revenue':'total','Unit Entry (UE)':'ue','Revenue Core':'core','Part & Bahan Non-PKB':'direct'}[_cm]
            _dv=1.0 if _cm=='Unit Entry (UE)' else 1e6; _un='UE' if _cm=='Unit Entry (UE)' else 'Jt'
            _vv=[m_[_fk]/_dv for m_ in _MM]; _gg=[None]+[RK.growth_pct(_MM[i][_fk],_MM[i-1][_fk]) for i in range(1,len(_MM))]
            _fc=go.Figure()
            _fc.add_trace(go.Bar(x=_ms_,y=_vv,name=_cm,marker_color=['#cbd5e1']*(len(_vv)-1)+['#2563eb'],text=[(f'{v:,.0f}' if _un=='UE' else f'{v:,.1f}') for v in _vv],textposition='outside',cliponaxis=False))
            _fc.add_trace(go.Scatter(x=_ms_,y=_gg,name='Pertumbuhan vs bulan sebelumnya (%)',mode='lines+markers',yaxis='y2',line=dict(color='#f59e0b',width=3),connectgaps=False))
            _fc.update_layout(font=dict(color='#0f172a',size=12),height=360,margin=dict(l=8,r=8,t=24,b=8),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',hovermode='x unified',legend=dict(orientation='h',y=-.15),
                yaxis=dict(title=_un,gridcolor='#e8eef6'),yaxis2=dict(title='%',overlaying='y',side='right',showgrid=False,zeroline=False))
            st.plotly_chart(fixfig(_fc),use_container_width=True,config=PCFG)
            st.markdown(RK.head_card(f'Matriks Perbandingan {_nm} Bulan Terakhir',f'Cakupan: {_html.escape(str(active_branch))} • Total Revenue = Revenue Core (PKB) + Part Non-PKB + Bahan Non-PKB','bar_chart','blue'),unsafe_allow_html=True)
            _mrows=[]
            for i,m_ in enumerate(_MM):
                _dU=(f'{m_["ue"]-_MM[i-1]["ue"]:+,}' if i else '-')
                _gbadge=('-' if _mg[i] is None else f'<span class="rk-badge {"green" if _mg[i]>=0 else "red"}">{RK.fmt_growth(_mg[i])}</span>')
                _mrows.append([f'<b>{_html.escape(_ml[i])}</b>',f'{m_["ue"]:,}',_dU,fmt(m_['core']),fmt(m_['part']),fmt(m_['bahan']),f'<b>{fmt(m_["total"])}</b>',_gbadge,fmt(m_['rev_ue'])])
            st.markdown(RK.table_html(['Periode','UE','Δ UE vs bln lalu','Revenue Core','Part Non-PKB','Bahan Non-PKB','Total Revenue','Pertumbuhan MoM','Rev / UE'],_mrows,num_cols=tuple(range(1,9)),hl_cols={6:'hl2'}),unsafe_allow_html=True)
            if active_branch=='Semua Cabang':
                _bm=[revenue_components_by_branch(*d,BRANCHES) for d in _MD]
                _bnames=[c for c in _bm[-1]['Cabang'] if c!='TOTAL SEMUA CABANG']
                if len(_bnames)>1:
                    st.markdown(RK.head_card(f'Total Revenue per Cabang — {_nm} Bulan Terakhir','Perbandingan kontribusi tiap cabang pada setiap bulan','apartment','purple'),unsafe_allow_html=True)
                    _brow=[]
                    for c in _bnames:
                        vals=[float(b_[b_['Cabang']==c]['Total Revenue'].iloc[0]) if not b_[b_['Cabang']==c].empty else 0.0 for b_ in _bm]
                        gg=RK.growth_pct(vals[-1],vals[0])
                        _brow.append([f'<b>{_html.escape(str(c))}</b>']+[fmt(v) for v in vals]+[f'<span class="rk-badge {"green" if (gg or 0)>=0 else "red"}">{RK.fmt_growth(gg)}</span>'])
                    st.markdown(RK.table_html(['Cabang']+_ms_+[f'Pertumbuhan {_ms_[0]} → {_ms_[-1]}'],_brow,num_cols=tuple(range(1,len(_ms_)+2))),unsafe_allow_html=True)
            # ---- temuan kunci ----
            _tots=[m_['total'] for m_ in _MM]; _ues=[m_['ue'] for m_ in _MM]
            _ib=max(range(len(_tots)),key=lambda i:_tots[i]); _iw=min(range(len(_tots)),key=lambda i:_tots[i])
            _gall=RK.growth_pct(_tots[-1],_tots[0]); _uall=RK.growth_pct(_ues[-1],_ues[0])
            _sh=(_m2['direct']/_m2['total']*100) if _m2['total'] else 0.0
            _avg_ry=(sum(_tots)/max(sum(_ues),1))
            _t1=(f'Revenue tertinggi <b>{_html.escape(_ml[_ib])}</b> ({fmt(_tots[_ib])}), terendah <b>{_html.escape(_ml[_iw])}</b> ({fmt(_tots[_iw])}). Dari {_html.escape(_ml[0])} ke {_html.escape(_ml[-1])}: revenue {RK.fmt_growth(_gall)}.')
            _t2=(f'UE bulan terakhir <b>{_ues[-1]:,}</b> ({RK.fmt_growth(_uall)} dibanding {_html.escape(_ml[0])}); rata-rata <b>{sum(_ues)/len(_ues):,.0f} UE/bulan</b>.')
            _t3=(f'Revenue/UE bulan terakhir <b>{fmt(_m2["rev_ue"])}</b> vs rata-rata periode {fmt(_avg_ry)}. Non-PKB menyumbang {_sh:.1f}% dari total revenue bulan terakhir.')
            st.markdown(RK.insight_band(f'Temuan Kunci {_nm} Bulan Terakhir',[('Pola Revenue',_t1,'g'),('Volume Unit (UE)',_t2,'o'),('Yield & Non-PKB',_t3,'c')]),unsafe_allow_html=True)
    else:
        _opts=RK.period_options(_alld,_g)
        if len(_opts)<2:
            st.info('Butuh minimal dua periode data untuk komparasi Periode 1 vs Periode 2.')
        else:
            _k1,_k2=f'rk_p1_{_g}',f'rk_p2_{_g}'
            if st.session_state.get(_k1) not in _opts: st.session_state[_k1]=_opts[-2]
            if st.session_state.get(_k2) not in _opts: st.session_state[_k2]=_opts[-1]
            _q1,_q2,_q3=st.columns([1,1,1.4])
            with _q1: st.selectbox('Periode 1',_opts,key=_k1,format_func=lambda k:RK.period_label(k,_g))
            with _q2: st.selectbox('Periode 2',_opts,key=_k2,format_func=lambda k:RK.period_label(k,_g))
            _pairs=RK.shortcut_pairs(_opts,3)
            if _pairs:
                _sc=st.columns(len(_pairs)+3)
                _sc[0].caption('Shortcut cepat:')
                for _ci,(a_,b_) in enumerate(_pairs):
                    _sc[_ci+1].button(f'{RK.period_label(a_,_g).split(" ")[0] if _g=="Bulan" else a_} vs {RK.period_label(b_,_g).split(" ")[0] if _g=="Bulan" else b_}',key=f'rk_sc_{_g}_{_ci}',use_container_width=True,on_click=lambda a=a_,b=b_,k1=_k1,k2=_k2:st.session_state.update({k1:a,k2:b}))
            _p1,_p2=st.session_state[_k1],st.session_state[_k2]; _l1,_l2=RK.period_label(_p1,_g),RK.period_label(_p2,_g)
            _d1=_cmp_data(_p1,_g); _d2=_cmp_data(_p2,_g)
            _r1=revenue_components(*_d1); _r2=revenue_components(*_d2)
            _m1,_m2=_mx(_d1,_r1),_mx(_d2,_r2)
            _gu,_tu=_gb(_m1['ue'],_m2['ue']); _gc,_tc=_gb(_m1['core'],_m2['core']); _gt,_tt=_gb(_m1['total'],_m2['total']); _gd,_td=_gb(_m1['direct'],_m2['direct'])
            st.markdown(RK.grid([
                RK.kpi_card('Volume Unit Entry (UE)',f'{_m2["ue"]:,} UE <small><s>({_m1["ue"]:,} UE)</s></small>',f'Perubahan: <b>{_m2["ue"]-_m1["ue"]:+,} Unit</b>','directions_car','blue',_gu),
                RK.kpi_card('Revenue Core (ada PKB)',fmt(_m2['core']),f'Delta: <b>{"+" if _m2["core"]-_m1["core"]>=0 else "-"}{fmt(abs(_m2["core"]-_m1["core"]))}</b>','build','blue',_gc),
                RK.kpi_card('Part & Bahan Non-PKB',fmt(_m2['direct']),f'Baseline: {fmt(_m1["direct"])} ({"+" if _m2["direct"]-_m1["direct"]>=0 else "-"}{fmt(abs(_m2["direct"]-_m1["direct"]))})','storefront','orange',_gd),
                RK.kpi_card('Total Revenue',fmt(_m2['total']),f'Rev/UE: <b>{fmt(_m2["rev_ue"])}</b> ({RK.fmt_growth(RK.growth_pct(_m2["rev_ue"],_m1["rev_ue"]))})','payments','purple',_gt)],230),unsafe_allow_html=True)
            _b1=revenue_components_by_branch(*_d1,BRANCHES); _b2=revenue_components_by_branch(*_d2,BRANCHES)
            # ---- diagram batang P1 vs P2 per cabang (gaya halaman Target) ----
            _bn=[c for c in _b2['Cabang'] if c!='TOTAL SEMUA CABANG']
            if active_branch!='Semua Cabang': _bn=[c for c in _bn if c==active_branch]
            if _bn:
                st.markdown(RK.head_card(f'Diagram Perbandingan: {_l1} vs {_l2}','Batang abu-abu = Periode 1, biru = Periode 2, garis oranye = pertumbuhan (%)','show_chart','blue'),unsafe_allow_html=True)
                _cm=st.segmented_control('Metrik diagram',['Total Revenue','Unit Entry (UE)','Revenue Core','Part & Bahan Non-PKB'],default='Total Revenue',key='rk_cmp_chart_metric',label_visibility='collapsed') or 'Total Revenue'
                def _val(df,c,mm_):
                    r_=df[df['Cabang']==c]
                    if r_.empty: return 0.0
                    r_=r_.iloc[0]
                    if mm_=='Unit Entry (UE)': return float(r_['UE'])
                    if mm_=='Revenue Core': return float(r_['Total Revenue']-r_['Part tanpa PKB']-r_['Bahan tanpa PKB'])
                    if mm_=='Part & Bahan Non-PKB': return float(r_['Part tanpa PKB']+r_['Bahan tanpa PKB'])
                    return float(r_['Total Revenue'])
                _dv=1.0 if _cm=='Unit Entry (UE)' else 1e6; _un='UE' if _cm=='Unit Entry (UE)' else 'Jt'
                _v1=[_val(_b1,c,_cm) for c in _bn]; _v2=[_val(_b2,c,_cm) for c in _bn]
                _gr=[((b_-a_)/a_*100 if a_>0 else None) for a_,b_ in zip(_v1,_v2)]
                _lb=[str(c).replace('Kharisma ','') for c in _bn]
                _fc=go.Figure()
                _fc.add_trace(go.Bar(x=_lb,y=[v/_dv for v in _v1],name=_l1,marker_color='#cbd5e1',text=[(f'{v/_dv:,.0f}' if _un=='UE' else f'{v/_dv:,.1f}') for v in _v1],textposition='outside',cliponaxis=False))
                _fc.add_trace(go.Bar(x=_lb,y=[v/_dv for v in _v2],name=_l2,marker_color='#2563eb',text=[(f'{v/_dv:,.0f}' if _un=='UE' else f'{v/_dv:,.1f}') for v in _v2],textposition='outside',cliponaxis=False))
                _fc.add_trace(go.Scatter(x=_lb,y=_gr,name='Pertumbuhan %',mode='lines+markers',yaxis='y2',line=dict(color='#f59e0b',width=3),connectgaps=False))
                _fc.update_layout(barmode='group',font=dict(color='#0f172a',size=12),height=360,margin=dict(l=8,r=8,t=24,b=8),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',hovermode='x unified',legend=dict(orientation='h',y=-.15),
                    yaxis=dict(title=_un,gridcolor='#e8eef6'),yaxis2=dict(title='%',overlaying='y',side='right',showgrid=False,zeroline=False))
                st.plotly_chart(fixfig(_fc),use_container_width=True,config=PCFG)
            _hdr=['Cabang','UE (P1)','UE (P2)','Δ UE','Core P1','Core P2','Part Non-PKB','Bahan Non-PKB','Total Rev (P1)','Total Rev (P2)','Pertumbuhan','Rev / UE (P2)']
            _rows=[]
            def _core(r): return float(r['Total Revenue']-r['Part tanpa PKB']-r['Bahan tanpa PKB'])
            _names=[c for c in _b2['Cabang'] if c!='TOTAL SEMUA CABANG']
            if active_branch!='Semua Cabang': _names=[c for c in _names if c==active_branch]
            _t1=dict(ue=0,core=0.0,tot=0.0); _t2=dict(ue=0,core=0.0,part=0.0,bahan=0.0,tot=0.0)
            for c in _names:
                a=_b1[_b1['Cabang']==c]; b=_b2[_b2['Cabang']==c]
                if a.empty or b.empty: continue
                a=a.iloc[0]; b=b.iloc[0]; ca=_core(a); cb=_core(b); ga=RK.growth_pct(b['Total Revenue'],a['Total Revenue'])
                _t1['ue']+=int(a['UE']); _t1['core']+=ca; _t1['tot']+=float(a['Total Revenue']); _t2['ue']+=int(b['UE']); _t2['core']+=cb; _t2['part']+=float(b['Part tanpa PKB']); _t2['bahan']+=float(b['Bahan tanpa PKB']); _t2['tot']+=float(b['Total Revenue'])
                _rows.append([_html.escape(str(c)),f'{int(a["UE"]):,}',f'{int(b["UE"]):,}',f'{int(b["UE"])-int(a["UE"]):+,}',fmt(ca),fmt(cb),fmt(b['Part tanpa PKB']),fmt(b['Bahan tanpa PKB']),fmt(a['Total Revenue']),fmt(b['Total Revenue']),f'<span class="rk-badge {"green" if (ga or 0)>=0 else "red"}">{RK.fmt_growth(ga)}</span>',fmt(b['Revenue/UE'])])
            if _rows:
                _gt2=RK.growth_pct(_t2['tot'],_t1['tot'])
                _rows.append(['TOTAL'+(' SEMUA CABANG' if active_branch=='Semua Cabang' else f' ({_html.escape(str(active_branch).upper())})'),f'{_t1["ue"]:,}',f'{_t2["ue"]:,}',f'{_t2["ue"]-_t1["ue"]:+,}',fmt(_t1['core']),fmt(_t2['core']),fmt(_t2['part']),fmt(_t2['bahan']),fmt(_t1['tot']),fmt(_t2['tot']),RK.fmt_growth(_gt2),fmt(_t2['tot']/_t2['ue'] if _t2['ue'] else 0)])
                st.markdown(RK.head_card(f'Matriks Perbandingan Performa ({_l1} vs {_l2})',f'Cakupan: {_html.escape(str(active_branch))} • Total Revenue = Revenue Core (PKB) + Part Non-PKB + Bahan Non-PKB','bar_chart','blue'),unsafe_allow_html=True)
                st.markdown(RK.table_html(_hdr,_rows,num_cols=tuple(range(1,12)),bold_last=True,hl_cols={4:'hl1',5:'hl1',8:'hl2',9:'hl2'}),unsafe_allow_html=True)
            st.markdown(RK.insight_band(f'Temuan Kunci Komparatif ({_l1} vs {_l2})',RK.compare_insights(_m1,_m2,_l1,_l2)),unsafe_allow_html=True)
    # ---- Tren Bulanan: Target vs Pencapaian Aktual (dipindah dari halaman Target) ----
    _tgt_all=load_targets()
    try: _dall=daily_components(_cu['master'],_cu['parts'],_cu['issue'],_cu['cons'])
    except Exception: _dall=pd.DataFrame()
    _today_=datetime.date.today(); _ddv=_alld.dropna()
    try: _ty=int(st.session_state.get('flt_year'))
    except Exception: _ty=int(_ddv.max().year) if len(_ddv) else _today_.year
    _fm_=st.session_state.get('flt_month','Semua')
    if _fm_ in MONTH_NAMES: _tm=MONTH_NAMES.index(_fm_)+1
    else:
        _yd_=_ddv[_ddv.dt.year==_ty]; _tm=int(_yd_.max().month) if len(_yd_) else (_today_.month if _ty==_today_.year else 12)
    _mkey=f'{_ty}-{_tm:02d}'
    _ms=RK.monthly_series(_dall,_tgt_all,active_branch,BRANCHES,last_n=6,upto=_mkey)
    if not _ms.empty:
        st.markdown(RK.head_card('Tren Bulanan: Target vs Pencapaian Aktual',f'Target bulanan vs realisasi aktual, 6 bulan terakhir hingga {MONTH_NAMES[_tm-1]} {_ty} • Cabang: {_html.escape(str(active_branch))}','show_chart','red'),unsafe_allow_html=True)
        _mv=st.segmented_control('Tampilan tren bulanan',['Total Revenue','Unit Entry (UE)','Jasa, Part, Bahan'],default='Total Revenue',key='rk_tabs_cmp_month',label_visibility='collapsed') or 'Total Revenue'
        _f2=go.Figure()
        if _mv=='Jasa, Part, Bahan':
            for m,c in [('Jasa','#10b981'),('Part','#8b5cf6'),('Bahan','#f59e0b')]: _f2.add_trace(go.Bar(x=_ms['Label'],y=_ms[m+'_A']/1e6,name=f'{m} aktual (Jt)',marker_color=c))
            _f2.add_trace(go.Scatter(x=_ms['Label'],y=(_ms['Jasa_T']+_ms['Part_T']+_ms['Bahan_T'])/1e6,name='Target Jasa+Part+Bahan (Jt)',mode='lines+markers',line=dict(color='#ef4444',dash='dash')))
            _f2.update_layout(barmode='group',yaxis=dict(title='Jt',gridcolor='#e8eef6'))
        else:
            _mm_='UE' if _mv=='Unit Entry (UE)' else 'Revenue'; _dvm=1.0 if _mm_=='UE' else 1e6
            _f2.add_trace(go.Bar(x=_ms['Label'],y=_ms[_mm_+'_T']/_dvm,name='Target',marker_color='#cbd5e1'))
            _f2.add_trace(go.Bar(x=_ms['Label'],y=_ms[_mm_+'_A']/_dvm,name='Aktual',marker_color='#2563eb'))
            _pp_=[(a_/t_*100 if t_>0 else None) for a_,t_ in zip(_ms[_mm_+'_A'],_ms[_mm_+'_T'])]
            _f2.add_trace(go.Scatter(x=_ms['Label'],y=_pp_,name='% Pencapaian',mode='lines+markers',yaxis='y2',line=dict(color='#f59e0b',width=3),connectgaps=False))
            _f2.update_layout(barmode='group',yaxis=dict(title=('UE' if _mm_=='UE' else 'Jt'),gridcolor='#e8eef6'),yaxis2=dict(title='%',overlaying='y',side='right',showgrid=False,rangemode='tozero'))
        _f2.update_layout(font=dict(color='#0f172a',size=12),height=340,margin=dict(l=8,r=8,t=12,b=8),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',hovermode='x unified',legend=dict(orientation='h',y=-.15))
        st.plotly_chart(fixfig(_f2),use_container_width=True,config=PCFG)
        _wt=_ms[_ms['Revenue_T']>0]
        _avg_p=float(_wt['Pct'].mean()) if len(_wt) else 0.0
        if len(_wt): _bi=_wt['Pct'].idxmax(); _best=(str(_wt.loc[_bi,'Label']),float(_wt.loc[_bi,'Pct']))
        else: _best=('-',0.0)
        _cur_row=_ms[_ms['Key']==_mkey]
        _cur_rev=float(_cur_row['Revenue_A'].iloc[0]) if len(_cur_row) else 0.0
        st.markdown(RK.grid([RK.mini_card('Rata-rata pencapaian revenue',f'{_avg_p:.1f}%','dari target bulan yang terisi'),
            RK.mini_card('Pencapaian tertinggi',_best[0],f'({_best[1]:.1f}%)','blue'),
            RK.mini_card(f'Realisasi {MONTH_NAMES[_tm-1]}',RK.fmt_jt(_cur_rev),'revenue aktual bulan acuan','yellow')],260),unsafe_allow_html=True)

    # Pencapaian revenue per cabang per bulan (acuan: periode header Laporan Estimasi SA bila aktif)
    if not f.empty and 'Cabang' in f.columns:
        st.markdown('### :material/apartment: Pencapaian Revenue per Cabang per Bulan')
        st.caption('Revenue PKB (Jasa + Sparepart + Bahan) mengikuti Laporan Estimasi Service Advisor: bulan = periode header laporan, bukan tanggal PKB. Part tanpa PKB dan Consumable ditambahkan sesuai tanggal transaksinya.' if EST_REF_ACTIVE else 'Acuan Laporan Estimasi SA tidak aktif / belum diupload: bulan mengikuti tanggal PKB.')
        _g=f.groupby(['Cabang','Periode Bulan'],as_index=False).agg(UE=('UE','sum'),Jasa=('Jasa','sum'),Sparepart=('Sparepart','sum'),Bahan=('Bahan','sum'),**{'Revenue Core':('Revenue','sum')})
        def _direct(df,datecol,valcol,filt=None):
            if not isinstance(df,pd.DataFrame) or df.empty or datecol not in df.columns or valcol not in df.columns or 'Cabang' not in df.columns: return pd.DataFrame(columns=['Cabang','Periode Bulan','v'])
            z=df if filt is None else df[filt(df)]
            z=z.assign(**{'Periode Bulan':pd.to_datetime(z[datecol],errors='coerce').dt.to_period('M').astype(str)})
            z=z[z['Periode Bulan']!='NaT']
            return z.groupby(['Cabang','Periode Bulan'],as_index=False).agg(v=(valcol,'sum'))
        _dp=_direct(parts,'Tanggal','Sales Amount',lambda d:d['Sumber Penjualan'].eq('Billing Part') if 'Sumber Penjualan' in d else d['Sales Amount'].notna()).rename(columns={'v':'Part tanpa PKB'})
        _db=_direct(cons,'Tanggal','Harga Net').rename(columns={'v':'Bahan tanpa PKB'})
        _g=_g.merge(_dp,on=['Cabang','Periode Bulan'],how='outer').merge(_db,on=['Cabang','Periode Bulan'],how='outer').fillna(0)
        _g['Total Revenue']=_g['Revenue Core']+_g['Part tanpa PKB']+_g['Bahan tanpa PKB']
        _g['Revenue/UE']=np.where(_g['UE']>0,_g['Total Revenue']/_g['UE'],0)
        _g=_g.sort_values(['Periode Bulan','Cabang'])
        pv=_g.pivot_table(index='Cabang',columns='Periode Bulan',values='Total Revenue',aggfunc='sum',fill_value=0)
        if active_branch=='Semua Cabang': pv=pv.reindex(list(dict.fromkeys(list(BRANCHES)+list(pv.index))),fill_value=0)  # cabang tanpa data / cabang baru tetap tampil
        pv['TOTAL']=pv.sum(axis=1); pv.loc['SEMUA CABANG']=pv.sum()
        st.dataframe(pv.reset_index().style.format({c:'{:,.0f}' for c in pv.columns}),use_container_width=True,hide_index=True)
        with st.expander('Rincian per cabang & bulan'):
            st.dataframe(_g,use_container_width=True,hide_index=True)
        download_table(_g,':material/download: Download Excel — Pencapaian Revenue per Cabang per Bulan','pencapaian_revenue_cabang_bulan.xlsx','dl_cab_bulan')


if page == '🧠 Management Insight':
    # ================= V23: Management Insight — status target, diagnosis gap, strategi & rencana aksi UE + Revenue =================
    _mi_today=datetime.date.today()
    _mi_m,_mi_p,_mi_i,_mi_c=UNF['master'],UNF['parts'],UNF['issue'],UNF['cons']   # periode target mengikuti Filter Waktu, tanpa memotong per hari
    _mi_dd=pd.to_datetime(_mi_m['Tanggal'],errors='coerce').dropna() if isinstance(_mi_m,pd.DataFrame) and not _mi_m.empty and 'Tanggal' in _mi_m.columns else pd.Series(dtype='datetime64[ns]')
    _mi_dy=int(_mi_dd.max().year) if len(_mi_dd) else _mi_today.year
    try: _mi_y=int(st.session_state.get('flt_year'))
    except Exception: _mi_y=_mi_dy
    _mi_fm=st.session_state.get('flt_month','Semua')
    if _mi_fm in MONTH_NAMES: _mi_mo=MONTH_NAMES.index(_mi_fm)+1
    else:
        _mi_yd=_mi_dd[_mi_dd.dt.year==_mi_y]
        _mi_mo=int(_mi_yd.max().month) if len(_mi_yd) else (_mi_today.month if _mi_y==_mi_today.year else 12)
    _mi_name=MONTH_NAMES[_mi_mo-1]; _mi_key=f'{_mi_y}-{_mi_mo:02d}'; _mi_days=calendar.monthrange(_mi_y,_mi_mo)[1]
    _mi_tg=load_targets().get(_mi_key,{}) or {}
    _mi_scope=active_branch
    _mi_bset=list(dict.fromkeys(list(BRANCHES)+[b_ for b_ in _mi_tg if b_ not in BRANCHES]))
    def _mi_t(m): return float(sum((_mi_tg.get(b_,{}).get(m,0) or 0) for b_ in _mi_bset)) if _mi_scope=='Semua Cabang' else float((_mi_tg.get(_mi_scope,{}) or {}).get(m,0) or 0)
    _mi_T={m:_mi_t(m) for m in TARGET_METRICS}
    def _mi_in(df):
        if not isinstance(df,pd.DataFrame) or df.empty or 'Tanggal' not in df.columns: return df
        d_=pd.to_datetime(df['Tanggal'],errors='coerce'); return df[((d_.dt.year==_mi_y)&(d_.dt.month==_mi_mo)).values].copy()
    _mm,_pp,_ii,_cc=_mi_in(_mi_m),_mi_in(_mi_p),_mi_in(_mi_i),_mi_in(_mi_c)
    _mi_daily=daily_components(_mm,_pp,_ii,_cc).reindex(pd.date_range(pd.Timestamp(_mi_y,_mi_mo,1),periods=_mi_days,freq='D'),fill_value=0.0)
    _mi_R={m:float(_mi_daily[m].sum()) for m in TARGET_METRICS}
    if pd.Timestamp(_mi_y,_mi_mo,1).date()>_mi_today: _mi_el=0
    elif pd.Timestamp(_mi_y,_mi_mo,_mi_days).date()<_mi_today: _mi_el=_mi_days
    else: _mi_el=_mi_today.day
    # ---- data pendukung strategi (bulan acuan) ----
    _mi_rc=revenue_components(_mm,_pp,_ii,_cc); _mi_tot=float(_mi_rc['total_revenue'])
    _mi_ctx={'direct_part':float(_mi_rc['direct_part']),'direct_bahan':float(_mi_rc['direct_bahan']),'nonpkb_share':((float(_mi_rc['direct_part']+_mi_rc['direct_bahan'])/_mi_tot*100) if _mi_tot else None),'no_sales':int(len(qc_no_sales)) if isinstance(qc_no_sales,pd.DataFrame) else 0}
    if isinstance(_mm,pd.DataFrame) and not _mm.empty:
        _ue_m=float(_mm['UE'].sum()) if 'UE' in _mm.columns else 0.0
        if 'Booking Type' in _mm.columns and _ue_m>0:
            _btm=_mm['Booking Type'].map(clean).str.upper()
            _mi_ctx['booking_pct']=float(_mm.loc[_btm=='BOOKING','UE'].sum())/_ue_m*100
        _wl_m=RK.dow_workload(_mm)
        if not _wl_m.empty:
            _lw=_wl_m.loc[_wl_m['Rata2'].idxmin()]; _mi_ctx['weak_day']=(str(_lw['Hari']),float(_lw['Rata2'])); _mi_ctx['avg_day_ue']=float(_wl_m['Rata2'].mean())
        if 'Service Advisor' in _mm.columns:
            _sam=_mm[_mm['Service Advisor'].map(clean)!=''].copy()
            if not _sam.empty:
                _sam['Service Advisor']=sa_disp(_sam['Service Advisor'])
                _sag=_sam.groupby('Service Advisor').agg(UE=('UE','sum'),Rev=('Revenue','sum'))
                _mi_ctx['sa_n']=int(len(_sag)); _sag=_sag[_sag['UE']>=3] if (_sag['UE']>=3).sum()>=2 else _sag
                if len(_sag)>=1:
                    _sag['RPU']=np.where(_sag['UE']>0,_sag['Rev']/_sag['UE'],0.0)
                    _hi=_sag['RPU'].idxmax(); _lo=_sag['RPU'].idxmin()
                    _mi_ctx['sa_high']=(_html.escape(str(_hi)),float(_sag.loc[_hi,'RPU'])); _mi_ctx['sa_low']=(_html.escape(str(_lo)),float(_sag.loc[_lo,'RPU']))
    MI=RK.mi_strategy(_mi_T,_mi_R,_mi_el,_mi_days,_mi_ctx)

    def _mi_sec(n,title,sub): return f'<div class="mi-sec"><span class="n">{n}</span><div><b>{_html.escape(title)}</b><small>{_html.escape(sub)}</small></div></div>'
    def _u(t): return _html.escape(_html.unescape(str(t)))
    st.markdown(RK.head_card('Management Insight: Strategi Menjaga UE & Revenue Sesuai Target',f'Periode analisa: <b>{_mi_name} {_mi_y}</b> ({_mi_el} dari {_mi_days} hari berjalan) &nbsp;•&nbsp; Cabang: <b>{_html.escape(str(_mi_scope))}</b>','psychology','purple'),unsafe_allow_html=True)
    if not MI['has_target']:
        st.info(f'Target {_mi_name} {_mi_y} untuk cabang ini belum diinput. Isi lewat menu **Target & Pencapaian → Input Target**. Strategi umum tetap ditampilkan di bawah.')

    # ---------- 1. Status pencapaian ----------
    st.markdown(_mi_sec(1,'Status Pencapaian Target','Posisi UE, Revenue, dan Revenue per UE terhadap target bulan ini dan pace harian'),unsafe_allow_html=True)
    def _mi_card(label,icon,tone,s,is_ue=False,is_rpu=False):
        if is_rpu:
            act=fmt(s['actual']); tgt=fmt(s['target']) if s['target']>0 else '-'
            body=(f'<div class="mi-kv"><span>Pencapaian terhadap target</span><b>{s["pct"]:.0f}%</b></div><div class="mi-kv"><span>Arti</span><b>Nilai rata-rata tiap unit</b></div>')
            bar=min(s['pct'],100) if s['target']>0 else 0; badge=f'{s["pct"]:.0f}%' if s['target']>0 else '-'
        else:
            f_=(lambda v:f'{v:,.0f}') if is_ue else (lambda v:RK.fmt_jt(v))
            act=f_(s['actual']); tgt=f_(s['target']) if s['target']>0 else '-'
            body=''.join(f'<div class="mi-kv"><span>{a_}</span><b>{b_}</b></div>' for a_,b_ in [('Sisa target',f_(s['remain']) if s['target']>0 else '-'),(f'Butuh / hari ({s["days_left"]} hari)',(f'{s["need"]:,.1f}' if is_ue else RK.fmt_jt(s['need'])) if s['target']>0 else '-'),('Rata-rata aktual / hari',(f'{s["run"]:,.1f}' if is_ue else RK.fmt_jt(s['run']))),('Proyeksi akhir bulan',(f'{f_(s["proj"])} ({s["proj_pct"]:.0f}%)' if s['target']>0 else f_(s['proj'])))])
            bar=min(s['pct'],100) if s['target']>0 else 0; badge=f'{s["pct"]:.1f}%' if s['target']>0 else 'belum ada target'
        exp_mark=(f'<div class="rk-foot2"><span>Pace seharusnya hari ini</span><b>{s["exp"]:.0f}%</b></div>' if (not is_rpu and s['target']>0) else '')
        return (f'<div class="mi-card {tone}"><div class="rk-top"><span class="rk-lab">{label}</span><span class="rk-ic {tone}"><span class="msr">{icon}</span></span></div>'
                f'<div class="mi-big">{act} <small>/ {tgt}</small></div>'
                f'<div class="rk-row"><span class="rk-badge {s["tone"]}">{s["label"]}</span><b class="rk-badge {tone if s["tone"]!="gray" else "gray"}">{badge}</b></div>'
                f'<div class="rk-bar"><i class="{tone if s["tone"]!="gray" else "gray"}" style="width:{bar:.1f}%"></i></div>{exp_mark}{body}</div>')
    st.markdown('<div class="mi-ach">'+_mi_card('Unit Entry (UE)','directions_car','blue',MI['s_ue'],is_ue=True)+_mi_card('Total Revenue','payments','green',MI['s_rev'])+_mi_card('Revenue per UE (Yield)','trending_up','purple',MI['s_rpu'],is_rpu=True)+'</div>',unsafe_allow_html=True)

    # ---------- 2. Diagnosis ----------
    st.markdown(_mi_sec(2,'Diagnosis Gap: Volume atau Yield?','Revenue = UE × Revenue per UE. Bagian mana yang menahan pencapaian?'),unsafe_allow_html=True)
    st.markdown(RK.dark_banner(MI['cause_text'],'DIAGNOSIS'),unsafe_allow_html=True)
    _cmp_rows=[]
    for m_,lab_ in [('UE','UE'),('Jasa','Jasa'),('Part','Part (sparepart)'),('Bahan','Bahan'),('Revenue','Total Revenue')]:
        t_=_mi_T[m_]; a_=_mi_R[m_]; p_=(a_/t_*100 if t_>0 else 0.0)
        _cmp_rows.append([f'<b>{lab_}</b>',(f'{a_:,.0f}' if m_=='UE' else fmt(a_)),((f'{t_:,.0f}' if m_=='UE' else fmt(t_)) if t_>0 else '-'),(f'<span class="rk-badge {RK.pct_tone(p_)}">{p_:.1f}%</span>' if t_>0 else '<span class="rk-badge gray">-</span>'),(RK.status_badge('Tertinggal','red') if (t_>0 and p_<MI['s_rev']['exp']*0.9) else RK.status_badge('Sesuai pace','green') if t_>0 else RK.status_badge('Target kosong','gray'))])
    st.markdown(RK.table_html(['Komponen','Realisasi','Target','% Pencapaian','Status vs Pace'],_cmp_rows,num_cols=(1,2,3)),unsafe_allow_html=True)
    if _mi_scope=='Semua Cabang' and MI['has_target']:
        _bb=revenue_components_by_branch(_mm,_pp,_ii,_cc,BRANCHES); _bb=_bb[_bb['Cabang']!='TOTAL SEMUA CABANG']
        _br_rows=[]
        for _,r_ in _bb.iterrows():
            tg_=_mi_tg.get(r_['Cabang'],{}) or {}; tu=float(tg_.get('UE',0) or 0); tr=float(tg_.get('Revenue',0) or 0)
            pu=(float(r_['UE'])/tu*100 if tu>0 else None); pr=(float(r_['Total Revenue'])/tr*100 if tr>0 else None)
            if pu is None and pr is None: foc='Target belum diisi'
            elif pu is not None and pr is not None and min(pu,pr)>=100: foc='Pertahankan'
            elif pu is None: foc='Naikkan Revenue'
            elif pr is None: foc='Naikkan Volume UE'
            else: foc=('Naikkan Volume UE' if pu<pr else 'Naikkan Revenue per UE') if abs(pu-pr)>=3 else 'Volume UE &amp; Revenue per UE'
            _bdg=lambda v:('<span class="rk-badge gray">-</span>' if v is None else f'<span class="rk-badge {RK.pct_tone(v)}">{v:.1f}%</span>')
            _br_rows.append([f'<b>{_html.escape(str(r_["Cabang"]))}</b>',f'{int(r_["UE"]):,}',_bdg(pu),fmt(r_['Total Revenue']),_bdg(pr),fmt(r_['Revenue/UE']),foc])
        if _br_rows:
            st.markdown(RK.head_card('Prioritas Fokus per Cabang','Cabang mana yang perlu dibantu, dan sisi mana (volume atau yield) yang harus didorong','apartment','teal'),unsafe_allow_html=True)
            paged_rows(['Cabang','UE','% Target UE','Revenue','% Target Revenue','Rev / UE','Fokus Utama'],_br_rows,'pg_mi_cabang',num_cols=(1,2,3,4,5),unit='cabang')

    # ---------- 3. Skenario ----------
    st.markdown(_mi_sec(3,'Skenario Mengejar Target Revenue','Apa yang harus dicapai per hari pada sisa bulan ini agar revenue tepat target'),unsafe_allow_html=True)
    if MI['scen']:
        st.markdown('<div class="mi-g" style="--mw:300px">'+''.join(f'<div class="mi-sc"><div class="h">{_html.escape(sc["title"])}</div><div class="v">{sc["value"]}</div><small>{sc["sub"]}</small></div>' for sc in MI['scen'])+'</div>',unsafe_allow_html=True)
        st.caption('Dihitung dari rata-rata realisasi per hari berjalan terhadap kebutuhan per hari untuk sisa bulan. Skenario C membagi beban peningkatan sama rata ke volume dan yield sehingga paling realistis.')
    else: st.info(MI['scen_note'])

    # ---------- 4. Strategi ----------
    st.markdown(_mi_sec(4,'Strategi Agar UE & Revenue Selalu Tercapai','Enam tuas strategi berbasis data bulan ini; urutan prioritas menyesuaikan hasil diagnosis'),unsafe_allow_html=True)
    _ord={'Prioritas Utama':0,'Pendukung':1,'Pertahankan':1,'Rekomendasi':2}
    _lev=sorted(MI['levers'],key=lambda l:_ord.get(l['prio'][0],3))
    _cards=[]
    for k_,lv in enumerate(_lev):
        _cards.append(f'<div class="mi-card {lv["tone"]}"><h4><span class="rk-ic {lv["tone"]}"><span class="msr">{lv["icon"]}</span></span>{_u(lv["title"])}</h4>'
            f'<div><span class="rk-badge {lv["prio"][1]}">{lv["prio"][0]}</span></div>'
            f'<div class="ev">{"<br>".join(lv["evidence"])}</div>'
            f'<ul>{"".join(f"<li>{_u(a_)}</li>" for a_ in lv["actions"])}</ul><div class="kp">{_u(lv["kpi"]).replace("Ukuran:","<b>Ukuran:</b>")}</div></div>')
    st.markdown('<div class="mi-g" style="--mw:340px">'+''.join(_cards)+'</div>',unsafe_allow_html=True)

    # ---------- 5. Rencana aksi & early warning ----------
    st.markdown(_mi_sec(5,'Rencana Aksi & Early Warning','Ritme kerja tim (siapa, kapan, ukuran sukses) dan pemicu eskalasi bila pace menurun'),unsafe_allow_html=True)
    paged_rows(['Aksi','PIC','Frekuensi','Ukuran Keberhasilan','Dampak'],MI['plan'],'pg_mi_plan',unit='aksi')
    st.markdown(RK.head_card('Early Warning: Jika Terjadi, Maka Lakukan','Pemicu yang perlu dipantau tiap hari/minggu beserta tindakan eskalasinya','warning','orange'),unsafe_allow_html=True)
    paged_rows(['Kondisi Pemicu','Tindakan Eskalasi'],MI['trig'],'pg_mi_trig',unit='pemicu')

    # ---------- 6. Insight otomatis dari data ----------
    if not f.empty:
        _ins=build_management_insights(f)
        if _ins:
            st.markdown(_mi_sec(6,'Catatan Otomatis dari Data','Temuan dari histori pada filter waktu sidebar'),unsafe_allow_html=True)
            st.markdown('<div class="insight-card">'+''.join(f'<div class="insight-row"><span class="insight-dot"><span class="msr">check</span></span><span>{x}</span></div>' for x in _ins)+'</div>',unsafe_allow_html=True)
if page == '🚗 UE':
    section('<span class="msr">directions_car</span> Unit Entry Analysis','Empat klasifikasi UE: jenis pekerjaan, type kendaraan, Service Advisor, dan Booking vs Walk In.')
    # ================= V17: ringkasan UE & kapasitas bengkel (sesuai referensi) =================
    _ue_tot=int(f['UE'].sum()) if 'UE' in f.columns else len(f)
    _md=RK.model_distribution(f,10)
    _rcu=revenue_components(f,parts,issue,cons); _rpu=(float(_rcu['total_revenue'])/_ue_tot) if _ue_tot else 0.0
    _bt=f['Booking Type'].map(clean).str.upper() if 'Booking Type' in f.columns else pd.Series('',index=f.index)
    _bk=int((_bt=='BOOKING').sum()); _bk_pct=(_bk/_ue_tot*100) if _ue_tot else 0.0
    if len(_md): _top_name=str(_md.iloc[0]['Model']); _top_sub=f'{_md.iloc[0]["Porsi"]:.1f}% dari seluruh unit terlayani'; _top_val=f'{_top_name} <small>({int(_md.iloc[0]["UE"]):,} UE)</small>'
    else: _top_val='-'; _top_sub='Belum ada data model kendaraan'
    st.markdown(RK.grid([RK.kpi_card('Total Unit Masuk',f'{_ue_tot:,} UE','Basis: laporan Estimasi SA / PKB periode terpilih','directions_car','blue'),
        RK.kpi_card('Model Terlaris (Volume)',_top_val,_top_sub,'emoji_events','purple'),
        RK.kpi_card('Rata-rata Revenue / UE',fmt(_rpu),'Yield rata-rata per transaksi','payments','green'),
        RK.kpi_card('Rasio Booking',f'{_bk_pct:.1f}%',f'{_bk:,} UE booking dari {_ue_tot:,} UE','event_available','teal')],230),unsafe_allow_html=True)
    if len(_md):
        st.markdown(RK.head_card('Distribusi UE Berdasarkan Model Kendaraan (Top 10 Models)','Volume unit, kontribusi persentase, total omzet per model, dan rata-rata belanja per kendaraan','directions_car','blue'),unsafe_allow_html=True)
        _rows=[[str(i+1),_html.escape(str(r['Model'])),f'<b style="color:#1d4ed8">{int(r["UE"]):,} UE</b>',f'{r["Porsi"]:.1f}%',fmt(r['Omzet']),fmt(r['Avg']),_html.escape(str(r['Pekerjaan Dominan']))] for i,r in _md.iterrows()]
        st.markdown(RK.table_html(['#','Model Kendaraan','Volume Masuk','Porsi UE','Total Omzet','Avg Spending / UE','Jenis Pekerjaan Dominan'],_rows,num_cols=(2,3,4,5)),unsafe_allow_html=True)
        download_table(_md,':material/download: Download Excel — Distribusi UE per Model','ue_distribusi_model.xlsx','dl_ue_model_v17')
    _sg=RK.customer_segments(f); _wl=RK.dow_workload(f)
    _u1,_u2=st.columns(2)
    with _u1:
        st.markdown(RK.head_card('UE Berdasarkan Segmentasi Pelanggan','Komposisi pelanggan perorangan, armada perusahaan, dan lainnya','groups','teal'),unsafe_allow_html=True)
        if _sg.empty: st.info('Segmentasi pelanggan belum tersedia (butuh laporan Pelanggan).')
        else:
            _cols=['#2563eb','#f59e0b','#10b981','#8b5cf6','#ef4444','#06b6d4']
            st.markdown(''.join(f'<div class="rk-seg"><div class="t"><span>{_html.escape(str(r["Segmen"]))}</span><span>{int(r["UE"]):,} UE ({r["Porsi"]:.0f}%)</span></div><div class="rk-bar"><i style="width:{min(r["Porsi"],100):.1f}%;background:{_cols[k%len(_cols)]}"></i></div><small style="color:#64748b">Omzet {fmt(r["Omzet"])}</small></div>' for k,r in _sg.iterrows()),unsafe_allow_html=True)
    with _u2:
        st.markdown(RK.head_card('Pola Beban Masuk Unit Harian (Day of Week Workload)','Rata-rata unit per hari kerja untuk penjadwalan & rotasi teknisi','calendar_month','purple'),unsafe_allow_html=True)
        if _wl.empty: st.info('Belum ada data tanggal.')
        else:
            _mxw=float(_wl['Rata2'].max()) or 1.0; _lc={'Sangat Padat':('#ef4444','red'),'Padat':('#f59e0b','orange'),'Sedang':('#2563eb','blue'),'Lancar':('#10b981','green')}
            st.markdown(''.join(f'<div class="rk-wl"><span class="d">{r["Hari"]}</span><span class="b"><i style="width:{r["Rata2"]/_mxw*100:.1f}%;background:{_lc[r["Level"]][0]}"></i></span><span class="v">~{r["Rata2"]:.0f} UE</span>{RK.status_badge(r["Level"],_lc[r["Level"]][1])}</div>' for _,r in _wl.iterrows()),unsafe_allow_html=True)
    # ===== V23: 4 klasifikasi UE dengan desain kartu yang sama seperti Dashboard (donat untuk SA & Booking) =====
    _bn_up=_html.escape(str(active_branch)).upper()
    _pal_ue=['#2563eb','#10b981','#f59e0b','#ef4444','#8b5cf6','#06b6d4','#ec4899','#84cc16','#f97316','#14b8a6','#6366f1','#a855f7']
    def ue_bar_card(ckey,icon,title,sub,col,fname):
        with st.container(border=True,key=ckey):
            st.markdown(f'<div class="rk-ct"><b><span class="msr">{icon}</span><span>{title} — <span class="a">{_bn_up}</span></span></b><span>{sub}</span></div>',unsafe_allow_html=True)
            z=f[f[col].map(clean)!=''].groupby(col,as_index=False).agg(UE=('UE','sum')).sort_values('UE',ascending=False).reset_index(drop=True)
            if z.empty or float(z['UE'].sum())<=0: st.info('Belum ada data pada periode ini.'); return
            _tot=float(z['UE'].sum()); _t10=z.head(10).iloc[::-1]
            fig=go.Figure(go.Bar(x=_t10['UE'],y=_t10[col],orientation='h',marker_color='#2563eb',text=[f'{int(v):,}' for v in _t10['UE']],textposition='outside',cliponaxis=False))
            fig.update_layout(height=max(260,34*len(_t10)+40),margin=dict(l=5,r=45,t=5,b=5),paper_bgcolor='#ffffff',plot_bgcolor='#ffffff',xaxis=dict(title='UE',gridcolor='#edf2f7'),yaxis=dict(title=''))
            st.plotly_chart(fixfig(fig),use_container_width=True,config=PCFG)
            z['% Porsi']=z['UE']/_tot*100; z.insert(0,'Rank',range(1,len(z)+1))
            paged_df(z,f'pg_{fname}',pct=['% Porsi'])
            download_table(z,f':material/download: Download Excel — {title}',f'{fname}.xlsx',f'dl_{fname}')
    c1,c2=st.columns(2)
    with c1: ue_bar_card('ue_c1','work','UE BERDASARKAN JENIS PEKERJAAN','Top 10 jenis pekerjaan','Jenis Pekerjaan','ue_jenis_pekerjaan')
    with c2: ue_bar_card('ue_c2','directions_car','UE BERDASARKAN TYPE KENDARAAN','Top 10 type kendaraan','Type Kendaraan','ue_type_kendaraan')
    c3,c4=st.columns(2)
    with c3:
        with st.container(border=True,key='dn_ue_sa'):
            _sa=f[f['Service Advisor'].astype(str).str.strip()!=''].copy()
            st.markdown(f'<div class="rk-ct"><b><span class="msr">groups</span><span>UE BERDASARKAN SERVICE ADVISOR — <span class="a">{_bn_up}</span></span></b><span>{_sa["Service Advisor"].nunique() if not _sa.empty else 0} Service Advisor</span></div>',unsafe_allow_html=True)
            if _sa.empty: st.info('Belum ada data Service Advisor pada periode ini.')
            else:
                _sa['Service Advisor']=sa_disp(_sa['Service Advisor'])
                _agg={'UE':('UE','sum')}
                if 'Revenue' in _sa.columns: _agg['Revenue']=('Revenue','sum')
                _sat=_sa.groupby('Service Advisor',as_index=False).agg(**_agg).sort_values('UE',ascending=False).reset_index(drop=True)
                _tu=float(_sat['UE'].sum()); _sat['% Porsi']=(_sat['UE']/_tu*100) if _tu>0 else 0.0
                _top=_sat.head(6); _rest=_sat.iloc[6:]
                _lb=_top['Service Advisor'].tolist()+(['Lainnya'] if len(_rest) else []); _vl=_top['UE'].tolist()+([float(_rest['UE'].sum())] if len(_rest) else [])
                _its=[(_pal_ue[k%len(_pal_ue)],str(r_['Service Advisor']),(RK.fmt_jt(r_['Revenue']) if 'Revenue' in _sat.columns else 'Service Advisor'),f'{int(r_["UE"]):,} UE',f'{r_["% Porsi"]:.0f}% Share') for k,(_,r_) in enumerate(_top.iterrows())]
                if len(_rest): _its.append(('#94a3b8',f'Lainnya ({len(_rest)} SA)','',f'{int(_rest["UE"].sum()):,} UE',f'{_rest["% Porsi"].sum():.0f}% Share'))
                donut_card_body(_lb,_vl,_pal_ue[:len(_lb)],'TOTAL SA',f'{len(_sat)} SA',f'{int(_tu):,} UE',_its)
                _sat_t=_sat.copy(); _sat_t.insert(0,'Rank',range(1,len(_sat_t)+1))
                paged_df(_sat_t,'pg_ue_sa',money=['Revenue'] if 'Revenue' in _sat_t.columns else (),pct=['% Porsi'])
                download_table(_sat_t,':material/download: Download Excel — UE per Service Advisor','ue_per_service_advisor.xlsx','dl_ue_per_sa')
    with c4:
        with st.container(border=True,key='dn_ue_bk'):
            st.markdown(f'<div class="rk-ct"><b><span class="msr">explore</span><span>UE BERDASARKAN LAYANAN — <span class="a">{_bn_up}</span></span></b><span>Booking vs Walk-In vs Field</span></div>',unsafe_allow_html=True)
            _bt2=f.copy(); _bt2['_L']=_bt2['Booking Type'].map(clean).str.upper().str.replace('-',' ',regex=False).str.strip()
            _lmap={'BOOKING':'Booking','WALK IN':'Walk-In','WALKIN':'Walk-In','FIELD SERVICE':'Field Service','FIELD':'Field Service'}
            _bt2['_L']=_bt2['_L'].map(lambda x:_lmap.get(x,'Lainnya' if x else 'Tidak diisi'))
            _lz=_bt2.groupby('_L',as_index=False).agg(UE=('UE','sum')).sort_values('UE',ascending=False).reset_index(drop=True)
            _lt=float(_lz['UE'].sum()); _lc={'Booking':'#2563eb','Walk-In':'#06b6d4','Field Service':'#f59e0b','Lainnya':'#94a3b8','Tidak diisi':'#cbd5e1'}
            if _lz.empty or _lt<=0: st.info('Belum ada data jenis layanan pada periode ini.')
            else:
                _lz['% Porsi']=_lz['UE']/_lt*100
                donut_card_body(_lz['_L'].tolist(),_lz['UE'].tolist(),[_lc.get(x,'#94a3b8') for x in _lz['_L']],'TOTAL UE',f'{int(_lt):,} UE','Unit Service',
                    [(_lc.get(r_['_L'],'#94a3b8'),r_['_L'],'Kedatangan Unit',f'{int(r_["UE"]):,} UE',RK.share_pct(r_['UE'],_lt)) for _,r_ in _lz.iterrows()])
                _lzt=_lz.rename(columns={'_L':'Layanan'}); _lzt.insert(0,'Rank',range(1,len(_lzt)+1))
                paged_df(_lzt,'pg_ue_layanan',pct=['% Porsi'])
                download_table(_lzt,':material/download: Download Excel — UE Booking Walk In','ue_booking_walkin.xlsx','dl_ue_booking_walkin')

if page == '💰 Revenue & Profit':
    section('<span class="msr">payments</span> Revenue & Profitability','Analisa profitabilitas berdasarkan sumber biaya yang tersedia. Revenue Core tetap dipisahkan dari OPL untuk mencegah double count.')

    # Profitability model: Service Penjualan supplies revenue only; Parts, Material and OPL
    # have cost fields. Therefore the app never invents a service cost and clearly labels
    # the resulting profit as "GP Terukur" rather than total business profit.
    rc=revenue_components(f,parts,issue,cons)
    revenue_core=float(rc['revenue_core'])
    core_rev=float(rc['total_revenue'])
    jasa=float(f['Jasa'].sum()); spare_core=float(f['Sparepart'].sum()); bahan_core=float(f['Bahan'].sum())
    part_sales=float(parts['Sales Amount'].sum()) if not parts.empty and 'Sales Amount' in parts else 0.0
    part_cost=float(parts['Total Cost'].sum()) if not parts.empty and 'Total Cost' in parts else 0.0
    part_gp=float(parts['Gross Profit'].sum()) if not parts.empty and 'Gross Profit' in parts else 0.0
    material_sales=0.0; material_cost=0.0
    for md in [issue,cons]:
        if not md.empty:
            material_sales += float(md['Harga Net'].sum()) if 'Harga Net' in md else 0.0
            material_cost += float(md['Cost'].sum()) if 'Cost' in md else 0.0
    material_gp=material_sales-material_cost
    opl_sales=float(opl['Sales OPL'].sum()) if not opl.empty else 0.0
    opl_cost=float(opl['Cost OPL'].sum()) if not opl.empty else 0.0
    opl_gp=float(opl['Gross Profit OPL'].sum()) if not opl.empty else 0.0
    measured_sales=part_sales+material_sales+opl_sales
    measured_cost=part_cost+material_cost+opl_cost
    measured_gp=part_gp+material_gp+opl_gp
    measured_margin=(measured_gp/measured_sales*100) if measured_sales else 0.0
    coverage=(measured_sales/core_rev*100) if core_rev else 0.0

    # Gunakan 3 + 3 kartu agar angka Rupiah panjang tidak terpotong di layar.
    st.markdown('<div class="revenue-profit-kpi">',unsafe_allow_html=True)
    c1,c2,c3=st.columns(3)
    c1.metric('Revenue Core',fmt(revenue_core))
    c2.metric('Part tanpa PKB',fmt(rc['direct_part']))
    c3.metric('Bahan tanpa PKB',fmt(rc['direct_bahan']))
    c4,c5,c6=st.columns(3)
    c4.metric('TOTAL REVENUE',fmt(core_rev))
    c5.metric('GP Terukur',fmt(measured_gp))
    c6.metric('Margin GP Terukur',f'{measured_margin:.1f}%')
    st.caption(f'Sales Berbiaya: {fmt(measured_sales)} • Coverage Biaya terhadap Total Revenue: {coverage:.1f}%')

    # Sumber Revenue tetap ditampilkan lengkap dan terpisah dari analisa profitability.
    st.markdown('### :material/payments: Sumber Revenue')
    revenue_source=pd.DataFrame([
        ['Jasa / Service dengan PKB',rc['jasa']],
        ['Sparepart dengan PKB — Billing Service',rc['service_sparepart']],
        ['Sparepart tanpa PKB — Billing Part',rc['direct_part']],
        ['Bahan dengan PKB — Issue To PKB',rc['service_bahan']],
        ['Bahan tanpa PKB — Consumable',rc['direct_bahan']],
    ],columns=['Sumber Revenue','Revenue'])
    revenue_source['% Total Revenue']=np.where(core_rev!=0,revenue_source['Revenue']/core_rev*100,0)
    st.dataframe(revenue_source.style.format({'Revenue':'Rp {:,.0f}','% Total Revenue':'{:.1f}%'}),use_container_width=True,hide_index=True)
    download_table(revenue_source,':material/download: Download Excel — Sumber Revenue','revenue_sources.xlsx','dl_revenue_sources')
    figsrc=go.Figure(go.Bar(x=revenue_source['Revenue'],y=revenue_source['Sumber Revenue'],orientation='h',text=revenue_source['Revenue'].map(fmt),textposition='outside',cliponaxis=False))
    figsrc.update_layout(height=330,margin=dict(l=8,r=95,t=10,b=10),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',xaxis=dict(tickprefix='Rp ',tickformat=',.0f',gridcolor='#edf2f7'))
    st.plotly_chart(figsrc,use_container_width=True,config={'displayModeBar':False})
    st.caption('GP Terukur hanya berasal dari sumber yang memiliki field cost: Penjualan Part, Material/Consumable, dan OPL. Cost jasa tidak tersedia di laporan Service Penjualan sehingga tidak dihitung/diasumsikan.')

    st.markdown('### :material/push_pin: Profitability Bridge')
    bridge=pd.DataFrame([
        ['Penjualan Part',part_sales,part_cost,part_gp,(part_gp/part_sales*100 if part_sales else 0)],
        ['Material / Consumable',material_sales,material_cost,material_gp,(material_gp/material_sales*100 if material_sales else 0)],
        ['OPL',opl_sales,opl_cost,opl_gp,(opl_gp/opl_sales*100 if opl_sales else 0)],
    ],columns=['Sumber','Sales','Cost','Gross Profit','Margin %'])
    st.dataframe(bridge.style.format({'Sales':'Rp {:,.0f}','Cost':'Rp {:,.0f}','Gross Profit':'Rp {:,.0f}','Margin %':'{:.1f}%'}),use_container_width=True,hide_index=True)
    download_table(bridge,':material/download: Download Excel — Profitability Bridge','profitability_bridge.xlsx','dl_profit_bridge')

    left,right=st.columns(2)
    with left:
        st.markdown('### :material/trending_up: Sales vs Cost vs GP')
        figp=go.Figure()
        for metric in ['Sales','Cost','Gross Profit']:
            figp.add_trace(go.Bar(x=bridge['Sumber'],y=bridge[metric],name=metric))
        figp.update_layout(barmode='group',height=360,margin=dict(l=8,r=8,t=10,b=10),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',yaxis=dict(tickprefix='Rp ',tickformat=',.0f',gridcolor='#e8eef6'))
        st.plotly_chart(figp,use_container_width=True,config={'displayModeBar':False})
    with right:
        st.markdown('### :material/track_changes: Margin by Source')
        figm=go.Figure(go.Bar(x=bridge['Margin %'],y=bridge['Sumber'],orientation='h',text=bridge['Margin %'].map(lambda x:f'{x:.1f}%'),textposition='outside',cliponaxis=False))
        figm.update_layout(height=360,margin=dict(l=8,r=45,t=10,b=10),paper_bgcolor='rgba(0,0,0,0)',plot_bgcolor='rgba(0,0,0,0)',xaxis=dict(title='Margin %',gridcolor='#edf2f7'))
        st.plotly_chart(figm,use_container_width=True,config={'displayModeBar':False})

    st.markdown('### :material/directions_car: Profitability by Vehicle — OPL')
    if opl.empty:
        st.info(':material/description: OPL belum tersedia. Analisa profit per kendaraan hanya dapat dibuat untuk sumber yang memiliki relasi kendaraan + cost.')
    else:
        veh=opl.groupby('Type Kendaraan',as_index=False).agg(OPL_Sales=('Sales OPL','sum'),OPL_Cost=('Cost OPL','sum'),OPL_GP=('Gross Profit OPL','sum'),OPL_Count=('OPL No','nunique')).sort_values('OPL_GP',ascending=False)
        veh['Margin %']=np.where(veh['OPL_Sales']!=0,veh['OPL_GP']/veh['OPL_Sales']*100,0)
        veh=veh.rename(columns={'Type Kendaraan':'Tipe Kendaraan','OPL_Sales':'Sales','OPL_Cost':'Cost','OPL_GP':'Gross Profit','OPL_Count':'Jumlah OPL'})
        st.dataframe(veh.style.format({'Sales':'Rp {:,.0f}','Cost':'Rp {:,.0f}','Gross Profit':'Rp {:,.0f}','Margin %':'{:.1f}%'}),use_container_width=True,hide_index=True)
        download_table(veh,':material/download: Download Excel — Profitability OPL per Kendaraan','profitability_opl_vehicle.xlsx','dl_profit_opl_vehicle')

    st.markdown('### :material/person: Profitability by Service Advisor — OPL')
    if opl.empty or 'Service Advisor' not in opl:
        st.info('OPL belum tersedia untuk analisa profit per Service Advisor.')
    else:
        sa_profit=opl[opl['Service Advisor'].map(clean)!=''].assign(**{'Service Advisor':lambda d:sa_disp(d['Service Advisor'])}).groupby('Service Advisor',as_index=False).agg(Sales=('Sales OPL','sum'),Cost=('Cost OPL','sum'),Gross_Profit=('Gross Profit OPL','sum'),OPL=('OPL No','nunique')).sort_values('Gross_Profit',ascending=False)
        sa_profit['Margin %']=np.where(sa_profit['Sales']!=0,sa_profit['Gross_Profit']/sa_profit['Sales']*100,0)
        sa_profit=sa_profit.rename(columns={'Gross_Profit':'Gross Profit'})
        st.dataframe(sa_profit.style.format({'Sales':'Rp {:,.0f}','Cost':'Rp {:,.0f}','Gross Profit':'Rp {:,.0f}','Margin %':'{:.1f}%'}),use_container_width=True,hide_index=True)
        download_table(sa_profit,':material/download: Download Excel — Profitability OPL per SA','profitability_opl_sa.xlsx','dl_profit_opl_sa')

    st.markdown('### :material/psychology: Profitability Insight')
    p_ins=[]
    if measured_sales>0:
        best=bridge.loc[bridge['Gross Profit'].idxmax()]
        worst=bridge.loc[bridge['Margin %'].idxmin()]
        p_ins.append(f"Kontributor GP terukur terbesar: {best['Sumber']} sebesar {fmt(best['Gross Profit'])}.")
        p_ins.append(f"Margin terendah di antara sumber berbiaya: {worst['Sumber']} ({worst['Margin %']:.1f}%).")
    if core_rev>0 and coverage<50:
        p_ins.append(f"Coverage biaya baru {coverage:.1f}% dari Total Revenue. Hindari menyebut angka GP Terukur sebagai total profit bengkel sebelum cost jasa tersedia.")
    if not opl.empty and opl_sales>0:
        if opl_gp<0: p_ins.append(f"OPL menghasilkan GP negatif {fmt(opl_gp)}. Perlu review vendor, harga jual atau cost OPL.")
        else: p_ins.append(f"OPL menghasilkan GP {fmt(opl_gp)} dengan margin {opl_gp/opl_sales*100:.1f}%.")
    if not p_ins:
        p_ins.append('Belum cukup data biaya untuk membangun profitability insight.')
    st.markdown('<div class="insight-card"><h3 style="margin:0 0 5px;color:#0f172a;"><span class="msr">lightbulb</span> Insight Profitability</h3>'+''.join(f'<div class="insight-row"><span class="insight-dot"><span class="msr">check</span></span><span>{x}</span></div>' for x in p_ins)+'</div>',unsafe_allow_html=True)

    st.markdown('### :material/warning: Profitability Data Limitation')
    st.warning('Revenue Core berasal dari PKB/Service Penjualan, sedangkan Total Revenue juga mencakup penjualan part dan bahan tanpa PKB. Karena laporan tersebut tidak menyediakan cost jasa, aplikasi tidak menghitung “Net Profit” atau “Total GP Bengkel”. Angka GP yang ditampilkan adalah GP yang dapat diukur dari sumber yang memiliki cost, dan ditampilkan terpisah agar tidak misleading.')
if page == '🔩 Sparepart':
    section('<span class="msr">settings</span> Sparepart','Rincian penjualan sparepart dengan PKB dan penjualan part langsung tanpa PKB.')
    if parts.empty:
        st.info('Upload Billing Service / Billing Part untuk membuka rincian sparepart.')
    else:
        ps=parts[parts['Sumber Penjualan']=='Billing Service'].copy() if 'Sumber Penjualan' in parts else pd.DataFrame()
        pdirect=parts[parts['Sumber Penjualan']=='Billing Part'].copy() if 'Sumber Penjualan' in parts else pd.DataFrame()
        with_pkb=float(ps['Sales Amount'].sum()) if not ps.empty else 0
        without_pkb=float(pdirect['Sales Amount'].sum()) if not pdirect.empty else 0
        total=with_pkb+without_pkb
        st.markdown(RK.grid([RK.kpi_card('Sparepart dengan PKB',fmt(with_pkb),'Suku cadang terpasang langsung pada pengerjaan service unit','build','blue'),
            RK.kpi_card('Part tanpa PKB (Counter Sales)',f'<span style="color:#ea580c">{fmt(without_pkb)}</span>','Penjualan retail suku cadang langsung ke konsumen / toko rekanan','storefront','orange'),
            RK.kpi_card('Total Omzet Sparepart',f'<span style="color:#7c3aed">{fmt(total)}</span>','Akumulasi penjualan part masuk ke revenue (PKB + tanpa PKB)','payments','purple')],300),unsafe_allow_html=True)
        _fm=RK.fast_moving(parts)
        if not _fm.empty:
            st.markdown(RK.head_card('Daftar Sparepart Terlaris (Top Fast-Moving Parts & Counter Sales)','Perbandingan jalur PKB vs Non-PKB, estimasi margin kotor, dan klasifikasi perputaran (analisis Pareto kuantitas)','sell','orange'),unsafe_allow_html=True)
            _s1,_s2=st.columns([3,1.4])
            with _s1: _q=st.text_input('Cari nama part / no part',key='rk_sp_q',placeholder='Cari nama part / no part...',label_visibility='collapsed')
            with _s2: _cat=st.selectbox('Kategori',['Semua Kategori','Fast Moving','Medium Moving','Slow Moving'],key='rk_sp_cat',label_visibility='collapsed')
            _z=_fm
            if _q.strip(): _qq=_q.strip().lower(); _z=_z[_z['No Part'].str.lower().str.contains(_qq,regex=False)|_z['Deskripsi'].str.lower().str.contains(_qq,regex=False)]
            if _cat!='Semua Kategori': _z=_z[_z['Kategori']==_cat]
            _kt={'Fast Moving':'green','Medium Moving':'blue','Slow Moving':'gray'}
            _rws=[[f'<span style="font-family:ui-monospace,monospace">{_html.escape(str(r["No Part"]))}</span>',_html.escape(str(r['Deskripsi'])),RK.status_badge(r['Kategori'],_kt[r['Kategori']]),f'{r["QtyPKB"]:,.0f}',f'<b style="color:#c2410c">{r["QtyCtr"]:,.0f}</b>',f'<b>{r["Qty"]:,.0f}</b>',fmt(r['Harga']),f'<b>{fmt(r["Omzet"])}</b>','-' if pd.isna(r['Margin']) else f'{r["Margin"]:.1f}%'] for _,r in _z.iterrows()]
            if _rws: paged_rows(['Part Number','Deskripsi Suku Cadang','Kategori','Qty (PKB)','Qty (Counter)','Total Qty','Harga Satuan','Total Omzet','Margin %'],_rws,'pg_sp_fast',num_cols=(3,4,5,6,7,8),unit='part'); st.caption('Kategori: Fast = 70% kuantitas teratas, Medium = 70–90%, Slow = sisanya.')
            else: st.info('Tidak ada part yang cocok dengan pencarian / kategori.')
            download_table(_fm.rename(columns={'QtyPKB':'Qty PKB','QtyCtr':'Qty Counter','Harga':'Harga Satuan','Omzet':'Total Omzet','Margin':'Margin %'}),':material/download: Download Excel — Sparepart Terlaris (lengkap)','sparepart_terlaris.xlsx','dl_sp_fast_v17')
        _bs=RK.branch_sparepart(parts,[b for b in BRANCHES if active_branch in ('Semua Cabang',b)])
        if not _bs.empty:
            st.markdown(RK.head_card(f'Rincian Penjualan Sparepart per Cabang ({_html.escape(str(active_branch))})','Perbandingan jalur PKB dan counter sales','apartment','purple'),unsafe_allow_html=True)
            _br=[[_html.escape(str(r['Cabang'])),fmt(r['PKB']),f'<span style="color:#c2410c">{fmt(r["Counter"])}</span>',f'<b>{fmt(r["Total"])}</b>',f'{r["Porsi"]:.1f}%'] for _,r in _bs.iterrows()]
            st.markdown(RK.table_html(['Cabang','Sparepart PKB','Part tanpa PKB (Counter)','Total Sparepart','% Porsi Non-PKB'],_br,num_cols=(1,2,3,4),hl_cols={3:'hl2'}),unsafe_allow_html=True)
        # TOP 10 PART TERBANYAK — gabungan Billing Service + Billing Part
        part_sources=[]
        if not ps.empty:
            part_sources.append(ps[['No Part','Deskripsi','Qty','Sales Amount','Gross Profit']].copy())
        if not pdirect.empty:
            part_sources.append(pdirect[['No Part','Deskripsi','Qty','Sales Amount','Gross Profit']].copy())
        if part_sources:
            top_parts=pd.concat(part_sources,ignore_index=True)
            top_parts['No Part']=top_parts['No Part'].map(clean)
            top_parts['Deskripsi']=top_parts['Deskripsi'].map(clean)
            top_parts=top_parts[(top_parts['No Part']!='') | (top_parts['Deskripsi']!='')].copy()
            top_parts['Qty']=pd.to_numeric(top_parts['Qty'],errors='coerce').fillna(0)
            top_parts['Sales Amount']=pd.to_numeric(top_parts['Sales Amount'],errors='coerce').fillna(0)
            top_parts['Gross Profit']=pd.to_numeric(top_parts['Gross Profit'],errors='coerce').fillna(0)
            top_parts=top_parts.groupby(['No Part','Deskripsi'],dropna=False,as_index=False).agg(Qty=('Qty','sum'),Sales=('Sales Amount','sum'),GP=('Gross Profit','sum')).sort_values(['Qty','Sales'],ascending=[False,False]).head(10).reset_index(drop=True)
            top_parts.insert(0,'Rank',range(1,len(top_parts)+1))
            st.markdown(RK.head_card('Top 10 Penjualan Part Terbanyak','Ranking berdasarkan total Qty terjual, gabungan Billing Service dan Billing Part.','emoji_events','orange'),unsafe_allow_html=True)
            paged_df(top_parts.rename(columns={'Sales':'Omzet','GP':'Gross Profit'}),'pg_sp_top10',money=['Omzet','Gross Profit'],code=['No Part'],unit='part')
            download_table(top_parts,':material/download: Download Excel — Top 10 Part Terbanyak','top10_part_terbanyak.xlsx','dl_top10_part')

        st.markdown(RK.head_card('Rincian Billing Service — Sparepart dengan PKB','Detail penjualan part per dokumen Billing Service (10 baris per halaman)','description','blue'),unsafe_allow_html=True)
        if ps.empty: st.info('Belum ada penjualan part Billing Service.')
        else:
            ps_rank=ps.groupby(['No Billing','No Referensi','No Part','Deskripsi'],dropna=False,as_index=False).agg(Qty=('Qty','sum'),Sales=('Sales Amount','sum'),GP=('Gross Profit','sum')).sort_values('Sales',ascending=False)
            paged_df(ps_rank.rename(columns={'Sales':'Omzet','GP':'Gross Profit'}),'pg_sp_bsvc',money=['Omzet','Gross Profit'],code=['No Part'],unit='baris'); download_table(ps_rank,':material/download: Download Excel — Part Billing Service','part_billing_service.xlsx','dl_part_billing_service')
        st.markdown(RK.head_card('Rincian Billing Part — Sparepart tanpa PKB','Detail penjualan part langsung (counter sales) per dokumen Billing Part (10 baris per halaman)','description','orange'),unsafe_allow_html=True)
        if pdirect.empty: st.info('Belum ada penjualan part langsung.')
        else:
            pd_rank=pdirect.groupby(['No Billing','No Referensi','No Part','Deskripsi'],dropna=False,as_index=False).agg(Qty=('Qty','sum'),Sales=('Sales Amount','sum'),GP=('Gross Profit','sum')).sort_values('Sales',ascending=False)
            paged_df(pd_rank.rename(columns={'Sales':'Omzet','GP':'Gross Profit'}),'pg_sp_bpart',money=['Omzet','Gross Profit'],code=['No Part'],unit='baris'); download_table(pd_rank,':material/download: Download Excel — Part Billing Part','part_billing_part.xlsx','dl_part_billing_part')
        st.caption(f'Total Revenue Sparepart = {fmt(with_pkb)} + {fmt(without_pkb)} = {fmt(total)}')
if page == '🧪 Bahan & Consumable':
    RK.BAHAN_MANUAL.clear(); RK.BAHAN_MANUAL.update(load_bahan_manual())
    # Angka headline memakai rumus yang SAMA dengan Dashboard (revenue_components).
    rc_b=revenue_components(master,parts,issue,cons)
    issue_sales=float(rc_b['service_bahan']); cons_sales=float(rc_b['direct_bahan']); total=float(rc_b['total_bahan'])
    _bh1,_bh2=st.columns([6,1.2])
    with _bh1: st.markdown(RK.head_card('Bahan & Consumable Bengkel',f'Monitoring bahan Issue To PKB, bahan tanpa PKB, dan omzet bahan per UE ({_html.escape(str(active_branch))})','science','orange'),unsafe_allow_html=True)
    with _bh2: st.markdown(f'<div class="rk-chip">{_html.escape(str(active_branch))}</div>',unsafe_allow_html=True)
    st.markdown(RK.grid([
        RK.kpi_card('Bahan Issue To PKB',fmt(issue_sales),'Material cat dan fluida yang dialokasikan khusus pada order perbaikan unit','science','blue'),
        RK.kpi_card('Bahan Tanpa PKB (Operasional / Direct)',fmt(cons_sales),'Pemakaian operasional umum bengkel, pembersihan stall, dan penjualan chemical langsung','storefront','orange'),
        RK.kpi_card('Total Omzet Bahan',fmt(total),'Akumulasi pemakaian dan penjualan bahan masuk ke revenue bulanan','payments','green')],300),unsafe_allow_html=True)

    # ---- data transaksi bahan (Issue To PKB + Consumable) ----
    def _sa_of(row):
        # SA ditentukan saat data dibaca (nomor PKB pada header dokumen Issue -> PKB di MaintainPKB / Estimasi SA)
        return sa_disp(row.get('SA Code','')) if clean(row.get('SA Code','')) else ''
    _tx=[]
    for _src_,_df_ in (('Issue To PKB',issue),('Tanpa PKB',cons)):
        if isinstance(_df_,pd.DataFrame) and not _df_.empty and 'Kode Bahan' in _df_.columns:
            z=_df_.copy()
            z['_Sumber']=_src_
            z['_SA']=[_sa_of(r_) for r_ in z.to_dict('records')] if _src_=='Issue To PKB' else ''
            _tx.append(z)
    _txd=pd.concat(_tx,ignore_index=True) if _tx else pd.DataFrame()
    if not _txd.empty:
        for c_ in ['Kode Bahan','Bahan']: _txd[c_]=_txd[c_].map(clean)
        for c_ in ['Qty','Harga Net','Cost']: _txd[c_]=pd.to_numeric(_txd[c_],errors='coerce').fillna(0.0)
        _txd['Kategori']=[RK.bahan_kategori(a_,b_) for a_,b_ in zip(_txd['Kode Bahan'],_txd['Bahan'])]
        if 'Satuan' not in _txd.columns: _txd['Satuan']=''
        _ue_all=int(f['UE'].sum()) if 'UE' in f.columns else 0
        # ---- blok: judul rincian + dropdown kategori bahan (satu kartu) ----
        with st.container(key='blk_bhn'):
            _ctop=st.columns([4,1.5],vertical_alignment='center')
            with _ctop[0]: st.markdown(RK.head_card('Rincian Penggunaan Bahan & Consumable Teratas','Konsumsi bahan per kode, kategori, qty PKB vs non-PKB, omzet, dan biaya per UE','list_alt','orange'),unsafe_allow_html=True)
            with _ctop[1]: _kat_sel=st.selectbox('Kategori Bahan',['Semua Kategori Bahan']+RK.BAHAN_KATS,key='bahan_kat_top',label_visibility='collapsed',help='Kategori ditentukan dari kode bahan (lihat Mapping Kode Bahan di bawah tabel).')
        _zz=_txd if _kat_sel=='Semua Kategori Bahan' else _txd[_txd['Kategori']==_kat_sel]
        _g=_zz.assign(_qp=np.where(_zz['_Sumber']=='Issue To PKB',_zz['Qty'],0.0),_qn=np.where(_zz['_Sumber']=='Tanpa PKB',_zz['Qty'],0.0)).groupby(['Kode Bahan','Bahan','Kategori'],as_index=False).agg(Satuan=('Satuan','first'),QtyP=('_qp','sum'),QtyN=('_qn','sum'),Omzet=('Harga Net','sum')).sort_values('Omzet',ascending=False)
        if _g.empty: st.info('Belum ada data bahan untuk kategori ini.')
        else:
            _rows=[[f'<span style="font-family:ui-monospace,monospace">{_html.escape(str(r_["Kode Bahan"]))}</span>',_html.escape(str(r_['Bahan'])),f'<span class="rk-badge gray">{_html.escape(str(r_["Kategori"]))}</span>',_html.escape(str(r_['Satuan'])),f'{r_["QtyP"]:,.0f}',f'<b style="color:#dc2626">{r_["QtyN"]:,.0f}</b>',f'<b>{r_["QtyP"]+r_["QtyN"]:,.0f}</b>',fmt(r_['Omzet']),(fmt(r_['Omzet']/_ue_all) if _ue_all else '-')] for _,r_ in _g.iterrows()]
            paged_rows(['Kode Bahan','Nama Material / Chemical','Kategori','Satuan','Qty (PKB)','Qty (Non-PKB)','Total Qty','Total Omzet','Biaya / UE'],_rows,'pg_bhn_top',num_cols=(4,5,6,7,8),hl_cols={7:'hl2'},unit='bahan')
        # ---- mapping kode bahan (tersembunyi dalam expander) ----
        with st.expander(':material/account_tree: Mapping Kode Bahan per Kategori',expanded=False):
            st.caption('Kategori ditentukan dari KODE bahan: mapping manual > daftar resmi > Bahan Lain-lain.')
            _nm_by_code={}
            for _c_,_n_ in _txd[['Kode Bahan','Bahan']].drop_duplicates('Kode Bahan').itertuples(index=False): _nm_by_code[RK.norm_kode(_c_)]=str(_n_)
            _mcols=st.columns(3)
            for _mc,_kn,_tone in zip(_mcols,RK.BAHAN_KATS,('blue','purple','gray')):
                with _mc:
                    _pool=list(dict.fromkeys([c_ for lst in RK.BAHAN_KODE.values() for c_ in lst]+sorted({str(c_) for c_ in _txd['Kode Bahan'] if c_})+[k_ for k_ in RK.BAHAN_MANUAL]))
                    _codes=[c_ for c_ in _pool if RK.bahan_kategori(c_)==_kn]
                    _note={'Oli & Pelumas':'Daftar resmi + mapping manual.','Chemical':'Daftar resmi + mapping manual.'}.get(_kn,'Kode di data yang tidak termasuk dua kategori lain.')
                    if _kn==RK.BAHAN_KODE_LAIN: _codes=[c_ for c_ in _codes if c_ in set(map(str,_txd['Kode Bahan'])) or RK.norm_kode(c_) in RK.BAHAN_MANUAL]
                    _chips=''.join(f'<span class="rk-chp" title="{_html.escape(_nm_by_code.get(RK.norm_kode(c_),""))}">{_html.escape(str(c_))}</span>' for c_ in _codes) or '<span style="color:#94a3b8;font-size:12px">Belum ada kode pada data aktif</span>'
                    st.markdown(f'<div class="rk-mapcol"><div class="t"><span>{_html.escape(_kn)}</span><span class="rk-badge {_tone}">{len(_codes)} kode</span></div><div class="rk-chips">{_chips}</div><small>{_note}</small></div>',unsafe_allow_html=True)
            st.markdown('<div class="rk-div"></div>',unsafe_allow_html=True)
            st.markdown('**Mapping manual** — pindahkan kode bahan ke kategori lain. Mapping manual diutamakan di atas daftar resmi dan tersimpan permanen.')
            _man=load_bahan_manual()
            _allc=sorted({str(c_) for c_ in _txd['Kode Bahan'] if c_}|{c_ for lst in RK.BAHAN_KODE.values() for c_ in lst})
            _mm1,_mm2,_mm3=st.columns([3,2,1.2],vertical_alignment='bottom')
            _sel_codes=_mm1.multiselect('Kode bahan',_allc,key='bmap_codes',placeholder='Pilih satu atau beberapa kode bahan',format_func=lambda c_:f'{c_} — {_nm_by_code.get(RK.norm_kode(c_),"")}'.rstrip(' —'))
            _sel_kat=_mm2.selectbox('Kategori tujuan',RK.BAHAN_KATS,key='bmap_kat')
            def _save_bmap():
                _d=load_bahan_manual()
                for c_ in st.session_state.get('bmap_codes',[]): _d[RK.norm_kode(c_)]=st.session_state.get('bmap_kat')
                save_bahan_manual(_d); st.session_state['bmap_codes']=[]; st.session_state['_bmap_ok']=True
            _mm3.button(':material/save: Simpan',key='bmap_save',type='primary',use_container_width=True,disabled=not _sel_codes,on_click=_save_bmap)
            if st.session_state.pop('_bmap_ok',None): st.success('Mapping manual tersimpan. Kategori pada tabel langsung diperbarui.')
            if _man:
                _mrows=[[f'<span style="font-family:ui-monospace,monospace">{_html.escape(k_)}</span>',_html.escape(_nm_by_code.get(k_,'-')),RK.status_badge(v_,'blue'),RK.status_badge(RK._KODE_TO_KAT.get(k_,RK.BAHAN_KODE_LAIN),'gray')] for k_,v_ in sorted(_man.items())]
                paged_rows(['Kode Bahan','Nama Bahan','Kategori Manual','Kategori Resmi (awal)'],_mrows,'pg_bmap',unit='mapping')
                _del=st.multiselect('Hapus mapping manual (kembali ke kategori resmi)',sorted(_man),key='bmap_del',placeholder='Pilih kode yang ingin dikembalikan')
                def _del_bmap():
                    _d=load_bahan_manual()
                    for k_ in st.session_state.get('bmap_del',[]): _d.pop(k_,None)
                    save_bahan_manual(_d); st.session_state['bmap_del']=[]
                st.button(':material/delete: Hapus mapping terpilih',key='bmap_delbtn',disabled=not _del,on_click=_del_bmap)
            else: st.caption('Belum ada mapping manual.')
        # rincian per cabang
        st.markdown(RK.head_card(f'Rincian Pemakaian & Penjualan Bahan per Cabang ({_html.escape(str(active_branch))})','Bahan Issue To PKB vs bahan tanpa PKB dan porsi non-PKB','apartment','purple'),unsafe_allow_html=True)
        _rbb=revenue_components_by_branch(master,parts,issue,cons,BRANCHES); _rbb=_rbb[_rbb['Cabang']!='TOTAL SEMUA CABANG']
        if active_branch!='Semua Cabang': _rbb=_rbb[_rbb['Cabang']==active_branch]
        _rr=[]; _ta=_tb=0.0
        for _,r_ in _rbb.iterrows():
            _ia=float(r_['Bahan (dengan PKB)']); _ib=float(r_['Bahan tanpa PKB']); _ta+=_ia; _tb+=_ib
            _rr.append([f'<b>{_html.escape(str(r_["Cabang"]))}</b>',fmt(_ia),f'<span style="color:#c2410c">{fmt(_ib)}</span>',fmt(_ia+_ib),RK.share_pct(_ib,_ia+_ib)])
        if _rr:
            _rr.append(['TOTAL',fmt(_ta),fmt(_tb),fmt(_ta+_tb),RK.share_pct(_tb,_ta+_tb)])
            st.markdown(RK.table_html(['Cabang','Bahan Issue To PKB','Bahan Tanpa PKB','Total Bahan','% Porsi Non-PKB'],_rr,num_cols=(1,2,3,4),bold_last=True,hl_cols={3:'hl2'}),unsafe_allow_html=True)
        # ---- Rekapitulasi Penjualan Bahan: saring berdasarkan kode bahan & SA; tampil kode, nama, qty ----
        st.markdown(RK.head_card('Rekapitulasi Penjualan Bahan','Rekap qty terjual per kode bahan. Sortir/saring berdasarkan Kode Bahan dan Nama SA.','filter_alt','blue'),unsafe_allow_html=True)
        _tbl=_txd.rename(columns={'Bahan':'Nama Bahan','_SA':'Nama SA'})
        _tbl['Nama SA']=_tbl['Nama SA'].replace('','-')
        _fc1,_fc2=st.columns(2)
        _f_kode=_fc1.multiselect('Kode Bahan',sorted(_tbl['Kode Bahan'].unique()),key='bt_kode',placeholder='Semua kode bahan')
        _f_sa=_fc2.multiselect('Nama SA',sorted(x for x in _tbl['Nama SA'].unique()),key='bt_sa',placeholder='Semua SA')
        _w=_tbl
        if _f_kode: _w=_w[_w['Kode Bahan'].isin(_f_kode)]
        if _f_sa: _w=_w[_w['Nama SA'].isin(_f_sa)]
        _w=_w.groupby(['Kode Bahan','Nama Bahan'],as_index=False).agg(Qty=('Qty','sum')).sort_values(['Qty','Kode Bahan'],ascending=[False,True]).reset_index(drop=True)
        if _w.empty: st.info('Tidak ada data yang cocok dengan saringan.')
        else:
            paged_df(_w,'pg_bhn_rekap',code=['Kode Bahan'],unit='kode bahan')
            st.caption(f'{len(_w):,} kode bahan • Total Qty {_w["Qty"].sum():,.0f}')
            download_table(_w,':material/download: Excel','rekapitulasi_penjualan_bahan.xlsx','dl_rekap_penjualan_bahan')
    else:
        st.info('Belum ada data Issue To PKB / Consumable. Upload laporan bulanan di menu Data Input.')
    # TOP 10 BAHAN TERBANYAK — gabungan Issue To PKB + Consumable
    bahan_sources=[]
    if not issue.empty:
        bahan_sources.append(issue[['Kode Bahan','Bahan','Qty','Harga Net','Cost']].copy())
    if not cons.empty:
        bahan_sources.append(cons[['Kode Bahan','Bahan','Qty','Harga Net','Cost']].copy())
    if bahan_sources:
        top_bahan=pd.concat(bahan_sources,ignore_index=True)
        top_bahan['Kode Bahan']=top_bahan['Kode Bahan'].map(clean)
        top_bahan['Bahan']=top_bahan['Bahan'].map(clean)
        top_bahan=top_bahan[(top_bahan['Kode Bahan']!='') | (top_bahan['Bahan']!='')].copy()
        top_bahan['Qty']=pd.to_numeric(top_bahan['Qty'],errors='coerce').fillna(0)
        top_bahan['Harga Net']=pd.to_numeric(top_bahan['Harga Net'],errors='coerce').fillna(0)
        top_bahan['Cost']=pd.to_numeric(top_bahan['Cost'],errors='coerce').fillna(0)
        top_bahan=top_bahan.groupby(['Kode Bahan','Bahan'],dropna=False,as_index=False).agg(Qty=('Qty','sum'),Net=('Harga Net','sum'),Cost=('Cost','sum'))
        top_bahan['GP']=top_bahan['Net']-top_bahan['Cost']
        top_bahan=top_bahan.sort_values(['Qty','Net'],ascending=[False,False]).head(10).reset_index(drop=True)
        top_bahan.insert(0,'Rank',range(1,len(top_bahan)+1))
        st.markdown(RK.head_card('Top 10 Penjualan Bahan Terbanyak','Ranking berdasarkan total Qty terjual, gabungan Issue To PKB dan Consumable.','emoji_events','orange'),unsafe_allow_html=True)
        paged_df(top_bahan.rename(columns={'Net':'Omzet','GP':'Gross Profit'}),'pg_bhn_top10',money=['Omzet','Cost','Gross Profit'],code=['Kode Bahan'],unit='bahan')
        download_table(top_bahan,':material/download: Download Excel — Top 10 Bahan Terbanyak','top10_bahan_terbanyak.xlsx','dl_top10_bahan')

    st.markdown(RK.head_card('Bahan Issue To PKB','Rincian item bahan dari laporan Issue To PKB (10 baris per halaman)','description','blue'),unsafe_allow_html=True)
    if issue.empty: st.info('Issue To PKB belum diupload.')
    else:
        ir=issue.groupby([c for c in ['Cabang','Sumber Cabang','No Dokumen','No Referensi','Kode Bahan','Bahan'] if c in issue.columns],dropna=False,as_index=False).agg(Qty=('Qty','sum'),Net=('Harga Net','sum'),Cost=('Cost','sum')).sort_values('Net',ascending=False)
        paged_df(ir.rename(columns={'Net':'Omzet'}),'pg_bhn_issue',money=['Omzet','Cost'],code=['Kode Bahan'],unit='baris'); download_table(ir,':material/download: Download Excel — Bahan Issue To PKB','bahan_issue_to_pkb.xlsx','dl_bahan_issue')
    st.markdown(RK.head_card('Bahan Consumable — Penjualan tanpa PKB','Rincian penjualan bahan consumable langsung (10 baris per halaman)','description','orange'),unsafe_allow_html=True)
    if cons.empty: st.info('Consumable belum diupload.')
    else:
        cr=cons.groupby([c for c in ['Cabang','Customer','No Dokumen','No Referensi','Kode Bahan','Bahan'] if c in cons.columns],dropna=False,as_index=False).agg(Qty=('Qty','sum'),Net=('Harga Net','sum'),Cost=('Cost','sum')).sort_values('Net',ascending=False)
        paged_df(cr.rename(columns={'Net':'Omzet'}),'pg_bhn_cons',money=['Omzet','Cost'],code=['Kode Bahan'],unit='baris'); download_table(cr,':material/download: Download Excel — Bahan Consumable','bahan_consumable.xlsx','dl_bahan_cons')
    st.caption(f'Total Revenue Bahan = {fmt(issue_sales)} + {fmt(cons_sales)} = {fmt(total)}')
if page == '🛠️ OPL':
    section('<span class="msr">build</span> OPL Performance','Monitor vendor, operation, sales, cost dan GP OPL.')
    if opl.empty:
        st.markdown('<div class="upload-note"><b>OPL belum tersedia</b><br>Upload laporan OPL melalui menu Data Input. Setelah tersimpan, data OPL langsung digunakan pada halaman ini dan Revenue & Profit.</div>',unsafe_allow_html=True)
    else:
        opl_operation=opl.groupby('Operation Description',as_index=False).agg(Count=('OPL No','nunique'),Sales=('Sales OPL','sum'),Cost=('Cost OPL','sum'),GP=('Gross Profit OPL','sum')).sort_values('Sales',ascending=False).head(30)
        st.dataframe(opl_operation,use_container_width=True)
        download_table(opl_operation,':material/download: Download Excel — OPL Operation','opl_per_operation.xlsx','dl_opl_operation')
        st.subheader('Vendor'); opl_vendor=opl.groupby('Vendor',as_index=False).agg(Sales=('Sales OPL','sum'),Cost=('Cost OPL','sum'),GP=('Gross Profit OPL','sum')).sort_values('GP',ascending=False).head(20); st.dataframe(opl_vendor,use_container_width=True)
        download_table(opl_vendor,':material/download: Download Excel — OPL Vendor','opl_per_vendor.xlsx','dl_opl_vendor')
if page == '👥 Customer':
    section('<span class="msr">group</span> Customer Intelligence','Frekuensi kunjungan, customer repeat dan customer baru.')
    if cust.empty: st.info('Upload RptPelanggan untuk analisa customer.')
    else:
        cc=cust.copy(); cc['Customer Clean']=cc['Customer'].map(clean); cc=cc[cc['Customer Clean']!=''].copy()
        visits=cc.groupby('Customer Clean',as_index=False).agg(Jumlah_Kunjungan=('Customer Clean','size'),Customer=('Customer','first'),Kendaraan=('Equipment No','nunique'),Last_Visit=('Tgl Service','max'),SA=('Service Advisor','last')).assign(SA=lambda d:sa_disp(d['SA'])).sort_values('Jumlah_Kunjungan',ascending=False)
        visits['Status Customer']=np.where(visits['Jumlah_Kunjungan']<=1,'Customer Baru','Customer Repeat')
        newcust=visits[visits['Jumlah_Kunjungan']==1].copy(); repeat=visits[visits['Jumlah_Kunjungan']>1].copy()
        a,b,c,d=st.columns(4); a.metric('Unique Customer',len(visits)); b.metric('Customer Repeat',len(repeat)); c.metric('Customer Baru',len(newcust)); d.metric('Total Kunjungan',int(visits['Jumlah_Kunjungan'].sum()))
        st.markdown('### :material/repeat: Customer yang sering melakukan kunjungan')
        st.dataframe(repeat[['Customer','Jumlah_Kunjungan','Kendaraan','Last_Visit','SA']],use_container_width=True,hide_index=True); download_table(repeat[['Customer','Jumlah_Kunjungan','Kendaraan','Last_Visit','SA']],':material/download: Download Excel — Customer Repeat','customer_repeat.xlsx','dl_customer_repeat')
        st.markdown('### 🆕 Customer baru — pertama kali melakukan kunjungan')
        st.dataframe(newcust[['Customer','Jumlah_Kunjungan','Kendaraan','Last_Visit','SA']],use_container_width=True,hide_index=True); download_table(newcust[['Customer','Jumlah_Kunjungan','Kendaraan','Last_Visit','SA']],':material/download: Download Excel — Customer Baru','customer_baru.xlsx','dl_customer_new')
        st.markdown('### :material/assignment: Seluruh frekuensi kunjungan customer')
        st.dataframe(visits[['Customer','Jumlah_Kunjungan','Status Customer','Kendaraan','Last_Visit','SA']],use_container_width=True,hide_index=True); download_table(visits[['Customer','Jumlah_Kunjungan','Status Customer','Kendaraan','Last_Visit','SA']],':material/download: Download Excel — Customer Visit Frequency','customer_visit_frequency.xlsx','dl_customer_frequency')
if page == '✅ Quality':
    rec=None
    section('<span class="msr">verified</span> Quality Control','PKB tanpa penjualan tetap dicantumkan untuk evaluasi, tetapi TIDAK dihitung sebagai UE maupun revenue.')
    _qcols=[c for c in ['Cabang','PKB','Tanggal','Tanggal PKB','Service Advisor','No Polisi','Customer','Type Kendaraan','PKB Status','Match Status'] if c in qc_no_sales.columns]
    quality_df=qc_no_sales[_qcols].copy()
    if 'Tanggal PKB' in quality_df.columns: quality_df['Tanggal']=quality_df['Tanggal PKB'].where(quality_df['Tanggal PKB'].notna(),quality_df['Tanggal']); quality_df=quality_df.drop(columns=['Tanggal PKB'])
    quality_df['Keterangan']='Tidak ada penjualan - tidak dihitung UE'
    q_a,q_b,q_c,q_d=st.columns(4)
    q_a.metric('PKB tanpa Sales (tidak dihitung UE)',f"{len(quality_df):,}"); q_b.metric('UE dihitung (ada penjualan)',f"{len(master):,}")
    q_c.metric('Sales PKB unique',f"{len(sales):,}"); q_d.metric('OPL dengan PKB',f"{int(opl.PKB.nunique()) if not opl.empty else 0:,}")
    status_sum=pd.DataFrame()
    if not quality_df.empty:
        st.markdown('**Rincian PKB tanpa penjualan per status**')
        status_sum=quality_df.groupby([c for c in ['Cabang','PKB Status'] if c in quality_df.columns],as_index=False).size().rename(columns={'size':'Jumlah PKB'}).sort_values('Jumlah PKB',ascending=False)
        st.dataframe(status_sum,use_container_width=True,hide_index=True)
        st.markdown('**Daftar PKB tanpa penjualan**')
    st.dataframe(quality_df,use_container_width=True,hide_index=True)
    download_table(quality_df,':material/download: Download Excel — PKB tanpa Sales','quality_pkb_tanpa_sales.xlsx','dl_quality')

    # Cek pemetaan cabang bahan Issue To PKB
    issue_unmapped=pd.DataFrame()
    if isinstance(issue,pd.DataFrame) and not issue.empty and 'Sumber Cabang' in issue.columns:
        st.markdown('### :material/rule: Pemetaan Cabang — Bahan Issue To PKB')
        _src=issue.groupby(['Cabang','Sumber Cabang'],as_index=False).agg(**{'Baris':('Harga Net','size'),'Harga Net':('Harga Net','sum')})
        st.dataframe(_src,use_container_width=True,hide_index=True)
        issue_unmapped=issue[issue['Sumber Cabang']=='Tidak terpetakan'][[c for c in ['Tanggal','No Dokumen','No Referensi','Customer','Kode Bahan','Bahan','Qty','Harga Net','Cabang'] if c in issue.columns]]
        if not issue_unmapped.empty:
            st.warning(f'{len(issue_unmapped):,} baris Issue To PKB tidak ditemukan PKB-nya (sementara masuk cabang upload). Periksa nomor referensinya.')
            st.dataframe(issue_unmapped,use_container_width=True,hide_index=True)
    if EST_REF_ACTIVE:
        st.markdown('### :material/fact_check: Rekonsiliasi Acuan Laporan Estimasi SA')
        q1,q2,q3,q4=st.columns(4)
        q1.metric('PKB cocok dengan laporan',f"{est_stats['matched']:,}"); q2.metric('Pindah bulan (tgl PKB ≠ periode laporan)',f"{est_stats['moved']:,}")
        q3.metric('Ada di laporan, belum di MaintainPKB',f"{est_stats['added']:,}"); q4.metric('Tidak di laporan (tidak dihitung)',f"{est_stats['excluded']:,}")
        _new=master.assign(**{'Bulan':master['Tanggal'].dt.to_period('M').astype(str)}).groupby(['Cabang','Bulan'],as_index=False).agg(**{'Revenue Core (acuan laporan)':('Revenue','sum')})
        _old=master_raw.assign(**{'Bulan':pd.to_datetime(master_raw['Tanggal'],errors='coerce').dt.to_period('M').astype(str)}).groupby(['Cabang','Bulan'],as_index=False).agg(**{'Revenue Core (cara lama: tgl PKB + Service Penjualan)':('Revenue','sum')})
        rec=_new.merge(_old,on=['Cabang','Bulan'],how='outer').fillna(0); rec['Selisih']=rec.iloc[:,2]-rec.iloc[:,3]
        st.dataframe(rec.sort_values(['Bulan','Cabang']),use_container_width=True,hide_index=True)
        download_table(rec,':material/download: Download Excel — Rekonsiliasi Acuan Estimasi','rekonsiliasi_acuan_estimasi.xlsx','dl_rec_est')
        if not est_excluded.empty:
            with st.expander(f'PKB di bulan yang sudah punya laporan tetapi tidak tercantum di laporan ({len(est_excluded):,})'):
                st.dataframe(est_excluded[[c for c in ['Cabang','PKB','Tanggal PKB','Service Advisor','Customer','PKB Status','Revenue'] if c in est_excluded.columns]],use_container_width=True,hide_index=True)
if page == '✅ Quality':
    _qs={'PKB tanpa Sales':quality_df,'Ringkasan per Status':status_sum}
    if 'rec' in globals() and isinstance(rec,pd.DataFrame): _qs['Rekonsiliasi Estimasi SA']=rec
    if isinstance(est_excluded,pd.DataFrame) and not est_excluded.empty: _qs['Tidak di Laporan Estimasi']=est_excluded[[c for c in ['Cabang','PKB','Tanggal PKB','Service Advisor','Customer','PKB Status','Revenue'] if c in est_excluded.columns]]
    if not issue_unmapped.empty: _qs['Issue PKB tak terpetakan']=issue_unmapped
    st.markdown('---')
    dl_button(':material/download: Download Laporan Quality Control (lengkap, semua sheet)',make_multi_sheet_excel(_qs),'laporan_quality_control.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='dl_quality_full',type='primary')
if page == '📋 Data':
    section('<span class="msr">upload_file</span> Data & Export','Pilih filter yang ingin ditampilkan. Semua export di halaman ini otomatis mengikuti hasil filter.')

    # =========================
    # DATA FILTER PANEL
    # =========================
    st.markdown('''<div class="dashboard-card"><h3><span class="msr">search</span> Filter Data</h3><p>Filter hanya berlaku pada halaman Data & Export dan seluruh file yang di-download dari halaman ini.</p></div>''',unsafe_allow_html=True)
    ff=f.copy()
    if not ff.empty:
        # Periode analisis: satu rentang atau beberapa periode sekaligus.
        date_min=ff['Tanggal'].min().date() if 'Tanggal' in ff and ff['Tanggal'].notna().any() else datetime.date.today()
        date_max=ff['Tanggal'].max().date() if 'Tanggal' in ff and ff['Tanggal'].notna().any() else datetime.date.today()
        c1,c2,c3=st.columns(3)
        with c1:
            period_mode=st.radio(':material/calendar_month: Mode Periode',['Satu Rentang','Multi Periode'],horizontal=True,key='data_period_mode')
        with c2:
            selected_dates=None
            selected_periods=[]
            period_gran='Bulan'
            if period_mode=='Satu Rentang':
                selected_dates=st.date_input('Rentang tanggal',value=(date_min,date_max),min_value=date_min,max_value=date_max,key='data_filter_dates')
            else:
                period_gran=st.selectbox('Bandingkan berdasarkan',['Bulan','Kuartal','Tahun'],key='data_period_granularity')
                field={'Bulan':'Periode Bulan','Kuartal':'Periode Kuartal','Tahun':'Periode Tahun'}[period_gran]
                opts=sorted([x for x in ff[field].dropna().astype(str).unique()])
                selected_periods=st.multiselect('Pilih periode (bisa lebih dari satu)',opts,format_func=lambda x:period_key_label(x,period_gran),key='data_filter_periods')
        with c3:
            vehicle_opts=sorted([x for x in ff['Type Kendaraan'].dropna().map(clean).unique() if x]) if 'Type Kendaraan' in ff else []
            selected_vehicle=st.multiselect(':material/directions_car: Tipe Kendaraan',vehicle_opts,key='data_filter_vehicle')
        with c3:
            sa_opts=sorted([x for x in ff['Service Advisor'].dropna().map(clean).unique() if x]) if 'Service Advisor' in ff else []
            selected_sa=st.multiselect(':material/person: Service Advisor',sa_opts,key='data_filter_sa')

        c4,c5,c6=st.columns(3)
        with c4:
            entry_opts=sorted([x for x in ff['Booking Type'].dropna().map(clean).unique() if x]) if 'Booking Type' in ff else []
            selected_entry=st.multiselect(':material/push_pin: Jenis Entry',entry_opts,key='data_filter_entry')
        with c5:
            job_opts=sorted([x for x in ff['Jenis Pekerjaan'].dropna().map(clean).unique() if x]) if 'Jenis Pekerjaan' in ff else []
            selected_job=st.multiselect(':material/build: Jenis Pekerjaan',job_opts,key='data_filter_job')
        with c6:
            status_opts=sorted([x for x in ff['PKB Status'].dropna().map(clean).unique() if x]) if 'PKB Status' in ff else []
            selected_status=st.multiselect(':material/assignment: Status PKB',status_opts,key='data_filter_status')

        c7,c8,c9=st.columns(3)
        with c7:
            match_opts=sorted([x for x in ff['Match Status'].dropna().map(clean).unique() if x]) if 'Match Status' in ff else []
            selected_match=st.multiselect(':material/link: Status Matching',match_opts,key='data_filter_match')
        with c8:
            customer_opts=sorted([x for x in ff['Customer Name Report'].dropna().map(clean).unique() if x]) if 'Customer Name Report' in ff else []
            selected_customer=st.multiselect(':material/group: Customer',customer_opts,key='data_filter_customer')
        with c9:
            police_search=st.text_input(':material/search: Cari No Polisi / PKB',placeholder='Contoh: DD 1234 / PKB',key='data_filter_search').strip().upper()

        # Apply filters
        if period_mode=='Satu Rentang' and isinstance(selected_dates,(tuple,list)) and len(selected_dates)==2 and 'Tanggal' in ff:
            start_date,end_date=selected_dates
            ff=ff[(ff['Tanggal'].dt.date>=start_date)&(ff['Tanggal'].dt.date<=end_date)]
        elif period_mode=='Multi Periode' and selected_periods:
            field={'Bulan':'Periode Bulan','Kuartal':'Periode Kuartal','Tahun':'Periode Tahun'}[period_gran]
            ff=ff[ff[field].astype(str).isin(selected_periods)]
        if selected_vehicle: ff=ff[ff['Type Kendaraan'].map(clean).isin(selected_vehicle)]
        if selected_sa: ff=ff[ff['Service Advisor'].map(clean).isin(selected_sa)]
        if selected_entry: ff=ff[ff['Booking Type'].map(clean).isin(selected_entry)]
        if selected_job: ff=ff[ff['Jenis Pekerjaan'].map(clean).isin(selected_job)]
        if selected_status: ff=ff[ff['PKB Status'].map(clean).isin(selected_status)]
        if selected_match: ff=ff[ff['Match Status'].map(clean).isin(selected_match)]
        if selected_customer: ff=ff[ff['Customer Name Report'].map(clean).isin(selected_customer)]
        if police_search:
            pkb_s=ff['PKB'].map(clean).str.upper() if 'PKB' in ff else pd.Series('',index=ff.index)
            pol_s=ff['No Polisi'].map(clean).str.upper() if 'No Polisi' in ff else pd.Series('',index=ff.index)
            ff=ff[pkb_s.str.contains(re.escape(police_search),na=False)|pol_s.str.contains(re.escape(police_search),na=False)]

    # Filter summary on Data & Export stays compact; KPI cards are reserved for Dashboard.
    total_filtered_ue=int(ff['UE'].sum()) if 'UE' in ff else len(ff)
    total_filtered_rev=float(ff['Revenue'].sum()) if 'Revenue' in ff else 0
    st.markdown(f'''<div class="upload-progress"><b><span class="msr">search</span> {len(ff):,} baris / {len(f):,} baris</b><span>{total_filtered_ue:,} UE • Revenue {fmt(total_filtered_rev)} • semua export mengikuti filter</span></div>''',unsafe_allow_html=True)

    st.dataframe(ff,use_container_width=True,height=560)

    # Semua tombol export menggunakan ff, bukan f.
    download_table(ff,':material/download: Download Excel — Data Sesuai Filter','data_terfilter.xlsx','dl_data_filtered_v112')

    # Workbook ringkas: Master dan Sales hanya untuk PKB yang lolos filter.
    selected_pkbs=set(ff['PKB'].map(clean)) if 'PKB' in ff else set()
    sales_filtered=sales[sales['PKB'].map(clean).isin(selected_pkbs)].copy() if not sales.empty and 'PKB' in sales else sales.copy()
    excel_bytes=make_excel_download(ff,sales_filtered,parts,issue,cons,opl,cust)
    dl_button(':material/download: Download Excel Rapi — Sesuai Filter',excel_bytes,'bengkel_analytics_filtered_export.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='dl_full_filtered_v112')

    template_bytes=make_template_unit_service(ff)
    dl_button(':material/description: Export Template Unit Service — Sesuai Filter',template_bytes,'Template_Unit_Service_KHARISMA_Filtered.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',help='Template diisi hanya dengan data yang lolos filter.',key='dl_template_filtered_v112')
    st.caption(':material/verified: Template Unit Service, Data Excel, Master CSV, dan Sales pada workbook mengikuti data yang sudah difilter. Data sumber tanpa relasi PKB yang valid (Parts/Issue/Consumable/OPL/Customer) tetap ditampilkan sebagai sheet sumber dan tidak dipaksakan match ke filter.')
    dl_button(':material/download: Download Master CSV — Sesuai Filter',ff.to_csv(index=False).encode('utf-8-sig'),'bengkel_analytics_filtered.csv','text/csv',key='dl_csv_filtered_v112')
